#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ALLOCATORE DI CAPITALE — distribuzione goals-based di una somma liquida
=======================================================================

Risponde alla domanda: "ho una somma da investire (es. liquidando un vecchio
gestito) e voglio costruire un portafoglio da zero. In base ai miei obiettivi,
al mio orizzonte e al mio profilo di rischio, quanto del capitale dedico a ogni
secchiello, e qual è la mia allocazione AGGREGATA tra azioni, obbligazioni e
liquidità?"

Logica (goals-based, per valore attuale):
  - Le riserve PERMANENTI (es. fondo di emergenza) si finanziano al VALORE PIENO
    in liquidità OGGI: non sono in crescita verso una data futura, le vuoi
    disponibili adesso. Capitale richiesto = importo di oggi (meno il già
    accantonato).
  - Gli obiettivi TERMINANTI (università, casa…) ricevono il loro PRESENT VALUE:
    il capitale che, al rendimento atteso del profilo, cresce fino al target
    nominale entro la scadenza.  PV = target_nominale / Π(1 + r_anno), dove r_anno
    segue il glide path (quota del profilo, poi discesa negli ultimi anni verso
    l'atterraggio, come il versamento del foglio); senza de-risking è
    target_oggi * ((1+infl)/(1+r))^anni.
  - Si alloca la somma in ordine di PRIORITÀ (essenziale > importante >
    desiderabile), poi per orizzonte crescente, fino a esaurimento.
  - Se la somma AVANZA, l'eccedenza è "capitale libero": va trattata come un
    obiettivo di lungo termine (da formalizzare!) a un profilo configurabile.
  - Se la somma NON BASTA, si finanzia per priorità e il resto si copre con i
    versamenti mensili (lo vede il Monte Carlo).

Poi calcola la VISTA AGGREGATA su TRE classi:
  azioni    = somma su ogni secchiello di (capitale * % azionario del profilo)
  liquidità = la parte non azionaria dei secchielli a profilo Preservazione
              (% azionario 0): riserve tenute liquide
  obbligaz. = il resto della parte non azionaria (sleeve difensivo degli
              obiettivi investiti)

Questo aggregato è l'INPUT che il pianificatore Excel e it-rebalance usano come
allocazione target.

NON è consulenza di investimento. È uno strumento di metodo: propone una
ripartizione coerente e ne mostra le conseguenze; le ipotesi (rendimenti,
profili, trattamento del capitale libero) sono esplicite ed editabili.

Uso:  python allocatore-capitale.py [percorso/pianificatore-obiettivi.xlsx] [capitale]
"""

import os
import sys

from motore_bootstrap import (applica_override, crescita_glide, leggi_finestra_derisk, percorso_json,
                              scrivi_json, verifica_intestazioni_obiettivi)

try:
    from openpyxl import load_workbook
except ImportError:
    print("Serve openpyxl:  pip install openpyxl")
    sys.exit(1)

# ── PARAMETRI ───────────────────────────────────────────────────────────────
CAPITALE_DA_ALLOCARE = 130000.0   # la somma liquida da distribuire (€)
PROFILO_CAPITALE_LIBERO = "Crescita"  # profilo dell'eccedenza se NON si usa lo split sotto
# Se il capitale libero (eccedenza) alimenta più obiettivi di lungo termine "il più possibile"
# (es. integrazione pensione + eredità), specificarne qui la ripartizione: (nome, peso, profilo).
# I pesi si normalizzano. Lista vuota -> un unico bucket a PROFILO_CAPITALE_LIBERO (retro-compatibile).
CAPITALE_LIBERO_SPLIT = [
    ("Libertà finanziaria", 0.30, "Crescita"),
    ("Eredità",             0.70, "Azionario totale"),
]
TRATTA_PERMANENTI_A_VALORE_PIENO = True  # riserve permanenti = valore pieno in liquidità oggi
FINESTRA_DERISK = None   # glide path dei terminanti: None = dal foglio (Parametri!B7); l'atterraggio è
                         # il profilo con la quota azionaria più bassa (Preservazione, 0%)
# ────────────────────────────────────────────────────────────────────────────

# layout del foglio (coerente con pianificatore-obiettivi.xlsx)
INFL_ROW = 6
PROF_START, PROF_N = 10, 6
GOAL_START, GOAL_N = 5, 10

RANK_PRIORITA = {"essenziale": 0, "importante": 1, "desiderabile": 2}


def leggi_excel(path):
    wb = load_workbook(path, data_only=True)
    verifica_intestazioni_obiettivi(wb["Obiettivi"])
    par = wb["Parametri"]
    obj = wb["Obiettivi"]

    inflazione = par.cell(row=INFL_ROW, column=2).value or 0.0
    finestra = leggi_finestra_derisk(par)

    profili = {}
    for i in range(PROF_N):
        r = PROF_START + i
        nome = par.cell(row=r, column=1).value
        rend = par.cell(row=r, column=2).value
        vol = par.cell(row=r, column=3).value
        peso = par.cell(row=r, column=4).value
        if nome:
            profili[nome] = (float(rend), float(vol or 0), float(peso) if peso is not None else 0.0)

    obiettivi = []
    for i in range(GOAL_N):
        r = GOAL_START + i
        nome = obj.cell(row=r, column=1).value
        imp = obj.cell(row=r, column=2).value
        oriz = obj.cell(row=r, column=3).value
        prof = obj.cell(row=r, column=4).value
        gia = obj.cell(row=r, column=6).value
        prio = obj.cell(row=r, column=7).value
        tipo = obj.cell(row=r, column=14).value
        if nome and imp and oriz and prof:
            obiettivi.append({
                "nome": nome,
                "target_oggi": float(imp),
                "orizzonte": int(oriz),
                "profilo": prof,
                "gia": float(gia) if gia else 0.0,
                "priorita": (prio or "").strip().lower(),
                "tipo": (tipo or "terminante").strip().lower(),
            })
    return inflazione, finestra, profili, obiettivi


def present_value(target_oggi, inflazione, rend, anni, rend_atterraggio=None, finestra=0):
    """Capitale da mettere OGGI perché raggiunga il target nominale a scadenza. Con
    rend_atterraggio la crescita segue il glide path (obiettivi terminanti); senza, il
    rendimento r del profilo per tutti gli anni."""
    if rend_atterraggio is None:
        return target_oggi * ((1 + inflazione) / (1 + rend)) ** anni
    return target_oggi * (1 + inflazione) ** anni / crescita_glide(rend, rend_atterraggio, anni, finestra)


def capitale_richiesto(g, inflazione, profili, finestra):
    """Capitale richiesto OGGI per finanziare l'obiettivo, al netto del già accantonato."""
    rend, _, peso = profili[g["profilo"]]
    permanente = g["tipo"].startswith("perman")
    if permanente and TRATTA_PERMANENTI_A_VALORE_PIENO:
        lordo = g["target_oggi"]            # valore pieno, tenuto liquido oggi
    else:
        rend_l, _, peso_l = min(profili.values(), key=lambda p: p[2])   # atterraggio del glide path
        glide = not permanente and peso > peso_l
        lordo = present_value(g["target_oggi"], inflazione, rend, g["orizzonte"],
                              rend_l if glide else None, finestra)
    return max(0.0, lordo - g["gia"]), permanente


def ordina(obiettivi):
    """Per priorità (essenziale prima), poi per orizzonte crescente."""
    return sorted(obiettivi, key=lambda g: (RANK_PRIORITA.get(g["priorita"], 3), g["orizzonte"]))


def alloca(obiettivi, capitale, inflazione, profili, finestra):
    """Allocazione greedy del capitale per priorità/orizzonte. Ritorna le righe e l'eccedenza."""
    righe = []
    residuo = capitale
    for g in ordina(obiettivi):
        rich, permanente = capitale_richiesto(g, inflazione, profili, finestra)
        dato = min(rich, residuo)
        residuo -= dato
        righe.append({
            "nome": g["nome"], "profilo": g["profilo"], "tipo": g["tipo"],
            "permanente": permanente, "richiesto": rich, "allocato": dato,
            "pieno": dato >= rich - 1e-6 and rich > 0,
            "peso_az": profili[g["profilo"]][2],
        })
    return righe, residuo


def buckets_capitale_libero(surplus, profili):
    """Divide il capitale libero tra gli obiettivi di lungo termine (pesi normalizzati).
    Ritorna [(nome, importo, profilo, peso_az)]. Senza split, un unico bucket."""
    if surplus <= 1e-6:
        return []
    if CAPITALE_LIBERO_SPLIT:
        tot = sum(p for _, p, _ in CAPITALE_LIBERO_SPLIT) or 1.0
        return [(nome, surplus * peso / tot, prof, profili.get(prof, (0, 0, 0.65))[2])
                for nome, peso, prof in CAPITALE_LIBERO_SPLIT]
    waz = profili.get(PROFILO_CAPITALE_LIBERO, (0, 0, 0.65))[2]
    return [("Capitale libero", surplus, PROFILO_CAPITALE_LIBERO, waz)]


def aggrega(righe, surplus, profili):
    """Vista aggregata su tre classi: azioni / obbligazioni / liquidità."""
    az = ob = cassa = 0.0
    for r in righe:
        a, w = r["allocato"], r["peso_az"]
        az += a * w
        if w == 0:
            cassa += a * (1 - w)     # riserva tenuta liquida
        else:
            ob += a * (1 - w)        # sleeve difensivo = obbligazioni
    # il capitale libero (diviso tra gli obiettivi di lungo termine): non azionario -> obbligazioni
    for _, imp, _, waz in buckets_capitale_libero(surplus, profili):
        az += imp * waz
        if waz == 0:
            cassa += imp * (1 - waz)
        else:
            ob += imp * (1 - waz)
    return az, ob, cassa


def euro(x):
    return f"€{x:,.0f}".replace(",", ".")


def main():
    applica_override(globals())
    json_out = percorso_json()
    path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "pianificatore-obiettivi.xlsx")
    capitale = float(sys.argv[2]) if len(sys.argv) > 2 else CAPITALE_DA_ALLOCARE
    if not os.path.exists(path):
        print(f"File non trovato: {path}")
        sys.exit(1)

    inflazione, finestra_foglio, profili, obiettivi = leggi_excel(path)
    finestra = finestra_foglio if FINESTRA_DERISK is None else int(FINESTRA_DERISK)
    if PROFILO_CAPITALE_LIBERO not in profili:
        print(f"⚠️  Profilo capitale libero '{PROFILO_CAPITALE_LIBERO}' non in tabella.")
    if not obiettivi:
        print("Nessun obiettivo nel foglio. Compila il pianificatore prima.")
        sys.exit(1)

    righe, residuo = alloca(obiettivi, capitale, inflazione, profili, finestra)
    tot_richiesto = sum(r["richiesto"] for r in righe)
    surplus = max(0.0, residuo)
    deficit = max(0.0, tot_richiesto - capitale)

    print("=" * 78)
    print("ALLOCATORE DI CAPITALE — come distribuire la somma liquida")
    print("=" * 78)
    print(f"Capitale da allocare: {euro(capitale)}  |  inflazione attesa {inflazione*100:.1f}%")
    if CAPITALE_LIBERO_SPLIT:
        dett = ", ".join(f"{n} {int(round(p/(sum(x[1] for x in CAPITALE_LIBERO_SPLIT) or 1)*100))}%" for n, p, _ in CAPITALE_LIBERO_SPLIT)
        print(f"Capitale libero diviso tra: {dett}")
    else:
        print(f"Capitale libero (eccedenza) trattato come profilo: {PROFILO_CAPITALE_LIBERO}")
    print("-" * 78)

    print("\nRIPARTIZIONE PER OBIETTIVO (ordine di priorità, poi orizzonte):")
    print(f"  {'Obiettivo':<22}{'Profilo':<16}{'Richiesto oggi':>16}{'Allocato':>12}  Stato")
    for r in righe:
        base = "valore pieno (liquido)" if r["permanente"] and TRATTA_PERMANENTI_A_VALORE_PIENO else "present value"
        if r["allocato"] <= 1e-6:
            stato = "✗ da versamenti"
        elif r["pieno"]:
            stato = "✓ pieno"
        else:
            stato = "◐ parziale"
        print(f"  {r['nome']:<22}{r['profilo']:<16}{euro(r['richiesto']):>16}{euro(r['allocato']):>12}  {stato}")
        print(f"  {'':<22}{'('+base+')':<16}")

    # eccedenza / deficit
    if surplus > 1e-6:
        for nome, imp, prof, waz in buckets_capitale_libero(surplus, profili):
            print(f"  {nome:<22}{prof:<16}{'—':>16}{euro(imp):>12}  capitale libero")
        if CAPITALE_LIBERO_SPLIT:
            print(f"\nCAPITALE LIBERO: {euro(surplus)} ({surplus/capitale*100:.0f}% del capitale), diviso tra gli")
            print("  obiettivi di lungo termine qui sopra. È la parte che pesa di più sull'aggregato:")
            print("  cambiarne i pesi o i profili sposta l'allocazione complessiva più di ogni altra leva.")
        else:
            print(f"\nCAPITALE LIBERO: {euro(surplus)} eccede gli obiettivi dichiarati ({surplus/capitale*100:.0f}% del capitale).")
            print(f"  Trattato come lungo termine al profilo {PROFILO_CAPITALE_LIBERO}. Definisci un obiettivo")
            print("  reale per questa somma, o usa CAPITALE_LIBERO_SPLIT per dividerla tra più obiettivi.")
    elif deficit > 1e-6:
        print(f"\n⚠️  Il capitale NON copre tutti gli obiettivi: mancano {euro(deficit)} di capitale.")
        print("  Gli obiettivi a priorità più bassa restano da finanziare con i VERSAMENTI mensili")
        print("  (quanto serve al mese lo calcola montecarlo-obiettivi.py). Qui si è dato precedenza")
        print("  agli essenziali e agli orizzonti più vicini.")
    else:
        print("\n  Il capitale copre esattamente gli obiettivi dichiarati (nessuna eccedenza).")

    # vista aggregata
    az, ob, cassa = aggrega(righe, surplus, profili)
    tot = az + ob + cassa
    print("\n" + "=" * 78)
    print("ALLOCAZIONE AGGREGATA — la tua vista finale sul capitale")
    print("-" * 78)
    if tot > 0:
        print(f"  Azioni:        {euro(az):>12}   {az/tot*100:5.1f}%")
        print(f"  Obbligazioni:  {euro(ob):>12}   {ob/tot*100:5.1f}%")
        print(f"  Liquidità:     {euro(cassa):>12}   {cassa/tot*100:5.1f}%")
        print(f"  {'Totale':<14}{euro(tot):>12}   100.0%")
    print("\nÈ questa l'allocazione target che il pianificatore Excel e it-rebalance")
    print("useranno come riferimento.")

    # note di onestà
    print("\n" + "-" * 78)
    print("NOTE")
    print("- Vista 'lump sum': il present value assume di finanziare l'obiettivo col solo")
    print("  capitale. Se versi anche mensilmente, il capitale richiesto SCENDE e l'eccedenza")
    print("  cresce — i versamenti sono la leva separata che gestisce montecarlo-obiettivi.py.")
    print(f"- Terminanti col glide path: il present value sconta il de-risking degli ultimi {finestra}")
    print("  anni (rendimento più basso vicino alla scadenza), come il versamento del foglio.")
    print("- Riserve permanenti al valore pieno: una riserva (fondo emergenza) la vuoi")
    print("  disponibile OGGI, non in crescita verso una scadenza. Spegni con")
    print("  TRATTA_PERMANENTI_A_VALORE_PIENO=False per trattarla a present value come le altre.")
    print("- Obbligazioni vs liquidità: la parte non azionaria dei secchielli a profilo")
    print("  Preservazione è liquidità; quella degli obiettivi investiti è obbligazionaria.")
    print("  Su orizzonti molto brevi valuta di tenere quella quota in liquidità.")
    print("- Liquidare il vecchio gestito (es. fondi) REALIZZA plus/minus: i fondi scontano")
    print("  il 26% sulle plusvalenze (12,5% sulla quota in titoli di Stato), e la plus al")
    print("  rimborso è REDDITO DI CAPITALE, non compensabile con le minus dello zainetto.")
    print("  Guarda con it-tax-optimization l'ordine di smontaggio prima di reinvestire.")
    print("- NON è consulenza di investimento. Ipotesi e profili sono editabili.")

    if json_out:
        scrivi_json(json_out, {
            "capitale": capitale,
            "aggregato": {"azioni": round(az, 2), "obbligazioni": round(ob, 2),
                          "liquidita": round(cassa, 2), "totale": round(tot, 2)},
            "surplus_capitale_libero": round(surplus, 2),
            "deficit": round(deficit, 2),
            "obiettivi": righe,
        })


if __name__ == "__main__":
    main()
