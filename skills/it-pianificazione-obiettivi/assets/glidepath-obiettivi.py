#!/usr/bin/env python3
"""
Glide path automatico per il pianificatore obiettivi.

Genera il CALENDARIO DI DE-RISKING anno per anno: per ogni obiettivo, la quota
azionaria target che dovrebbe avere a ciascun anno fino alla scadenza, e
l'allocazione azionaria complessiva del portafoglio (aggregata, pesata per il
capitale proiettato). È il calendario che alimenta it-rebalance: gli dice come
spostare i target nel tempo, eseguendo il de-risking progressivo.

Logica del glide path (per ogni obiettivo):
  - tiene la quota azionaria di partenza (del profilo) finché mancano più anni
    della "finestra di de-risking"
  - poi scende LINEARMENTE fino alla quota di atterraggio entro la scadenza
  - la finestra effettiva è min(finestra, orizzonte): un obiettivo corto de-riska
    sull'intera vita, uno lungo solo negli ultimi anni

Aggregato di portafoglio:
  - per ogni anno, pesa la quota azionaria di ogni obiettivo ATTIVO per il suo
    capitale proiettato a quell'anno (crescita al rendimento del profilo +
    versamenti); un obiettivo "esce" dal portafoglio quando raggiunge la scadenza

Dipendenze: openpyxl  (pip install openpyxl)
Uso:        python glidepath-obiettivi.py [percorso/pianificatore-obiettivi.xlsx]

NON è consulenza di investimento. Il glide path è una regola metodologica, le %
sono esempi; la proiezione del capitale usa il rendimento medio del profilo (non
modella la volatilità). Vedi glide-path.md per i limiti e l'interpretazione.
"""

import sys
import os
from openpyxl import load_workbook

from motore_bootstrap import applica_override, verifica_intestazioni_obiettivi

# ---- Parametri del glide path (editabili) ----
FINESTRA_DERISK = 7      # anni prima della scadenza in cui inizia il de-risking
ATTERRAGGIO_AZ = 0.00    # quota azionaria target alla scadenza (0% = tutto liquido/sicuro)
CURVATURA = 1.0          # forma della rampa: 1.0 = lineare; <1 tiene la quota azionaria
                         #   più alta all'inizio della finestra e scende ripida vicino alla scadenza
                         #   (più crescita, de-risking tardivo); >1 de-riska prima e più dolce dopo

# Posizioni note nel foglio (coerenti con build_obiettivi.py)
CAP_ROW = 5
INFL_ROW = 6
PROF_START = 10
PROF_N = 6
GOAL_START = 5
GOAL_N = 10


def leggi_excel(path):
    wb = load_workbook(path, data_only=True)
    verifica_intestazioni_obiettivi(wb["Obiettivi"])
    par = wb["Parametri"]
    obj = wb["Obiettivi"]
    inflazione = par.cell(row=INFL_ROW, column=2).value or 0
    profili = {}
    for i in range(PROF_N):
        r = PROF_START + i
        nome = par.cell(row=r, column=1).value
        rend = par.cell(row=r, column=2).value
        eq = par.cell(row=r, column=4).value  # % azionario (colonna D)
        if nome:
            profili[nome] = (float(rend), float(eq))
    obiettivi = []
    for i in range(GOAL_N):
        r = GOAL_START + i
        nome = obj.cell(row=r, column=1).value
        imp = obj.cell(row=r, column=2).value
        oriz = obj.cell(row=r, column=3).value
        prof = obj.cell(row=r, column=4).value
        gia = obj.cell(row=r, column=6).value
        tipo = obj.cell(row=r, column=14).value  # colonna N: terminante / permanente
        if nome and oriz and prof:
            obiettivi.append({
                "nome": nome,
                # target opzionale: vuoto = "il più possibile" (si proietta il solo capitale)
                "target_oggi": float(imp) if imp else None,
                "orizzonte": int(oriz),
                "profilo": prof, "gia": float(gia) if gia else 0.0,
                # default: terminante (caso comune); permanente solo se indicato
                "permanente": (str(tipo).strip().lower() == "permanente"),
            })
    return inflazione, profili, obiettivi


def pmt_mensile(target_nom, gia, rend, anni):
    if rend == 0:
        return max(0.0, (target_nom - gia) / (anni * 12))
    fattore = (((1 + rend) ** anni) - 1) / ((1 + rend) ** (1 / 12) - 1)
    return max(0.0, (target_nom - gia * (1 + rend) ** anni) / fattore)


def equity_glide(anni_residui, eq_start, finestra, atterraggio):
    """Quota azionaria target dato il tempo residuo all'obiettivo."""
    if eq_start <= atterraggio:
        return eq_start                      # già conservativo: niente da de-riskare
    if anni_residui <= 0:
        return atterraggio
    if anni_residui >= finestra:
        return eq_start                      # ancora lontano: tiene la quota piena
    # rampa dentro la finestra: frazione di azionario (sopra l'atterraggio) ancora trattenuta
    frazione = (anni_residui / finestra) ** CURVATURA
    return atterraggio + (eq_start - atterraggio) * frazione


def pct(x):
    return f"{x*100:.0f}%"


def main():
    applica_override(globals())
    if len(sys.argv) > 1:
        path = sys.argv[1]
    else:
        here = os.path.dirname(os.path.abspath(__file__))
        path = os.path.join(here, "pianificatore-obiettivi.xlsx")
    if not os.path.exists(path):
        print(f"File non trovato: {path}")
        sys.exit(1)

    inflazione, profili, obiettivi = leggi_excel(path)
    if not obiettivi:
        print("Nessun obiettivo trovato nel foglio.")
        sys.exit(0)

    orizzonte_max = max(g["orizzonte"] for g in obiettivi)

    # pre-calcolo per ogni obiettivo: pmt, finestra, capitale proiettato anno per anno
    for g in obiettivi:
        rend, eq = profili.get(g["profilo"], (0.0, 0.0))
        g["rend"], g["eq_start"] = rend, eq
        g["finestra"] = min(FINESTRA_DERISK, g["orizzonte"])
        if g["target_oggi"] is None:
            # "il più possibile": nessun versamento "necessario" da calcolare — si proietta
            # il solo capitale assegnato (eventuali versamenti liberi sono una scelta, non un obbligo)
            g["pmt"] = 0.0
        else:
            target_nom = g["target_oggi"] * (1 + inflazione) ** g["orizzonte"]
            g["pmt"] = pmt_mensile(target_nom, g["gia"], rend, g["orizzonte"])
        # proiezione capitale fino a orizzonte_max:
        #  - fino alla scadenza: cresce al rendimento del profilo + versamenti
        #  - oltre la scadenza: TERMINANTE = irrilevante (esce); PERMANENTE = tenuto costante (riserva piena)
        cap = [g["gia"]]
        for t in range(1, orizzonte_max + 1):
            if t <= g["orizzonte"]:
                cap.append(cap[-1] * (1 + rend) + 12 * g["pmt"])
            else:
                cap.append(cap[-1] if g["permanente"] else 0.0)
        g["cap"] = cap

    print("=" * 78)
    print("GLIDE PATH — CALENDARIO DI DE-RISKING")
    print("=" * 78)
    print(f"Finestra de-risking: {FINESTRA_DERISK} anni  |  Atterraggio: {pct(ATTERRAGGIO_AZ)} azionario alla scadenza")
    print(f"Inflazione: {inflazione*100:.1f}%")
    print("-" * 78)

    # intestazione tabella per-obiettivo
    nomi = [g["nome"] + (" (P)" if g["permanente"] else "") + (" (max)" if g["target_oggi"] is None else "")
            for g in obiettivi]
    intest = "Anno  " + "  ".join(f"{n[:16]:>16}" for n in nomi) + "   |  PORTAFOGLIO"
    print("\nQUOTA AZIONARIA TARGET per obiettivo e aggregata (per anno).  (P) = permanente\n")
    print(intest)
    print("-" * len(intest))

    for t in range(orizzonte_max + 1):
        celle = []
        num = 0.0   # capitale azionario aggregato
        den = 0.0   # capitale totale aggregato
        for g in obiettivi:
            attivo = (t <= g["orizzonte"]) or g["permanente"]
            if not attivo:
                celle.append(f"{'concl.':>16}")
                continue
            if g["permanente"]:
                # nessun glide: una riserva permanente tiene la quota del profilo (non ha scadenza di spesa)
                eqt = g["eq_start"]
                etich = pct(eqt) + (" (piena)" if t == g["orizzonte"] else "")
            else:
                ar = g["orizzonte"] - t
                eqt = equity_glide(ar, g["eq_start"], g["finestra"], ATTERRAGGIO_AZ)
                etich = pct(eqt)
            celle.append(f"{etich:>16}")
            capt = g["cap"][t]
            num += capt * eqt
            den += capt
        agg = (num / den) if den > 0 else 0.0
        riga = f"{t:>4}  " + "  ".join(celle) + f"   |  {pct(agg):>10}"
        print(riga)

    # transizioni chiave dell'aggregato
    print("\n" + "-" * 78)
    print("LETTURA:")
    print("- Ogni colonna obiettivo è il target che it-rebalance mantiene per QUEL secchiello.")
    print("- La colonna PORTAFOGLIO è l'allocazione azionaria complessiva attesa anno per anno,")
    print("  pesata per il capitale proiettato di ogni obiettivo ancora attivo.")
    print("- 'concl.' = obiettivo TERMINANTE arrivato a scadenza (si spende → esce dal portafoglio).")
    print("- (P) = obiettivo PERMANENTE: NON de-riska (tiene la quota del profilo) e NON esce a scadenza")
    print("  (è una riserva, es. fondo emergenza). '(piena)' = anno in cui è pienamente finanziato.")
    print("  Restando in portafoglio, smussa l'aggregato invece di farlo saltare.")
    print("- (max) = obiettivo SENZA target («il più possibile»): si proietta il solo capitale")
    print("  assegnato; nessun versamento 'necessario' viene calcolato. Se terminante, de-riska")
    print("  e a scadenza si spende come gli altri; se permanente, resta e non de-riska mai.")
    print("\nNON è consulenza di investimento. Proiezione al rendimento medio del profilo")
    print("(non modella la volatilità). Le % azionario per profilo e i parametri del glide")
    print("path sono editabili. Vedi glide-path.md.")


if __name__ == "__main__":
    main()
