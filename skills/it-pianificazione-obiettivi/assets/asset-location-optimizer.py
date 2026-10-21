#!/usr/bin/env python3
"""
Ottimizzatore quantitativo dell'asset location.

Calcola il valore finale AL NETTO DI TASSE E COSTI in ciascun contenitore fiscale
italiano, per due decisioni dove il beneficio è quantificabile:

  1. FONDO PENSIONE vs DOSSIER (per l'obiettivo pensione): la deduzione IRPEF è il
     driver, ma i rendimenti tassati ogni anno al 20% e i costi più alti la erodono.
  2. PIR vs DOSSIER (per obiettivi a lungo termine): l'esenzione (0% sulle plus)
     contro il TER più alto del PIR. Il risultato è spesso controintuitivo: su
     orizzonti lunghi il costo del PIR può superare il risparmio fiscale.

Mostra anche il TER di pareggio del PIR: sotto quel costo il PIR conviene, sopra no.

Rendimento LORDO: i profili del pianificatore sono rendimenti NETTI di costi e imposte
(netto ≈ (lordo − costi) × (1 − 26%)); qui invece costi e imposte si tolgono contenitore
per contenitore, quindi lo script parte dal rendimento lordo di mercato. Se
RENDIMENTO_LORDO è None lo ricava dal netto del profilo con la conversione inversa
(lordo_da_netto); si può fissare esplicitamente con --set RENDIMENTO_LORDO=0.07.
La conversione tratta l'imposta come pagata ogni anno: dà il lordo più alto compatibile col
netto, e un lordo più alto aumenta il valore delle esenzioni, quindi favorisce i contenitori
esenti. Per questo, per ogni obiettivo, lo script mostra anche la variante con l'imposta pagata
alla vendita (ETF ad accumulazione, lordo_da_netto_differito) e il pareggio del PIR che ne esce.

Dipendenze: openpyxl  (pip install openpyxl)
Uso:        python asset-location-optimizer.py [percorso/pianificatore-obiettivi.xlsx]

⚠️ NON è consulenza di investimento né fiscale. Calcoli illustrativi a parametri
editabili. La comparazione PIR/dossier assume LO STESSO rendimento: NON cattura il
costo di concentrazione del PIR (PMI italiane, ≈2% del mondo), che è un rischio
aggiuntivo che il risparmio fiscale deve compensare. Il fondo pensione è vincolato
fino alla pensione. Verificare con un consulente abilitato e un commercialista.
Vedi asset-location.md.
"""

import sys
import os
from openpyxl import load_workbook

from motore_bootstrap import applica_override, verifica_intestazioni_obiettivi

# ============================================================
# PARAMETRI EDITABILI
# ============================================================
ALIQUOTA_IRPEF_MARGINALE = 0.33   # aliquota IRPEF marginale (per la deduzione fondo pensione)

# Costi annui (TER / ISC)
TER_DOSSIER = 0.0020              # ETF efficiente nel dossier (~0,2%)
ISC_FONDO_PENSIONE = 0.0050       # fondo pensione (negoziale ~0,2-0,5%, aperto 0,8-1,5%, PIP 1,5-3%)
TER_PIR = 0.0130                  # prodotti PIR tipici (1,2-1,35%)
BOLLO = 0.0020                   # imposta di bollo 0,2%/anno (dossier e PIR; il fondo pensione ne è esente)

# Tassazione
TASSA_RENDIMENTI_FP = 0.20        # imposta sostitutiva annua sui rendimenti del fondo pensione
TASSA_PLUS_DOSSIER = 0.26         # plusvalenze nel dossier (alla realizzazione)
TASSA_PLUS_PIR = 0.00             # plusvalenze nel PIR dopo 5 anni

# Scenario PENSIONE (il fondo pensione è vincolato → si valuta sull'obiettivo pensione, non sui goal a scadenza)
CONTRIBUTO_ANNUO_FP = 5300        # versamento annuo deducibile (max deducibile 2026: €5.300)
ANNI_A_PENSIONE = 30             # anni al pensionamento

# Rendimento LORDO di mercato (prima di costi e imposte), uguale per tutti i contenitori.
# None = ricavato dal netto del profilo (Crescita per la pensione, il profilo dell'obiettivo
# per il PIR): lordo = netto / (1 − TASSA_PLUS_DOSSIER) + TER_DOSSIER + BOLLO.
RENDIMENTO_LORDO = None

# Soglia: gli obiettivi con orizzonte >= questa sono valutati per il PIR
ORIZZONTE_MIN_PIR = 7
REINVESTI_RISPARMIO_IRPEF = True  # True: il risparmio IRPEF è reinvestito nel dossier (realistico);
                                  # False: contato come cassa nominale non reinvestita

# Posizioni note nel foglio
INFL_ROW = 6
PROF_START = 10
PROF_N = 6
GOAL_START = 5
GOAL_N = 10


def leggi_excel(path):
    wb = load_workbook(path, data_only=True)
    verifica_intestazioni_obiettivi(wb["Obiettivi"])
    par = wb["Parametri"]
    obj = wb["Obiettivi"]
    inflazione = par.cell(row=INFL_ROW, column=2).value or 0
    profili = {}
    for i in range(PROF_N):
        r = PROF_START + i
        nome = par.cell(row=r, column=1).value
        rend = par.cell(row=r, column=2).value
        if nome:
            profili[nome] = float(rend)
    obiettivi = []
    for i in range(GOAL_N):
        r = GOAL_START + i
        nome = obj.cell(row=r, column=1).value
        imp = obj.cell(row=r, column=2).value
        oriz = obj.cell(row=r, column=3).value
        prof = obj.cell(row=r, column=4).value
        gia = obj.cell(row=r, column=6).value
        if nome and imp and oriz and prof:
            obiettivi.append({
                "nome": nome, "orizzonte": int(oriz), "profilo": prof,
                "gia": float(gia) if gia else 0.0,
            })
    return inflazione, profili, obiettivi


def lordo_da_netto(netto):
    """Inverso della convenzione dei profili, netto ≈ (lordo − costi) × (1 − 26%): il
    rendimento lordo di mercato da cui l'ottimizzatore toglie costi e imposte di ogni
    contenitore. Tratta il netto come tassato ogni anno: per il versamento è la scelta
    prudente, qui no, perché dà il lordo più alto compatibile col netto e quindi favorisce i
    contenitori esenti (vedi lordo_da_netto_differito)."""
    return netto / (1 - TASSA_PLUS_DOSSIER) + TER_DOSSIER + BOLLO


def lordo_da_netto_differito(netto, anni):
    """Lordo equivalente al netto con l'imposta pagata alla vendita (ETF ad accumulazione nel
    dossier): il lordo per cui, dopo costi e plusvalenza tassata alla fine, il capitale cresce
    come al netto composto per `anni`. È più basso di lordo_da_netto: lo scenario prudente per
    il confronto con i contenitori esenti."""
    finale = ((1 + netto) ** anni - TASSA_PLUS_DOSSIER) / (1 - TASSA_PLUS_DOSSIER)
    return finale ** (1 / anni) - 1 + TER_DOSSIER + BOLLO


def rendimento_lordo(netto_profilo):
    return RENDIMENTO_LORDO if RENDIMENTO_LORDO is not None else lordo_da_netto(netto_profilo)


def aliquota_finale_fp(anni_iscrizione):
    """15% fino al 15° anno, poi -0,3%/anno fino al minimo 9%."""
    return max(0.09, 0.15 - 0.003 * max(0, anni_iscrizione - 15))


def montante_fondo_pensione(c_annuo, anni, r):
    """Montante del fondo pensione: rendimenti tassati ogni anno al 20%, niente bollo."""
    saldo = 0.0
    for _ in range(anni):
        crescita = saldo * (r - ISC_FONDO_PENSIONE)
        tassa = max(0.0, crescita) * TASSA_RENDIMENTI_FP
        saldo = saldo + crescita - tassa + c_annuo
    return saldo


def montante_dossier(c_annuo, anni, r, gia=0.0):
    """Montante lordo nel dossier (ETF accumulazione: tassa differita alla fine), con bollo."""
    saldo = gia
    for _ in range(anni):
        saldo = saldo * (1 + r - TER_DOSSIER - BOLLO) + c_annuo
    return saldo


def dossier_netto(c_annuo, anni, r, gia=0.0):
    m = montante_dossier(c_annuo, anni, r, gia)
    versato = gia + c_annuo * anni
    plus = max(0.0, m - versato)
    return m - plus * TASSA_PLUS_DOSSIER


def pir_netto(gia, anni, r, ter):
    """Montante PIR netto (plus 0% dopo 5 anni), con bollo e TER del PIR."""
    saldo = gia
    for _ in range(anni):
        saldo = saldo * (1 + r - ter - BOLLO)
    # plus esente (0%) se >= 5 anni
    return saldo  # tassa plus 0


def ter_pareggio_pir(gia, anni, r):
    """Trova (bisezione) il TER del PIR a cui PIR netto = dossier netto."""
    target = dossier_netto(0, anni, r, gia)
    lo, hi = 0.0, 0.10
    for _ in range(60):
        mid = (lo + hi) / 2
        if pir_netto(gia, anni, r, mid) > target:
            lo = mid   # PIR ancora meglio: posso permettermi TER più alto
        else:
            hi = mid
    return (lo + hi) / 2


def euro(x):
    s = "-" if x < 0 else ""
    return f"{s}€{abs(x):,.0f}".replace(",", ".")


def main():
    applica_override(globals())
    if len(sys.argv) > 1:
        path = sys.argv[1]
    else:
        here = os.path.dirname(os.path.abspath(__file__))
        path = os.path.join(here, "pianificatore-obiettivi.xlsx")

    inflazione, profili, obiettivi = (0.02, {}, [])
    if os.path.exists(path):
        inflazione, profili, obiettivi = leggi_excel(path)
    # rendimento di riferimento per lo scenario pensione: profilo Crescita se presente, portato
    # al LORDO (i profili sono netti di costi e imposte, qui si tolgono contenitore per contenitore)
    netto_crescita = profili.get("Crescita", 0.055)
    r_pensione = rendimento_lordo(netto_crescita)

    print("=" * 80)
    print("OTTIMIZZATORE ASSET LOCATION — beneficio fiscale netto per contenitore")
    print("=" * 80)
    print(f"Aliquota IRPEF marginale: {ALIQUOTA_IRPEF_MARGINALE*100:.0f}%  |  "
          f"TER dossier {TER_DOSSIER*100:.1f}%  ·  ISC fondo pensione {ISC_FONDO_PENSIONE*100:.1f}%  ·  TER PIR {TER_PIR*100:.1f}%")
    if RENDIMENTO_LORDO is None:
        print(f"Rendimento lordo di mercato {r_pensione*100:.1f}% = netto del profilo Crescita "
              f"{netto_crescita*100:.1f}% / (1 − {TASSA_PLUS_DOSSIER*100:.0f}%) + TER {TER_DOSSIER*100:.1f}% "
              f"+ bollo {BOLLO*100:.1f}%")
        print("  (i profili sono netti di costi e imposte; qui si tolgono contenitore per contenitore)")
    else:
        print(f"Rendimento lordo di mercato {r_pensione*100:.1f}% (RENDIMENTO_LORDO esplicito)")

    # ---------------------------------------------------------
    # 1) FONDO PENSIONE vs DOSSIER (pensione)
    # ---------------------------------------------------------
    print("\n" + "-" * 80)
    print(f"1) FONDO PENSIONE vs DOSSIER  —  €{CONTRIBUTO_ANNUO_FP:,}/anno per {ANNI_A_PENSIONE} anni "
          f"@ {r_pensione*100:.1f}% lordo".replace(",", "."))
    print("-" * 80)

    c = CONTRIBUTO_ANNUO_FP
    H = ANNI_A_PENSIONE
    versato = c * H

    m_fp = montante_fondo_pensione(c, H, r_pensione)
    aliq_fin = aliquota_finale_fp(H)
    tassa_finale_fp = versato * aliq_fin            # base ~ contributi dedotti
    fp_netto = m_fp - tassa_finale_fp
    risparmio_annuo = c * ALIQUOTA_IRPEF_MARGINALE
    if REINVESTI_RISPARMIO_IRPEF:
        # il risparmio IRPEF annuo viene reinvestito nel dossier (cresce, plus 26% alla fine)
        risparmio_irpef = dossier_netto(risparmio_annuo, H, r_pensione)
        etichetta_irpef = f"risparmio IRPEF reinvestito nel dossier ({ALIQUOTA_IRPEF_MARGINALE*100:.0f}% della deduzione)"
    else:
        risparmio_irpef = risparmio_annuo * H   # cassa nominale non reinvestita
        etichetta_irpef = f"risparmio IRPEF cumulato ({ALIQUOTA_IRPEF_MARGINALE*100:.0f}% della deduzione)"
    fp_totale = fp_netto + risparmio_irpef

    doss = dossier_netto(c, H, r_pensione)

    print(f"  Versato nominale: {euro(versato)}  ({euro(c)}/anno)")
    print(f"  FONDO PENSIONE:")
    print(f"    montante lordo (rend. tassati 20%/anno, ISC {ISC_FONDO_PENSIONE*100:.1f}%): {euro(m_fp)}")
    print(f"    − tassa finale {aliq_fin*100:.1f}% sui contributi dedotti:               −{euro(tassa_finale_fp)}")
    print(f"    + {etichetta_irpef}:")
    print(f"      {' '*54}+{euro(risparmio_irpef)}")
    print(f"    = TOTALE netto fondo pensione:                              {euro(fp_totale)}")
    print(f"  DOSSIER (ETF accumulazione, plus 26% alla fine, bollo 0,2%):  {euro(doss)}")
    vantaggio = fp_totale - doss
    print(f"  → VANTAGGIO FONDO PENSIONE: {euro(vantaggio)}  ({vantaggio/doss*100:+.0f}% vs dossier)")
    print(f"  Driver: la deduzione IRPEF ({euro(risparmio_irpef)}). ⚠️ Capitale VINCOLATO fino alla pensione.")
    nota_reinv = "reinvestito nel dossier" if REINVESTI_RISPARMIO_IRPEF else "contato come cassa non reinvestita (reinvestendolo il vantaggio cresce)"
    print(f"  Nota: il risparmio IRPEF è {nota_reinv}.")

    # ---------------------------------------------------------
    # 2) PIR vs DOSSIER (obiettivi a lungo termine)
    # ---------------------------------------------------------
    print("\n" + "-" * 80)
    print(f"2) PIR vs DOSSIER  —  obiettivi a lungo termine (orizzonte ≥ {ORIZZONTE_MIN_PIR} anni)")
    print("-" * 80)

    lunghi = [g for g in obiettivi if g["orizzonte"] >= ORIZZONTE_MIN_PIR and g["gia"] > 0]
    if not lunghi:
        # scenario di default se non ci sono goal idonei
        lunghi = [{"nome": "Esempio: €10.000 a 15 anni", "orizzonte": 15, "profilo": "Crescita", "gia": 10000.0}]
        print("  (nessun obiettivo idoneo nel foglio: uso uno scenario di esempio)")

    for g in lunghi:
        r = rendimento_lordo(profili.get(g["profilo"], 0.055))
        gia = g["gia"]
        H = g["orizzonte"]
        doss = dossier_netto(0, H, r, gia)
        pir = pir_netto(gia, H, r, TER_PIR)
        netto = pir - doss
        ter_par = ter_pareggio_pir(gia, H, r)
        print(f"\n  ▸ {g['nome']}  ({euro(gia)} di capitale, {H} anni @ {r*100:.1f}% lordo, profilo {g['profilo']})")
        print(f"      DOSSIER (plus 26%, TER {TER_DOSSIER*100:.1f}%):      {euro(doss)}")
        print(f"      PIR (plus 0%, TER {TER_PIR*100:.1f}%):           {euro(pir)}")
        segno = "VANTAGGIO PIR" if netto >= 0 else "SVANTAGGIO PIR"
        print(f"      → {segno}: {euro(netto)}")
        print(f"      TER di pareggio del PIR: {ter_par*100:.2f}%  "
              f"(sotto questo costo il PIR conviene, sopra no — TER tipici PIR: 1,2-1,35%)")
        if RENDIMENTO_LORDO is None:
            r_diff = lordo_da_netto_differito(profili.get(g["profilo"], 0.055), H)
            ter_par_diff = ter_pareggio_pir(gia, H, r_diff)
            print(f"      Con l'imposta alla vendita (ETF ad accumulazione) il lordo equivalente a {H} anni "
                  f"è {r_diff*100:.2f}%:")
            print(f"      il pareggio scende a {ter_par_diff*100:.2f}% (la conversione annua favorisce "
                  f"il PIR; scenario prudente: --set RENDIMENTO_LORDO={r_diff:.4f})")

    print("\n" + "-" * 80)
    print("LETTURA:")
    print("- Il fondo pensione domina per la PENSIONE grazie alla deduzione, ma è vincolato:")
    print("  non è utilizzabile per obiettivi che si spendono prima (casa, università).")
    print("- Il PIR azzera le tasse sulle plus, ma il TER più alto le erode: su orizzonti lunghi")
    print("  il costo può SUPERARE il risparmio fiscale (vantaggio PIR negativo). Conta il TER")
    print("  effettivo del prodotto rispetto al TER di pareggio.")
    print("- ⚠️ La comparazione PIR/dossier assume lo STESSO rendimento: NON include il costo di")
    print("  concentrazione del PIR (PMI italiane). È un rischio in più che il fisco deve ripagare.")
    print("- Il verdetto PIR dipende dal rendimento LORDO: più è alto, più vale l'esenzione delle")
    print("  plus. Trattare come lordo il rendimento netto dei profili lo ribalterebbe; per un'altra")
    print("  ipotesi: --set RENDIMENTO_LORDO=0.07. Il lordo ricavato dal netto tratta l'imposta come")
    print("  annua ed è il più alto compatibile col profilo: favorisce i contenitori esenti (PIR e")
    print("  fondo pensione); la variante con l'imposta alla vendita è lo scenario prudente.")
    print("\nNON è consulenza. Calcoli illustrativi, parametri editabili in cima allo script.")
    print("Vedi asset-location.md. Verificare con consulente abilitato e commercialista.")


if __name__ == "__main__":
    main()
