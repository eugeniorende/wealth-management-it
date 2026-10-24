#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
genera-ips.py — genera l'Investment Policy Statement (IPS) come documento Word.

L'IPS è il deliverable finale del flusso: traduce l'analisi (obiettivi, allocazione,
glide path, pensione) nel documento operativo che l'investitore esegue e rilegge.
Questo script lo produce da un unico blocco DATI qui sotto, con testi di policy
standard (ribilanciamento, guardrail, fiscalità italiana, checklist) come default
sovrascrivibili: compila i dati del TUO caso e il documento esce completo.

Uso:
    python genera-ips.py [output.docx]        # default: IPS.docx nella directory corrente

Dipendenze: python-docx  (pip install python-docx)

Struttura del documento: copertina, 12 sezioni (scopo, profilo, obiettivi, allocazione,
glide path, ribilanciamento, versamenti, strumenti, fiscalità, previdenza, guardrail,
revisione), Appendice A (checklist operativa), Appendice B (broker, opzionale), nota finale.

ONESTÀ METODOLOGICA
- La nota di rischio sull'allocazione è calcolata dalla quota azionaria aggregata
  (drawdown atteso in un orso tipico ≈ metà della quota azionaria): il documento
  non nasconde mai il rischio dietro il piano.
- Lo script valida le coerenze di base (allocazione = 100%, capitali = totale) e
  AVVISA a video se non tornano — ma genera comunque: decidere resta compito tuo.
Questo strumento è educativo e non costituisce consulenza finanziaria.
"""

import argparse
import datetime
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

def numero(x, decimali=0):
    """Numero in formato italiano: punto per le migliaia, virgola per i decimali (come
    motore_bootstrap.numero nelle skill previdenziali: questa skill non ha il motore)."""
    testo = f"{x:,.{decimali}f}"
    return testo.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


# ═══════════════════════════════ DATI ═══════════════════════════════
# Compila questo blocco con i risultati del flusso (allocatore, glide path,
# proiezione pensione). Le chiavi assenti usano i default di policy più sotto.

DATI = {
    "investitore": "Mario Rossi",
    "data": "",                        # vuota = mese/anno correnti
    "capitale": 100_000,
    "sottotitolo": "Impostazione di un portafoglio da zero, su base obiettivi",

    "profilo": [
        ("Età", "35 anni"),
        ("Reddito (RAL)", "€32.000 lordi/anno"),
        ("Tolleranza al rischio (auto-valutata)", "5 su una scala 1–8"),
        ("Capacità di risparmio", "€300–600/mese"),
        ("Capitale da dispiegare", "€100.000"),
        ("Liquidità operativa esclusa", "€3.000 sul conto corrente (cuscinetto, non investito)"),
    ],

    # nome, orizzonte/tipo, profilo (% azioni), target, capitale assegnato (€)
    "obiettivi": [
        ("Fondo emergenza", "permanente", "Preservazione (0%)", "€12.000", 12_000),
        ("Casa", "10 anni · terminante", "Moderato (40%)", "€60.000", 45_000),
        ("Lungo termine", "30 anni · permanente", "Crescita massima (90%)", "il più possibile", 43_000),
    ],
    "obiettivi_nota": ("Il fondo emergenza è tenuto a valore pieno in liquidità. Gli obiettivi con "
                       "target sono finanziati al **valore attuale**; quelli \u201cil più possibile\u201d "
                       "ricevono il capitale residuo."),

    # classe, importo (€), quota (%)
    "allocazione": [
        ("Azioni", 56_700, 56.7),
        ("Obbligazioni", 31_300, 31.3),
        ("Liquidità", 12_000, 12.0),
    ],
    "allocazione_nota": ("Il blocco più grande e privo di scadenza (\u201clungo termine\u201d) guida "
                         "l'aggregato più di ogni altro input: è la leva su cui tornare se il rischio "
                         "complessivo non convince."),

    # fase, anni, azioni aggregate, cosa succede
    "glide_path": [
        ("Partenza", "oggi", "57%", "allocazione iniziale"),
        ("Deriva \u2191", "1–3", "57 \u2192 60%", "le azioni compongono più della liquidità"),
        ("De-risk casa", "4–10", "60 \u2192 45%", "l'obiettivo casa si mette al sicuro"),
        ("Casa esce", "11+", "~75%", "resta il lungo termine (90%) + riserve"),
    ],

    # (facoltativi: se assenti si usano i default di policy qui sotto)
    # "ribilanciamento": [...], "guardrail": [...], "revisione": [...],
    # "fiscalita": [...], "checklist": [...],

    "versamenti": [
        ("**Obiettivo casa**: ~€250/mese portano la confidenza sopra l'85%; stanno nella "
         "capacità dichiarata."),
        ("**Il resto della capacità** alimenta il lungo termine, che compone per 30 anni, "
         "oppure resta come buffer."),
        ("**Fondo pensione**: versare almeno la quota minima a proprio carico per sbloccare "
         "il contributo del datore (se il CCNL lo prevede); deducibilità fino a €5.300/anno, "
         "contributo del datore compreso e TFR escluso (Legge di Bilancio 2026). I contributi non "
         "dedotti vanno comunicati al fondo entro il 31/12 dell'anno successivo."),
    ],

    "strumenti": [
        ("Azioni", "Azionario globale diversificato (sviluppati + emergenti) via ETF UCITS ad "
                   "accumulazione, ampia capitalizzazione, TER basso, replica fisica, masse elevate."),
        ("Obbligazioni", "La parte stabilizzante, quindi con copertura del cambio: global aggregate "
                         "bond EUR-hedged (tassato 26%) o governativo euro (titoli di Stato al 12,5%). "
                         "Evitare il global bond non coperto: il rischio valutario vanifica il ruolo."),
        ("Liquidità", "Il fondo emergenza: conto deposito, fondi monetari o titoli di Stato a "
                      "brevissima scadenza. Non è capitale \u201cda far rendere\u201d: è ciò che permette "
                      "al resto di restare investito nei momenti sbagliati."),
    ],

    "previdenza": [
        ("Il pilastro pubblico e quello complementare coprono il tenore di vita dai requisiti di "
         "vecchiaia in poi; il portafoglio copre tutto ciò che viene prima. Aggiornare la proiezione "
         "(montante INPS, fondo pensione) almeno a ogni revisione annuale con gli script della "
         "skill it-pianificazione: importatore + proiezione parametrica + Monte Carlo."),
    ],

    "broker": None,  # oppure: {"nome": "…", "sezioni": [("Titolo", "testo…"), …]}
}

# ══════════════════════ DEFAULT DI POLICY (sovrascrivibili) ══════════════════════

POLICY_RIBILANCIAMENTO = [
    ("**Bande 5/25.** Si interviene su una classe quando devia di oltre 5 punti percentuali in "
     "assoluto, oppure di oltre il 25% in relativo rispetto al peso target (vale la soglia che "
     "scatta prima)."),
    ("**Attenzione ai costi.** Prima si ribilancia coi flussi (versamenti e prelievi), solo poi "
     "vendendo, per non realizzare plusvalenze e costi evitabili. Una vendita che paga imposta si "
     "fa solo oltre il doppio della soglia."),
    ("**Rientro.** A target se la correzione è a costo zero (flussi, minusvalenze, plusvalenze "
     "compensate dallo zainetto); fino a metà banda per la parte che paga imposta."),
    ("**Modulazione per volatilità (opzionale).** Spenta di default; se accesa, bande un po' più "
     "larghe per le componenti volatili e più strette per le stabili (fattore limitato 0,7–1,6)."),
    ("**Frequenza.** Controllo trimestrale delle soglie; azione solo se una banda è violata. "
     "Niente ribilanciamenti \u201ca calendario\u201d fini a sé stessi."),
]

# Verificato: 2026-10 — plus ETF/fondi = redditi di capitale (art. 44 c. 1 lett. g TUIR, dal 2027 art. 46 c. 1 lett. h; circ. 33/E/2011),
# aliquote 26%/12,5% (art. 3 DL 66/2014), fondo pensione 20%/12,5% (art. 17 D.Lgs. 252/2005, dal 2027 art. 257 D.Lgs. 117/2026; L. 190/2014
# c. 622, dal 2027 art. 257 c. 10), riporto delle minus 4 anni (art. 68 c. 5 TUIR, dal 2027 art. 77 c. 6), comunicazione dei non dedotti (art. 8 c. 4
# D.Lgs. 252/2005). Non verificata su fonte primaria la compensabilità dei certificati.
POLICY_FISCALITA = [
    ("**La \u201ctrappola degli ETF\u201d.** In Italia le plusvalenze da ETF e fondi sono redditi di "
     "capitale: NON compensabili con le minusvalenze dello zainetto fiscale. Le minusvalenze "
     "(redditi diversi) si compensano con le plus da azioni singole, certificati e obbligazioni "
     "(per queste solo la differenza di prezzo: cedole e scarto di emissione sono redditi di capitale)."),
    ("**Aliquote.** 26% sulla generalità delle rendite finanziarie; 12,5% sui titoli di Stato "
     "italiani e white-list."),
    ("**Asset location.** Il fondo pensione tassa i rendimenti al 20% (12,5% sulla quota in titoli "
     "di Stato) ed è poco liquido fino alla pensione, salvo anticipazioni e riscatti: è il contenitore "
     "della crescita di lunghissimo periodo. Il conto titoli ospita i secchielli investibili."),
    ("**Zainetto.** Le minusvalenze scadono (anno di realizzo + 4): coordinare plus e minus nello "
     "stesso anno fiscale."),
    ("**Fondo pensione.** Entro il 31/12 dell'anno successivo al versamento (o prima, se nel "
     "frattempo matura il diritto alla prestazione) comunicare al fondo i contributi non dedotti: "
     "senza comunicazione, in uscita sono tassati come se fossero stati dedotti."),
]

POLICY_GUARDRAIL = [
    ("**Non vendere nel panico.** Il piano funziona solo se si resta investiti: le perdite "
     "diventano permanenti solo quando si vende."),
    ("**Il cuscinetto serve a questo.** Il fondo emergenza va ricostituito se intaccato: è ciò che "
     "evita di smobilizzare l'azionario nel momento peggiore."),
    ("**Leva di sicurezza.** Se il rischio aggregato pesa, ridurre il profilo del secchiello di "
     "lungo termine di un gradino: scelta reversibile e legittima."),
]

POLICY_REVISIONE = [
    ("**Cadenza annuale**: deviazione dalle bande, milestone del glide path, variazioni di "
     "capacità di risparmio e di reddito."),
    ("**Eventi di vita**: cambio di lavoro o RAL, famiglia, spese straordinarie \u2192 rivedere "
     "obiettivi, capitale assegnato e versamenti."),
    ("**Dati fiscali**: tetti e aliquote cambiano (verificare ogni anno con le reference del "
     "plugin, aggiornate alle leggi di bilancio)."),
]

POLICY_CHECKLIST = [
    ("Fase 0 — Preparazione", [
        "Aprire il **conto titoli** e scegliere il regime fiscale (amministrato o dichiarativo).",
        "Recuperare il **montante INPS esatto** da \u201cLa mia pensione futura\u201d.",
        "Verificare l'attivazione del **contributo datore** sul fondo pensione di categoria.",
    ]),
    ("Fase 1 — Costituire il cuscinetto PRIMA di investire", [
        "Accantonare il **fondo emergenza** su strumenti liquidi e sicuri, separato dal capitale investibile.",
        "Confermare che la liquidità operativa di conto resti fuori dal piano.",
    ]),
    ("Fase 2 — Dispiegamento", [
        "Definire i prodotti per classe secondo lo strato strumenti (selezione a propria discrezione).",
        "Investire verso l'allocazione target, rispettando il profilo di ciascun secchiello.",
        "Decidere **ingresso unico o scaglionato** (l'unico ha rendimento atteso più alto; lo "
        "scaglionato riduce il rischio di timing e il rimpianto).",
        "Annotare **prezzi e quantità di carico** per il tracciamento fiscale.",
    ]),
    ("Fase 3 — Disciplina", [
        "Registrare allocazione target e **bande 5/25** nello strumento di ribilanciamento.",
        "Attivare i **versamenti ricorrenti** decisi nella politica dei versamenti.",
        "Promemoria **trimestrale** per il controllo delle bande; **revisione annuale** dell'IPS.",
        "Ogni anno entro il 31/12: comunicare al fondo pensione i contributi non dedotti dell'anno precedente.",
    ]),
]

# ═══════════════════════════ MOTORE DOCUMENTO ═══════════════════════════

ACCENT, DARK, GREY, RED = "2E75B6", "1F3864", "595959", "B03A2E"


def _rgb(hexcolor):
    return RGBColor.from_string(hexcolor)


def runs_markup(par, testo, size=None, color=None, italic=False):
    """Aggiunge testo con mini-markup **grassetto** al paragrafo."""
    for i, pezzo in enumerate(testo.split("**")):
        if not pezzo:
            continue
        r = par.add_run(pezzo)
        r.bold = (i % 2 == 1)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = _rgb(color)
    return par


def bordo_sotto(par, colore=ACCENT, spessore=8):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(spessore))
    bottom.set(qn("w:color"), colore)
    bottom.set(qn("w:space"), "4")
    pBdr.append(bottom)
    pPr.append(pBdr)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def h1(doc, testo):
    p = doc.add_heading(testo, level=1)
    for r in p.runs:
        r.font.color.rgb = _rgb(DARK)
        r.font.size = Pt(14)
    bordo_sotto(p)
    return p


def h2(doc, testo):
    p = doc.add_heading(testo, level=2)
    for r in p.runs:
        r.font.color.rgb = _rgb(ACCENT)
        r.font.size = Pt(11.5)
    return p


def para(doc, testo, **kw):
    return runs_markup(doc.add_paragraph(), testo, **kw)


def bullets(doc, voci):
    for v in voci:
        runs_markup(doc.add_paragraph(style="List Bullet"), v)


def checks(doc, voci):
    for v in voci:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.add_run("\u2610   ")
        runs_markup(p, v)


def tabella(doc, intestazioni, righe, larghezze_cm=None, bold_prima_colonna=True):
    t = doc.add_table(rows=1, cols=len(intestazioni))
    t.style = "Table Grid"
    for j, testo in enumerate(intestazioni):
        c = t.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(testo)
        r.bold = True
        r.font.color.rgb = _rgb("FFFFFF")
        shade(c, ACCENT)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for i, riga in enumerate(righe):
        cells = t.add_row().cells
        for j, val in enumerate(riga):
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(str(val))
            if j == 0 and bold_prima_colonna:
                r.bold = True
            if i % 2 == 1:
                shade(cells[j], "EAF2FA")
            cells[j].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if larghezze_cm:
        for j, w in enumerate(larghezze_cm):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return t


def footer_con_pagina(doc, testo):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{testo}   ·   Pag. ")
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(GREY)
    r2 = p.add_run()
    r2.font.size = Pt(8)
    r2.font.color.rgb = _rgb(GREY)
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.text = "PAGE"
        r2._r.append(el)


def nota_rischio(quota_azioni):
    """Nota di rischio ONESTA calibrata sulla quota azionaria aggregata."""
    drawdown = quota_azioni * 0.5  # orso azionario tipico ≈ −50% → portafoglio ≈ −metà quota
    if quota_azioni >= 70:
        giudizio = "**aggressiva**"
    elif quota_azioni >= 50:
        giudizio = "dinamica"
    else:
        giudizio = "prudente"
    return (f"**Onestà sul rischio.** Un'allocazione al {numero(quota_azioni)}% azionario è {giudizio}: "
            f"in un mercato orso tipico il portafoglio complessivo può perdere circa il "
            f"{numero(drawdown)}% del valore. La strategia regge solo se non si vende nel panico "
            f"(vedi Guardrail).")


# ═══════════════════════════════ COSTRUZIONE ═══════════════════════════════

MESI = ("gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio",
        "agosto", "settembre", "ottobre", "novembre", "dicembre")  # strftime("%B") dipende dal locale


def genera(dati, path_out):
    d = dict(DATI)
    d.update(dati or {})

    # validazioni oneste (avvisano, non bloccano)
    tot_pct = sum(p for _, _, p in d["allocazione"])
    if abs(tot_pct - 100) > 0.6:
        print(f"⚠ l'allocazione somma a {numero(tot_pct, 1)}% (attesi 100%)")
    tot_cap = sum(o[4] for o in d["obiettivi"])
    if abs(tot_cap - d["capitale"]) > max(1, d["capitale"] * 0.005):
        print(f"⚠ i capitali per obiettivo sommano a €{numero(tot_cap)} vs capitale €{numero(d['capitale'])}")
    # tutte le classi azionarie ("Azionario sviluppati", "ETF azionario globale", "Equity"),
    # non solo la prima; \b esclude "obbligazionario", che contiene "azion" a metà parola
    quota_az = sum(p for cl, _, p in d["allocazione"] if re.search(r"\b(azion|equity)", cl, re.I))
    if not quota_az:
        print("⚠ nessuna classe azionaria riconosciuta nell'allocazione: la nota di rischio usa 0%")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    footer_con_pagina(doc, f"Investment Policy Statement — {d['investitore']}")

    # copertina
    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph()
    runs_markup(p, "INVESTMENT POLICY STATEMENT", size=24, color=DARK)
    p.runs[0].bold = True
    bordo_sotto(p, spessore=14)
    para(doc, f"Piano di dispiegamento e gestione di un capitale di €{numero(d['capitale'])}",
         size=13, color=GREY)
    para(doc, d["sottotitolo"], size=12, color=GREY, italic=True)
    doc.add_paragraph()
    para(doc, d["investitore"], size=15).runs[0].bold = True
    oggi = datetime.date.today()
    data = d["data"] or f"{MESI[oggi.month - 1].capitalize()} {oggi.year}"
    para(doc, data, size=12, color=GREY)
    for _ in range(3):
        doc.add_paragraph()
    para(doc, "Documento di pianificazione personale. Non costituisce consulenza finanziaria "
              "o di investimento (vedi nota finale).", size=9, color=GREY, italic=True)
    doc.add_page_break()

    # 1–2
    h1(doc, "1. Scopo del documento")
    para(doc, "Questo Investment Policy Statement traduce l'analisi di pianificazione in un piano "
              "operativo: definisce cosa fare quando il capitale verrà investito e quali regole "
              "seguire dopo. È il riferimento contro cui misurare ogni decisione futura, da "
              "**rileggere almeno una volta l'anno** e a ogni evento di vita rilevante.")
    h1(doc, "2. Profilo dell'investitore")
    tabella(doc, ["Voce", "Valore"], d["profilo"], [6.2, 10.3])

    # 3
    h1(doc, "3. Obiettivi")
    tabella(doc, ["Obiettivo", "Orizzonte / tipo", "Profilo (% azioni)", "Target", "Capitale"],
            [(n, o, pr, t, f"€{numero(c)}") for n, o, pr, t, c in d["obiettivi"]],
            [3.6, 3.4, 3.6, 2.6, 3.3])
    if d.get("obiettivi_nota"):
        para(doc, d["obiettivi_nota"])

    # 4
    h1(doc, "4. Allocazione strategica target")
    righe = [(cl, f"€{numero(imp)}", f"{numero(quota, 1)}%") for cl, imp, quota in d["allocazione"]]
    righe.append(("Totale", f"€{numero(sum(i for _, i, _ in d['allocazione']))}", f"{numero(tot_pct)}%"))
    tabella(doc, ["Classe", "Importo", "Quota"], righe, [5.5, 5.5, 5.5])
    if d.get("allocazione_nota"):
        para(doc, "**Variabile dominante.** " + d["allocazione_nota"])
    para(doc, nota_rischio(quota_az), color=None)

    # 5
    h1(doc, "5. Glide path — calendario di de-risking")
    para(doc, "Ogni obiettivo terminante si mette progressivamente al sicuro verso la scadenza; "
              "i permanenti non de-riskano mai. La quota azionaria aggregata non scende in modo "
              "lineare — segue queste fasi:")
    tabella(doc, ["Fase", "Anni", "Azioni aggreg.", "Cosa succede"], d["glide_path"], [3.6, 2.2, 3.0, 7.7])

    # 6–7
    h1(doc, "6. Politica di ribilanciamento")
    bullets(doc, d.get("ribilanciamento", POLICY_RIBILANCIAMENTO))
    h1(doc, "7. Politica dei versamenti")
    bullets(doc, d["versamenti"])

    # 8
    h1(doc, "8. Strato strumenti — traduzione in classi")
    para(doc, "Le percentuali macro si traducono in tipologie di strumento; la selezione dei "
              "prodotti specifici resta all'investitore. Qui si fissa il mandato per classe.")
    for titolo, testo in d["strumenti"]:
        h2(doc, titolo)
        para(doc, testo)

    # 9–10
    h1(doc, "9. Struttura e ottimizzazione fiscale")
    bullets(doc, d.get("fiscalita", POLICY_FISCALITA))
    h1(doc, "10. Pilastro previdenziale")
    for testo in d["previdenza"]:
        para(doc, testo)

    # 11–12
    h1(doc, "11. Guardrail e disciplina")
    bullets(doc, d.get("guardrail", POLICY_GUARDRAIL))
    h1(doc, "12. Revisione del piano")
    bullets(doc, d.get("revisione", POLICY_REVISIONE))

    # Appendice A
    doc.add_page_break()
    h1(doc, "Appendice A — Checklist operativa")
    para(doc, "La sequenza per mettere in pratica il piano. L'ordine conta: il cuscinetto viene "
              "PRIMA del dispiegamento. Spunta man mano.")
    for fase, voci in d.get("checklist", POLICY_CHECKLIST):
        h2(doc, fase)
        checks(doc, voci)

    # Appendice B (opzionale)
    if d.get("broker"):
        doc.add_page_break()
        h1(doc, f"Appendice B — Attuazione su {d['broker']['nome']}")
        for titolo, testo in d["broker"]["sezioni"]:
            h2(doc, titolo)
            if isinstance(testo, (list, tuple)):
                bullets(doc, testo)
            else:
                para(doc, testo)

    # nota finale
    doc.add_paragraph()
    p = doc.add_paragraph()
    bordo_sotto(p, spessore=6)
    para(doc, "**Nota.** Documento di pianificazione personale costruito sui dati e le scelte "
              "dell'investitore. Non costituisce consulenza finanziaria, di investimento, fiscale "
              "o legale, né sollecitazione all'investimento. Le proiezioni si basano su ipotesi: "
              "i risultati reali possono differire. I dati fiscali possono cambiare. Per decisioni "
              "vincolanti rivolgersi a un professionista abilitato.", size=8, color=GREY, italic=True)

    doc.save(path_out)
    print(f"IPS generato: {path_out}")
    print(f"  investitore: {d['investitore']} | capitale €{numero(d['capitale'])} | "
          f"azioni {numero(quota_az)}% | obiettivi: {len(d['obiettivi'])}"
          + (f" | Appendice B: {d['broker']['nome']}" if d.get("broker") else ""))


def main(argv=None):
    for flusso in (sys.stdout, sys.stderr):   # UTF-8 anche su Windows (cp1252)
        try:
            flusso.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(
        description="Genera l'Investment Policy Statement come documento Word dal blocco DATI in testa allo script. "
                    "Codici d'uscita: 0 riuscito, 1 errore nei dati o nella scrittura, 2 riga di comando errata.")
    ap.add_argument("output", nargs="?", default="IPS.docx",
                    help="file .docx da scrivere (default: IPS.docx nella cartella corrente)")
    args = ap.parse_args(argv)
    try:
        genera(None, args.output)
    except OSError as e:
        print(f"ERRORE: impossibile scrivere {args.output}: {e}", file=sys.stderr)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
