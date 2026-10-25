#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
run_tests.py — test suite del plugin wealth-management-it.

Uso (dalla radice del plugin o da qualunque directory):
    python tests/run_tests.py

Filosofia: nessuna dipendenza oltre a quelle del plugin (numpy, openpyxl).
Ogni test è veloce e verifica una proprietà ROBUSTA (invariante, valore chiuso,
regressione di un bug reale), non l'output cosmetico. La suite protegge i
refactor: se passa, la meccanica di calcolo non è cambiata.
"""

import csv
import glob
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SKILLS = os.path.join(ROOT, "skills")
FIXTURES = os.path.join(ROOT, "tests", "fixtures")

RISULTATI = []  # (nome, ok, dettaglio)


def asset(skill, nome):
    return os.path.join(SKILLS, skill, "assets", nome)


def carica_modulo(path):
    """Importa uno script (anche con trattini nel nome) come modulo; la sua cartella entra nel path
    per il tempo dell'import, così trova motore_bootstrap come quando viene lanciato."""
    spec = importlib.util.spec_from_file_location(
        os.path.basename(path).replace("-", "_").replace(".py", ""), path)
    mod = importlib.util.module_from_spec(spec)
    cartella = os.path.dirname(os.path.abspath(path))
    sys.path.insert(0, cartella)
    try:
        spec.loader.exec_module(mod)
    finally:
        sys.path.remove(cartella)
    return mod


def run_script(path, *args, cwd=None, env=None):
    """Esegue uno script e ritorna (returncode, output). L'output è letto come UTF-8 (gli script lo
    forzano su stdout), così il test non dipende dalla console: su Windows sarebbe cp1252."""
    r = subprocess.run([sys.executable, path, *args],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
                       cwd=cwd or os.path.dirname(path), env=env)
    return r.returncode, r.stdout + r.stderr


def run_script_split(path, *args, cwd=None, env=None):
    """Come run_script, ma con stdout e stderr separati: (returncode, stdout, stderr)."""
    r = subprocess.run([sys.executable, path, *args],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
                       cwd=cwd or os.path.dirname(path), env=env)
    return r.returncode, r.stdout, r.stderr


def euro(testo, etichetta):
    """Estrae il primo importo in formato italiano ('€ 1.234' o '€ 1.234,56') che segue l'etichetta."""
    m = re.search(re.escape(etichetta) + r"[^€\n]*€\s*([\d.,]+)", testo)
    if not m:
        raise AssertionError(f"etichetta '{etichetta}' non trovata nell'output")
    return float(m.group(1).replace(".", "").replace(",", "."))


class Salta(Exception):
    """Il test non si può eseguire in questo ambiente (es. LibreOffice assente)."""


def test(nome):
    def deco(fn):
        def wrapped():
            try:
                fn()
                RISULTATI.append((nome, True, ""))
            except Salta as e:
                RISULTATI.append((nome, None, str(e)))
            except SystemExit as e:  # errore() di uno script importato in-process: è un test fallito, non la fine della suite
                RISULTATI.append((nome, False, f"SystemExit: {e}"))
            except Exception as e:  # noqa: BLE001 — un test fallito non ferma la suite
                RISULTATI.append((nome, False, f"{type(e).__name__}: {e}"))
        wrapped.__test__ = True
        wrapped.__nome__ = nome
        return wrapped
    return deco


# ═══════════════════════════ IGIENE DEL REPO ═══════════════════════════

@test("igiene: marketplace.json valido e coerente col plugin.json")
def t_marketplace():
    mk = json.load(open(os.path.join(ROOT, ".claude-plugin", "marketplace.json"), encoding="utf-8"))
    pl = json.load(open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8"))
    assert mk.get("name") and mk.get("owner", {}).get("name"), "marketplace: name/owner mancanti"
    assert len(mk.get("plugins", [])) >= 1, "marketplace: nessun plugin elencato"
    voce = mk["plugins"][0]
    assert voce["name"] == pl["name"], "marketplace e plugin.json dichiarano nomi diversi"
    assert voce.get("source") == "./", "il plugin del marketplace deve puntare a './'"
    assert voce.get("version") == pl["version"], \
        f"versioni diverse: marketplace {voce.get('version')!r}, plugin.json {pl['version']!r}"


@test("igiene: plugin.json valido, SKILL.md < 500 righe, nessun tag cite nei deliverable")
def t_igiene():
    d = json.load(open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8"))
    assert re.match(r"^\d+\.\d+\.\d+$", d.get("version", "")), "version non semver"
    assert isinstance(d.get("author"), dict) and d["author"].get("name"), \
        "author deve essere un oggetto {name, ...} (claude plugin validate lo richiede)"
    assert str(d.get("repository", "")).startswith("https://github.com/"), \
        "repository mancante o non è un URL GitHub"
    for skill in sorted(os.listdir(SKILLS)):
        p = os.path.join(SKILLS, skill, "SKILL.md")
        if os.path.exists(p):
            n = sum(1 for _ in open(p, encoding="utf-8"))
            assert n < 500, f"{skill}/SKILL.md ha {n} righe (limite 500)"
    for base, _, files in os.walk(ROOT):
        for f in files:
            if f.endswith(".md"):
                testo = open(os.path.join(base, f), encoding="utf-8").read()
                assert "antml:cite" not in testo, f"tag cite in {f}"


@test("igiene CI: i workflow che usano una pipe nei run girano con pipefail")
def t_pipefail():
    # Senza "shell: bash" GitHub usa `bash -e {0}`, senza pipefail: in
    # `comando | tee log` l'exit code è quello di tee e un fallimento passa verde.
    cartella = os.path.join(ROOT, ".github", "workflows")
    for nome in sorted(os.listdir(cartella)):
        righe = open(os.path.join(cartella, nome), encoding="utf-8").read().splitlines()
        con_pipe, blocco = [], None
        for riga in righe:
            rientro = len(riga) - len(riga.lstrip())
            if blocco is not None and riga.strip() and rientro <= blocco:
                blocco = None
            m = re.match(r"^(\s*)(?:- )?run:\s*(.*)$", riga)
            if m:
                if m.group(2)[:1] in ("|", ">"):
                    blocco = len(m.group(1))
                    continue
                comando = m.group(2)
            elif blocco is not None:
                comando = riga
            else:
                continue
            if re.search(r"(?<!\|)\|(?!\|)", comando):
                con_pipe.append(comando.strip())
        if con_pipe:
            testo = "\n".join(righe)
            assert re.search(r"^defaults:\s*\n\s+run:\s*\n\s+shell:\s*bash\s*$", testo, re.M), \
                f"{nome}: pipe senza pipefail ({con_pipe[0]!r}); aggiungere defaults.run.shell: bash"


@test("igiene: il README dichiara un solo conteggio dei test, uguale ai decoratori @test del runner")
def t_conteggio_test():
    # Lo storico (CHANGELOG.md, dalla 0.11.0 fuori dal README) cita i conteggi di allora
    # («Suite: 24 test»): ogni «N test» del README deve essere il numero dei decoratori
    # @test di questo file. Il conteggio vive in «## Test».
    sorgente = open(os.path.abspath(__file__), encoding="utf-8").read()
    decoratori = len(re.findall(r"^@test\(", sorgente, re.M))
    registrati = sum(1 for v in globals().values() if getattr(v, "__test__", False))
    assert decoratori == registrati, f"decoratori {decoratori} ≠ test registrati {registrati}"
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    conteggi = [int(n) for n in re.findall(r"\b(\d+) test\b", readme)]
    assert conteggi, "il README deve dichiarare il numero dei test (sezione «Test»)"
    assert all(n == decoratori for n in conteggi), \
        f"il README dice {sorted(set(conteggi))} test, i decoratori @test sono {decoratori}"
    assert len(conteggi) == 1, f"un solo conteggio dei test nel README, trovati {len(conteggi)}"


@test("igiene CI: i workflow mensile e di release usano una versione di Python della matrice di tests.yml")
def t_python_workflow():
    cartella = os.path.join(ROOT, ".github", "workflows")
    testi = {n: open(os.path.join(cartella, n), encoding="utf-8").read() for n in sorted(os.listdir(cartella))}
    m = re.search(r"python-version:\s*\[([^\]]+)\]", testi["tests.yml"])
    assert m, "tests.yml: matrice python-version non trovata"
    matrice = set(re.findall(r"[\d.]+", m.group(1)))
    assert "3.11" in matrice, "la matrice deve includere il minimo dichiarato nel README (3.11)"
    for nome, testo in testi.items():
        for v in re.findall(r'python-version:\s*"([\d.]+)"', testo):
            assert v in matrice, f"{nome}: Python {v} fuori dalla matrice {sorted(matrice)} di tests.yml"


@test("igiene CI: il check «test» chiesto dal ruleset riassume la matrice e Windows e fallisce se uno fallisce")
def t_check_ruleset():
    # Il ruleset di main e develop (24/10/2026) chiede i check «test» e «validate». I job della matrice si chiamano
    # «suite (3.11)» e «suite (3.14)»: senza un job che si chiami esattamente «test» la PR resterebbe bloccata
    # anche con la CI verde, e con un «test» saltato (needs falliti, senza !cancelled()) passerebbe con la CI rossa.
    ci = open(os.path.join(ROOT, ".github", "workflows", "tests.yml"), encoding="utf-8").read()
    jobs = re.findall(r"^  ([\w-]+):\n", ci.split("\njobs:\n")[1], re.M)
    assert {"suite", "test-windows", "test", "validate"} <= set(jobs), jobs
    test = ci.split("\n  test:\n")[1].split("\n\n")[0]
    assert "if: ${{ !cancelled() }}" in test and "needs: [suite, test-windows]" in test, test
    righe = [r.strip() for r in test.split("run: |")[1].strip().splitlines()]
    # l'ultimo comando decide l'esito: entrambi i job devono essere «success», in AND e senza vie d'uscita
    assert righe[-1] == '[ "${{ needs.suite.result }}" = success ] && [ "${{ needs.test-windows.result }}" = success ]', righe
    assert not any("||" in r or r == "true" for r in righe), righe
    suite = ci.split("\n  suite:\n")[1].split("\n\n  #")[0]
    assert "matrix:" in suite and "RICHIEDI_LIBREOFFICE" in suite, "la matrice con LibreOffice è il job suite"
    # T17: il cruscotto si verifica in Chromium, obbligatorio nel job suite
    assert "RICHIEDI_BROWSER" in suite and "playwright==1.55.0" in suite, "Playwright fissato e obbligatorio nel job suite"
    assert "playwright install --with-deps chromium" in suite, "Chromium installato da Playwright nel job suite"


@test("igiene CI: le action dei workflow sono fissate per SHA completo")
def t_action_sha():
    # Supply chain: un tag (@v4) può essere spostato da chi controlla l'action, uno SHA no.
    # Dependabot aggiorna SHA e commento di versione.
    cartella = os.path.join(ROOT, ".github", "workflows")
    for nome in sorted(os.listdir(cartella)):
        for n, riga in enumerate(open(os.path.join(cartella, nome), encoding="utf-8"), 1):
            m = re.search(r"uses:\s*([^\s#]+)", riga)
            if m:
                assert re.search(r"@[0-9a-f]{40}$", m.group(1)), f"{nome}:{n}: action non fissata per SHA ({m.group(1)})"


@test("compilazione: tutti gli script Python compilano")
def t_compila():
    import py_compile
    trovati = 0
    for base, _, files in os.walk(SKILLS):
        for f in files:
            if f.endswith(".py"):
                py_compile.compile(os.path.join(base, f), doraise=True)
                trovati += 1
    assert trovati >= 10, f"attesi ≥10 script, trovati {trovati}"


# ══════════════════════ it-pianificazione (previdenza) ══════════════════════

@test("importa-previdenza: parser XML INPS su fixture (settimane, retribuzioni)")
def t_inps_parser():
    m = carica_modulo(asset("it-pianificazione", "importa-previdenza.py"))
    anag, righe = m.leggi_estratto_inps(os.path.join(FIXTURES, "estratto-inps-test.xml"))
    assert anag["Cognome"] == "ROSSI" and "nascita" in anag
    assert len(righe) == 2 and sum(r["settimane"] for r in righe) == 104
    assert abs(sum(r["retribuzione"] for r in righe) - 61000) < 0.01


@test("importa-previdenza: ricostruzione montante = formula replicata a mano")
def t_inps_montante():
    m = carica_modulo(asset("it-pianificazione", "importa-previdenza.py"))
    retrib = {2023: 30000.0, 2024: 31000.0}
    montante, nominali, _ = m.ricostruisci_montante(retrib, 2026)
    # replica indipendente della meccanica L.335/1995 con la stessa tabella tassi
    atteso = 0.0
    for anno in range(2023, 2027):
        atteso *= m.TASSI_CAPITALIZZAZIONE.get(anno, m.TASSO_CAP_DEFAULT)
        atteso += retrib.get(anno, 0.0) * m.ALIQUOTA_COMPUTO
    assert abs(montante - atteso) < 0.01, f"{montante} != {atteso}"
    assert abs(nominali - 61000 * m.ALIQUOTA_COMPUTO) < 0.01


@test("importa-previdenza: fondo su fixture + REGRESSIONE substring (obbligazionario ≠ azionario)")
def t_fondo():
    m = carica_modulo(asset("it-pianificazione", "importa-previdenza.py"))
    quote, nav, vers = m.leggi_storico_fondo(os.path.join(FIXTURES, "storico-fondo-test.csv"))
    val = {c: quote[c] * nav[c][1] for c in quote}
    assert abs(sum(val.values()) - 2000.0) < 0.01, f"montante fondo {sum(val.values())}"
    # regressione del bug reale: 'azionar' è sottostringa di 'obbligAZIONARio'
    pct_ob, _ = m.quota_azionaria("FONDO TEST COMPARTO OBBLIGAZIONARIO")
    pct_az, _ = m.quota_azionaria("FONDO TEST COMPARTO AZIONARIO")
    assert pct_ob == 0.05, f"obbligazionario → {pct_ob} (bug substring?)"
    assert pct_az == 0.90
    assert abs(sum(vers[2025].values()) - 2000.0) < 0.01


@test("proiezione-inps: valore chiuso via --set (formula annuità indipendente, PIL ≠ 0)")
def t_proiezione():
    rc, out = run_script(asset("it-pianificazione", "proiezione-inps.py"),
                         "--set", "MONTANTE_ATTUALE=100_000.0", "--set", "RAL_ATTUALE=30_000.0",
                         "--set", "ETA_ATTUALE=37.0", "--set", "ETA_PENSIONAMENTO=67.0",
                         "--set", "CRESCITA_REALE_RAL=0.0", "--set", "PIL_REALE=0.02",
                         "--set", "COEFF_ESPLICITO=0.05")
    assert rc == 0, out[-400:]
    assert "[override]" in out, "gli override applicati devono essere loggati"
    # verifica incrociata: lo script usa un loop, qui la formula chiusa dell'annuità.
    # Il PIL ≠ 0 rende osservabile la meccanica "l'ultimo anno non si rivaluta":
    # montante = 100k×1.02^30 + 9.900×(1.02^30 − 1)/0.02, pensione = montante × 5%.
    atteso = (100_000 * 1.02 ** 30 + 30_000 * 0.33 * ((1.02 ** 30 - 1) / 0.02)) * 0.05
    pens = euro(out, "PENSIONE LORDA ANNUA")
    assert abs(pens - atteso) <= 2, f"pensione {pens} != {atteso:.0f}"


@test("override --set: refuso su costante inesistente respinto con errore chiaro")
def t_override_refuso():
    rc, out = run_script(asset("it-pianificazione", "proiezione-inps.py"),
                         "--set", "COSTANTE_INESISTENTE=1")
    assert rc != 0, "un --set su costante inesistente deve fermare lo script"
    assert "sconosciuta" in out, out[-200:]


@test("montecarlo-pensione: run sui default, mediana simulata ≈ check deterministico")
def t_mc_pensione():
    rc, out = run_script(asset("it-pianificazione", "montecarlo-pensione.py"))
    assert rc == 0, out[-400:]
    det = euro(out, "montante FP deterministico")
    med = euro(out, "mediana simulata")
    dev = abs(det - med) / det
    assert dev < 0.20, f"mediana devia dal deterministico del {dev:.0%} (>20%)"


@test("serie storica: gli script la trovano anche lanciati dalla cartella di lavoro dell'utente")
def t_cwd_esterna():
    # REGRESSIONE F005: ponte e montecarlo-pensione aprivano rendimenti-storici.csv
    # rispetto alla cwd e andavano in FileNotFoundError fuori da assets/.
    tmp = tempfile.mkdtemp()
    try:
        for skill, nome, *args in (("it-pianificazione", "ponte-pensione.py"),
                                   ("it-pianificazione", "montecarlo-pensione.py"),
                                   ("it-pianificazione-obiettivi", "montecarlo-obiettivi.py",
                                    "--set", "MODELLO=bootstrap"),
                                   ("it-pianificazione-obiettivi", "decumulo-obiettivi.py")):
            rc, out = run_script(asset(skill, nome), *args, cwd=tmp)
            assert rc == 0, f"{nome} da una cartella esterna: {out[-300:]}"
            assert "serie storica non trovata" not in out, f"{nome}: serie storica non trovata"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ═══════════════ T9: interfaccia comune degli script (F056, Windows, formati) ═══════════════

SCRIPT_CLI = [
    ("it-pianificazione", "proiezione-inps.py"),
    ("it-pianificazione", "montecarlo-pensione.py"),
    ("it-pianificazione", "netto-pensione.py"),
    ("it-pianificazione", "ponte-pensione.py"),
    ("it-pianificazione", "importa-previdenza.py"),
    ("it-pianificazione", "riscatto-laurea.py"),
    ("it-pianificazione-obiettivi", "allocatore-capitale.py"),
    ("it-pianificazione-obiettivi", "asset-location-optimizer.py"),
    ("it-pianificazione-obiettivi", "capitale-umano.py"),
    ("it-pianificazione-obiettivi", "importa-composizione-etf.py"),
    ("it-pianificazione-obiettivi", "decumulo-obiettivi.py"),
    ("it-pianificazione-obiettivi", "glidepath-obiettivi.py"),
    ("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"),
    ("it-rebalance", "ribilanciamento-multibroker.py"),
    ("it-revisione", "genera-cruscotto.py"),
    ("it-revisione", "genera-ips.py"),
    ("it-tax-optimization", "importa-estratto-broker.py"),
]
# argomenti con cui ogni script gira sui default (quelli con un input obbligatorio usano le fixture)
ARGOMENTI_DEFAULT = {
    "importa-previdenza.py": (os.path.join(FIXTURES, "estratto-inps-test.xml"),
                              os.path.join(FIXTURES, "storico-fondo-test.csv")),
    "importa-estratto-broker.py": (os.path.join(FIXTURES, "estratto-broker-test.csv"),),
    "genera-cruscotto.py": ("--from", os.path.join(FIXTURES, "cruscotto-allocatore.json")),
    "importa-composizione-etf.py": ("--file", os.path.join(FIXTURES, "composizione-ishares-en.csv"),
                                    "--strumento", "ETF MSCI World"),
}
_OUTPUT_CP1252 = {}


def _tutti_gli_script_cp1252():
    """Ogni script lanciato una volta da una cartella esterna con una console cp1252 (come Windows);
    un solo giro per esecuzione della suite, riusato dai test sui formati."""
    if not _OUTPUT_CP1252:
        tmp = tempfile.mkdtemp()
        env = dict(os.environ, PYTHONIOENCODING="cp1252", PYTHONUTF8="0")
        try:
            for skill, nome in SCRIPT_CLI:
                rc, out = run_script(asset(skill, nome), *ARGOMENTI_DEFAULT.get(nome, ()), cwd=tmp, env=env)
                _OUTPUT_CP1252[nome] = (rc, out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    return _OUTPUT_CP1252


@test("T9 interfaccia: ogni script ha --help (rc 0) e respinge un argomento ignoto con rc 2")
def t_cli_help():
    tutti = [asset(s, n) for s, n in SCRIPT_CLI]
    assert len(tutti) == 17 and all(os.path.exists(p) for p in tutti), tutti
    sorgenti = {os.path.basename(p): open(p, encoding="utf-8").read() for p in tutti}
    for path in tutti:
        nome = os.path.basename(path)
        rc, out = run_script(path, "--help")
        assert rc == 0 and "usage:" in out and "Traceback" not in out, f"{nome} --help: rc {rc}, {out[-300:]}"
        assert "-h, --help" in out, f"{nome}: --help senza l'aiuto di argparse"
        if "motore_bootstrap import" in sorgenti[nome] and "cli(" in sorgenti[nome]:
            assert "--set NOME=VALORE" in out and "costanti sovrascrivibili" in out, f"{nome}: --help senza l'elenco delle costanti"
        rc, out = run_script(path, "--argomento-ignoto", *ARGOMENTI_DEFAULT.get(nome, ()))
        assert rc == 2 and "usage:" in out and "Traceback" not in out, f"{nome} con un flag ignoto: rc {rc}, {out[-300:]}"
    # codice 1 per i dati, con «ERRORE:» su stderr e niente su stdout: file mancante (ogni script che legge un
    # file), costante --set sconosciuta o non numerica, file di output non scrivibile, foglio senza obiettivi
    def errore_pulito(path, *args, atteso="ERRORE:"):
        rc, so, se = run_script_split(path, *args)
        assert rc == 1, f"{os.path.basename(path)} {args}: rc {rc}, {so[-200:]} {se[-200:]}"
        assert se.lstrip().startswith("ERRORE:") and "ERRORE:" not in so and "Traceback" not in se, \
            f"{os.path.basename(path)} {args}: stdout «{so[-150:]}» stderr «{se[-300:]}»"
        assert atteso in se, f"{os.path.basename(path)} {args}: manca «{atteso}» in «{se[-300:]}»"
    OB = "it-pianificazione-obiettivi"
    for nome in ("glidepath-obiettivi.py", "montecarlo-obiettivi.py", "allocatore-capitale.py", "asset-location-optimizer.py"):
        errore_pulito(asset(OB, nome), "/non/esiste.xlsx", atteso="non trovato")
    vuoto = _excel_test([])
    for nome in ("glidepath-obiettivi.py", "montecarlo-obiettivi.py", "allocatore-capitale.py", "asset-location-optimizer.py"):
        errore_pulito(asset(OB, nome), vuoto, atteso="nessun obiettivo")
    shutil.rmtree(os.path.dirname(vuoto), ignore_errors=True)
    errore_pulito(asset("it-pianificazione", "importa-previdenza.py"), "/non/esiste.xml", atteso="non trovato")
    errore_pulito(asset("it-pianificazione", "importa-previdenza.py"), os.path.join(FIXTURES, "estratto-inps-test.xml"), "/non/esiste.csv", atteso="non trovato")
    errore_pulito(asset("it-pianificazione", "netto-pensione.py"), "--set", "NON_ESISTE=1", atteso="sconosciuta")
    errore_pulito(asset("it-pianificazione", "proiezione-inps.py"), "--set", "RAL_ATTUALE=abc", atteso="non è un numero")
    errore_pulito(asset("it-pianificazione", "proiezione-inps.py"), "--json", "/non/esiste/x.json", atteso="impossibile scrivere")
    errore_pulito(asset("it-revisione", "genera-ips.py"), os.path.join(tempfile.gettempdir(), "cartella-inesistente", "x.docx"), atteso="impossibile scrivere")
    errore_pulito(asset("it-tax-optimization", "importa-estratto-broker.py"), "/non/esiste.csv", atteso="non trovato")
    errore_pulito(asset("it-pianificazione-obiettivi", "importa-composizione-etf.py"), "--file", "/non/esiste.csv",
                  "--strumento", "X", atteso="non trovato")
    # niente abbreviazioni dei flag: «--se» non è «--set»
    rc, out = run_script(asset("it-pianificazione", "netto-pensione.py"), "--se", "RAL_FINALE_LORDA=1")
    assert rc == 2 and "usage:" in out, (rc, out[-200:])


@test("T9 interfaccia: tutti gli script girano da una cartella esterna con una console cp1252 (Windows)")
def t_cli_utf8():
    # Prima della 0.11 undici script su tredici andavano in UnicodeEncodeError su «€», «→» e «≈» su qualsiasi
    # output, gli altri due (genera-ips, importa-estratto-broker) su percorsi specifici
    # con stdout in cp1252: ogni script forza l'UTF-8 (sys.stdout.reconfigure) prima di stampare.
    for nome, (rc, out) in _tutti_gli_script_cp1252().items():
        assert rc == 0, f"{nome} da una cartella esterna con cp1252: rc {rc}, {out[-400:]}"
        assert "UnicodeEncodeError" not in out and "Traceback" not in out, f"{nome}: {out[-400:]}"
        assert "serie storica non trovata" not in out, f"{nome}: serie storica non trovata"
        # lo stdout resta UTF-8: ogni script stampa almeno un «€» (l'importatore della composizione, che non
        # stampa importi, una «à»), e senza reconfigure il carattere uscirebbe in cp1252 (qui decodificato
        # come «\ufffd») anche dove non va in crash
        carattere = "à" if nome == "importa-composizione-etf.py" else "€"
        assert carattere in out and "\ufffd" not in out, f"{nome}: stdout non UTF-8 ({out[:200]!r})"
    # nessuna open() senza encoding negli script e nel runner: su Windows leggerebbero in cp1252
    cartelle = [os.path.join(SKILLS, s, "assets") for s in sorted(os.listdir(SKILLS))] + \
               [os.path.join(ROOT, "scripts"), os.path.join(ROOT, "tests")]
    senza = []
    for cartella in cartelle:
        for f in sorted(os.listdir(cartella)) if os.path.isdir(cartella) else []:
            if f.endswith(".py"):
                righe = open(os.path.join(cartella, f), encoding="utf-8").read().splitlines()
                for n, riga in enumerate(righe, 1):
                    chiamata = riga + " " + (righe[n] if n < len(righe) else "")   # argomenti anche a capo
                    if re.search(r"\bopen\(", riga) and "encoding=" not in chiamata and "rb" not in chiamata \
                            and "def " not in riga and not riga.strip().startswith(("#", "assert")):
                        senza.append(f"{f}:{n}")
    assert not senza, f"open() senza encoding (su Windows leggerebbero in cp1252): {senza}"


@test("T9 formati: importi e percentuali in formato italiano in tutti gli script (punto migliaia, virgola decimali)")
def t_formati_italiani():
    # Nei sorgenti nessun format spec anglosassone («{x:,.0f}», «{x:.1%}»): si passa da numero()/pct()
    # («{x:.1f}» con decimali è anglosassone anche senza «%»: «12.6 punti»; «.0f» non ha decimali e resta)
    anglo = re.compile(r":[<>^]?\+?\d*,\.?\d*[fd]?\}|:[<>^]?\+?\d*\.\d%\}|:[<>^]?\+?\d*\.[1-9]\d*f\}")
    for skill, nome in SCRIPT_CLI + [("it-pianificazione", "motore_bootstrap.py")]:
        for n, riga in enumerate(open(asset(skill, nome), encoding="utf-8"), 1):
            if anglo.search(riga) and ".replace(" not in riga and "--set" not in riga:
                raise AssertionError(f"{nome}:{n}: format spec anglosassone: {riga.strip()[:100]}")
    # Negli output nessun «€ 1,234» né «7.0%»; valori chiusi sui default
    for nome, (rc, out) in _tutti_gli_script_cp1252().items():
        assert rc == 0, nome
        m = re.search(r"€\s?\d{1,3},\d{3}\b", out)
        assert not m, f"{nome}: importo anglosassone «{m.group(0)}»"
        m = re.search(r"\d\.\d+%", out)
        assert not m, f"{nome}: percentuale col punto «{m.group(0)}»"
    out = _tutti_gli_script_cp1252()["proiezione-inps.py"][1]
    assert "RAL: € 36.000 oggi" in out and "+1,5%/anno" in out and "PIL reale +0,5%" in out, out[:600]
    out = _tutti_gli_script_cp1252()["ribilanciamento-multibroker.py"][1]
    assert "VENDI €8.000 azionario/diverso @ Fineco" in out and "68,0%" in out and "patrimonio €100.000" in out, out[:900]
    out = _tutti_gli_script_cp1252()["decumulo-obiettivi.py"][1]
    assert re.search(r"€\s*\d{1,3}\.\d{3}\s+\d+,\d%\s+\d+,\d%", out), "decumulo: tabella di sensibilità non in formato italiano"
    motore = carica_modulo(asset("it-pianificazione", "motore_bootstrap.py"))
    assert motore.numero(1234567.891, 2) == "1.234.567,89" and motore.numero(-0.5, 1) == "-0,5"
    assert motore.euro(-1234.4) == "-€1.234" and motore.euro(999.5) == "€1.000" and motore.euro(0) == "€0"
    assert motore.pct(0.125) == "12,5%" and motore.pct(0.0712, 1, segno=True) == "+7,1%" and motore.pct(1, 0) == "100%"


@test("T9 formati: CSV degli ordini apribile in Excel italiano (BOM, «;», virgola decimale), importi invariati")
def t_csv_italiano():
    import csv
    m, p = _ribilancia()
    tmp = tempfile.mkdtemp()
    try:
        (path, n), = m.scrivi_csv_per_broker(m.genera_ordini(p), tmp)
        grezzo = open(path, "rb").read()
        assert grezzo.startswith(b"\xef\xbb\xbf"), "senza BOM Excel apre il CSV in cp1252 e sporca «€»"
        testo = grezzo.decode("utf-8-sig")
        assert testo.splitlines()[0] == ";".join(m.CAMPI_CSV), testo.splitlines()[0]
        assert re.search(r"^Fineco;VENDITA;azionario;reddito_diverso;;8000,00;3200,00;0,00;3200,00;", testo, re.M), testo
        assert not re.search(r"\d\.\d\d;", testo), "importi con il punto decimale"
        with open(path, encoding="utf-8-sig", newline="") as f:
            righe = list(csv.DictReader(f, delimiter=";"))
        assert [r["operazione"] for r in righe] == ["VENDITA", "ACQUISTO", "ACQUISTO", "NOTA"], righe
        assert sum(float(r["importo_eur"].replace(",", ".")) for r in righe[:-1]) == 16000, righe
        assert "€3.200" in righe[0]["nota"], righe[0]["nota"]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ══════════════════════ previdenza (T5, T4, T6) ══════════════════════

def _modulo_previdenza(nome):
    """Importa uno script di it-pianificazione/assets (importano motore_bootstrap dalla loro cartella)."""
    cartella = os.path.join(SKILLS, "it-pianificazione", "assets")
    sys.path.insert(0, cartella)
    try:
        return carica_modulo(os.path.join(cartella, nome))
    finally:
        sys.path.remove(cartella)


def _per_etichetta(ws, testo, col="B"):
    """Valore della colonna col nella riga la cui etichetta in A contiene testo (layout-indipendente)."""
    for r in range(1, ws.max_row + 1):
        if testo.lower() in str(ws[f"A{r}"].value or "").lower():
            return ws[f"{col}{r}"].value
    raise AssertionError(f"{ws.title}: nessuna riga con etichetta «{testo}»")


_VARIANTI = {}


def _varianti(skill, nome_xlsx, varianti):
    """Copie ricalcolate di un xlsx con input modificati, una sola volta per suite.

    varianti: {nome: {(foglio, cella): valore}} → {nome: workbook data_only}.
    Le copie sono temporanee: l'originale non viene mai salvato con openpyxl."""
    from openpyxl import load_workbook
    chiave = (skill, nome_xlsx, repr(sorted((k, sorted(v.items())) for k, v in varianti.items())))
    if chiave not in _VARIANTI:
        mod = ricalcolatore()
        tmp = tempfile.mkdtemp()
        try:
            percorsi = []
            for nome, celle in varianti.items():
                wb = load_workbook(asset(skill, nome_xlsx))
                for (foglio, cella), v in celle.items():
                    wb[foglio][cella] = v
                percorsi.append(os.path.join(tmp, f"{nome}.xlsx"))
                wb.save(percorsi[-1])
            prodotti = mod.ricalcola(percorsi, os.path.join(tmp, "out"))
            _VARIANTI[chiave] = {n: load_workbook(p, data_only=True) for n, p in zip(varianti, prodotti)}
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    return _VARIANTI[chiave]


def _pensione_contributiva(m0, ral, g, r, n, massimale, aliquota, coeff):
    """Replica indipendente della meccanica L. 335/1995 in euro di oggi: il montante si rivaluta
    al PIL reale r, il contributo dell'anno entra senza rivalutazione, la retribuzione oltre il
    massimale non versa contributi."""
    m = m0
    for t in range(n):
        m = m * (1 + r) + min(ral * (1 + g) ** t, massimale) * aliquota
    return m * coeff


def _coeff_excel(wb, eta):
    c = wb["Coefficienti"]
    return next(c[f"C{r}"].value for r in range(5, 20) if c[f"A{r}"].value == eta)


def _montante_stimato(ral, g, r, anzianita, massimale, aliquota):
    """Replica indipendente della stima di Pubblica!E5 (montante al 31/12 dell'anno scorso):
    la retribuzione dell'anno −k è RAL/(1+g)^k, il massimale è costante in euro di oggi
    (L. 335/1995 art. 2 c. 18: rivalutato sui prezzi) e il contributo dell'anno −k si è
    rivalutato k−1 volte (il contributo dell'anno scorso non ancora)."""
    return sum(min(ral / (1 + g) ** k, massimale) * aliquota * (1 + r) ** (k - 1)
               for k in range(1, int(anzianita) + 1))


PP = ("it-pianificazione", "proiezione-pensione.xlsx")
VARIANTI_PP = {"ral150k": {("Input", "B10"): 150_000, ("Input", "B12"): 0.0},
               "ral150k_g": {("Input", "B10"): 150_000, ("Input", "B12"): 0.015},
               "ral8k": {("Input", "B10"): 8_000},
               # coefficiente congelato: le colonne degli Scenari restano ordinate (offset dalla Base)
               "decadimento_zero": {("Input", "B18"): 0.0}}


@test("proiezione-pensione.xlsx: Pubblica in euro di oggi = proiezione-inps.py (PIL reale, massimale, contributo dell'anno)")
def t_pubblica_vs_script():
    # REGRESSIONE F013/F002: Pubblica!F rivalutava anche il contributo dell'anno, il PIL era
    # nominale (1,8%) su una RAL reale, e né l'Excel né lo script applicavano il massimale.
    from openpyxl import load_workbook
    inps = _modulo_previdenza("proiezione-inps.py")
    base = load_workbook(xlsx_ricalcolati()[asset(*PP)], data_only=True)
    assert base["Input"]["B16"].value == inps.PIL_REALE, \
        f"PIL dell'Excel {base['Input']['B16'].value} ≠ PIL_REALE di proiezione-inps.py {inps.PIL_REALE}"
    var = _varianti(*PP, VARIANTI_PP)
    for nome, wb in (("default", base), ("ral150k", var["ral150k"]), ("ral150k_g", var["ral150k_g"])):
        i, p = wb["Input"], wb["Pubblica"]
        massimale = _per_etichetta(i, "Massimale contributivo")
        n = int(i["B6"].value - i["B5"].value)
        # REGRESSIONE (revisione T5): E5 applicava il massimale alla RAL di oggi e poi lo
        # deflazionava con la RAL, e rivalutava anche il contributo dell'anno scorso
        assert not i["B14"].value, f"{nome}: la variante deve usare la stima del montante"
        stima = _montante_stimato(i["B10"].value, i["B12"].value, i["B16"].value, i["B11"].value,
                                  massimale, i["B13"].value)
        assert abs(p["E5"].value - stima) <= 1e-6 * stima, \
            f"{nome}: Pubblica!E5 = {p['E5'].value:,.2f}, stima anno per anno col massimale = {stima:,.2f}"
        atteso = _pensione_contributiva(p["E5"].value, i["B10"].value, i["B12"].value, i["B16"].value,
                                        n, massimale, i["B13"].value, p["B60"].value)
        assert abs(p["B61"].value - atteso) <= 1e-6 * atteso, \
            f"{nome}: Pubblica!B61 = {p['B61'].value:,.2f}, meccanica L. 335/1995 = {atteso:,.2f}"
        sets = (f"MONTANTE_ATTUALE={p['E5'].value!r}", f"RAL_ATTUALE={float(i['B10'].value)!r}",
                f"ETA_ATTUALE={float(i['B5'].value)!r}", f"ETA_PENSIONAMENTO={float(i['B6'].value)!r}",
                f"CRESCITA_REALE_RAL={float(i['B12'].value)!r}", f"PIL_REALE={float(i['B16'].value)!r}",
                f"COEFF_ESPLICITO={p['B60'].value!r}", f"MASSIMALE_CONTRIBUTIVO={float(massimale)!r}")
        rc, out = run_script(asset("it-pianificazione", "proiezione-inps.py"),
                             *[x for s in sets for x in ("--set", s)])
        assert rc == 0, out[-300:]
        pens = euro(out, "PENSIONE LORDA ANNUA")
        assert abs(pens - p["B61"].value) <= 1.0, f"{nome}: script {pens:,.0f} ≠ Excel {p['B61'].value:,.0f}"
    p = _varianti(*PP, VARIANTI_PP)["ral150k"]["Pubblica"]
    assert abs(p["D5"].value - 122_295 * 0.33) < 0.01, f"RAL 150k: contributo {p['D5'].value:,.2f} senza massimale"
    # lo script da solo, valore chiuso: 30 anni sopra il massimale, PIL 0, coefficiente 5%
    rc, out = run_script(asset("it-pianificazione", "proiezione-inps.py"),
                         "--set", "MONTANTE_ATTUALE=0.0", "--set", "RAL_ATTUALE=150_000.0",
                         "--set", "ETA_ATTUALE=37.0", "--set", "ETA_PENSIONAMENTO=67.0",
                         "--set", "CRESCITA_REALE_RAL=0.0", "--set", "PIL_REALE=0.0",
                         "--set", "COEFF_ESPLICITO=0.05")
    assert rc == 0, out[-300:]
    assert abs(euro(out, "PENSIONE LORDA ANNUA") - 30 * 122_295 * 0.33 * 0.05) <= 1, out[-500:]


@test("proiezione-pensione.xlsx: IRPEF, addizionali e netto di Pubblica = netto_pensionato() di netto-pensione.py")
def t_pubblica_netto():
    # Deliverable T5: Pubblica!B63:B66 ricostruite dallo script sugli stessi input, anche sotto
    # la no tax area (addizionali dovute solo se c'è IRPEF, come nello script).
    from openpyxl import load_workbook
    netto = _modulo_previdenza("netto-pensione.py")
    base = load_workbook(xlsx_ricalcolati()[asset(*PP)], data_only=True)
    for nome, wb in (("default", base), ("ral8k", _varianti(*PP, VARIANTI_PP)["ral8k"])):
        p = wb["Pubblica"]
        lorda = p["B61"].value
        n, irpef, addiz = netto.netto_pensionato(lorda, 0.02, 0.0)
        assert abs(max(0.0, p["B63"].value - p["B64"].value) - irpef) < 0.01, f"{nome}: IRPEF netta"
        assert abs(p["B65"].value - addiz) < 0.01, f"{nome}: addizionali {p['B65'].value:.2f} ≠ {addiz:.2f} (lorda {lorda:,.0f})"
        assert abs(p["B66"].value - n) < 0.01, f"{nome}: netto {p['B66'].value:,.2f} ≠ {n:,.2f}"


@test("proiezione-pensione.xlsx: Scenari in euro di oggi, Base = fogli di calcolo, parametri di colonna usati davvero")
def t_scenari_reali():
    # REGRESSIONE F013: spesa target inflazionata contro pensione in euro di oggi (il gap cambiava
    # segno) e righe crescita RAL / PIL / rendimento decorative. Anche sopra il massimale
    # (revisione T5): con i default il tetto di B11/D11 non scatta mai.
    from openpyxl import load_workbook
    base = load_workbook(xlsx_ricalcolati()[asset(*PP)], data_only=True)
    netto = _modulo_previdenza("netto-pensione.py")
    var = _varianti(*PP, VARIANTI_PP)
    for nome, wb in (("default", base), ("ral150k", var["ral150k"]), ("decadimento_zero", var["decadimento_zero"])):
        s, i, p = wb["Scenari"], wb["Input"], wb["Pubblica"]
        for col in "BCD":
            assert abs(s[f"{col}18"].value - i["B39"].value) < 0.01, \
                f"{nome}: Scenari!{col}18 = {s[f'{col}18'].value:,.0f}: la spesa target va in euro di oggi ({i['B39'].value:,.0f})"
        for riga, (foglio, cella) in ((11, ("Pubblica", "B61")), (12, ("Pubblica", "B66")),
                                      (13, ("Complementare", "B58")), (14, ("Complementare", "B61")),
                                      (15, ("Complementare", "B62"))):
            assert abs(s[f"C{riga}"].value - wb[foglio][cella].value) < 0.01, \
                f"{nome}: Base Scenari!C{riga} ≠ {foglio}!{cella}"
        n = int(i["B6"].value - i["B5"].value)
        massimale = _per_etichetta(i, "Massimale contributivo")
        coeff = _coeff_excel(wb, i["B6"].value)
        for col in "BD":
            g, r, d, rend = (s[f"{col}{k}"].value for k in (5, 6, 8, 9))
            atteso = _pensione_contributiva(p["E5"].value, i["B10"].value, g, r, n, massimale, i["B13"].value,
                                            coeff * (1 - d) ** (n / 2))
            assert abs(s[f"{col}11"].value - atteso) <= 1e-6 * atteso, \
                f"{nome}: Scenari!{col}11 = {s[f'{col}11'].value:,.0f}, ricalcolo con i parametri di colonna = {atteso:,.0f}"
            assert abs(s[f"{col}12"].value - netto.netto_pensionato(atteso, 0.02, 0.0)[0]) < 0.01, \
                f"{nome}: Scenari!{col}12"
            fondo = _fondo_scenario(wb, g, rend, n)
            assert abs(s[f"{col}13"].value - fondo) <= 1e-6 * fondo, \
                f"{nome}: Scenari!{col}13 = {s[f'{col}13'].value:,.0f}, fondo col rendimento di colonna = {fondo:,.0f}"
        assert s["B11"].value < s["C11"].value < s["D11"].value and s["B13"].value < s["C13"].value < s["D13"].value, nome


def _fondo_scenario(wb, g, rend, n):
    """Montante del fondo (foglio Complementare) ricalcolato con crescita RAL g e rendimento rend."""
    i = wb["Input"]
    dovuto = _per_etichetta(i, "Contributo dovuto lavoratore")
    m = i["B22"].value * i["B21"].value + i["B25"].value   # il saldo TFR entra solo se conferito
    for t in range(n):
        ral = i["B10"].value * (1 + g) ** t
        m = m * (1 + rend) + ral * (0.0691 * i["B21"].value + i["B27"].value + dovuto) + i["B26"].value
    return m


def _estratto_inps_xml(retrib_per_anno, path):
    """Estratto conto INPS sintetico: una riga di 52 settimane per anno."""
    righe = "".join(
        f"<RigaContributi><PeriodoDal><Giorno>01</Giorno><Mese>01</Mese><Anno>{a}</Anno></PeriodoDal>"
        f"<PeriodoAl><Giorno>31</Giorno><Mese>12</Mese><Anno>{a}</Anno></PeriodoAl>"
        f"<TipoContribuzione>Lavoro dipendente</TipoContribuzione><ContributiUtiliCalcolo>52</ContributiUtiliCalcolo>"
        f"<RetribuzioneEuro>{r}</RetribuzioneEuro><Descrizione>AZIENDA TEST SRL</Descrizione></RigaContributi>"
        for a, r in sorted(retrib_per_anno.items()))
    with open(path, "w", encoding="utf-8") as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?><EstrattoConto><DatiAnagrafici><Cognome>ROSSI</Cognome>'
                '<Nome>MARIO</Nome></DatiAnagrafici><RegimeGenerale><Contributi>' + righe +
                '</Contributi></RegimeGenerale></EstrattoConto>')


@test("importa-previdenza: tassi di capitalizzazione ufficiali 1996-2025, solo contributi dal 1996, avvisi veri")
def t_tassi_capitalizzazione():
    # REGRESSIONE F062: tabella limitata al 2021-2025, default 1,01 per gli anni precedenti
    # (sottostima ~8% per chi ha iniziato a fine anni '90) e anche per l'anno in corso, non pubblicato.
    # Revisione T5: i contributi prima del 1996 entravano nel montante e si rivalutavano (×2,31)
    # mentre l'avviso diceva il contrario, e l'avviso perdeva il primo anno (carriera dal 1995: nessuno).
    ufficiali = {  # nota ISTAT 16/10/2025 (lavoro.gov.it); 2014 e 2021 a 1, 2022 con recupero (DL 65/2015 art. 5)
        1996: 1.062054,
        1997: 1.055871, 1998: 1.053597, 1999: 1.056503, 2000: 1.051781, 2001: 1.047781, 2002: 1.043698,
        2003: 1.041614, 2004: 1.039272, 2005: 1.040506, 2006: 1.035386, 2007: 1.033937, 2008: 1.034625,
        2009: 1.033201, 2010: 1.017935, 2011: 1.016165, 2012: 1.011344, 2013: 1.001643, 2014: 1.000000,
        2015: 1.005058, 2016: 1.004684, 2017: 1.005205, 2018: 1.013478, 2019: 1.018254, 2020: 1.019199,
        2021: 1.000000, 2022: 1.009756, 2023: 1.023082, 2024: 1.036622, 2025: 1.040445}
    m = carica_modulo(asset("it-pianificazione", "importa-previdenza.py"))
    for anno, tasso in ufficiali.items():
        assert m.TASSI_CAPITALIZZAZIONE.get(anno) == tasso, f"tasso {anno}: {m.TASSI_CAPITALIZZAZIONE.get(anno)} ≠ {tasso}"
    retrib = {a: 30_000.0 for a in range(1998, 2026)}
    montante, _, mancanti = m.ricostruisci_montante(retrib, 2026)
    atteso = 0.0
    for anno in range(1998, 2027):  # il 2026 non ha ancora un tasso: il montante al 31/12/2025 resta lì
        atteso = atteso * ufficiali.get(anno, 1.0) + retrib.get(anno, 0.0) * 0.33
    assert abs(montante - atteso) < 0.01, f"montante {montante:,.2f} ≠ {atteso:,.2f}"
    assert mancanti == [2026], f"anni senza tasso pubblicato: {mancanti}"
    # posizione mista: le retribuzioni prima del 1996 non entrano nel montante contributivo
    # (L. 335/1995 art. 1 c. 12-13: quota retributiva), né come nominali né rivalutate
    dal96 = {a: 30_000.0 for a in range(1996, 2026)}
    solo96, nom96, _ = m.ricostruisci_montante(dal96, 2026)
    for prima in ({1995: 25_000.0}, {1994: 30_000.0}, {a: 20_000.0 for a in range(1990, 1996)}):
        misto, nom, _ = m.ricostruisci_montante({**prima, **dal96}, 2026)
        assert abs(misto - solo96) < 0.01 and abs(nom - nom96) < 0.01, \
            f"{sorted(prima)}: montante {misto:,.2f} ≠ {solo96:,.2f} dei soli contributi dal 1996"
    rc, out = run_script(asset("it-pianificazione", "importa-previdenza.py"),
                         os.path.join(FIXTURES, "estratto-inps-test.xml"))
    assert rc == 0, out[-300:]
    assert "sono i primi anni" not in out and "non ancora pubblicat" in out, out[-600:]
    assert "prima del 1996" not in out, "carriera dal 2023: nessun avviso sugli anni prima del 1996"
    tmp = tempfile.mkdtemp()
    try:
        for inizio, intervallo in ((1995, "(1995,"), (1990, "(1990-1995,")):
            xml = os.path.join(tmp, f"da{inizio}.xml")
            _estratto_inps_xml({a: 30_000 for a in range(inizio, 2026)}, xml)
            rc, out = run_script(asset("it-pianificazione", "importa-previdenza.py"), xml)
            assert rc == 0, out[-300:]
            avviso = next((r for r in out.splitlines() if "prima del 1996" in r), "")
            assert intervallo in avviso and "retributiv" in avviso, f"carriera dal {inizio}: avviso «{avviso.strip()}»"
            assert "non sono rivalutati" not in out, f"carriera dal {inizio}: avviso falso sulla rivalutazione"
            atteso = m.ricostruisci_montante({a: 30_000.0 for a in range(1996, 2026)},
                                             __import__("datetime").date.today().year)[0]
            assert abs(euro(out, "MONTANTE CONTRIBUTIVO RICOSTRUITO") - atteso) <= 1, \
                f"carriera dal {inizio}: il montante stampato deve contare solo i contributi dal 1996"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("importa-previdenza: con --anonimo non stampa intestatario né data di nascita")
def t_importa_anonimo():
    # Pre-pubblicazione, privacy: tutto ciò che lo script stampa entra nel contesto del modello.
    xml = os.path.join(FIXTURES, "estratto-inps-test.xml")
    rc, out = run_script(asset("it-pianificazione", "importa-previdenza.py"), xml)
    assert rc == 0 and "ROSSI" in out and "15/03/1990" in out, "senza il flag lo script resta com'era"
    for args in ((xml, "--anonimo"), ("--anonimo", xml)):
        rc, out = run_script(asset("it-pianificazione", "importa-previdenza.py"), *args)
        assert rc == 0, out[-300:]
        for dato in ("ROSSI", "MARIO", "15/03/1990", "1990)"):
            assert dato not in out, f"--anonimo: stampato «{dato}»"
        eta = re.search(r"ETA_ATTUALE\s*=\s*([\d.]+)", out)
        assert eta and "." not in eta.group(1), "con --anonimo l'età resta, in anni interi"
        assert "MONTANTE_ATTUALE" in out and "omesso" in out
    # età arrotondata, non troncata: proiezione-inps.py usa round(ETA_PENSIONAMENTO − ETA_ATTUALE)
    m = carica_modulo(asset("it-pianificazione", "importa-previdenza.py"))
    assert m.eta_stampata(36.6, True) == 37 and m.eta_stampata(36.4, True) == 36 and m.eta_stampata(36.6, False) == 36.6
    # un flag sbagliato non deve stampare i dati personali (revisione delle note del § 4)
    csv_fondo = os.path.join(FIXTURES, "storico-fondo-test.csv")
    for args in ((xml, "--anonymous"), (xml, "-anonimo"), (xml, "—anonimo"), (xml, csv_fondo, "--anonimi"),
                 (xml, csv_fondo, "altro.csv")):
        rc, out = run_script(asset("it-pianificazione", "importa-previdenza.py"), *args)
        assert rc == 2 and "ROSSI" not in out and "15/03/1990" not in out and "usage:" in out, \
            f"argomenti {args[1:]}: rc {rc}, output «{out.strip()[:120]}»"


@test("netto-pensione: netto dello stipendio con le misure L. 207/2024 c. 4-6 e il contributo IVS aggiuntivo")
def t_netto_stipendio():
    netto = _modulo_previdenza("netto-pensione.py")
    # RAL 30.000: reddito 27.243 → ulteriore detrazione di 1.000 € (c. 6, reddito tra 20.000 e 32.000)
    _, contr, irpef, _ = netto.netto_lavoratore(30_000, netto.ALIQ_CONTRIBUTI_DIP, 0.0, 0.0)
    imponibile = 30_000 - 30_000 * 0.0919
    atteso = imponibile * 0.23 - (1_910 + 1_190 * (28_000 - imponibile) / 13_000) - 1_000
    assert abs(irpef - atteso) < 0.01, f"RAL 30k: IRPEF {irpef:,.2f} ≠ {atteso:,.2f}"
    # RAL 36.000: reddito 32.692,40 → detrazione decrescente 1.000 × (40.000 − R) / 8.000
    _, _, irpef, _ = netto.netto_lavoratore(36_000, netto.ALIQ_CONTRIBUTI_DIP, 0.0, 0.0)
    r = 36_000 * (1 - 0.0919)
    atteso = 6_440 + (r - 28_000) * 0.33 - 1_910 * (50_000 - r) / 22_000 - 1_000 * (40_000 - r) / 8_000
    assert abs(irpef - atteso) < 0.01, f"RAL 36k: IRPEF {irpef:,.2f} ≠ {atteso:,.2f}"
    # RAL 18.000: reddito 16.345,80 ≤ 20.000 → somma esente del 4,8% del reddito di lavoro (c. 4)
    n18, _, irpef18, _ = netto.netto_lavoratore(18_000, netto.ALIQ_CONTRIBUTI_DIP, 0.0, 0.0)
    r = 18_000 * (1 - 0.0919)
    assert abs(n18 - (r - irpef18 + r * 0.048)) < 0.01, f"RAL 18k: netto {n18:,.2f} senza la somma del c. 4"
    # RAL 80.000: 1% aggiuntivo sulla parte oltre la prima fascia di €56.224
    _, contr, _, _ = netto.netto_lavoratore(80_000, netto.ALIQ_CONTRIBUTI_DIP, 0.0, 0.0)
    assert abs(contr - (80_000 * 0.0919 + (80_000 - 56_224) * 0.01)) < 0.01, f"RAL 80k: contributi {contr:,.2f}"
    # RAL 150.000: niente contributi oltre il massimale di €122.295 (revisione T5, mutante «base = ral»)
    _, contr, _, _ = netto.netto_lavoratore(150_000, netto.ALIQ_CONTRIBUTI_DIP, 0.0, 0.0)
    assert abs(contr - (122_295 * 0.0919 + (122_295 - 56_224) * 0.01)) < 0.01, f"RAL 150k: contributi {contr:,.2f}"
    # fasce del c. 4: 7,1% fino a 8.500 di reddito, 5,3% fino a 15.000 (limiti inclusi)
    for ral, quota in ((9_000, 0.071), (12_000, 0.053)):
        n, _, irpef, _ = netto.netto_lavoratore(ral, netto.ALIQ_CONTRIBUTI_DIP, 0.0, 0.0)
        r = ral * (1 - 0.0919)
        assert abs(n - (r - irpef + r * quota)) < 0.01, f"RAL {ral}: netto {n:,.2f}, somma esente al {quota:.1%} attesa"
    assert netto.netto_lavoratore(9_000, netto.ALIQ_CONTRIBUTI_DIP, 0.0, 0.0)[2] == 0, "RAL 9k: IRPEF dovuta"
    for r, quota in ((8_500, 0.071), (15_000, 0.053), (20_000, 0.048)):
        assert abs(netto.cuneo_l207(r)[0] - r * quota) < 0.01, f"reddito {r}: «non superiore a» include il limite"


def _file_testo(*basi):
    for base in basi:
        for cartella, _, files in os.walk(base):
            for f in sorted(files):
                if f.endswith((".md", ".py")):
                    p = os.path.join(cartella, f)
                    yield os.path.relpath(p, ROOT).replace(os.sep, "/"), open(p, encoding="utf-8").read()


@test("coerenza tra skill: massimale, prima fascia, plafond e assegno sociale 2026 uguali in tutte le copie")
def t_coerenza_parametri_2026():
    # Decisione 5 della roadmap: copie dei dati normativi in ogni skill, con un test di coerenza.
    from openpyxl import load_workbook
    regole = {
        "massimale": r"(?:[Mm]assimale(?: annuo)?(?: contributivo)?(?: 2026)?|[Tt]etto contributivo(?: annuo)?(?: 2026)?)"
                     r"\**:?\**\s*(?:di\s*)?€\s*(\d{3}\.\d{3})|massimale_(\d{6})",
        "prima fascia": r"prima fascia[^€\n]{0,60}€\s*(\d{2}\.\d{3})(?![\d.,])",
        "plafond": r"(?:[Dd]educibil\w*|plafond|[Tt]etto di deducibilità)[^€\n]{0,40}€\s*(\d\.\d{3}(?:,\d{2})?)(?![\d.])",
        "assegno sociale": r"assegno sociale[^€\n]{0,60}€\s*(\d{3},\d{2}) al mese",
    }
    trovati = {k: {} for k in regole}
    for rel, testo in _file_testo(SKILLS, os.path.join(ROOT, "commands")):
        for riga in testo.splitlines():
            if "120.607" in riga:
                assert "2025" in riga, f"{rel}: 120.607 (massimale 2025) senza l'anno: {riga.strip()[:120]}"
            for nome, rx in regole.items():
                for m in re.finditer(rx, riga):
                    v = next(g for g in m.groups() if g)
                    v = v.replace(".", "").replace(",", ".")
                    trovati[nome].setdefault(float(v), set()).add(rel)
    for nome, valori in trovati.items():
        assert len(valori) == 1, f"{nome}: valori diversi tra le copie {valori}"
    massimale = next(iter(trovati["massimale"]))
    prima_fascia = next(iter(trovati["prima fascia"]))
    assert _modulo_previdenza("proiezione-inps.py").MASSIMALE_CONTRIBUTIVO == massimale
    np_ = _modulo_previdenza("netto-pensione.py")
    assert np_.MASSIMALE_CONTRIBUTIVO == massimale and np_.PRIMA_FASCIA_PENSIONABILE == prima_fascia
    wb = load_workbook(asset(*PP))
    assert _per_etichetta(wb["Input"], "Massimale contributivo") == massimale, "massimale dell'Excel dipendenti"
    assert _per_etichetta(load_workbook(asset(*PA))["Input"], "Massimale contributivo") == massimale, \
        "massimale dell'Excel autonomi"
    plafond = next(iter(trovati["plafond"]))
    for skill, nome in (PP, ("it-pianificazione-autonomi", "proiezione-pensione-autonomi.xlsx")):
        assert _per_etichetta(load_workbook(asset(skill, nome))["Input"], "Plafond deducibilità") == plafond, \
            f"{nome}: plafond diverso da €{plafond:,.0f}"


def _colonna(ws, testo, riga=4):
    """Lettera della colonna la cui intestazione (riga 4) contiene testo."""
    from openpyxl.utils import get_column_letter
    for c in range(1, ws.max_column + 1):
        if testo.lower() in str(ws.cell(riga, c).value or "").lower():
            return get_column_letter(c)
    raise AssertionError(f"{ws.title}: nessuna colonna «{testo}» in riga {riga}")


def _irpef(r):
    return r * 0.23 if r <= 28_000 else 6_440 + (r - 28_000) * 0.33 if r <= 50_000 else 13_700 + (r - 50_000) * 0.43


def _base_uscita_fondo(i, g, n, infl):
    """Base imponibile in uscita del fondo dei dipendenti, in euro di oggi: TFR conferito e contributi
    dedotti (lavoratore + datore entro il plafond), al valore nominale riportato a oggi. Esclusi i
    rendimenti già tassati e i contributi non dedotti comunicati (art. 11 c. 6 D.Lgs. 252/2005)."""
    plafond = _per_etichetta(i, "Plafond deducibilità")
    dovuto = _per_etichetta(i, "Contributo dovuto lavoratore")
    base = (i["B22"].value * i["B21"].value + i["B25"].value) * (1 + infl) ** -n
    for t in range(n):
        ral = i["B10"].value * (1 + g) ** t
        dedotti = min(plafond, i["B26"].value + ral * (dovuto + i["B27"].value))
        base += (ral * 0.0691 * i["B21"].value + dedotti) * (1 + infl) ** (t - n)
    return base


VARIANTI_T4 = {
    "ral60k": {("Input", "B10"): 60_000, ("Input", "B27"): 0.02, ("Input", "B26"): 5_000},
    "ral60k_dovuto": {("Input", "B10"): 60_000, ("Input", "B27"): 0.02, ("Input", "B47"): 0.01},
    "rend6": {("Input", "B28"): 0.06},
    "tfr_azienda": {("Input", "B21"): 0},
    "tfr40k": {("Input", "B5"): 37, ("Input", "B6"): 67, ("Input", "B10"): 40_000, ("Input", "B12"): 0.0,
               ("Input", "B22"): 0, ("Input", "B48"): 0, ("Input", "B29"): 0, ("Input", "B17"): 0.02,
               ("Input", "B28"): 0.025, ("Input", "B21"): 1},
    "tfr_azienda_senza_saldo": {("Input", "B21"): 0, ("Input", "B22"): 0},
}


@test("proiezione-pensione.xlsx: plafond condiviso col datore, quota deducibile e spazio residuo (RAL 60k, datore 2% → €4.100)")
def t_plafond_excel():
    # REGRESSIONE F007: Complementare!E5 tappava solo il volontario a 5.300 e B64 deduceva tutto il
    # volontario anche quando il datore aveva già usato parte del plafond.
    v = _varianti(*PP, VARIANTI_T4)
    for nome, residuo, deducibile, non_dedotto in (("ral60k", 4_100, 4_100, 900), ("ral60k_dovuto", 3_500, 2_000, 0)):
        wb = v[nome]
        i, c = wb["Input"], wb["Complementare"]
        assert abs(_per_etichetta(c, "Spazio deducibile residuo") - residuo) < 0.01, \
            f"{nome}: spazio residuo {_per_etichetta(c, 'Spazio deducibile residuo')} ≠ {residuo}"
        q, nd, e = _colonna(c, "Quota deducibile"), _colonna(c, "Non dedotto"), _colonna(c, "Contrib. lavoratore")
        assert abs(c[f"{q}5"].value - deducibile) < 0.01, f"{nome}: quota deducibile {c[f'{q}5'].value}"
        assert abs(c[f"{nd}5"].value - non_dedotto) < 0.01, f"{nome}: non dedotto {c[f'{nd}5'].value}"
        attesi = i["B26"].value + i["B10"].value * _per_etichetta(i, "Contributo dovuto lavoratore")
        assert abs(c[f"{e}5"].value - attesi) < 0.01, f"{nome}: il contributo del lavoratore va tutto a montante"
        # replica indipendente anno per anno (revisione T4: le righe dopo la prima non erano vincolate)
        plafond = _per_etichetta(i, "Plafond deducibilità")
        dovuto = _per_etichetta(i, "Contributo dovuto lavoratore")
        deducibili = non_dedotti = 0.0
        for t in range(int(i["B36"].value)):
            ral = i["B10"].value * (1 + i["B12"].value) ** t
            versato = i["B26"].value + ral * (dovuto + i["B27"].value)
            deducibili += min(plafond, versato) - min(plafond, ral * (dovuto + i["B27"].value))
            non_dedotti += max(0.0, versato - plafond)
        assert abs(_per_etichetta(c, "Risparmio IRPEF") - deducibili * i["B37"].value) < 0.01, \
            f"{nome}: risparmio IRPEF {_per_etichetta(c, 'Risparmio IRPEF'):,.2f} ≠ {deducibili * i['B37'].value:,.2f}"
        assert abs(_per_etichetta(c, "Non dedotto da comunicare") - non_dedotti) < 0.01, \
            f"{nome}: non dedotto {_per_etichetta(c, 'Non dedotto da comunicare'):,.2f} ≠ {non_dedotti:,.2f}"
    formula = load_workbook_formule(asset(*PP))["Complementare"][f"{_colonna(v['ral60k']['Complementare'], 'Quota deducibile')}5"].value
    assert "Input!$B$44" in formula, f"la quota deducibile deve usare la cella del plafond: {formula}"


def load_workbook_formule(path):
    from openpyxl import load_workbook
    return load_workbook(path)


@test("proiezione-pensione.xlsx: tassa in uscita solo su TFR e contributi dedotti, TFR in azienda fuori dal fondo")
def t_uscita_fondo_excel():
    # REGRESSIONE F014/F037: B60 tassava l'intero montante (rendimenti già tassati compresi) e
    # G5/H5 sommavano al fondo il saldo TFR anche con il TFR lasciato in azienda.
    from openpyxl import load_workbook
    base = load_workbook(xlsx_ricalcolati()[asset(*PP)], data_only=True)
    v = _varianti(*PP, VARIANTI_T4)
    for nome, wb in (("default", base), ("rend6", v["rend6"]), ("tfr_azienda", v["tfr_azienda"]),
                     ("ral60k", v["ral60k"]), ("ral60k_dovuto", v["ral60k_dovuto"])):
        i, c = wb["Input"], wb["Complementare"]
        n = int(i["B36"].value)
        attesa = _base_uscita_fondo(i, i["B12"].value, n, i["B17"].value) * i["B38"].value
        assert abs(c["B60"].value - attesa) < 0.01, f"{nome}: tassa in uscita {c['B60'].value:,.2f} ≠ {attesa:,.2f}"
    # ral60k: 900 € l'anno oltre il plafond, non dedotti e quindi fuori dalla base (revisione T4, NV0)
    i = v["ral60k"]["Input"]
    n = int(i["B36"].value)
    tutti = (i["B22"].value * i["B21"].value + i["B25"].value) * (1 + i["B17"].value) ** -n + sum(
        (i["B10"].value * (1 + i["B12"].value) ** t * (0.0691 * i["B21"].value + i["B27"].value) + i["B26"].value)
        * (1 + i["B17"].value) ** (t - n) for t in range(n))
    assert v["ral60k"]["Complementare"]["B60"].value < tutti * i["B38"].value - 1, \
        "ral60k: i contributi non dedotti non vanno tassati in uscita"
    assert abs(v["rend6"]["Complementare"]["B60"].value - base["Complementare"]["B60"].value) < 0.01, \
        "il rendimento non deve cambiare la base imponibile in uscita"
    assert v["rend6"]["Complementare"]["B58"].value > base["Complementare"]["B58"].value
    c, i = v["tfr_azienda"]["Complementare"], v["tfr_azienda"]["Input"]
    assert c["D5"].value == 0 and abs(c["H5"].value - (i["B25"].value * (1 + i["B28"].value) + c["E5"].value + c["F5"].value)) < 0.01, \
        "con il TFR in azienda il saldo TFR non entra nel montante del fondo"
    f = load_workbook_formule(asset(*PP))["Complementare"]
    assert "Input!$B$22*Input!$B$21" in f["G5"].value and "Input!$B$22*Input!$B$21" in f["H5"].value
    # Scenari: la prestazione netta di ogni colonna usa la stessa base, con crescita RAL e inflazione della colonna
    for nome, wb in (("default", base), ("ral60k", v["ral60k"])):
        s, i = wb["Scenari"], wb["Input"]
        n = int(i["B36"].value)
        for col in "BCD":
            g, infl, rend = s[f"{col}5"].value, s[f"{col}7"].value, s[f"{col}9"].value
            netta = _fondo_scenario(wb, g, rend, n) - _base_uscita_fondo(i, g, n, infl) * i["B38"].value
            assert abs(s[f"{col}14"].value - netta) <= 1e-6 * netta, \
                f"{nome}: Scenari!{col}14 = {s[f'{col}14'].value:,.0f} ≠ {netta:,.0f}"


def _tfr_confronto(i):
    """Confronto TFR in azienda vs nel fondo, in euro di oggi, replicato dalle regole di legge:
    quota RAL × 6,91%; in azienda rivalutazione 1,5% + 75% dell'inflazione tassata al 17% e
    tassazione separata all'aliquota media sul reddito di riferimento (base × 12 / anni);
    nel fondo rendimento del comparto e 15%→9% sulle sole quote (art. 2120 c.c., art. 19 TUIR,
    art. 11 c. 6 D.Lgs. 252/2005)."""
    n, infl, rend = int(i["B36"].value), i["B17"].value, i["B28"].value
    r_tfr = (1 + (0.015 + 0.75 * infl) * 0.83) / (1 + infl) - 1
    # il TFR già maturato in azienda ci resta (art. 8 c. 7 D.Lgs. 252/2005: si conferisce il maturando):
    # se non è già nel fondo, la decisione riguarda solo le quote future (revisione T4)
    saldo = i["B22"].value * i["B21"].value
    azienda = fondo = saldo
    base = saldo * (1 + infl) ** -n
    for t in range(n):
        quota = i["B10"].value * (1 + i["B12"].value) ** t * 0.0691
        azienda = azienda * (1 + r_tfr) + quota
        fondo = fondo * (1 + rend) + quota
        base += quota * (1 + infl) ** (t - n)
    rif = base * 12 / (_per_etichetta(i, "Anni di servizio già maturati") * i["B21"].value + n)
    aliq_sep = _irpef(rif) / rif if rif > 0 else 0.0
    netto_azienda = azienda - base * aliq_sep
    netto_fondo = fondo - base * i["B38"].value
    return netto_azienda, netto_fondo


@test("proiezione-pensione.xlsx: confronto TFR in azienda vs fondo (RAL 40k per 30 anni) e TFR in azienda negli Scenari")
def t_tfr_confronto():
    # F037: il confronto era promesso dalla skill ma nessuno strumento lo calcolava.
    from openpyxl import load_workbook
    v = _varianti(*PP, VARIANTI_T4)
    base = load_workbook(xlsx_ricalcolati()[asset(*PP)], data_only=True)
    sopra = _varianti(*PP, VARIANTI_PP)["ral150k"]  # reddito di riferimento oltre €50.000: scaglione al 43%
    for nome, wb in (("tfr40k", v["tfr40k"]), ("default", base), ("tfr_azienda", v["tfr_azienda"]),
                     ("ral150k", sopra)):
        i, c = wb["Input"], wb["Complementare"]
        azienda, fondo = _tfr_confronto(i)
        assert abs(_per_etichetta(c, "TFR in azienda NETTO") - azienda) < 0.01, \
            f"{nome}: TFR in azienda netto {_per_etichetta(c, 'TFR in azienda NETTO'):,.2f} ≠ {azienda:,.2f}"
        assert abs(_per_etichetta(c, "TFR nel fondo NETTO") - fondo) < 0.01, f"{nome}: TFR nel fondo netto"
        assert abs(_per_etichetta(c, "Vantaggio netto del fondo") - (fondo - azienda)) < 0.01, f"{nome}: delta"
    assert _per_etichetta(sopra["Complementare"], "Aliquota media della tassazione separata") > 0.33, \
        "ral150k: il reddito di riferimento supera €50.000 e deve toccare lo scaglione al 43%"
    # il saldo già in azienda resta lì in entrambi i rami: il conto della decisione non dipende dal saldo
    d_saldo = _per_etichetta(v["tfr_azienda"]["Complementare"], "Vantaggio netto del fondo")
    d_zero = _per_etichetta(v["tfr_azienda_senza_saldo"]["Complementare"], "Vantaggio netto del fondo")
    assert abs(d_saldo - d_zero) < 0.01, f"TFR in azienda: delta {d_saldo:,.2f} con il saldo, {d_zero:,.2f} senza"
    # valore chiuso del caso RAL 40k × 30 anni (inflazione 2%, comparto 2,5% reale, nessun saldo)
    delta = _per_etichetta(v["tfr40k"]["Complementare"], "Vantaggio netto del fondo")
    # azienda: lordo 88.963,53, base 61.903,80 al 23% → netto 74.725,66; fondo: 121.347,07 − 10,5% della base
    assert abs(delta - DELTA_TFR40K) < 0.5, f"RAL 40k × 30 anni: vantaggio del fondo {delta:,.2f} ≠ {DELTA_TFR40K:,.2f}"
    # la forbice della reference contiene l'esempio calcolato (revisione T4)
    azienda, fondo = _tfr_confronto(v["tfr40k"]["Input"])
    pc = open(os.path.join(SKILLS, "it-pianificazione", "references", "previdenza-complementare.md"), encoding="utf-8").read()
    lo, hi = map(int, re.search(r"\+(\d+)-(\d+)% di prestazione netta", pc).groups())
    assert lo <= (fondo / azienda - 1) * 100 <= hi, f"forbice +{lo}-{hi}% senza l'esempio (+{(fondo / azienda - 1) * 100:.0f}%)"
    s = v["tfr_azienda"]["Scenari"]
    assert abs(s["C16"].value - _tfr_azienda_scenario(v["tfr_azienda"], "C")) < 0.01, \
        "Scenari: con il TFR in azienda il suo netto va mostrato come capitale fuori dalla rendita"
    for col in "BD":  # colonne con crescita RAL e inflazione proprie (revisione T4, NV2)
        assert abs(s[f"{col}16"].value - _tfr_azienda_scenario(v["tfr_azienda"], col)) < 0.01, f"Scenari!{col}16"
        assert abs(s[f"{col}34"].value - s[f"{col}16"].value * 0.055) < 0.01, f"Scenari!{col}34: rendita equivalente"
    assert base["Scenari"]["C16"].value == 0, "con il TFR nel fondo la riga del TFR in azienda è zero"


DELTA_TFR40K = 40_121.51


def _tfr_azienda_scenario(wb, col):
    """TFR effettivamente in azienda (saldo compreso) netto al pensionamento con crescita RAL e
    inflazione della colonna degli Scenari: rivalutazione 1,5% + 75% dell'inflazione tassata al 17%,
    tassazione separata sul reddito di riferimento (saldo e quote, anni già maturati + residui)."""
    i, s = wb["Input"], wb["Scenari"]
    n, g, infl = int(i["B36"].value), s[f"{col}5"].value, s[f"{col}7"].value
    r_tfr = (1 + (0.015 + 0.75 * infl) * 0.83) / (1 + infl) - 1
    lordo, base = i["B22"].value, i["B22"].value * (1 + infl) ** -n
    for t in range(n):
        quota = i["B10"].value * (1 + g) ** t * 0.0691
        lordo = lordo * (1 + r_tfr) + quota
        base += quota * (1 + infl) ** (t - n)
    rif = base * 12 / (_per_etichetta(i, "Anni di servizio già maturati") + n)
    return lordo - base * (_irpef(rif) / rif)


PA = ("it-pianificazione-autonomi", "proiezione-pensione-autonomi.xlsx")
VARIANTI_PA_T4 = {
    "forf_comunica": {("Input", "B28"): 3_000, ("Input", "B19"): 1},
    "forf_non_comunica": {("Input", "B28"): 3_000, ("Input", "B19"): 1, ("Input", "B47"): 0},
    "ordinario": {("Input", "B28"): 6_000, ("Input", "B19"): 2},
    "ordinario_montante": {("Input", "B28"): 6_000, ("Input", "B19"): 2, ("Input", "B31"): 20_000},
    "forf_montante": {("Input", "B28"): 3_000, ("Input", "B19"): 1, ("Input", "B31"): 20_000},
    "ordinario_plafond_2025": {("Input", "B28"): 6_000, ("Input", "B19"): 2, ("Input", "B48"): 5_164.57},
}


@test("proiezione-pensione-autonomi.xlsx: forfettario che comunica i non dedotti → tassa in uscita zero")
def t_uscita_fondo_autonomi():
    # REGRESSIONE F014: Complementare!B60 tassava l'intero versato anche per il forfettario puro, che
    # non deduce e, comunicando i non dedotti, in uscita non paga nulla su quella parte.
    v = _varianti(*PA, VARIANTI_PA_T4)
    for nome, dedotto_annuo, comunica in (("forf_comunica", 0, 1), ("forf_non_comunica", 0, 0), ("ordinario", 5_300, 1),
                                          ("ordinario_montante", 5_300, 1), ("forf_montante", 0, 1),
                                          ("ordinario_plafond_2025", 5_164.57, 1)):
        i, c = v[nome]["Input"], v[nome]["Complementare"]
        n, infl = int(i["B7"].value), i["B39"].value
        assert _per_etichetta(i, "Versamento dedotto") == dedotto_annuo, f"{nome}: dedotto {_per_etichetta(i, 'Versamento dedotto')}"
        tassato = dedotto_annuo + (0 if comunica else i["B28"].value - dedotto_annuo)
        attesa = (i["B31"].value * (1 + infl) ** -n + sum(tassato * (1 + infl) ** (t - n) for t in range(n))) * i["B38"].value
        tassa = _per_etichetta(c, "Tassa finale")
        assert abs(tassa - attesa) < 0.01, f"{nome}: tassa in uscita {tassa:,.2f} ≠ {attesa:,.2f}"
    assert _per_etichetta(v["forf_comunica"]["Complementare"], "Tassa finale") == 0
    # montante già accumulato: trattato come dedotto, al valore riportato a oggi (revisione T4, NV3)
    i = v["forf_montante"]["Input"]
    attesa = 20_000 * (1 + i["B39"].value) ** -int(i["B7"].value) * i["B38"].value
    assert abs(_per_etichetta(v["forf_montante"]["Complementare"], "Tassa finale") - attesa) < 0.01, "forf_montante"
    # un solo plafond nel foglio: Fiscalità deduce quanto Input!B46 (revisione T4)
    for nome in ("ordinario", "ordinario_plafond_2025"):
        f, i = v[nome]["Fiscalità"], v[nome]["Input"]
        assert abs(f["C8"].value - _per_etichetta(i, "Versamento dedotto")) < 0.01, \
            f"{nome}: Fiscalità!C8 = {f['C8'].value} ≠ versamento dedotto {_per_etichetta(i, 'Versamento dedotto')}"
        assert f["B8"].value == 0, f"{nome}: il forfettario non deduce il fondo"
    # nessun plafond cablato nelle formule dei due Excel
    for skill, nome_xlsx in (PP, PA):
        for ws in load_workbook_formule(asset(skill, nome_xlsx)).worksheets:
            for riga in ws.iter_rows():
                for cella in riga:
                    if isinstance(cella.value, str) and cella.value.startswith("="):
                        assert not re.search(r"(?<![\d.])(5300|5164)", cella.value), \
                            f"{nome_xlsx} {ws.title}!{cella.coordinate}: plafond cablato {cella.value[:80]}"


@test("proiezione-pensione.xlsx: rendimento del comparto dichiarato reale, default = montecarlo-pensione.py")
def t_rendimento_reale_fondo():
    # Revisione T5/T4: il foglio è in euro di oggi ma il campo era «Rendimento netto comparto» al 3,5%,
    # contro il 2,5% reale del Monte Carlo.
    from openpyxl import load_workbook
    wb = load_workbook(asset(*PP))
    i = wb["Input"]
    mc = _modulo_previdenza("montecarlo-pensione.py")
    assert "reale" in i["A28"].value.lower() and i["B28"].value == mc.FP_RENDIMENTO_ATTESO_REALE, \
        f"Input!A28 «{i['A28'].value}» = {i['B28'].value}, Monte Carlo {mc.FP_RENDIMENTO_ATTESO_REALE}"
    assert "inflazione" in i["C28"].value
    assert "reale" in wb["Scenari"]["A9"].value.lower()


@test("previdenza complementare: base in uscita, non dedotti, nuove erogazioni 2026, capitale al 50%, marcatore")
def t_testi_fondo():
    pc = open(os.path.join(SKILLS, "it-pianificazione", "references", "previdenza-complementare.md"), encoding="utf-8").read()
    for atteso in ("durata definita", "frazionata"):
        assert atteso in pc, f"previdenza-complementare.md: manca «{atteso}»"
    # ancorati alle frasi nuove (revisione T4: le stringhe sciolte erano già nel file di prima)
    assert re.search(r"non dedott[^\n]*31/12[^\n]*9-15%", pc), "perdita del 9-15% sui non dedotti non comunicati"
    assert "si applica al TFR conferito e ai contributi dedotti" in pc, "base in uscita sulle sole quote dedotte"
    # vantaggio del TFR nel fondo: niente stime vecchie, vantaggio fiscale separato da quello di rendimento
    for vietato in ("decine di migliaia", "15-25%", "27% medio"):
        assert vietato not in pc, f"previdenza-complementare.md: stima superata «{vietato}»"
    # nuove forme di erogazione con i vincoli (art. 11 c. 3-quater e 3-quinquies, delibera COVIP 25/6/2026)
    nuove = pc[pc.index("### Nuove forme di erogazione"):pc.index("## Heuristic per la skill")]
    for atteso in ("rate maturate", "irrevocabil", "RITA", "31/12/2026", "3-quinquies"):
        assert atteso in nuove, f"nuove forme di erogazione: manca il vincolo «{atteso}»"
    # § 2.7: vincoli di legge riletti su Normattiva e vincoli della delibera COVIP 25/6/2026 letti in
    # GU n. 155 del 7/7/2026 (revisione T14: il testo è in Gazzetta), separati per fonte
    legge, covip = nuove.split("Vincoli delle istruzioni COVIP")
    assert "in luogo della rendita vitalizia" in legge and "rate maturate e non riscosse" in legge, "vincoli di legge"
    for solo_covip in ("irrevocabil", "RITA", "non si combinano", "revoca", "non si versa", "transitorio"):
        assert solo_covip not in legge, f"vincolo della delibera COVIP presentato come vincolo di legge: «{solo_covip}»"
    assert "GU n. 155 del 7/7/2026" in covip and "non riletti" not in covip, "vincoli COVIP senza la fonte in Gazzetta"
    assert "nuovo rapporto di lavoro" in covip, "durante l'erogazione si versa ancora il TFR di un nuovo rapporto (par. 4)"
    assert "art. 258 c. 5" in legge, "c. 6-bis senza la corrispondenza 2027"
    marcatore = next(r for r in pc.splitlines() if "201-204" in r and "Verificato: 2026-10" in r)
    verificato = marcatore.split("— Verificato")[0]
    assert "delibera del 25/6/2026, letta in GU n. 155" in verificato and "restano fuori" not in marcatore, marcatore[-300:]
    # la copia in it-rebalance (sequenza-liquidazione.md) cita la stessa fonte
    sl_testo = open(os.path.join(SKILLS, "it-rebalance", "references", "sequenza-liquidazione.md"), encoding="utf-8").read()
    riga_sl = next(r for r in sl_testo.splitlines() if "irrevocabil" in r)
    assert "COVIP" in riga_sl and "GU n. 155" in riga_sl, riga_sl[:200]
    assert "mix 50/50" not in pc, "regola pratica capitale/rendita come default consulenziale"
    # esempio RAL 60k: gli ultimi ~850 € della deduzione cadono nello scaglione al 33%
    netto = _modulo_previdenza("netto-pensione.py")
    imp = 60_000 - netto.contributi_lavoratore(60_000, netto.ALIQ_CONTRIBUTI_DIP)
    risparmio = netto.irpef_lorda(imp) - netto.irpef_lorda(imp - 5_300)
    assert f"€{round(risparmio):,}".replace(",", ".") in pc, f"esempio RAL 60k: risparmio atteso €{risparmio:,.0f}"
    assert "€2.279" not in pc
    # doppia citazione dal 2027 della base in uscita (art. 258 c. 4 TU imposte sui redditi, D.Lgs. 117/2026)
    for rel in ("it-pianificazione/references/previdenza-complementare.md", "it-pianificazione/SKILL.md",
                "it-pianificazione/references/proiezione-tenore-vita.md",
                "it-pianificazione-autonomi/references/fiscalita-forfettario.md"):
        testo = open(os.path.join(SKILLS, rel), encoding="utf-8").read()
        for par in testo.split("\n\n"):
            if "art. 11 c. 6 D.Lgs. 252/2005" in par:
                assert "258" in par, f"{rel}: art. 11 c. 6 D.Lgs. 252/2005 senza «dal 2027 art. 258»: {par.strip()[:90]}"
    # stessa regola nelle altre skill: it-rebalance e l'esempio J1
    sl = open(os.path.join(SKILLS, "it-rebalance", "references", "sequenza-liquidazione.md"), encoding="utf-8").read()
    for vietato in ("almeno il 50% va in rendita vitalizia", "solo sulla parte di capitale", "non a chiamata"):
        assert vietato not in sl, f"sequenza-liquidazione.md: «{vietato}»"
    for atteso in ("durata definita", "prelievi", "contributi dedotti"):
        assert atteso in sl, f"sequenza-liquidazione.md: manca «{atteso}»"
    joint = open(os.path.join(SKILLS, "it-pianificazione", "references", "integrazione-fiscale-joint.md"), encoding="utf-8").read()
    assert "€86.700" not in joint and "11,4%" in joint, "J1: tassa in uscita sull'intero montante"
    sk = open(os.path.join(SKILLS, "it-pianificazione", "SKILL.md"), encoding="utf-8").read()
    assert "max 50% montante in molte forme" not in sk
    assert "Verificato: 2026-06" not in pc, "marcatore di giugno non aggiornato dopo la riverifica"
    assert re.search(r"201-204[^\n]*Verificato: 2026-10|Verificato: 2026-10[^\n]*201-204", pc), "marcatore senza i commi 201-204"
    assert "contributi + rendimenti già tassati" not in pc, "base imponibile in uscita con i rendimenti già tassati"
    assert not re.search(r"(?:fino al|massimo) 60%", pc), "il capitale resta al 50% (DL 62/2026)"
    ff = open(os.path.join(SKILLS, "it-pianificazione-autonomi", "references", "fiscalita-forfettario.md"), encoding="utf-8").read()
    assert "12.800" not in ff and "15→9% sulla prestazione" not in ff, "fiscalita-forfettario.md: uscita del forfettario ancora tassata"
    for rel in ("it-pianificazione/SKILL.md", "it-pianificazione-autonomi/SKILL.md"):
        testo = open(os.path.join(SKILLS, rel), encoding="utf-8").read()
        assert "tassazione agevolata della prestazione finale (15→9%)" not in testo, f"{rel}: uscita del forfettario"


VARIANTI_PA_T6 = {
    "ord150k": {("Input", "B19"): 2, ("Input", "B21"): 150_000, ("Input", "B22"): 0},
    # affitti tassati a IRPEF ordinaria: nessuna detrazione propria, il risparmio è l'IRPEF lorda
    "forf_altri12k": {("Input", "B19"): 1, ("Input", "B28"): 3_000, ("Input", "B52"): 12_000},
    "forf_altri2k": {("Input", "B19"): 1, ("Input", "B28"): 3_000, ("Input", "B52"): 2_000},
    "forf_altri40k": {("Input", "B19"): 1, ("Input", "B28"): 8_000, ("Input", "B52"): 40_000},
    "fondo3k": {("Input", "B28"): 3_000, ("Input", "B31"): 10_000, ("Input", "B32"): 30},
    # cassa con coefficiente proprio, scritto sopra Input!B43 come dice il foglio Coefficienti
    "cassa_coeff": {("Input", "B11"): 2, ("Input", "B43"): 0.045, ("Input", "B28"): 3_000, ("Input", "B31"): 10_000},
    # cassa che imposta solo la soggettiva: l'aliquota di computo la segue (revisione delle correzioni T6)
    "cassa_inarcassa": {("Input", "B11"): 2, ("Input", "B12"): 0.145},
    "gs_24": {("Input", "B12"): 0.24},
    "decadimento_zero": {("Input", "B15"): 0.0},
    "forf8k": {("Input", "B21"): 8_000},
}


def _pensione_autonomi(i, g, r, n, massimale, coeff):
    """Previdenza autonomi: stessa meccanica L. 335/1995 dei dipendenti, sul reddito professionale.
    A montante va l'aliquota di computo (Gestione Separata: 25% IVS), non quella versata (26,07%)."""
    return _pensione_contributiva(i["B13"].value, i["B23"].value, g, r, n, massimale,
                                  _per_etichetta(i, "Aliquota di computo"), coeff)


@test("proiezione-pensione-autonomi.xlsx: massimale (150k in ordinario → versati €31.882,31, a montante €30.573,75) e rivalutazione L. 335/1995")
def t_autonomi_massimale():
    # REGRESSIONE F015: nessuna cella applicava il massimale (contributo 39.105 invece di 31.882,
    # pensione +38,8%); Previdenza!F rivalutava anche il contributo dell'anno.
    from openpyxl import load_workbook
    base = load_workbook(xlsx_ricalcolati()[asset(*PA)], data_only=True)
    v = _varianti(*PA, VARIANTI_PA_T6)
    p, f = v["ord150k"]["Previdenza"], v["ord150k"]["Fiscalità"]
    # a montante il 25% IVS (Circ. INPS 8/2026: 26,07% = 25% IVS + 0,72% + 0,35% ISCRO); si deduce tutto il versato
    assert abs(p["D5"].value - 122_295 * 0.25) < 0.01, f"Previdenza!D5 = {p['D5'].value:,.2f}"
    assert abs(f["C7"].value - 122_295 * 0.2607) < 0.01, f"Fiscalità!C7 = {f['C7'].value:,.2f} senza massimale"
    assert _per_etichetta(base["Input"], "Aliquota di computo") == 0.25
    assert base["Input"]["B14"].value == _modulo_previdenza("proiezione-inps.py").PIL_REALE, \
        "rivalutazione di default degli autonomi ≠ PIL_REALE di proiezione-inps.py"
    formule = load_workbook_formule(asset(*PA))
    assert "Input!$B$51" in formule["Previdenza"]["D5"].value and "MIN(" in formule["Previdenza"]["D5"].value
    netto = _modulo_previdenza("netto-pensione.py")
    for nome in ("ord150k", "forf8k"):  # rami della detrazione oltre 28.000 e sotto la no tax area
        p = v[nome]["Previdenza"]
        assert abs(_per_etichetta(p, "NETTA annua") - netto.netto_pensionato(p["B60"].value, 0.02, 0.0)[0]) < 0.01, \
            f"{nome}: netto della pensione (lorda {p['B60'].value:,.0f})"
    assert v["forf8k"]["Previdenza"]["B60"].value <= 8_500 < 28_000 < v["ord150k"]["Previdenza"]["B60"].value
    for nome, wb in (("default", base), ("ord150k", v["ord150k"])):
        i, p = wb["Input"], wb["Previdenza"]
        n = int(i["B7"].value)
        atteso = _pensione_autonomi(i, i["B24"].value, i["B14"].value, n, _per_etichetta(i, "Massimale contributivo"), 1.0)
        assert abs(p["B58"].value - atteso) <= 1e-6 * atteso, f"{nome}: montante {p['B58'].value:,.0f} ≠ {atteso:,.0f}"


@test("proiezione-pensione-autonomi.xlsx: il forfettario deduce il fondo solo fino agli altri redditi IRPEF")
def t_autonomi_altri_redditi():
    # REGRESSIONE F039: Fiscalità!B8 = MIN(versamento; 5.300) × 0, cablato: l'eccezione del forfettario
    # con affitti o lavoro dipendente part-time non era calcolabile.
    from openpyxl import load_workbook
    base = load_workbook(xlsx_ricalcolati()[asset(*PA)], data_only=True)
    v = _varianti(*PA, VARIANTI_PA_T6)
    for nome, wb, dedotto in (("altri 12k", v["forf_altri12k"], 3_000), ("altri 2k", v["forf_altri2k"], 2_000),
                              ("altri 40k", v["forf_altri40k"], 5_300), ("default", base, 0)):
        i, f = wb["Input"], wb["Fiscalità"]
        altri = _per_etichetta(i, "Altri redditi IRPEF")
        assert _per_etichetta(i, "Versamento dedotto") == min(dedotto, i["B28"].value), f"{nome}: versamento dedotto"
        risparmio = _irpef(altri) - _irpef(altri - min(i["B28"].value, 5_300, altri))
        assert abs(f["B15"].value - risparmio) < 0.01, f"{nome}: risparmio forfettario {f['B15'].value:,.2f} ≠ {risparmio:,.2f}"
        # colonna Ordinario: gli altri redditi si sommano al reddito professionale (revisione T6)
        x = f["C6"].value - f["C7"].value + altri
        ded = min(i["B28"].value, _per_etichetta(i, "Plafond deducibilità"))
        assert abs(f["C8"].value - ded) < 0.01 and abs(f["C9"].value - max(0, x - ded)) < 0.01, f"{nome}: Ordinario C8/C9"
        assert abs(f["C14"].value - _irpef(max(0, x))) < 0.01 and abs(f["C10"].value - _irpef(f["C9"].value)) < 0.01
        assert abs(f["C11"].value - (x - f["C10"].value)) < 0.01, f"{nome}: Ordinario C11"
        assert abs(f["C15"].value - (_irpef(x) - _irpef(max(0, x - ded)))) < 0.01, f"{nome}: risparmio ordinario"
    assert abs(v["forf_altri12k"]["Fiscalità"]["B15"].value - 690) < 0.01
    assert abs(v["forf_altri12k"]["Fiscalità"]["C15"].value - 990) < 0.01, "ordinario con 12k di altri redditi: 3.000 al 33%"
    assert abs(v["forf_altri40k"]["Fiscalità"]["B15"].value - 1_749) < 0.01, "plafond e secondo scaglione"
    # il modello vale per redditi senza detrazioni proprie: col lavoro dipendente il risparmio è minore
    c52 = load_workbook_formule(asset(*PA))["Input"]["C52"].value
    assert "detrazion" in c52 and "part-time: danno capienza" not in c52, f"Input!C52: {c52}"


def _tassa_uscita_autonomi(i):
    """Tassa in uscita del fondo degli autonomi, in euro di oggi: montante già accumulato, versamenti
    dedotti e non dedotti NON comunicati, al nominale riportato a oggi (art. 11 c. 6 D.Lgs. 252/2005)."""
    n, infl = int(i["B7"].value), i["B39"].value
    ded = _per_etichetta(i, "Versamento dedotto")
    tassato = ded + (0 if i["B47"].value == 1 else i["B28"].value - ded)
    return (i["B31"].value * (1 + infl) ** -n + sum(tassato * (1 + infl) ** (t - n) for t in range(n))) * i["B38"].value


def _scenari_autonomi(wb, nome):
    """Tabella Scenari del Riepilogo autonomi contro repliche indipendenti, colonna per colonna:
    pensione lorda (massimale, aliquota di computo, coefficiente effettivo di Input!B43), netta con
    netto_pensionato(), rendita del fondo al netto della tassa in uscita, totale e gap mai letti dal foglio."""
    i, p, c, r = wb["Input"], wb["Previdenza"], wb["Complementare"], wb["Riepilogo"]
    netto = _modulo_previdenza("netto-pensione.py")
    righe = {str(r[f"A{k}"].value): k for k in range(1, r.max_row + 1) if r[f"A{k}"].value}
    def riga(testo):
        return next(k for e, k in righe.items() if e.startswith(testo))
    intest = riga("Voce")
    assert [r[f"{x}{intest}"].value for x in "BCD"] == ["Conservativo", "Base", "Ottimistico"], nome
    # Base = fogli di calcolo, anche con il coefficiente della cassa scritto in Input!B43
    assert abs(r[f"C{riga('Pensione obbligatoria lorda')}"].value - p["B60"].value) < 0.01, f"{nome}: Base lorda"
    assert abs(r[f"C{riga('Pensione obbligatoria netta')}"].value - _per_etichetta(p, "NETTA annua")) < 0.01, nome
    assert abs(r[f"C{riga('Rendita fondo')}"].value - c["B62"].value) < 0.01, f"{nome}: Base ≠ Complementare!B62"
    spesa = _per_etichetta(i, "Spesa mensile attuale") * 12 * _per_etichetta(i, "Fattore riduzione spesa")
    n, massimale = int(i["B7"].value), _per_etichetta(i, "Massimale contributivo")
    dec_base = i["B15"].value
    # coefficiente replicato dalla tabella (o quello della cassa scritto nella variante), mai letto da Input!B43
    coeff = VARIANTI_PA_T6["cassa_coeff"][("Input", "B43")] if nome == "cassa_coeff" else \
        _coeff_excel(wb, i["B6"].value) * (1 - dec_base) ** (n / 2)
    assert abs(p["B59"].value - coeff) < 1e-12, f"{nome}: coefficiente proiettato {p['B59'].value} ≠ {coeff}"
    tassa = _tassa_uscita_autonomi(i)
    lorde = []
    for col in "BCD":
        g, rv, d, rend = (r[f"{col}{riga(t)}"].value for t in ("Crescita reddito", "Rivalutazione montante",
                                                                "Decadimento", "Rendimento netto fondo"))
        coeff_col = coeff * ((1 - d) / (1 - dec_base)) ** (n / 2)
        attesa = _pensione_autonomi(i, g, rv, n, massimale, coeff_col)
        lorda = r[f"{col}{riga('Pensione obbligatoria lorda')}"].value
        assert abs(lorda - attesa) <= 1e-6 * attesa, f"{nome} {col}: pensione {lorda:,.2f} ≠ {attesa:,.2f}"
        netta = netto.netto_pensionato(attesa, 0.02, 0.0)[0]
        assert abs(r[f"{col}{riga('Pensione obbligatoria netta')}"].value - netta) < 0.01, f"{nome} {col}: netta"
        montante = i["B31"].value * (1 + rend) ** n + sum(i["B28"].value * (1 + rend) ** (n - t) for t in range(n))
        rendita = (montante - tassa) * coeff_col
        nel_foglio = r[f"{col}{riga('Rendita fondo')}"].value
        assert abs(nel_foglio - rendita) <= 1e-6 * max(rendita, 1), \
            f"{nome} {col}: rendita del fondo {nel_foglio:,.2f} ≠ {rendita:,.2f}"
        tot = netta + rendita
        assert abs(r[f"{col}{riga('Pensione totale netta')}"].value - tot) < 0.01, f"{nome} {col}: totale netto"
        assert abs(r[f"{col}{riga('Spesa target')}"].value - spesa) < 0.01, f"{nome} {col}: spesa target in euro di oggi"
        assert abs(r[f"{col}{riga('GAP')}"].value - (spesa - tot)) < 0.01, f"{nome} {col}: gap = spesa − pensione netta"
        lorde.append(attesa)
    assert lorde[0] < lorde[1] < lorde[2], f"{nome}: scenari non ordinati {lorde}"


@test("proiezione-pensione-autonomi.xlsx: pensione netta e tre scenari sul gap rispetto alla spesa target")
def t_autonomi_scenari():
    # REGRESSIONE F039: nessun netto, gap solo rispetto al reddito, «Scenari» = due righe.
    # Revisione T6: totale e gap verificati senza leggere il foglio; sopra il massimale; coefficiente
    # della cassa in Input!B43; tassa in uscita su dedotti e non dedotti non comunicati.
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_ricalcolati()[asset(*PA)], data_only=True)
    netto = _modulo_previdenza("netto-pensione.py")
    assert abs(_per_etichetta(wb["Previdenza"], "NETTA annua")
               - netto.netto_pensionato(wb["Previdenza"]["B60"].value, 0.02, 0.0)[0]) < 0.01, "pensione netta"
    v6, v4 = _varianti(*PA, VARIANTI_PA_T6), _varianti(*PA, VARIANTI_PA_T4)
    i = v6["ord150k"]["Input"]
    assert i["B23"].value > _per_etichetta(i, "Massimale contributivo"), "ord150k: il massimale deve scattare dall'anno 0"
    for nome, w in (("default", wb), ("ord150k", v6["ord150k"]), ("fondo3k", v6["fondo3k"]),
                    ("cassa_coeff", v6["cassa_coeff"]), ("ordinario", v4["ordinario"]),
                    ("forf_non_comunica", v4["forf_non_comunica"]), ("decadimento_zero", v6["decadimento_zero"])):
        _scenari_autonomi(w, nome)
    # aliquota di computo: segue la soggettiva della cassa; in Gestione Separata è l'IVS (25%, o 24%)
    for nome, attesa in (("cassa_inarcassa", 0.145), ("gs_24", 0.24)):
        i, p = v6[nome]["Input"], v6[nome]["Previdenza"]
        assert abs(_per_etichetta(i, "Aliquota di computo") - attesa) < 1e-12, f"{nome}: aliquota di computo"
        assert abs(p["D5"].value - p["C5"].value * attesa) < 0.01, f"{nome}: contributo a montante"
    assert v6["fondo3k"]["Complementare"]["B62"].value > 0 and v4["forf_non_comunica"]["Complementare"]["B60"].value > 0


@test("casse professionali: ogni riga della tabella ha una fonte ufficiale e il marcatore, dati 2026 corretti")
def t_casse_fonti():
    # F038/T6: tabella «indicativa, da verificare» senza fonti; Inarcassa con un 3% oltre il massimale
    # abolito dal 2013, Forense con il volontario all'1-10% (è 1-20%).
    testo = open(os.path.join(SKILLS, "it-pianificazione-autonomi", "references", "casse-professionali.md"),
                 encoding="utf-8").read()
    righe = [r for r in testo.splitlines() if r.startswith("| **")]
    assert len(righe) >= 6, f"tabella casse: {len(righe)} righe"
    for riga in righe:
        nome = riga.split("**")[1]
        link = re.findall(r"\]\((https://[^)]+)\)", riga)
        assert link and all(re.match(r"https://[^/]+/\S+", u) for u in link), \
            f"{nome}: ogni link deve puntare a una pagina ufficiale, non alla homepage: {link}"
        assert "Verificato: 2026-10" in riga, f"{nome}: manca il marcatore"
    per_nome = {r.split("**")[1]: r for r in righe}
    assert "2.785" in per_nome["Inarcassa"] and "3%" not in per_nome["Inarcassa"], per_nome["Inarcassa"][:160]
    assert "reddito medio" in per_nome["Inarcassa"] and "non oltre i 35" in per_nome["Inarcassa"], "Inarcassa: limiti dell'agevolazione under 35"
    # RGP art. 4.4 verificato: soglia = reddito medio del biennio precedente, contribuzione figurativa a 25 anni
    assert "da verificare sul regolamento" not in per_nome["Inarcassa"] and "25 anni" in per_nome["Inarcassa"], \
        "Inarcassa: condizioni under 35 dal RGP art. 4.4 (soglia di reddito, integrazione figurativa)"
    assert "20%" in per_nome["Cassa Forense"] and "131.800" in per_nome["Cassa Forense"], per_nome["Cassa Forense"][:160]
    assert "209.300" in per_nome["CNPADC"] and "856" in per_nome["ENPAP"], "CNPADC/ENPAP"


@test("previdenza: requisiti 2027-2028 completi nelle due reference, tassi storici veri, nessun dominio inesistente")
def t_requisiti_e_link_inps():
    # Revisione T5. Requisiti: Circ. INPS 28/2026, All. 2 (anticipata ordinaria 43a1m/42a1m dal 2028,
    # contributi dell'anticipata contributiva adeguati anche loro, finestra di 3 mesi).
    dip = open(os.path.join(SKILLS, "it-pianificazione", "references", "previdenza-obbligatoria.md"),
               encoding="utf-8").read()
    gs = open(os.path.join(SKILLS, "it-pianificazione-autonomi", "references", "previdenza-gestione-separata.md"),
              encoding="utf-8").read()
    for nome, testo in (("previdenza-obbligatoria.md", dip), ("previdenza-gestione-separata.md", gs)):
        for req in ("42 anni e 11 mesi", "43 anni e 1 mese", "42 anni e 1 mese", "20 anni e 1 mese", "20 anni e 3 mesi"):
            assert req in testo, f"{nome}: manca «{req}» tra i requisiti 2027-2028"
    assert not re.search(r"Dal 2027: 42 anni e 11 mesi", dip), "anticipata ordinaria: «dal 2027» copre anche il 2028"
    for req in ("71 anni e 3 mesi", "41 anni e 3 mesi"):
        assert req in dip, f"previdenza-obbligatoria.md: manca «{req}»"
    ponte = open(asset("it-pianificazione", "ponte-pensione.py"), encoding="utf-8").read()
    inizio = ponte.index('"""') + 3
    doc = " ".join(ponte[inizio:ponte.index('"""', inizio)].split())
    assert "20 anni e 1 mese" in doc and "20 anni e 3 mesi" in doc and re.search(r"(tre|3) mesi dopo", doc), \
        "ponte-pensione: anticipata contributiva senza i contributi adeguati o senza la finestra di 3 mesi"
    # tassi nominali di rivalutazione: medie dalla tabella ufficiale dello script, non «1,5-2,5%»
    tassi = carica_modulo(asset("it-pianificazione", "importa-previdenza.py")).TASSI_CAPITALIZZAZIONE
    def media(da, a):
        prodotto = 1.0
        for anno in range(da, a + 1):
            prodotto *= tassi[anno]
        return prodotto ** (1 / (a - da + 1)) - 1
    for da in (1996, 2006):
        cifra = f"~{media(da, 2025) * 100:.1f}%".replace(".", ",")
        assert f"{cifra} annuo nel {da}-2025" in dip, f"previdenza-obbligatoria.md: media {da}-2025 attesa {cifra}"
    assert "1,5-2,5% annuo negli ultimi decenni" not in dip
    # mypensione.it non è un dominio INPS: le credenziali SPID vanno solo sul portale INPS
    for rel, testo in _file_testo(SKILLS, os.path.join(ROOT, "commands")):
        assert "mypensione" not in testo.lower(), f"{rel}: rimando a «mypensione», dominio che non è dell'INPS"
        assert "PIN/CIE/SPID" not in testo, f"{rel}: il PIN INPS dei cittadini non esiste più dal 2021"
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    assert "mypensione" not in readme.lower(), "README: rimando a «mypensione»"
    # cuneo L. 207/2024: il c. 6 (ulteriore detrazione) vale fino al 2026, poi il TUIR 2027
    netto = open(asset("it-pianificazione", "netto-pensione.py"), encoding="utf-8").read()
    netto = " ".join(netto.split())
    assert re.search(r"c\. 6 vale fino al 2026", netto) and "D.Lgs. 117/2026" in netto, \
        "netto-pensione: il c. 6 della L. 207/2024 senza la doppia citazione dal 2027"
    # comma verificato sul testo in GU (S.O. 26 alla GU 152/2026): art. 13 c. 10, abrogazione art. 376 c. 1 lett. pppppp
    assert "art. 13 c. 10" in netto and "comma da verificare" not in netto, \
        "netto-pensione: l'ulteriore detrazione dal 2027 è l'art. 13 c. 10 D.Lgs. 117/2026"


# ══════════════════ it-pianificazione-obiettivi (goals-based) ══════════════════

def _excel_test(obiettivi):
    """Copia del pianificatore con gli obiettivi dati; ritorna il path."""
    from openpyxl import load_workbook
    tmp = tempfile.mkdtemp()
    dst = os.path.join(tmp, "p.xlsx")
    shutil.copy(asset("it-pianificazione-obiettivi", "pianificatore-obiettivi.xlsx"), dst)
    wb = load_workbook(dst)
    o = wb["Obiettivi"]
    for i, (nome, target, oriz, prof, gia, prio, tipo) in enumerate(obiettivi):
        r = 5 + i
        for col, v in zip("ABCDFGN", (nome, target, oriz, prof, gia, prio, tipo)):
            o[f"{col}{r}"] = v
    # F022: azzera TUTTE le righe residue (5..14): prima restava la riga demo 7 con un solo obiettivo
    for r in range(5 + len(obiettivi), 15):
        for col in "ABCDFGN":
            o[f"{col}{r}"] = None
    wb.save(dst)
    return dst


@test("glidepath: REGRESSIONE obiettivi senza target (entrano, marcati (max), aggregato coerente)")
def t_glidepath_senza_target():
    xlsx = _excel_test([
        ("Con target", 50000, 10, "Moderato", 10000, "importante", "terminante"),
        ("Senza target", None, 20, "Crescita massima", 20000, "desiderabile", "permanente"),
    ])
    rc, out = run_script(asset("it-pianificazione-obiettivi", "glidepath-obiettivi.py"), xlsx)
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    assert rc == 0, out[-400:]
    assert "(max" in out, "marcatore (max) assente: obiettivo senza target saltato?"
    assert "PORTAFOGLIO" in out
    # il permanente senza target non de-riska: la sua colonna resta al 90% anche all'ultimo anno
    ultima = [r for r in out.splitlines() if re.match(r"\s*20\s\s", r)]
    assert ultima and "90%" in ultima[-1], "il permanente senza target deve restare al 90%"


def _pianificatore_modificato(modifica):
    """Copia del pianificatore d'esempio a cui `modifica(wb)` applica un caso limite; ritorna il path."""
    from openpyxl import load_workbook
    tmp = tempfile.mkdtemp()
    dst = os.path.join(tmp, "p.xlsx")
    shutil.copy(asset("it-pianificazione-obiettivi", "pianificatore-obiettivi.xlsx"), dst)
    wb = load_workbook(dst)
    modifica(wb)
    wb.save(dst)
    return dst


@test("T9 lettore unico del pianificatore: profilo minuscolo, target vuoto, 1,5 anni, foglio mancante, già coperto")
def t_lettore_pianificatore():
    # F028: quattro lettori con filtri diversi; «crescita» dava 0% al glide path, KeyError all'allocatore e
    # «Salto» al Monte Carlo; il target vuoto spariva in silenzio; 1,5 anni diventava 1; la riserva già
    # piena risultava «da versamenti». Ora un solo lettore con le stesse regole per i quattro script.
    OB = "it-pianificazione-obiettivi"
    script = {n: asset(OB, n) for n in ("glidepath-obiettivi.py", "montecarlo-obiettivi.py",
                                        "allocatore-capitale.py", "asset-location-optimizer.py")}
    motore = carica_modulo(asset(OB, "motore_bootstrap.py"))
    avvisi = []
    d = motore.leggi_pianificatore(asset(OB, "pianificatore-obiettivi.xlsx"), avvisa=avvisi.append)
    assert [g["nome"] for g in d["obiettivi"]] == ["Fondo di emergenza", "Anticipo casa", "Lungo termine / libertà"]
    assert d["capacita"] == 1200 and d["inflazione"] == 0.02 and d["finestra"] == 7 and d["finestra_nel_foglio"]
    assert d["profili"]["Crescita"] == {"rendimento": 0.055, "volatilita": 0.13, "azionario": 0.65}
    casa = d["obiettivi"][1]
    assert casa == {"nome": "Anticipo casa", "target_oggi": 40000.0, "orizzonte": 6, "profilo": "Moderato",
                    "gia": 5000.0, "priorita": "importante", "tipo": "terminante", "permanente": False, "coperto": False}
    assert d["obiettivi"][0]["permanente"] and d["obiettivi"][0]["priorita"] == "essenziale"
    assert any("dentro il plugin" in a for a in avvisi), "il modello dentro il plugin va segnalato"
    # (a) profilo in minuscolo e con spazi: canonico della tabella, stessi numeri di «Crescita» in tutti e quattro
    xlsx = _excel_test([("Lungo", 100000, 20, " crescita ", 0, "desiderabile", "terminante")])
    d = motore.leggi_pianificatore(xlsx, avvisa=avvisi.append)
    assert d["obiettivi"][0]["profilo"] == "Crescita", d["obiettivi"][0]
    assert not any("dentro il plugin" in a for a in avvisi[1:]), "una copia fuori dal plugin non va segnalata"
    for nome, path in script.items():
        rc, out = run_script(path, xlsx)
        assert rc == 0 and "Traceback" not in out and "Salto" not in out and "non in tabella" not in out, f"{nome}: {out[-300:]}"
    rc, out = run_script(script["glidepath-obiettivi.py"], xlsx)
    assert re.search(r"^\s+1\s+65%", out, re.M) and "Lungo: €384/mese" in out, \
        f"glide path con «crescita»: attesi il 65% del profilo Crescita all'anno 1 e €384/mese (prima dava 0%): {out[:900]}"
    rc, out = run_script(script["montecarlo-obiettivi.py"], xlsx)
    assert "Rendimento 5,5% · volatilità 13,0%" in out, out[-600:]
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    # (b) target vuoto: il glide path e l'asset location lo tengono; Monte Carlo e allocatore lo escludono dicendolo
    xlsx = _excel_test([("Con target", 50000, 10, "Moderato", 10000, "importante", "terminante"),
                        ("Senza target", None, 15, "Crescita", 20000, "desiderabile", "terminante")])
    assert motore.leggi_pianificatore(xlsx)["obiettivi"][1]["target_oggi"] is None
    rc, out = run_script(script["glidepath-obiettivi.py"], xlsx)
    assert rc == 0 and "(max)" in out, out[-300:]
    for nome in ("montecarlo-obiettivi.py", "allocatore-capitale.py"):
        rc, out = run_script(script[nome], xlsx)
        assert rc == 0 and "Senza target: senza importo target" in out and "escluso" in out, f"{nome}: {out[-400:]}"
        assert "Con target" in out
    # revisione T9: all'asset location il target non serve (contano capitale, orizzonte e profilo),
    # quindi l'obiettivo «il più possibile» a lungo termine entra nel confronto PIR/dossier
    rc, out = run_script(script["asset-location-optimizer.py"], xlsx)
    assert rc == 0 and "▸ Senza target" in out and "scenario di esempio" not in out, out[-500:]
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    # revisione T9: due obiettivi con lo stesso nome mescolavano i calendari del glide path (chiave = nome)
    xlsx = _excel_test([("Anticipo casa", 40000, 6, "Moderato", 0, "importante", "terminante"),
                        (" anticipo CASA", 20000, 3, "Conservativo", 0, "importante", "terminante")])
    for nome, path in script.items():
        rc, out = run_script(path, xlsx)
        assert rc == 1 and "ERRORE:" in out and "ripetuto" in out and "Obiettivi!A6" in out and "Obiettivi!A5" in out, f"{nome}: {out[-300:]}"
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    # (c) orizzonte 1,5 anni: errore esplicito, non «1 anni»; (f) profilo ignoto: errore con la tabella
    for caso, atteso in ((("Casa", 40000, 1.5, "Moderato", 0, "importante", "terminante"), "anni interi"),
                         (("Casa", 40000, 6, "Boh", 0, "importante", "terminante"), "non è nella tabella dei profili")):
        xlsx = _excel_test([caso])
        for nome, path in script.items():
            rc, out = run_script(path, xlsx)
            assert rc == 1 and "ERRORE:" in out and atteso in out and "Traceback" not in out, f"{nome} {caso[2:4]}: rc {rc}, {out[-300:]}"
        shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    # (d) riserva già coperta: «già coperto», non «da versamenti»
    xlsx = _excel_test([("Fondo di emergenza", 12000, 2, "Preservazione", 12000, "essenziale", "permanente"),
                        ("Casa", 40000, 6, "Moderato", 5000, "importante", "terminante")])
    assert motore.leggi_pianificatore(xlsx)["obiettivi"][0]["coperto"] is True
    rc, out = run_script(script["allocatore-capitale.py"], xlsx, "50000")
    riga = next(r for r in out.splitlines() if r.strip().startswith("Fondo di emergenza"))
    assert rc == 0 and "già coperto" in riga and "da versamenti" not in riga, riga
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    # (e) foglio mancante: errore, non KeyError; importo scritto come testo: errore con la cella
    xlsx = _pianificatore_modificato(lambda wb: wb.remove(wb["Parametri"]))
    for nome, path in script.items():
        rc, out = run_script(path, xlsx)
        assert rc == 1 and "foglio «Parametri» mancante" in out and "KeyError" not in out, f"{nome}: {out[-300:]}"
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    # file esistente ma non xlsx (revisione T9a, codice-test 3): ERRORE, non BadZipFile
    finto = os.path.join(tempfile.mkdtemp(), "finto.xlsx")
    open(finto, "w", encoding="utf-8").write("non sono un foglio")
    for nome, path in script.items():
        rc, out = run_script(path, finto)
        assert rc == 1 and "ERRORE:" in out and "non è un pianificatore xlsx leggibile" in out and "Traceback" not in out, f"{nome}: {out[-300:]}"
    shutil.rmtree(os.path.dirname(finto), ignore_errors=True)
    # revisione T9: anche il testo che sembra un numero è un errore. Prima «1.5» diventava 15 anni,
    # «0.02» un'inflazione del 200% e «40.000» passava: ogni punto era preso per le migliaia.
    for cella, testo in (("Obiettivi!B6", "quarantamila"), ("Obiettivi!B6", "40.000"), ("Obiettivi!C6", "1.5"),
                         ("Parametri!B6", "0.02"), ("Parametri!B11", "0.025")):
        foglio, rif = cella.split("!")
        xlsx = _pianificatore_modificato(lambda wb: wb[foglio].__setitem__(rif, testo))
        for nome, path in script.items():
            rc, out = run_script(path, xlsx)
            assert rc == 1 and "ERRORE:" in out and cella in out and "testo" in out and "Traceback" not in out, \
                f"{nome}, {cella} = {testo!r}: rc {rc}, {out[-300:]}"
        shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    # (h) riga col nome ma senza orizzonte: scartata con il motivo; riga senza nome: scartata
    def incompleta(wb):
        o = wb["Obiettivi"]
        o["A8"], o["B8"], o["D8"] = "Monca", 1000, "Moderato"
        o["B9"], o["C9"], o["D9"] = 5000, 3, "Moderato"
    xlsx = _pianificatore_modificato(incompleta)
    avvisi = []
    d = motore.leggi_pianificatore(xlsx, avvisa=avvisi.append)
    assert [g["nome"] for g in d["obiettivi"]] == ["Fondo di emergenza", "Anticipo casa", "Lungo termine / libertà"]
    assert any("«Monca» scartato, manca orizzonte" in a for a in avvisi) and any("riga 9" in a and "senza nome" in a for a in avvisi), avvisi
    rc, out = run_script(script["montecarlo-obiettivi.py"], xlsx)
    assert rc == 0 and "«Monca» scartato" in out, out[-400:]
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)


@test("allocatore-capitale: invariante — la somma allocata è il capitale dato")
def t_allocatore():
    xlsx = _excel_test([
        ("Emergenza", 10000, 5, "Preservazione", 0, "essenziale", "permanente"),
        ("Casa", 40000, 12, "Moderato", 0, "importante", "terminante"),
    ])
    out_json = os.path.join(os.path.dirname(xlsx), "a.json")
    rc, out = run_script(asset("it-pianificazione-obiettivi", "allocatore-capitale.py"),
                         xlsx, "50000", "--json", out_json)
    assert rc == 0, out[-400:]
    d = json.load(open(out_json, encoding="utf-8"))
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    # F032: prima si parsava lo stdout con una regex; qui il JSON, al centesimo
    agg = d["aggregato"]
    assert agg["totale"] == 50000 and abs(agg["azioni"] + agg["obbligazioni"] + agg["liquidita"] - 50000) <= 0.01, agg
    assert abs(sum(o["allocato"] for o in d["obiettivi"]) + d["surplus_capitale_libero"] - d["capitale"]) <= 0.01, d
    assert d["capitale"] == 50000 and d["deficit"] == 0
    emergenza = next(o for o in d["obiettivi"] if o["nome"] == "Emergenza")
    assert emergenza["richiesto"] == 10000 and emergenza["allocato"] == 10000 and emergenza["pieno"], emergenza
    # l'aggregato dello stdout coincide col JSON
    tot = sum(float(x.replace(".", "").replace(",", "."))
              for x in re.findall(r"€\s*([\d.]+,?\d*)\s", out.split("AGGREGATA")[-1])[:3])
    assert abs(tot - 50000) < 3, f"somma aggregata stampata {tot} != 50000"


def _prob_obiettivi(out):
    """{nome: probabilità di successo col versamento deterministico} dall'output di montecarlo-obiettivi."""
    return {nome: int(re.search(r"Probabilità di successo con questo versamento: (\d+)%", testo).group(1))
            for nome, (_, _, testo) in _sezioni_obiettivi(out).items()}


@test("montecarlo-obiettivi: obiettivo ampiamente finanziato → probabilità alta, controllo a bassa volatilità ≈ 50%")
def t_mc_obiettivi():
    # F022: la fixture lasciava la riga demo 7 e il regex prendeva la quota azionaria dell'OPZIONE
    # (~100%) invece della probabilità: il mutante «già accantonato ignorato» passava.
    xlsx = _excel_test([("Sovra-finanziato", 10000, 10, "Moderato", 14000, "importante", "terminante"),
                        ("Controllo", 20000, 10, "Preservazione", 6000, "importante", "terminante")])
    rc, out = run_script(asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"), xlsx)
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    assert rc == 0, out[-400:]
    prob = _prob_obiettivi(out)
    assert set(prob) == {"Sovra-finanziato", "Controllo"}, f"obiettivi letti: {sorted(prob)} (riga demo residua?)"
    assert prob["Sovra-finanziato"] >= 95, f"14.000 accantonati su 12.190 nominali: probabilità {prob['Sovra-finanziato']}%"
    # versamento deterministico a volatilità 1,5%: ~50%. Se il già accantonato è ignorato nella
    # simulazione la probabilità crolla; se è ignorato nel versamento sale verso il 100%.
    assert 40 <= prob["Controllo"] <= 55, f"controllo Preservazione: probabilità {prob['Controllo']}%, attesa 40-55%"


def _sezioni_obiettivi(out):
    """{nome: (anni, priorità, testo della sezione)} dall'output di montecarlo-obiettivi."""
    sezioni = {}
    for m in re.finditer(r"▸ (.+?)  \[[^,]+, (\d+) anni, priorità: ([^\]]*)\]\n(.*?)(?=\n▸ |\n={10,})", out, re.S):
        sezioni[m.group(1)] = (int(m.group(2)), m.group(3).strip().lower(), m.group(4))
    return sezioni


@test("montecarlo-obiettivi: la leva di rischio è un'OPZIONE solo per i desiderabili oltre la finestra di de-risking")
def t_leva_rischio():
    # REGRESSIONE F017: sui default proponeva ~95% azionario per l'anticipo casa a 6 anni.
    rc, out = run_script(asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"))
    assert rc == 0, out[-400:]
    sezioni = _sezioni_obiettivi(out)
    assert "Anticipo casa" in sezioni, f"obiettivo d'esempio non trovato: {list(sezioni)}"
    assert "INADATTA" in sezioni["Anticipo casa"][2], "anticipo casa a 6 anni: la leva deve essere INADATTA"
    for nome, (anni, prio, testo) in sezioni.items():
        if anni <= 7 or not prio.startswith("desider"):
            assert "OPZIONE:" not in testo, f"{nome} ({anni} anni, {prio}): OPZIONE fuori dal gate"
    # il gate non spegne la leva: un desiderabile lungo e sotto-finanziato la riceve ancora
    xlsx = _excel_test([("Barca", 400000, 20, "Moderato", 0, "desiderabile", "terminante")])
    rc, out = run_script(asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"), xlsx)
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    assert rc == 0, out[-400:]
    assert "OPZIONE:" in _sezioni_obiettivi(out)["Barca"][2], "desiderabile a 20 anni: OPZIONE attesa"


@test("decumulo-obiettivi: il bootstrap a blocchi campiona ogni anno della serie, ultimo compreso")
def t_decumulo_indici():
    # REGRESSIONE F029: start = integers(0, n - BLOCCO) escludeva l'ultimo anno (2025).
    import numpy as np
    cartella = os.path.join(SKILLS, "it-pianificazione-obiettivi", "assets")
    sys.path.insert(0, cartella)
    try:
        mod = carica_modulo(os.path.join(cartella, "decumulo-obiettivi.py"))
    finally:
        sys.path.remove(cartella)
    n = 98
    rng = np.random.default_rng(0)
    visti = np.zeros(n, dtype=int)
    for _ in range(3000):
        for i in mod.percorso_bootstrap(np.arange(n, dtype=float), rng):
            visti[int(i)] += 1
    assert visti.min() > 0, f"anni mai campionati: {np.flatnonzero(visti == 0).tolist()}"


@test("smoke: decumulo, asset-location, ribilanciamento, importa-estratto-broker girano")
def t_smoke_altri():
    rc, out = run_script(asset("it-pianificazione-obiettivi", "decumulo-obiettivi.py"))
    assert rc == 0 and "%" in out, "decumulo: " + out[-200:]
    rc, out = run_script(asset("it-pianificazione-obiettivi", "asset-location-optimizer.py"))
    assert rc == 0, "asset-location: " + out[-200:]
    tmp = tempfile.mkdtemp()  # il ribilanciamento scrive ordini/ nella cwd: non deve sporcare gli assets
    rc, out = run_script(asset("it-rebalance", "ribilanciamento-multibroker.py"), cwd=tmp)
    shutil.rmtree(tmp, ignore_errors=True)
    assert rc == 0, "ribilanciamento: " + out[-200:]
    rc, out = run_script(asset("it-tax-optimization", "importa-estratto-broker.py"),
                         os.path.join(FIXTURES, "estratto-broker-test.csv"), "--json", "-")
    assert rc == 0 and json.loads(out)["analisi"]["plus_latenti_compensabili"] == 300, "importatore broker: " + out[-200:]


@test("postura: ogni comando chiede in fondo il disclaimer «non è consulenza»")
def t_disclaimer_comandi():
    cartella = os.path.join(ROOT, "commands")
    for nome in sorted(os.listdir(cartella)):
        testo = open(os.path.join(cartella, nome), encoding="utf-8").read()
        assert re.search(r"[Dd]isclaimer obbligatorio[^:\n]*:[^\n]*consulenza", testo), \
            f"commands/{nome}: manca «Disclaimer obbligatorio …: … consulenza …»"


@test("ribilanciamento: CSV degli ordini e stdout portano il disclaimer, intestazione invariata")
def t_disclaimer_csv():
    import csv
    import glob
    tmp = tempfile.mkdtemp()
    try:
        rc, out = run_script(asset("it-rebalance", "ribilanciamento-multibroker.py"), cwd=tmp)
        assert rc == 0, out[-300:]
        assert "non costituisce consulenza" in out, "stdout senza disclaimer"
        files = glob.glob(os.path.join(tmp, "ordini", "ordini_*.csv"))
        assert files, "nessun CSV degli ordini"
        for path in files:
            with open(path, encoding="utf-8-sig", newline="") as f:
                righe = list(csv.DictReader(f, delimiter=";"))
            assert righe and righe[-1]["operazione"] == "NOTA", f"{os.path.basename(path)}: manca la riga NOTA finale"
            assert "non costituisce consulenza" in righe[-1]["nota"], f"{os.path.basename(path)}: nota senza disclaimer"
            assert all(r["operazione"] in ("VENDITA", "ACQUISTO", "PRELIEVO") for r in righe[:-1]), "righe d'ordine alterate"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ══════════════════════ ribilanciamento tax-aware (T7) ══════════════════════

def _ribilancia(**parametri):
    """Importa ribilanciamento-multibroker.py, imposta i parametri e ritorna (modulo, piano)."""
    m = carica_modulo(asset("it-rebalance", "ribilanciamento-multibroker.py"))
    for k, v in parametri.items():
        setattr(m, k, v)
    return m, m.pianifica()


def _stampato(m, piano):
    import contextlib
    import io
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        m.stampa(piano)
    return buf.getvalue()


def _cassa_per_broker(m, piano):
    """Per ogni broker: acquisti + prelievi ≤ vendite nette d'imposta + versamenti che vi atterrano."""
    ordini = m.genera_ordini(piano)
    entrate = {}
    for o in ordini:
        segno = {"VENDITA": 1, "ACQUISTO": -1, "PRELIEVO": -1}[o["operazione"]]
        netto = o["importo_eur"] - (o["imposta_eur"] or 0) if o["operazione"] == "VENDITA" else o["importo_eur"]
        entrate[o["broker"]] = entrate.get(o["broker"], 0) + segno * netto
    entrate[m.BROKER_VERSAMENTI] = entrate.get(m.BROKER_VERSAMENTI, 0) + m.NUOVI_VERSAMENTI + m.FLUSSI_ATTESI
    return ordini, entrate


@test("T9 input esterno: ribilanciamento con --dati, --zainetto dell'importatore, --out, sinonimi del vocabolario")
def t_dati_ribilanciamento():
    # F033: portafoglio, zainetto e target erano cablati; F026: tipo_fiscale/categoria_fiscale e
    # reddito_diverso/diverso non erano accettati dallo script (tipo fuori vocabolario → piano degradato in silenzio).
    import csv
    script = asset("it-rebalance", "ribilanciamento-multibroker.py")
    m = carica_modulo(script)
    tmp = tempfile.mkdtemp()
    try:
        piano = os.path.join(tmp, "piano.json")
        # lo stesso caso di default, scritto col vocabolario del plugin (categoria_fiscale, plus_minus, classe)
        pos = [{"broker": "Fineco", "classe": "azionario", "categoria_fiscale": "reddito_capitale", "controvalore": 40000, "plus_minus": 10000},
               {"broker": "Fineco", "classe": "azionario", "categoria_fiscale": "reddito_diverso", "controvalore": 10000, "plus_minus": 4000},
               {"broker": "Fineco", "classe": "obbligazionario", "tipo_fiscale": "capitale", "valore": 15000, "plus": -500},
               {"broker": "Directa", "classe": "azionario", "tipo": "reddito_capitale", "valore": 18000, "plus": 3000},
               {"broker": "Directa", "classe": "obbligazionario", "tipo": "diverso", "valore": 9000, "plus": 0},
               {"broker": "Directa", "classe": "monetario", "tipo": "capitale", "valore": 8000, "plus": 0}]
        json.dump({"posizioni": pos, "zainetto": {"Fineco": 5000, "Directa": 0},
                   "target": {"azionario": 0.6, "obbligazionario": 0.3, "monetario": 0.1}}, open(piano, "w", encoding="utf-8"))
        rc, out = run_script(script, "--dati", piano, "--out", os.path.join(tmp, "csv"), cwd=tmp)
        assert rc == 0 and "Dati di esempio" not in out, out[-400:]
        assert "VENDI €8.000 azionario/diverso @ Fineco" in out and "COSTO FISCALE TOTALE: €0" in out, out[-900:]
        assert os.path.exists(os.path.join(tmp, "csv", "ordini_Fineco.csv")) and not os.path.exists(os.path.join(tmp, "ordini")), os.listdir(tmp)
        with open(os.path.join(tmp, "csv", "ordini_Fineco.csv"), encoding="utf-8-sig", newline="") as f:
            righe = list(csv.DictReader(f, delimiter=";"))
        # rilettura degli output (§ 2.6): nel CSV il valore canonico del vocabolario, non la forma breve interna
        assert righe[0]["categoria_fiscale"] == "reddito_diverso" and "tipo_fiscale" not in righe[0], righe[0]
        assert "esempio" not in righe[-1]["nota"].lower(), righe[-1]
        assert "un CSV per broker in " + os.path.join(tmp, "csv") + os.sep in out and "./" + tmp not in out, \
            "intestazione della sezione 5 con il percorso di --out così com'è"
        # senza --dati: dati di esempio dichiarati, anche nella riga NOTA del CSV
        rc, out = run_script(script, cwd=tmp)
        assert rc == 0 and "Dati di esempio" in out, out[:300]
        with open(os.path.join(tmp, "ordini", "ordini_Fineco.csv"), encoding="utf-8-sig", newline="") as f:
            nota = list(csv.DictReader(f, delimiter=";"))[-1]["nota"]
        assert nota.startswith("DATI DI ESEMPIO") and "non costituisce consulenza" in nota, nota
        # chiave ignota, posizione fuori perimetro, tipo sconosciuto, target sbagliato → rc 1 con il motivo
        for dati, atteso in (({"posizioni": pos, "zainetto": {}, "target": {"azionario": 1.0}, "prelievi": 1}, "non riconosciute"),
                             ({"posizioni": [dict(pos[0], categoria_fiscale="da_classificare")], "zainetto": {}, "target": {"azionario": 1.0}}, "classifica prima"),
                             ({"posizioni": [dict(pos[0], categoria_fiscale="reddito_capitale_irpef")], "zainetto": {}, "target": {"azionario": 1.0}}, "IRPEF"),
                             ({"posizioni": [dict(pos[0], categoria_fiscale="boh")], "zainetto": {}, "target": {"azionario": 1.0}}, "reddito_diverso o reddito_capitale"),
                             ({"posizioni": pos, "zainetto": {}, "target": {"azionario": "60%"}}, "numeri")):
            json.dump(dati, open(piano, "w", encoding="utf-8"))
            rc, out = run_script(script, "--dati", piano, cwd=tmp)
            assert rc == 1 and "ERRORE" in out and atteso in out and "Traceback" not in out, (atteso, rc, out[-300:])
        # --zainetto dal JSON dell'importatore: zainetto_26_equivalente per broker
        rc, out = run_script(asset("it-tax-optimization", "importa-estratto-broker.py"),
                             os.path.join(FIXTURES, "estratto-multibroker.csv"), os.path.join(FIXTURES, "zainetto-multibroker.csv"),
                             "--anno", "2026", "--json", os.path.join(tmp, "imp.json"))
        assert rc == 0, out[-300:]
        imp = json.load(open(os.path.join(tmp, "imp.json"), encoding="utf-8"))
        atteso = {b: v["zainetto_26_equivalente"] for b, v in imp["analisi"]["per_broker"].items()}
        json.dump({"posizioni": [dict(p, broker=list(atteso)[0]) for p in pos], "zainetto": {"X": 1},
                   "target": {"azionario": 0.6, "obbligazionario": 0.3, "monetario": 0.1}}, open(piano, "w", encoding="utf-8"))
        rc, out = run_script(script, "--dati", piano, "--zainetto", os.path.join(tmp, "imp.json"), "--out", os.path.join(tmp, "z"), cwd=tmp)
        assert rc == 0, out[-400:]
        m2 = carica_modulo(script)
        m2.carica_dati(piano)
        assert m2.carica_zainetto_importatore(os.path.join(tmp, "imp.json")) == atteso and m2.ZAINETTO == atteso
        rc, out = run_script(script, "--dati", piano, "--zainetto", piano, cwd=tmp)
        assert rc == 1 and "non è il JSON di importa-estratto-broker" in out, out[-300:]
        # revisione T9: dal JSON dell'importatore al piano copiando le posizioni così come sono (più la
        # classe). L'importatore scrive l'aliquota in percento (12.5): prima il BTP veniva respinto, o
        # tassato al 26% togliendo il campo; ora vale il 12,5% e consuma lo zainetto al 48,08%.
        open(os.path.join(tmp, "e.csv"), "w", encoding="utf-8").write(
            "broker,isin,descrizione,quantita,pmc,prezzo_corrente\n"
            "Fineco,IT0005436693,BTP 0.6% 2031,10000,0.80,0.95\n"
            "Fineco,IE00B4L5Y983,ISHARES CORE MSCI WORLD UCITS ETF,100,50,50\n")
        open(os.path.join(tmp, "z.csv"), "w", encoding="utf-8").write("broker,anno,importo,aliquota\nFineco,2024,1000,26\n")
        rc, out = run_script(asset("it-tax-optimization", "importa-estratto-broker.py"), os.path.join(tmp, "e.csv"),
                             os.path.join(tmp, "z.csv"), "--anno", "2026", "--json", os.path.join(tmp, "imp2.json"))
        assert rc == 0, out[-300:]
        imp2 = json.load(open(os.path.join(tmp, "imp2.json"), encoding="utf-8"))
        assert sorted(p["aliquota"] for p in imp2["posizioni"]) == [12.5, 26.0], imp2["posizioni"]
        classi = {"IT0005436693": "obbligazionario", "IE00B4L5Y983": "azionario"}
        json.dump({"posizioni": [dict(p, classe=classi[p["isin"]]) for p in imp2["posizioni"]], "zainetto": {},
                   "target": {"azionario": 0.6, "obbligazionario": 0.4}}, open(piano, "w", encoding="utf-8"))
        rc, out = run_script(script, "--dati", piano, "--zainetto", os.path.join(tmp, "imp2.json"), "--out", os.path.join(tmp, "z2"), cwd=tmp)
        assert rc == 0 and "in percento" in out, out[-600:]
        vendita = next(r for r in out.splitlines() if "VENDI" in r and "obbligazionario" in r)
        assert "imposta €0" in vendita and "48,08%" in out, out[-900:]
        # zainetto con il nome del broker scritto diversamente: allineato con un avviso, non perso;
        # zainetto di un broker senza posizioni: avviso «non usato» (con il suggerimento di --broker)
        json.dump({"posizioni": pos, "zainetto": {"FINECO": 5000, "(broker non indicato)": 700, "Directa": 0},
                   "target": {"azionario": 0.6, "obbligazionario": 0.3, "monetario": 0.1}}, open(piano, "w", encoding="utf-8"))
        rc, out = run_script(script, "--dati", piano, "--out", os.path.join(tmp, "z3"), cwd=tmp)
        assert rc == 0 and "«FINECO» → «Fineco»" in out and "compensata €3.200 dallo zainetto Fineco" in out, out[-900:]
        assert "(broker non indicato)" in out and "non usato" in out and "--broker" in out, out[-900:]
        # revisione T14: un --dati senza zainetto o senza target usava in silenzio quelli d'esempio
        # (e il CSV si dichiarava dato reale); lo zainetto può arrivare anche da --zainetto
        base = {"posizioni": pos, "zainetto": {"Fineco": 5000}, "target": {"azionario": 0.6, "obbligazionario": 0.3, "monetario": 0.1}}
        for manca in ("zainetto", "target", "posizioni"):
            json.dump({k: v for k, v in base.items() if k != manca}, open(piano, "w", encoding="utf-8"))
            rc, out = run_script(script, "--dati", piano, cwd=tmp)
            assert rc == 1 and "ERRORE" in out and f"«{manca}»" in out and "Traceback" not in out, (manca, rc, out[-300:])
        json.dump({k: v for k, v in base.items() if k != "zainetto"}, open(piano, "w", encoding="utf-8"))
        rc, out = run_script(script, "--dati", piano, "--zainetto", os.path.join(tmp, "imp.json"), "--out", os.path.join(tmp, "z4"), cwd=tmp)
        assert rc == 0, out[-300:]
        # numeri scritti come testo nel JSON: ERRORE con rc 1, mai un Traceback
        for chiave, valore in (("target", {"azionario": "0.6", "obbligazionario": 0.3, "monetario": 0.1}),
                               ("nuovi_versamenti", "1000"), ("zainetto", {"Fineco": "5.000"})):
            json.dump({"posizioni": pos, "zainetto": {"Fineco": 5000}, "target": {"azionario": 0.6, "obbligazionario": 0.3,
                       "monetario": 0.1}, chiave: valore}, open(piano, "w", encoding="utf-8"))
            rc, out = run_script(script, "--dati", piano, cwd=tmp)
            assert rc == 1 and "ERRORE" in out and "Traceback" not in out, (chiave, rc, out[-300:])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    # i sinonimi normalizzati producono lo stesso piano dei dati canonici
    canon = [m.normalizza_posizione(p) for p in pos]
    assert canon[0] == {"broker": "Fineco", "asset_class": "azionario", "tipo": "capitale", "valore": 40000, "plus": 10000}, canon[0]
    assert [p["tipo"] for p in canon] == ["capitale", "diverso", "capitale", "capitale", "diverso", "capitale"]
    m.POSITIONS = pos
    piano_sin = m.pianifica()
    m.POSITIONS = carica_modulo(script).POSITIONS
    piano_can = m.pianifica()
    assert [(v["broker"], v["classe"], round(v["vendi"])) for v in piano_sin["vendite"]] == \
        [(v["broker"], v["classe"], round(v["vendi"])) for v in piano_can["vendite"]] == [("Fineco", "azionario", 8000)]


@test("T9 input esterno: genera-ips con --dati dalla cartella di lavoro, esempio dichiarato senza, JSON errato → rc 1")
def t_dati_ips():
    try:
        from docx import Document
    except ImportError:
        raise AssertionError("python-docx non installato: pip install python-docx")
    script = asset("it-revisione", "genera-ips.py")
    tmp = tempfile.mkdtemp()
    try:
        dati = os.path.join(tmp, "ips.json")
        json.dump({"investitore": "Giulia Test", "capitale": 50000, "sottotitolo": "Prova",
                   "profilo": [["Età", "40 anni"]],
                   "obiettivi": [["Fondo emergenza", "permanente", "Preservazione (0%)", "€10.000", 10000],
                                 ["Casa", "8 anni · terminante", "Moderato (40%)", "€40.000", 40000]],
                   "allocazione": [["Azioni", 16000, 32.0], ["Obbligazioni", 24000, 48.0], ["Liquidità", 10000, 20.0]],
                   "glide_path": [["Partenza", "oggi", "32%", "allocazione iniziale"]],
                   "versamenti": ["€200/mese sulla casa"], "strumenti": [["Azioni", "ETF globale"]],
                   "previdenza": ["Aggiornare la proiezione ogni anno."]}, open(dati, "w", encoding="utf-8"), ensure_ascii=False)
        out_docx = os.path.join(tmp, "ips.docx")
        rc, out = run_script(script, out_docx, "--dati", dati, cwd=tmp)
        assert rc == 0 and "Dati di esempio" not in out and "Giulia Test" in out, out[-400:]
        testo = "\n".join(p.text for p in Document(out_docx).paragraphs)
        assert "Giulia Test" in testo and "Mario Rossi" not in testo and "€50.000" in testo, testo[:600]
        m = re.search(r"perdere circa il (\d+)%", testo)
        assert m and int(m.group(1)) == 16, "nota di rischio dalla quota azionaria del JSON (32% → 16%)"
        assert "DATI DI ESEMPIO" not in testo, "IPS con i dati del JSON marcato come esempio"
        # revisione T9: con --dati nulla viene dall'esempio di Mario Rossi (prima: note, profilo e
        # glide path dell'esempio riempivano in silenzio le chiavi mancanti)
        for estraneo in ("fondo emergenza è tenuto", "35 anni", "€32.000", "lungo termine\u201d) guida"):
            assert estraneo not in testo, f"IPS con --dati contiene un pezzo dell'esempio: «{estraneo}»"
        completo = json.load(open(dati, encoding="utf-8"))
        for parziale, attesi in (({"investitore": "X", "capitale": 50000}, ("chiavi mancanti", "profilo", "glide_path")),
                                 (dict(completo, capitale="50000"), ("capitale",)),
                                 (dict(completo, allocazione=[["Azioni", "16000", 32.0]]), ("allocazione",))):
            json.dump(parziale, open(dati, "w", encoding="utf-8"))
            rc, out = run_script(script, os.path.join(tmp, "x.docx"), "--dati", dati, cwd=tmp)
            assert rc == 1 and "ERRORE:" in out and all(a in out for a in attesi) and "Traceback" not in out, (parziale, rc, out[-400:])
        json.dump(completo, open(dati, "w", encoding="utf-8"))
        rc, out = run_script(script, os.path.join(tmp, "demo.docx"), cwd=tmp)
        assert rc == 0 and "Dati di esempio" in out and "Mario Rossi" in out, out[-300:]
        # rilettura degli output (§ 2.6): il documento può circolare da solo, lo dice anche lui
        demo = "\n".join(p.text for p in Document(os.path.join(tmp, "demo.docx")).paragraphs)
        assert "DATI DI ESEMPIO" in demo.split("INVESTMENT POLICY STATEMENT")[1][:600], "copertina senza l'avviso dei dati di esempio"
        assert not os.path.exists(os.path.join(tmp, "--dati")), "«--dati» non è un nome di file di output"
        json.dump({"capitale": 1}, open(dati, "w", encoding="utf-8"))
        rc, out = run_script(script, "--dati", dati, cwd=tmp)
        assert rc == 1 and "investitore" in out and "Traceback" not in out, (rc, out[-300:])
        json.dump({"investitore": "X", "capitali": 1}, open(dati, "w", encoding="utf-8"))
        rc, out = run_script(script, "--dati", dati, cwd=tmp)
        assert rc == 1 and "non riconosciute" in out and "capitali" in out, (rc, out[-300:])
        open(dati, "w", encoding="utf-8").write("{non json")
        rc, out = run_script(script, "--dati", dati, cwd=tmp)
        assert rc == 1 and "JSON valido" in out, (rc, out[-300:])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T9 «Dove lavorare»: ogni SKILL.md ha la sezione con ${CLAUDE_SKILL_DIR}, nessuna invocazione a nome nudo, command con ${CLAUDE_PLUGIN_ROOT}")
def t_dove_lavorare():
    # F059/F058: la cartella di lavoro era solo prosa nel README e gli script erano invocati a nome nudo,
    # cioè solo con cwd = assets/. Variabili verificate sulla documentazione ufficiale di Claude Code
    # (code.claude.com/docs/en/skills, «Available string substitutions»; plugins-reference, «Environment variables»).
    # gli script del repository (scripts/…) si lanciano dal checkout e non sono nello zip: restano fuori dalla regola
    nudo = re.compile(r"python3? (?!\")(?!scripts/)[\w./-]*[\w-]+\.py\b")
    for skill in sorted(os.listdir(SKILLS)):
        testo = open(os.path.join(SKILLS, skill, "SKILL.md"), encoding="utf-8").read()
        assert "## Dove lavorare" in testo, f"{skill}: manca la sezione «Dove lavorare»"
        sez = testo.split("## Dove lavorare")[1].split("\n## ")[0]
        assert "${CLAUDE_SKILL_DIR}" in sez and "cartella di lavoro" in sez and "${user_config.cartella_lavoro}" in sez, skill
        assert "mai dentro il plugin" in sez or "mai in `assets/`" in sez, skill
        if os.path.isdir(os.path.join(SKILLS, skill, "assets")) and any(f.endswith(".xlsx") for f in os.listdir(os.path.join(SKILLS, skill, "assets"))):
            assert "copiano" in sez, f"{skill}: i modelli Excel vanno copiati nella cartella di lavoro"
        m = nudo.search(testo)
        assert not m, f"{skill}/SKILL.md: invocazione a nome nudo «{m.group(0)}»: usare python \"${{CLAUDE_SKILL_DIR}}/assets/…\""
        # T14: i comandi sono alias della skill, quindi niente percorso «per i command» (revisione T14)
        assert "nei command vale" not in sez, f"{skill}: i comandi non lanciano script"
        assert not re.search(r"python3? scripts/", testo), f"{skill}: scripts/ si lancia con ${{CLAUDE_PLUGIN_ROOT}}"
    for nome in sorted(os.listdir(os.path.join(ROOT, "commands"))):
        testo = open(os.path.join(ROOT, "commands", nome), encoding="utf-8").read()
        m = nudo.search(testo)
        assert not m, f"commands/{nome}: invocazione a nome nudo «{m.group(0)}»"
        assert ".py" not in testo, f"commands/{nome}: un alias non lancia script, lo fa la skill"
    manifest = json.load(open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8"))
    cfg = manifest["userConfig"]["cartella_lavoro"]
    assert cfg["type"] == "directory" and cfg.get("default") and set(cfg) <= {"type", "title", "description", "default", "required"}, cfg
    # contestato in revisione T9: la riga delle dipendenze era la stessa in ogni skill. Ora elenca i
    # pacchetti esterni importati in testa agli script della skill (e del motore, se lo importano).
    pacchetti = {"numpy": "numpy", "openpyxl": "openpyxl", "docx": "python-docx"}
    for skill in sorted(os.listdir(SKILLS)):
        assets_dir = os.path.join(SKILLS, skill, "assets")
        sez = open(os.path.join(SKILLS, skill, "SKILL.md"), encoding="utf-8").read().split("## Dove lavorare")[1].split("\n## ")[0]
        script = [f for f in os.listdir(assets_dir) if f.endswith(".py")] if os.path.isdir(assets_dir) else []
        if not script:
            assert "non ha script" in sez and "pip install" not in sez, f"{skill}: senza script, niente dipendenze"
            continue
        usati = set()
        for f in script:
            testo_f = open(os.path.join(assets_dir, f), encoding="utf-8").read()
            for m in re.finditer(r"^(?:import|from) (\w+)", testo_f, re.M):
                if m.group(1) in pacchetti:
                    usati.add(pacchetti[m.group(1)])
            if f != "motore_bootstrap.py" and "leggi_pianificatore" in testo_f:
                usati.add("openpyxl")   # il lettore del motore importa openpyxl quando serve
        riga_dip = re.search(r"Dipendenz[ae][^\n]*?(?:\)|standard)\.", sez).group(0)
        for nome in pacchetti.values():
            assert (nome in riga_dip) == (nome in usati), f"{skill}: dipendenze dichiarate «{riga_dip}», usate {sorted(usati)}"
        if not usati:
            assert "solo libreria standard" in riga_dip, f"{skill}: {riga_dip}"
    # il rimando al ricalcolo da riga di comando dice che scripts/ non è nello zip di Cowork
    rev = open(os.path.join(SKILLS, "it-revisione", "SKILL.md"), encoding="utf-8").read()
    assert re.search(r"ricalcola-excel\.py[^\n]*", rev) and "non lo zip per Cowork" in rev, "it-revisione: ricalcolo condizionale"


@test("T9 vocabolario: categoria fiscale con gli stessi nomi e valori in README, skill, importatore e ribilanciamento")
def t_lessico():
    # F026: lo stesso campo si chiamava categoria_fiscale / tipo_fiscale / tipo con valori reddito_diverso / diverso.
    imp = carica_modulo(asset("it-tax-optimization", "importa-estratto-broker.py"))
    valori = set(imp.CATEGORIE)
    assert valori == {"reddito_diverso", "reddito_capitale", "reddito_capitale_irpef", "da_classificare"}, valori
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    gloss = readme.split("## Glossario")[1].split("\n## ")[0]
    for v in valori:
        assert f"`{v}`" in gloss, f"glossario del README senza «{v}»"
    for termine in ("Zainetto fiscale", "Categoria fiscale", "PMC", "Secchiello", "Glide path", "Profilo", "Asset class", "Tasso di sostituzione"):
        assert f"**{termine}**" in gloss, f"glossario senza «{termine}»"
    assert gloss.count("\n| **") >= 8, "glossario con meno di 8 termini"
    assert "broker,isin,descrizione" in readme and "broker,anno_realizzo" in readme, "README «Input richiesto»: schema senza broker"
    tax = open(os.path.join(SKILLS, "it-tax-optimization", "SKILL.md"), encoding="utf-8").read()
    reb = open(os.path.join(SKILLS, "it-rebalance", "SKILL.md"), encoding="utf-8").read()
    for v in valori:
        assert f"`{v}`" in tax, f"it-tax-optimization/SKILL.md senza «{v}»"
    assert "`categoria_fiscale`" in reb and not re.search(r"^- `tipo_fiscale`", reb, re.M), "it-rebalance: il campo si chiama categoria_fiscale"
    assert "`categoria_fiscale`" in tax and "`broker`" in tax
    m = carica_modulo(asset("it-rebalance", "ribilanciamento-multibroker.py"))
    assert "categoria_fiscale" in m.CAMPI_CSV and "tipo_fiscale" not in m.CAMPI_CSV
    for sinonimo, canonico in (("reddito_diverso", "diverso"), ("reddito_capitale", "capitale"), ("Reddito di capitale", "capitale")):
        assert m.normalizza_posizione({"tipo_fiscale": sinonimo})["tipo"] == canonico
    assert set(m.SINONIMI_CHIAVI) >= {"categoria_fiscale", "tipo_fiscale", "classe", "plus_minus", "controvalore"}
    # revisione T9: i sinonimi del PMC nel glossario sono quelli del prezzo unitario dell'importatore,
    # mai quelli del costo totale della posizione (prima `costo`, che l'importatore legge come totale)
    riga_pmc = next(r for r in gloss.splitlines() if r.startswith("| **PMC**"))
    sinonimi_pmc = [s.strip("`«» ") for s in riga_pmc.rstrip(" |").split("|")[-1].split(",")]
    for s in sinonimi_pmc:
        assert s.lower() in imp.SINONIMI["prezzo_carico"] and s.lower() not in imp.SINONIMI["valore_carico"], \
            f"glossario: «{s}» non è un sinonimo del prezzo medio di carico nell'importatore"


@test("T7 ribilanciamento: caso di validazione con asserzioni (vendita €8.000 a imposta zero, CSV, cassa per broker)")
def t_t7_default():
    # REGRESSIONE F047: lo smoke test controllava solo rc == 0, e il caso ATTESO era scaduto
    # (monetario «FUORI banda» con la modulazione per volatilità accesa di default).
    import csv
    m, p = _ribilancia()
    assert m.MODULA_VOLATILITA is False, "modulazione per volatilità spenta di default (decisione 5)"
    assert [(v["broker"], v["classe"], v["tipo"], round(v["vendi"])) for v in p["vendite"]] == \
        [("Fineco", "azionario", "diverso", 8000)], p["vendite"]
    assert p["imposta"] == 0 and round(p["zainetto"]["Fineco"]) == 1800, (p["imposta"], p["zainetto"])
    assert not p["drift"]["monetario"]["fuori_banda"], p["drift"]["monetario"]
    assert {c: round(w, 6) for c, w in p["risultante"].items()} == {"azionario": 0.6, "obbligazionario": 0.3, "monetario": 0.1}
    ordini, saldo = _cassa_per_broker(m, p)
    assert all(v >= -0.01 for v in saldo.values()), saldo
    tmp = tempfile.mkdtemp()
    try:
        scritti = m.scrivi_csv_per_broker(ordini, tmp)
        assert [(os.path.basename(f), n) for f, n in scritti] == [("ordini_Fineco.csv", 3)], scritti
        with open(scritti[0][0], encoding="utf-8-sig", newline="") as f:
            assert next(csv.reader(f, delimiter=";")) == m.CAMPI_CSV
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    testo = _stampato(m, p)
    assert "COSTO FISCALE TOTALE: €0" in testo and "nessuna vendita necessaria" not in testo, testo[-600:]


@test("T7 ribilanciamento: versamenti nel totale, minus che ricaricano lo zainetto, aliquota 12,5% per posizione")
def t_t7_flussi_minus_aliquota():
    # REGRESSIONE F004: con €20.000 in entrata lo script vendeva comunque €8.000 di azionario e
    # lasciava la cassa fuori piano; una minus realizzata non ricaricava lo zainetto del broker.
    m, p = _ribilancia(NUOVI_VERSAMENTI=20000)
    assert p["vendite"] == [] and round(sum(p["acquisti"].values())) == 20000, (p["vendite"], p["acquisti"])
    assert {c: round(w, 6) for c, w in p["risultante"].items()} == {"azionario": 0.6, "obbligazionario": 0.3, "monetario": 0.1}
    pos = [{"broker": "X", "asset_class": "azionario", "tipo": "capitale", "valore": 20000, "plus": -6000},
           {"broker": "X", "asset_class": "azionario", "tipo": "diverso", "valore": 10000, "plus": 4000},
           {"broker": "X", "asset_class": "obbligazionario", "tipo": "capitale", "valore": 10000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={"azionario": 0.125, "obbligazionario": 0.875})
    assert round(sum(v["vendi"] for v in p["vendite"])) == 25000 and p["imposta"] == 0, p["vendite"]
    assert round(p["zainetto"]["X"]) == 4000, p["zainetto"]          # 6.000 di minus − 2.000 compensati
    # REGRESSIONE F044: BTP al 12,5%, non al 26%; il 12,5% consuma lo zainetto al 48,08%
    btp = [{"broker": "X", "asset_class": "obbligazionario", "tipo": "diverso", "valore": 13000, "plus": 3000,
            "aliquota": 0.125},
           {"broker": "X", "asset_class": "azionario", "tipo": "capitale", "valore": 7000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=btp, ZAINETTO={"X": 0}, TARGET={"azionario": 1.0, "obbligazionario": 0.0})
    assert abs(p["imposta"] - 375) < 0.01, p["imposta"]
    m, p = _ribilancia(POSITIONS=btp, ZAINETTO={"X": 5000}, TARGET={"azionario": 1.0, "obbligazionario": 0.0})
    assert p["imposta"] == 0 and abs(5000 - p["zainetto"]["X"] - 3000 * 12.5 / 26) < 0.01, p["zainetto"]
    # zainetto di €1.000 (al 26%): copre €2.080 di plus al 12,5%; restano €920 tassati al 12,5%
    m, p = _ribilancia(POSITIONS=btp, ZAINETTO={"X": 1000}, TARGET={"azionario": 1.0, "obbligazionario": 0.0})
    assert abs(p["imposta"] - (3000 - 1000 * 26 / 12.5) * 0.125) < 0.01 and abs(p["zainetto"]["X"]) < 0.01, p


@test("T7 ribilanciamento: quota a costo zero venduta, residuo rinviato stampato, regola di rientro mista")
def t_t7_gate_rientro():
    # REGRESSIONE F019: con zainetto €2.000 il gate saltava l'intera posizione e stampava
    # «nessuna vendita necessaria»; F045: «a target» contro «in banda» senza regola unica.
    m, p = _ribilancia(ZAINETTO={"Fineco": 2000, "Directa": 0})
    assert [(v["broker"], v["tipo"], round(v["vendi"])) for v in p["vendite"]] == [("Fineco", "diverso", 5000)], p["vendite"]
    assert p["imposta"] == 0 and round(p["drift"]["azionario"]["residuo_non_corretto"]) == 3000, p["drift"]["azionario"]
    assert round(p["acquisti"]["obbligazionario"]) == 5000 and round(p["acquisti"]["monetario"]) == 0, p["acquisti"]
    testo = _stampato(m, p)
    assert "nessuna vendita necessaria" not in testo and re.search(r"rinviat[^\n]*3\.000|3\.000[^\n]*rinviat", testo), testo[-800:]
    # drift +18% oltre la banda larga, nessuna vendita gratis: si vende con imposta fino a metà banda
    target = {"azionario": 0.50, "obbligazionario": 0.40, "monetario": 0.10}
    m, p = _ribilancia(ZAINETTO={"Fineco": 0, "Directa": 0}, TARGET=target)
    assert m.OBIETTIVO_RIENTRO == "misto" and m.MARGINE_RIENTRO == 0.5
    assert round(sum(v["vendi"] for v in p["vendite"])) == 15500, p["vendite"]       # 68% → 52,5%
    assert abs(p["imposta"] - 15500 * 3000 / 18000 * 0.26) < 0.01, p["imposta"]     # prima l'ETF Directa (costo minore)
    az = p["drift"]["azionario"]
    assert abs(az["costo_a_target"] - 780) < 0.01 and abs(az["costo_in_banda"] - p["imposta"]) < 0.01, az
    testo = _stampato(m, p)
    assert "780" in testo and "672" in testo, testo[-800:]
    m, p = _ribilancia(ZAINETTO={"Fineco": 0, "Directa": 0}, TARGET=target, OBIETTIVO_RIENTRO="target")
    assert round(sum(v["vendi"] for v in p["vendite"])) == 18000 and abs(p["imposta"] - 780) < 0.01, p["vendite"]


@test("T7 ribilanciamento: prelievo in decumulo (vendite = prelievo + imposta) e validazione dell'input")
def t_t7_prelievo_input():
    # REGRESSIONE F046: sequenza-liquidazione prometteva il calcolo del prelievo, che lo script non aveva.
    m, p = _ribilancia(PRELIEVO=10000)
    venduto = sum(v["vendi"] for v in p["vendite"])
    assert venduto >= 10000 + p["imposta"] - 0.01, (venduto, p["imposta"])
    assert abs(venduto - p["imposta"] - sum(p["acquisti"].values()) - 10000) < 0.01, p
    assert p["prelievo"] == 10000
    ordini, saldo = _cassa_per_broker(m, p)
    assert all(v >= -0.01 for v in saldo.values()) and \
        round(sum(o["importo_eur"] for o in ordini if o["operazione"] == "PRELIEVO")) == 10000, ordini
    testo = _stampato(m, p)
    assert re.search(r"netto prelevato[^\n]*10\.000", testo), testo[-600:]
    # portafoglio già in banda: il prelievo da solo obbliga a vendere, imposta compresa
    target = {"azionario": 0.68, "obbligazionario": 0.24, "monetario": 0.08}
    m, p = _ribilancia(PRELIEVO=3000, TARGET=target)
    assert not any(d["fuori_banda"] for d in p["drift"].values()), p["drift"]
    venduto = sum(v["vendi"] for v in p["vendite"])
    assert venduto >= 3000 - 0.01 and abs(venduto - p["imposta"] - sum(p["acquisti"].values()) - 3000) < 0.01, p
    assert all(v["forzata"] for v in p["vendite"]), p["vendite"]
    assert all(abs(p["risultante"][c] - t) < 0.001 for c, t in target.items()), p["risultante"]
    # instradamento: il broker trattiene l'imposta, il prelievo non può superare la cassa netta
    piano = {"vendite": [{"broker": "A", "classe": "azionario", "tipo": "capitale", "aliquota": 0.26, "vendi": 10000,
                          "plus_real": 7692.31, "comp": 0, "tassa": 2000, "forzata": True},
                         {"broker": "B", "classe": "azionario", "tipo": "capitale", "aliquota": 0.26, "vendi": 9000,
                          "plus_real": 0, "comp": 0, "tassa": 0, "forzata": True}],
             "entrate": 0, "prelievo": 17000, "acquisti": {}}
    prelievi = {o["broker"]: o["importo_eur"] for o in m.genera_ordini(piano) if o["operazione"] == "PRELIEVO"}
    assert prelievi == {"A": 8000, "B": 9000}, prelievi
    for errato in ({"TARGET": {"azionario": 0.6, "obbligazionario": 0.3}},
                   {"POSITIONS": [{"broker": "X", "asset_class": "cripto", "tipo": "diverso", "valore": 100, "plus": 0}]},
                   {"POSITIONS": [{"broker": "X", "asset_class": "azionario", "tipo": "etf", "valore": 100, "plus": 0}]},
                   {"POSITIONS": [{"broker": "X", "asset_class": "azionario", "tipo": "diverso", "valore": 100,
                                   "plus": 0, "aliquota": 0.20}]},
                   {"OBIETTIVO_RIENTRO": "boh"}):
        m = carica_modulo(asset("it-rebalance", "ribilanciamento-multibroker.py"))
        for k, v in errato.items():
            setattr(m, k, v)
        try:
            m.pianifica()
        except m.ErroreInput:
            continue
        raise AssertionError(f"input non valido accettato: {errato}")


@test("T7 ribilanciamento: Leva 3 di tutte le classi prima della Leva 4, piano indipendente dall'ordine di TARGET")
def t_t7_ordine_leve():
    # REGRESSIONE (revisione T7): la minus di una classe elaborata dopo non compensava la plus
    # venduta prima sullo stesso broker, e piano e imposta dipendevano dall'ordine delle chiavi.
    pos = [{"broker": "X", "asset_class": "azionario", "tipo": "diverso", "valore": 68000, "plus": 20000},
           {"broker": "X", "asset_class": "obbligazionario", "tipo": "capitale", "valore": 24000, "plus": -3000},
           {"broker": "X", "asset_class": "monetario", "tipo": "capitale", "valore": 8000, "plus": 0}]
    classi = ("azionario", "obbligazionario", "monetario")
    for pesi, imposta, residuo in (((0.60, 0.15, 0.25), 0.0, 4175), ((0.55, 0.15, 0.30), 6675 * 20000 / 68000 * 0.26, 2500)):
        esiti = []
        for ordine in (classi, ("obbligazionario", "azionario", "monetario"), ("monetario", "obbligazionario", "azionario")):
            pesi_cl = dict(zip(classi, pesi))
            m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={c: pesi_cl[c] for c in ordine})
            assert abs(p["imposta"] - imposta) < 0.01, (ordine, p["imposta"], p["vendite"])
            assert abs(p["drift"]["azionario"]["residuo_non_corretto"] - residuo) < 0.01, (ordine, p["drift"]["azionario"])
            vendite = {v["classe"]: (round(v["vendi"], 2), round(v["comp"], 2)) for v in p["vendite"]}
            assert vendite["obbligazionario"] == (9000, 0) and vendite["azionario"][1] == 1125, (ordine, vendite)
            esiti.append((sorted(vendite.items()), round(p["imposta"], 2), round(p["zainetto"]["X"], 2),
                          {c: round(w, 6) for c, w in p["risultante"].items()}))
            # nel CSV di X la vendita in minus viene prima di quella in plus, e la nota lo dice
            vendite_x = [o for o in m.genera_ordini(p) if o["operazione"] == "VENDITA"]
            assert [o["classe"] for o in vendite_x] == ["obbligazionario", "azionario"], vendite_x
            assert "prima delle vendite in plus" in vendite_x[0]["nota"], vendite_x[0]["nota"]
        assert esiti[0] == esiti[1] == esiti[2], esiti
    # la minus venduta per riempire una sottopesata (dopo la plus decisa per la sovrappesata) va
    # comunque eseguita per prima, e l'imposta si calcola in quell'ordine
    pos = [{"broker": "X", "asset_class": "azionario", "tipo": "diverso", "valore": 70000, "plus": 21000},
           {"broker": "X", "asset_class": "obbligazionario", "tipo": "capitale", "valore": 25000, "plus": -2500},
           {"broker": "X", "asset_class": "monetario", "tipo": "capitale", "valore": 5000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={"azionario": 0.55, "obbligazionario": 0.21, "monetario": 0.24})
    assert [(v["classe"], round(v["vendi"], 2)) for v in p["vendite"]] == [("obbligazionario", 4000), ("azionario", 13833.33)], p["vendite"]
    assert abs(p["imposta"] - 12500 * 0.3 * 0.26) < 0.01 and abs(p["zainetto"]["X"]) < 0.01, (p["imposta"], p["zainetto"])
    # due sottopesate fuori banda in concorrenza per la stessa vendita a costo zero: prima la più grave
    pos = [{"broker": "X", "asset_class": c, "tipo": "capitale", "valore": v, "plus": pl}
           for c, v, pl in (("azionario", 64000, 16000), ("obbligazionario", 26000, 0), ("monetario", 6000, 0), ("oro", 4000, 0))]
    pesi = {"azionario": 0.60, "obbligazionario": 0.22, "monetario": 0.10, "oro": 0.08}
    for ordine in (list(pesi), ["monetario", "oro", "azionario", "obbligazionario"]):
        m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={c: pesi[c] for c in ordine})
        assert round(p["acquisti"]["oro"]) == 4000 and round(p["drift"]["monetario"]["residuo_non_corretto"]) == 4000, \
            (ordine, p["acquisti"], p["drift"]["monetario"])
    # la cassa che non basta per tutte le sottopesate fuori banda si divide in proporzione al fabbisogno
    pos = [{"broker": "X", "asset_class": c, "tipo": "capitale", "valore": v, "plus": pl}
           for c, v, pl in (("azionario", 70500, 0), ("obbligazionario", 19500, 4000), ("monetario", 6000, 0), ("oro", 4000, 0))]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={"azionario": 0.60, "obbligazionario": 0.16, "monetario": 0.14, "oro": 0.10})
    assert {c: round(x) for c, x in p["acquisti"].items()} == {"azionario": 0, "obbligazionario": 0, "monetario": 6000, "oro": 4500}, p["acquisti"]


@test("T7 ribilanciamento: una sottopesata fuori banda si riempie vendendo dalle classi sopra il target, anche in banda")
def t_t7_sottopesata():
    # REGRESSIONE (revisione T7): con il solo monetario fuori banda (sottopeso) lo script non
    # vendeva nulla e stampava «nessuna vendita necessaria (… o tutto in banda)».
    target = {"azionario": 0.64, "obbligazionario": 0.21, "monetario": 0.15}
    m, p = _ribilancia(TARGET=target)
    assert [d["fuori_banda"] for d in p["drift"].values()] == [False, False, True], p["drift"]
    assert sorted((v["broker"], v["classe"], v["tipo"], round(v["vendi"])) for v in p["vendite"]) == \
        [("Fineco", "azionario", "diverso", 4000), ("Fineco", "obbligazionario", "capitale", 3000)], p["vendite"]
    assert p["imposta"] == 0 and round(p["acquisti"]["monetario"]) == 7000, (p["imposta"], p["acquisti"])
    assert all(abs(p["risultante"][c] - t) < 1e-6 for c, t in target.items()), p["risultante"]
    assert "nessuna vendita necessaria" not in _stampato(m, p)
    pos = [{"broker": "X", "asset_class": c, "tipo": "capitale", "valore": v, "plus": 0}
           for c, v in (("azionario", 44000), ("obbligazionario", 44000), ("monetario", 12000))]
    target = {"azionario": 0.40, "obbligazionario": 0.40, "monetario": 0.20}
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET=target)
    assert sorted((v["classe"], round(v["vendi"])) for v in p["vendite"]) == [("azionario", 4000), ("obbligazionario", 4000)]
    assert all(abs(p["risultante"][c] - t) < 1e-6 for c, t in target.items()), p["risultante"]
    # sottopesata entro la banda larga e niente a costo zero: si rinvia, con il costo dichiarato
    m, p = _ribilancia(ZAINETTO={"Fineco": 0, "Directa": 0})
    ob = p["drift"]["obbligazionario"]
    assert round(ob["residuo_non_corretto"]) == 6000 and abs(ob["costo_a_target"] - 6000 * 3000 / 18000 * 0.26) < 0.01, ob
    # dopo la cassa delle vendite già decise la sottopesata è in banda: nessuna vendita tassata in più
    pos = [{"broker": "X", "asset_class": "azionario", "tipo": "capitale", "valore": 30000, "plus": 6000},
           {"broker": "Y", "asset_class": "azionario", "tipo": "diverso", "valore": 40000, "plus": 12000, "aliquota": 0.125},
           {"broker": "X", "asset_class": "obbligazionario", "tipo": "capitale", "valore": 30000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0, "Y": 0}, TARGET={"azionario": 0.5, "obbligazionario": 0.5})
    assert [(v["broker"], v["aliquota"], round(v["vendi"], 2)) for v in p["vendite"]] == [("Y", 0.125, 17500)], p["vendite"]
    assert abs(p["imposta"] - 17500 * 0.3 * 0.125) < 0.01, p["imposta"]      # Leva 4 per imposta effettiva (F044)
    # sottopesata ancora oltre la banda larga dopo cassa e costo zero: vendite tassate dalle classi
    # sopra il target (anche in banda), le più economiche per prime, fino a metà banda (o a target)
    pos = [{"broker": "X", "asset_class": c, "tipo": "capitale", "valore": v, "plus": pl}
           for c, v, pl in (("azionario", 34000, 6800), ("obbligazionario", 24000, 2400), ("oro", 24000, 4800), ("monetario", 18000, 0))]
    target = {"azionario": 0.30, "obbligazionario": 0.20, "oro": 0.20, "monetario": 0.30}
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET=target)
    lordo = (9500 - 4000 * (0.052 - 0.026)) / (1 - 0.052)        # obbligazionario al 2,6% per euro, poi il 5,2%
    assert abs(p["acquisti"]["monetario"] - 9500) < 0.01 and abs(sum(v["vendi"] for v in p["vendite"]) - lordo) < 0.01, p["vendite"]
    assert [v["classe"] for v in p["vendite"]][0] == "obbligazionario" and abs(p["imposta"] - (lordo - 9500)) < 0.01, p["vendite"]
    mon = p["drift"]["monetario"]
    assert "rientro a metà banda" in mon["motivo"] and f"così €{m.numero(p['imposta'])}" in mon["motivo"], mon["motivo"]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET=target, OBIETTIVO_RIENTRO="target")
    assert round(sum(v["vendi"] for v in p["vendite"])) == 12000 and abs(p["imposta"] - 520) < 0.01, p["vendite"]
    assert "metà banda" not in p["drift"]["monetario"]["motivo"], p["drift"]["monetario"]
    # nessuna vendita, classi fuori banda corrette dai versamenti: il messaggio lo dice
    m, p = _ribilancia(NUOVI_VERSAMENTI=20000)
    testo = _stampato(m, p)
    assert "si correggono con versamenti e flussi" in testo and "tutto in banda" not in testo, testo


@test("T7 ribilanciamento: minus al 12,5% nello zainetto al 48,08%, quota a costo zero scalata, etichette del consumo")
def t_t7_coefficiente():
    # REGRESSIONE (revisione T7, F044): tre mutanti che cambiano l'imposta passavano i test.
    pos = [{"broker": "X", "asset_class": "obbligazionario", "tipo": "diverso", "valore": 20000, "plus": -2000, "aliquota": 0.125},
           {"broker": "X", "asset_class": "azionario", "tipo": "diverso", "valore": 20000, "plus": 4000},
           {"broker": "X", "asset_class": "monetario", "tipo": "capitale", "valore": 60000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={"azionario": 0.0, "obbligazionario": 0.0, "monetario": 1.0})
    zero = 2000 * 12.5 / 26 / 0.2                                 # euro di azionario coperti dalla minus del BTP
    assert abs(p["imposta"] - (20000 - zero) * 0.2 * 0.26) < 0.01, p["imposta"]
    az = [v for v in p["vendite"] if v["classe"] == "azionario"][0]
    assert abs(az["comp"] - 2000 * 12.5 / 26) < 0.01, az
    testo = _stampato(m, p)
    assert re.search(r"minus €2\.000[^\n]*€962 di zainetto: minus al 12,5% computata al 48,08%", testo), testo
    assert "VENDI €20.000 obbligazionario/diverso (12,5%) @ X" in testo, testo
    pos = [{"broker": "X", "asset_class": "obbligazionario", "tipo": "diverso", "valore": 38000, "plus": 3800, "aliquota": 0.125},
           {"broker": "X", "asset_class": "azionario", "tipo": "capitale", "valore": 62000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 300}, TARGET={"obbligazionario": 0.30, "azionario": 0.70})
    assert [round(v["vendi"], 2) for v in p["vendite"]] == [round(300 / (12.5 / 26 * 0.1), 2)] and p["imposta"] == 0, p["vendite"]
    assert abs(p["zainetto"]["X"]) < 0.01, p["zainetto"]
    btp = [{"broker": "X", "asset_class": "obbligazionario", "tipo": "diverso", "valore": 13000, "plus": 3000, "aliquota": 0.125},
           {"broker": "X", "asset_class": "azionario", "tipo": "capitale", "valore": 7000, "plus": 0}]
    # l'imposta pagata esce dal portafoglio: non è un sottopeso dell'azionario (100% del patrimonio finale)
    m, p = _ribilancia(POSITIONS=btp, ZAINETTO={"X": 0}, TARGET={"azionario": 1.0, "obbligazionario": 0.0})
    assert p["drift"]["azionario"]["residuo_non_corretto"] == 0 and p["risultante"]["azionario"] == 1.0, p["drift"]["azionario"]
    m, p = _ribilancia(POSITIONS=btp, ZAINETTO={"X": 5000}, TARGET={"azionario": 1.0, "obbligazionario": 0.0})
    testo = _stampato(m, p)
    assert re.search(r"compensata €3\.000[^\n]*€1\.442 di zainetto: plus al 12,5% computata al 48,08%", testo), testo
    nota = [o for o in m.genera_ordini(p) if o["operazione"] == "VENDITA"][0]["nota"]
    assert "1.442" in nota and "48,08%" in nota, nota


@test("T7 ribilanciamento: OBIETTIVO_RIENTRO «banda» e residui dichiarati (nessuna vendita, sezione 5 senza ordini)")
def t_t7_rientro_banda_residui():
    # REGRESSIONE (revisione T7): «rinviata per imposta €0» in modalità «banda»; caso EC3 (zainetto 0,
    # nessuna vendita) mai testato; sezione 5 «ORDINI ESPORTATI» stampata senza ordini.
    m, p = _ribilancia(OBIETTIVO_RIENTRO="banda")
    assert [round(v["vendi"]) for v in p["vendite"]] == [5500] and p["imposta"] == 0, p["vendite"]
    # la cassa oltre il rientro va prima alla classe fuori banda, poi alle altre
    assert {c: round(x) for c, x in p["acquisti"].items()} == {"azionario": 0, "obbligazionario": 5500, "monetario": 0}, p["acquisti"]
    az = p["drift"]["azionario"]
    assert round(az["residuo_non_corretto"]) == 2500 and "rinviat" not in az["motivo"] and "OBIETTIVO_RIENTRO" in az["motivo"], az
    m, p = _ribilancia(OBIETTIVO_RIENTRO="banda", ZAINETTO={"Fineco": 2000, "Directa": 0})
    az = p["drift"]["azionario"]
    assert "€22 " in az["motivo"] and "€130" not in az["motivo"], az["motivo"]
    m, p = _ribilancia(ZAINETTO={"Fineco": 0, "Directa": 0})
    assert p["vendite"] == [] and p["imposta"] == 0, (p["vendite"], p["imposta"])
    assert round(p["drift"]["azionario"]["residuo_non_corretto"]) == 8000, p["drift"]["azionario"]
    testo = _stampato(m, p)
    assert "nessuna vendita necessaria" not in testo and re.search(r"rinviat[^\n]*8\.000|8\.000[^\n]*rinviat", testo), testo[-800:]
    import contextlib
    import io
    tmp = tempfile.mkdtemp()
    try:
        m.CARTELLA_OUTPUT = os.path.join(tmp, "ordini")
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            m.esporta_ordini(p)
        assert "NESSUN ORDINE" in buf.getvalue() and "ORDINI ESPORTATI" not in buf.getvalue(), buf.getvalue()
        assert not os.path.exists(m.CARTELLA_OUTPUT), "cartella ordini/ creata senza ordini"
        assert m.scrivi_csv_per_broker([], m.CARTELLA_OUTPUT) == [] and not os.path.exists(m.CARTELLA_OUTPUT)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    # MARGINE_RIENTRO diverso da 0,5: la stampa non dice «metà banda»
    m, p = _ribilancia(ZAINETTO={"Fineco": 0, "Directa": 0}, MARGINE_RIENTRO=0.2,
                       TARGET={"azionario": 0.50, "obbligazionario": 0.40, "monetario": 0.10})
    assert round(sum(v["vendi"] for v in p["vendite"])) == 17000, p["vendite"]
    testo = _stampato(m, p)
    assert "metà banda" not in testo and "20% della banda stretta" in testo, testo


@test("T7 ribilanciamento: prelievo con imposta in una sola vendita per posizione, residui ricalcolati, prelievo non finanziabile")
def t_t7_prelievo_imposta():
    # REGRESSIONE (revisione T7): il gross-up del prelievo aggiungeva una riga a ogni giro (fino a
    # €0,01), vendeva per intero la posizione con una piccola quota a costo zero, lasciava il residuo
    # calcolato prima delle vendite forzate e dichiarava eseguito un prelievo non finanziabile.
    zero = {"Fineco": 0, "Directa": 0}
    for prelievo, target in ((2000, None), (3000, {"azionario": 0.68, "obbligazionario": 0.24, "monetario": 0.08})):
        parametri = {"PRELIEVO": prelievo, "ZAINETTO": zero}
        if target:
            parametri["TARGET"] = target
        m, p = _ribilancia(**parametri)
        venduto = sum(v["vendi"] for v in p["vendite"])
        assert p["imposta"] > 0 and abs(venduto - p["imposta"] - sum(p["acquisti"].values()) - prelievo) < 0.01, p
        chiavi = [(v["broker"], v["classe"], v["tipo"]) for v in p["vendite"]]
        assert len(chiavi) == len(set(chiavi)), chiavi
        ordini, saldo = _cassa_per_broker(m, p)
        assert all(o["importo_eur"] >= 1 for o in ordini if o["operazione"] == "VENDITA"), ordini
        assert all(v >= -0.01 for v in saldo.values()), saldo
        testo = _stampato(m, p)
        assert "il resto è reinvestito" not in testo, testo
        if target:
            assert all(abs(p["risultante"][c] - t) < 0.001 for c, t in target.items()), p["risultante"]
            # per euro lordo: 68% dall'azionario, di cui la parte coperta dalla minus dell'ETF obbligazionario
            # Fineco (24% × 500/15.000, sul diverso Fineco con plus al 40%) a costo zero, il resto
            # dall'ETF Directa (plus 3.000/18.000 al 26%)
            k = (0.68 - 0.24 * 500 / 15000 / 0.4) * 3000 / 18000 * 0.26
            assert abs(p["imposta"] - prelievo * k / (1 - k)) < 0.01, p["imposta"]
        else:
            az = p["drift"]["azionario"]
            atteso = 9200 - sum(v["vendi"] for v in p["vendite"] if v["classe"] == "azionario")
            assert abs(az["residuo_non_corretto"] - atteso) < 0.01, (az, atteso)
            assert m.numero(atteso) in testo and "9.200" not in testo and "prelievo" in az["motivo"], testo
    # la quota a costo zero va per prima, il resto della posizione dopo le tassate più economiche
    imposte = []
    for z in (0, 200, 1000):
        m, p = _ribilancia(PRELIEVO=5000, ZAINETTO={"Fineco": z, "Directa": 0},
                           TARGET={"azionario": 0.68, "obbligazionario": 0.24, "monetario": 0.08})
        imposte.append(p["imposta"])
        if z == 200:
            div = [v for v in p["vendite"] if v["tipo"] == "diverso" and v["classe"] == "azionario"]
            assert div and div[0]["tassa"] == 0 and div[0]["vendi"] < 1000, p["vendite"]
    assert imposte[0] > imposte[1] > imposte[2], imposte
    # una classe appena sopra il livello del water-filling: nessuna vendita da centesimi
    # (o una minus minuscola che renderebbe a costo zero pochi centesimi di plus)
    for indice, meno, prelievi in ((5, 10, (112, 116, 123)), (4, 20, (100,))):
        pos = [dict(x) for x in carica_modulo(asset("it-rebalance", "ribilanciamento-multibroker.py")).POSITIONS]
        pos[indice]["valore"] -= meno
        for prelievo in prelievi:
            m, p = _ribilancia(PRELIEVO=prelievo, POSITIONS=pos, ZAINETTO={"Fineco": 0, "Directa": 0},
                               TARGET={"azionario": 0.68, "obbligazionario": 0.24, "monetario": 0.08})
            assert all(o["importo_eur"] >= 1 for o in m.genera_ordini(p) if o["operazione"] == "VENDITA"), (prelievo, p["vendite"])
    pos = [{"broker": "X", "asset_class": "azionario", "tipo": "capitale", "valore": 70000, "plus": 0},
           {"broker": "X", "asset_class": "azionario", "tipo": "capitale", "valore": 0.6, "plus": -0.1},
           {"broker": "X", "asset_class": "obbligazionario", "tipo": "capitale", "valore": 30000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={"azionario": 0.6, "obbligazionario": 0.4})
    assert [(v["vendi"] > 1, round(v["vendi"])) for v in p["vendite"]] == [(True, 10000)], p["vendite"]
    m, p = _ribilancia(PRELIEVO=20000, ZAINETTO={"Fineco": 1000, "Directa": 0})
    ordini, saldo = _cassa_per_broker(m, p)
    assert all(o["importo_eur"] >= 1 for o in ordini if o["operazione"] == "VENDITA"), ordini
    chiavi = [(v["broker"], v["classe"], v["tipo"]) for v in p["vendite"]]
    assert len(chiavi) == len(set(chiavi)), chiavi
    # prelievo oltre il netto di liquidazione: errore; pari al netto: tutto venduto, pesi a zero
    m = carica_modulo(asset("it-rebalance", "ribilanciamento-multibroker.py"))
    m.PRELIEVO = 99000
    try:
        m.pianifica()
        raise AssertionError("prelievo non finanziabile accettato")
    except m.ErroreInput as e:
        assert "96.620" in str(e), e
    m, p = _ribilancia(PRELIEVO=96620)
    assert all(0 <= w <= 1 for w in p["risultante"].values()), p["risultante"]
    assert round(sum(o["importo_eur"] for o in m.genera_ordini(p) if o["operazione"] == "PRELIEVO"), 2) == 96620
    assert "interamente liquidato" in _stampato(m, p)


@test("T7 ribilanciamento: aliquota mista di ETF e fondi su titoli di Stato (DM 13/12/2011, circ. 19/E/2014)")
def t_t7_aliquota_mista():
    # Per gli OICR la parte di proventi e perdite riferibile ai titoli di Stato e white list (quota
    # media dell'attivo) conta per il 48,08%: aliquota effettiva = 12,5% × p + 26% × (1 − p), e le
    # minus entrano nello zainetto per (1 − 51,92% × p). Esempio della circ. 19/E/2014, par. 5.1:
    # provento 300, p = 26,5% → ritenuta 67,27.
    etf = [{"broker": "X", "asset_class": "obbligazionario", "tipo": "capitale", "valore": 1300, "plus": 300,
            "quota_titoli_stato": 0.265},
           {"broker": "X", "asset_class": "azionario", "tipo": "capitale", "valore": 8700, "plus": 0}]
    m, p = _ribilancia(POSITIONS=etf, ZAINETTO={"X": 0}, TARGET={"azionario": 1.0, "obbligazionario": 0.0})
    assert abs(p["imposta"] - 67.27) < 0.01, p["imposta"]
    assert abs(m._aliquota(etf[0]) - (0.125 * 0.265 + 0.26 * 0.735)) < 1e-12
    nota = [o for o in m.genera_ordini(p) if o["operazione"] == "VENDITA"][0]["nota"]
    assert "plus €300 TASSATA ad aliquota mista 22,4% (reddito di capitale: non si compensa" in nota, nota
    # quota 1 → 12,5% pieno; vicino agli estremi l'etichetta passa a due decimali
    for quota, etichetta, imposta in ((1.0, "obbligazionario/capitale (12,5%) @", 37.5),
                                      (0.003, "(aliquota mista 25,96%)", 300 * (0.125 * 0.003 + 0.26 * 0.997)),
                                      (0.9995, "(aliquota mista 12,51%)", 300 * (0.125 * 0.9995 + 0.26 * 0.0005))):
        m, p = _ribilancia(POSITIONS=[dict(etf[0], quota_titoli_stato=quota), etf[1]], ZAINETTO={"X": 0},
                           TARGET={"azionario": 1.0, "obbligazionario": 0.0})
        assert abs(p["imposta"] - imposta) < 0.01 and etichetta in _stampato(m, p), (quota, p["imposta"])
    # la minus dell'ETF misto entra nello zainetto ridotta, e compensa così la plus di un'azione
    pos = [{"broker": "X", "asset_class": "obbligazionario", "tipo": "capitale", "valore": 10000, "plus": -1000,
            "quota_titoli_stato": 0.6},
           {"broker": "X", "asset_class": "azionario", "tipo": "diverso", "valore": 10000, "plus": 5000},
           {"broker": "X", "asset_class": "monetario", "tipo": "capitale", "valore": 80000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={"obbligazionario": 0.0, "azionario": 0.0, "monetario": 1.0})
    credito = 1000 * (1 - 0.6 * (26 - 12.5) / 26)
    az = [v for v in p["vendite"] if v["classe"] == "azionario"][0]
    assert abs(az["comp"] - credito) < 0.01 and abs(p["imposta"] - (5000 - credito) * 0.26) < 0.01, (az, p["imposta"])
    testo = _stampato(m, p)
    assert re.search(r"VENDI €10\.000 obbligazionario/capitale \(aliquota mista 17,9%\) @ X", testo), testo
    assert re.search(r"minus €1\.000[^\n]*€688 di zainetto: minus ad aliquota mista 17,9% computata al 68,85%",
                     testo), testo                                      # 12,5% × 0,6 + 26% × 0,4
    # in alternativa alla quota si può dare l'aliquota effettiva tra 12,5% e 26%
    pos[0] = dict(pos[0], aliquota=0.179)
    del pos[0]["quota_titoli_stato"]
    m, p2 = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={"obbligazionario": 0.0, "azionario": 0.0, "monetario": 1.0})
    assert abs(p2["imposta"] - p["imposta"]) < 0.01, (p2["imposta"], p["imposta"])
    # Leva 4: a parità di plus, prima l'ETF con più titoli di Stato (imposta per euro più bassa)
    pos = [{"broker": "X", "asset_class": "obbligazionario", "tipo": "capitale", "valore": 20000, "plus": 4000,
            "quota_titoli_stato": q} for q in (0.0, 0.9)] + \
          [{"broker": "X", "asset_class": "monetario", "tipo": "capitale", "valore": 60000, "plus": 0}]
    m, p = _ribilancia(POSITIONS=pos, ZAINETTO={"X": 0}, TARGET={"obbligazionario": 0.25, "monetario": 0.75})
    venduta = [v for v in p["vendite"] if v["classe"] == "obbligazionario"]
    assert len(venduta) == 1 and abs(venduta[0]["aliquota"] - (0.125 * 0.9 + 0.26 * 0.1)) < 1e-9, venduta
    # validazione: la quota vale solo per fondi ed ETF (capitale), l'aliquota effettiva tra 12,5% e 26%
    for errato in ({"tipo": "diverso", "quota_titoli_stato": 0.5}, {"tipo": "capitale", "quota_titoli_stato": 1.2},
                   {"tipo": "capitale", "quota_titoli_stato": -0.1},
                   {"tipo": "capitale", "aliquota": 0.30}, {"tipo": "capitale", "aliquota": 0.10},
                   {"tipo": "capitale", "aliquota": 0.2, "quota_titoli_stato": 0.5},
                   {"tipo": "capitale", "aliquota": "12,5%"}, {"tipo": "capitale", "quota_titoli_stato": "60%"},
                   {"tipo": "capitale", "quota_titoli_stato": True}):
        pos = [dict({"broker": "X", "asset_class": "azionario", "valore": 100, "plus": 0}, **errato)]
        m = carica_modulo(asset("it-rebalance", "ribilanciamento-multibroker.py"))
        m.POSITIONS, m.TARGET = pos, {"azionario": 1.0}
        try:
            m.pianifica()
        except m.ErroreInput:
            continue
        raise AssertionError(f"input non valido accettato: {errato}")


@test("T7 ribilanciamento: una sola regola di rientro e di banda in skill e reference")
def t_t7_testi():
    # REGRESSIONE F045/F046/F030: «a target» contro «in banda», prelievo promesso e non fatto,
    # «banda fissa ±5%» come regola di base.
    reb = os.path.join(SKILLS, "it-rebalance")
    testi = {rel: open(os.path.join(reb, rel), encoding="utf-8").read()
             for rel in ("SKILL.md", "references/gerarchia-leve.md", "references/bande-dinamiche.md",
                         "references/sequenza-liquidazione.md", "references/ribilanciamento-multibroker.md",
                         "references/coordinamento-zainetto.md")}
    regola = re.compile(r"a target[^\n]*costo (netto )?zero[^\n]*metà (della )?banda", re.I)
    for rel in ("SKILL.md", "references/gerarchia-leve.md", "references/bande-dinamiche.md"):
        assert regola.search(testi[rel]), f"{rel}: manca la regola di rientro mista"
    vietate = [r"raccomandare quella in banda", r"\"In banda\" batte", r"banda fissa\*\* \(±5% assoluto\)",
               r"prevista per v0\.5", r"Lo strumento fa esattamente così", r"PMC \(o lotti\)",
               r"Ribilanciare \"a target esatto\" quando",
               # revisione T7: banda fissa, Leva 4 per plus latente, vendita tassata senza gate,
               # «capitale per ultime», OBIETTIVO_RIENTRO a due valori, frase citata che non esiste più
               r"banda ±5%", r"plus latente minore", r"accettando la tassa", r"per ultime le posizioni",
               r"permette di scegliere «target» o «banda»", r"ribilanciamento parziale o attesa\"",
               r"Risparmio dal coordinamento: €582",
               # aliquota mista: ETF e fondi su titoli di Stato non più forzati al 26%
               r"che lo script non rappresenta", r"imposta stimata è per eccesso", r"ultima certificazione"]
    for rel, testo in testi.items():
        for v in vietate:
            assert not re.search(v, testo), f"{rel}: «{v}»"
    # il prelievo in decumulo si passa con la chiave «prelievo» del JSON di --dati (revisione T9)
    assert "chiave `prelievo`" in testi["references/sequenza-liquidazione.md"] and "chiave `prelievo`" in testi["SKILL.md"]
    assert "MODULA_VOLATILITA = False" in testi["references/bande-dinamiche.md"] or \
        "spenta di default" in testi["references/bande-dinamiche.md"], "bande-dinamiche: modulazione spenta di default"
    assert "aliquota" in testi["SKILL.md"].split("### Step 1")[1].split("### Step 2")[0], "Step 1 senza l'aliquota per posizione"
    for rel in ("SKILL.md", "references/ribilanciamento-multibroker.md"):
        assert "quota_titoli_stato" in testi[rel], f"{rel}: manca l'aliquota mista di ETF e fondi su titoli di Stato"
    aliquote = open(os.path.join(SKILLS, "it-tax-optimization", "references", "aliquote-12-5-26.md"), encoding="utf-8").read()
    for frase in ("67,27", "Non dal KID", "art. 1 c. 3", "valida per il semestre della vendita"):
        assert frase in aliquote, f"aliquote-12-5-26.md: manca «{frase}»"
    assert "ultima certificazione" not in aliquote, "aliquote-12-5-26.md: conta la certificazione del semestre della vendita"
    assert re.search(r"\*\*perdite\*\* \(DM 13/12/2011, art\. 1 c\. 4\)[^\n]*ridotta del 51,92% \(circ\. AdE "
                     r"19/E/2014, par\. 5\.1", aliquote), \
        "aliquote-12-5-26.md: le minus dei fondi su titoli di Stato entrano ridotte nello zainetto"
    # revisione T7: Leva 4 per imposta per euro venduto, minus eseguite per prime, sottopesate,
    # tre valori di OBIETTIVO_RIENTRO, gate nel coordinamento di fine anno, unità dello zainetto
    for rel in ("SKILL.md", "references/ribilanciamento-multibroker.md", "references/sequenza-liquidazione.md",
                "references/gerarchia-leve.md"):
        assert re.search(r"imposta per euro venduto", testi[rel]), f"{rel}: Leva 4 senza l'imposta per euro venduto"
    for rel in ("SKILL.md", "references/gerarchia-leve.md", "references/ribilanciamento-multibroker.md"):
        assert re.search(r"minus[^\n]*(prima delle plus|eseguit[ao] per prim|per prim[ae][^\n]*CSV)", testi[rel]), \
            f"{rel}: manca l'ordine di esecuzione (le minus prima delle plus)"
    assert "art. 6 c. 5 D.Lgs. 461/97 → art. 306 c. 6 TU 2027" in testi["SKILL.md"]
    for rel in ("SKILL.md", "references/bande-dinamiche.md", "references/ribilanciamento-multibroker.md"):
        assert re.search(r"(?i)sottopesat[ae] fuori banda", testi[rel]), f"{rel}: manca la regola delle sottopesate fuori banda"
    for rel in ("SKILL.md", "references/bande-dinamiche.md"):
        assert all(x in testi[rel] for x in ("«misto»", "«target»", "«banda»", "MARGINE_RIENTRO")), f"{rel}: OBIETTIVO_RIENTRO"
    coord = testi["references/coordinamento-zainetto.md"]
    fase4 = coord.split("### Fase 4")[1].split("## Esempio")[0]
    assert "banda larga" in fase4 and "metà banda" in fase4, "coordinamento-zainetto: Leva 4 senza gate"
    assert "0,4808" in coord.split("### Fase 3")[1].split("### Fase 4")[0], "coordinamento-zainetto: Fase 3 senza il 48,08%"
    assert re.search(r"\+7%[^\n]*sotto la banda larga[^\n]*rinvia", coord), "coordinamento-zainetto: esempio senza gate"
    for rel in ("SKILL.md",):
        assert "26%-equivalenti" in testi[rel], f"{rel}: zainetto senza unità"
    assert re.search(r"reddito_diverso` \([^)]*obbligazioni", testi["SKILL.md"]), "SKILL.md: tipo_fiscale senza obbligazioni"

# ══════════════════════ obiettivi (T8) ══════════════════════

def _modulo_obiettivi(nome):
    """Importa uno script di it-pianificazione-obiettivi (con il suo motore_bootstrap)."""
    cartella = os.path.join(SKILLS, "it-pianificazione-obiettivi", "assets")
    sys.path.insert(0, cartella)
    try:
        return carica_modulo(os.path.join(cartella, nome))
    finally:
        sys.path.remove(cartella)


def _testo_obiettivi(rel):
    return open(os.path.join(SKILLS, "it-pianificazione-obiettivi", rel), encoding="utf-8").read()


def _euro_it(x):
    return f"€{x:,.0f}".replace(",", ".")


def _pct_it(x, decimali=1):
    return f"{x:.{decimali}f}".replace(".", ",")


def _da_it(testo):
    """Numero da un testo in formato italiano: '1.234,5' → 1234.5."""
    return float(testo.replace(".", "").replace(",", "."))


def _pmt_anno_per_anno(target, gia, rend, rend_atterraggio, anni, finestra):
    """Replica indipendente del versamento col calendario annuale del glide path: la quota
    dell'anno y è quella del calendario a inizio anno (anni residui = anni − y), il rendimento
    è lineare nella quota tra il profilo e l'atterraggio; 12 versamenti di fine mese."""
    w = min(finestra, anni)
    tassi = [rend_atterraggio + (rend - rend_atterraggio) * (1.0 if anni - y >= w else (anni - y) / w)
             for y in range(anni)]
    crescita, coeff = 1.0, 0.0
    for r in reversed(tassi):  # FV = gia·Π(1+r) + PMT·Σ a(r_y)·Π_{z>y}(1+r_z)
        coeff += (12.0 if r == 0 else r / ((1 + r) ** (1 / 12) - 1)) * crescita
        crescita *= 1 + r
    return (target - gia * crescita) / coeff, crescita


def _obiettivi_foglio():
    """Obiettivi e parametri del pianificatore d'esempio, dalle cache ricalcolate."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_ricalcolati()[asset("it-pianificazione-obiettivi", "pianificatore-obiettivi.xlsx")],
                       data_only=True)
    par, obj = wb["Parametri"], wb["Obiettivi"]
    profili = {par[f"A{r}"].value: (par[f"B{r}"].value, par[f"C{r}"].value, par[f"D{r}"].value)
               for r in range(10, 16) if par[f"A{r}"].value}
    obiettivi = []
    for r in range(5, 15):
        if obj[f"A{r}"].value:
            obiettivi.append({c: obj[f"{col}{r}"].value for c, col in (
                ("nome", "A"), ("target", "B"), ("anni", "C"), ("profilo", "D"), ("gia", "F"),
                ("priorita", "G"), ("versamento", "I"), ("tipo", "N"))})
    return par, profili, obiettivi


@test("obiettivi: versamento col glide path ≥ quota fissa, stesso numero in foglio, motore, Monte Carlo e glide path")
def t_versamento_glide():
    # F040: gli obiettivi terminanti de-riskano negli ultimi anni, ma il versamento usava il
    # rendimento pieno del profilo (shortfall ~10% a 20 anni). Il foglio approssima la rampa col
    # suo rendimento medio: sugli obiettivi d'esempio entro l'1,5% dal calcolo anno per anno, mai
    # sotto la quota fissa (l'errore per profilo e finestra è in t_errore_due_fasi).
    par, profili, obiettivi = _obiettivi_foglio()
    infl = par["B6"].value
    rend_l = min(profili.values(), key=lambda p: p[2])[0]   # rendimento all'atterraggio (0% azionario)
    finestra = 7
    attesi = {}
    for g in obiettivi:
        rend, _, quota = profili[g["profilo"]]
        target_nom = g["target"] * (1 + infl) ** g["anni"]
        fisso, _ = _pmt_anno_per_anno(target_nom, g["gia"] or 0, rend, rend, g["anni"], finestra)
        foglio = g["versamento"]
        if g["tipo"] == "permanente" or quota == 0:
            assert abs(foglio - fisso) < 0.01, f"{g['nome']}: permanente €{foglio:.2f} ≠ quota fissa €{fisso:.2f}"
        else:
            esatto, _ = _pmt_anno_per_anno(target_nom, g["gia"] or 0, rend, rend_l, g["anni"], finestra)
            assert foglio >= 0.985 * esatto, (f"{g['nome']}: il foglio chiede €{foglio:.0f}/mese, col glide path "
                                              f"ne servono €{esatto:.0f}: de-risking ignorato nel versamento")
            assert fisso < foglio <= esatto + 0.01, f"{g['nome']}: €{foglio:.2f} fuori da [{fisso:.2f}, {esatto:.2f}]"
        attesi[g["nome"]] = (foglio, target_nom, g, rend)
    assert "finestra" in str(par["A7"].value).lower() and par["B7"].value == finestra, \
        "Parametri!A7:B7 deve portare la finestra di de-risking (7 anni, come glidepath-obiettivi.py)"
    motore = _modulo_obiettivi("motore_bootstrap.py")
    for nome, (foglio, target_nom, g, rend) in attesi.items():
        glide = g["tipo"] != "permanente" and profili[g["profilo"]][2] > 0
        calcolo = motore.pmt_glide(target_nom, g["gia"] or 0, rend, rend_l if glide else rend, g["anni"], finestra)
        assert abs(calcolo - foglio) < 0.01, f"{nome}: motore €{calcolo:.2f} ≠ foglio €{foglio:.2f}"
    rc, out = run_script(asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"))
    assert rc == 0, out[-400:]
    for nome, (_, _, testo) in _sezioni_obiettivi(out).items():
        m = re.search(r"Versamento deterministico:\s*(€[\d.]+)/mese", testo)
        assert m and m.group(1) == _euro_it(attesi[nome][0]), f"Monte Carlo, {nome}: {m and m.group(1)} ≠ foglio"
    rc, out = run_script(asset("it-pianificazione-obiettivi", "glidepath-obiettivi.py"))
    assert rc == 0, out[-400:]
    assert "VERSAMENTI" in out, "glidepath: manca il riepilogo dei versamenti usati nella proiezione"
    for nome, (foglio, *_) in attesi.items():
        m = re.search(re.escape(nome) + r"[^\n]*?(€[\d.]+)/mese", out.split("VERSAMENTI")[-1])
        assert m and m.group(1) == _euro_it(foglio), f"glide path, {nome}: {m and m.group(1)} ≠ foglio"
    # allocatore: il capitale richiesto oggi da un terminante segue la stessa crescita col glide
    tmp = tempfile.mkdtemp()
    try:
        rc, out = run_script(asset("it-pianificazione-obiettivi", "allocatore-capitale.py"),
                             "--json", os.path.join(tmp, "a.json"))
        assert rc == 0, out[-400:]
        righe = {r["nome"]: r for r in json.load(open(os.path.join(tmp, "a.json"), encoding="utf-8"))["obiettivi"]}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    g = attesi["Lungo termine / libertà"][2]
    rend = profili[g["profilo"]][0]
    target_nom = g["target"] * (1 + infl) ** g["anni"]
    _, crescita = _pmt_anno_per_anno(target_nom, 0, rend, rend_l, g["anni"], finestra)
    atteso = target_nom / crescita - (g["gia"] or 0)
    assert abs(righe[g["nome"]]["richiesto"] - atteso) < 1, \
        f"allocatore: richiesto €{righe[g['nome']]['richiesto']:.0f}, col glide path €{atteso:.0f}"


@test("obiettivi: l'errore della formula in due fasi è quello documentato, per profilo e finestra")
def t_errore_due_fasi():
    # Revisione T8: «sbaglia per difetto di meno dell'1,5%» era falso già con la finestra di
    # default (Crescita massima 1,7%, Azionario totale 1,8%) e con la finestra di 10 anni, che
    # glide-path.md propone, arriva al 3%. Il caso peggiore è un orizzonte pari alla finestra.
    motore = _modulo_obiettivi("motore_bootstrap.py")
    _, profili, _ = _obiettivi_foglio()
    rend_l = min(profili.values(), key=lambda p: p[2])[0]

    def errore_max(rend, finestra):
        peggiore = 0.0
        for n in range(1, 41):
            target = 100_000 * 1.02 ** n
            esatto, _ = _pmt_anno_per_anno(target, 0, rend, rend_l, n, finestra)
            peggiore = max(peggiore, 1 - motore.pmt_glide(target, 0, rend, rend_l, n, finestra) / esatto)
        return peggiore * 100

    err = {(nome, f): errore_max(p[0], f) for nome, p in profili.items() if p[2] > 0 for f in (7, 10)}
    for (nome, f), e in err.items():
        assert e >= 0, f"{nome}, finestra {f}: la formula in due fasi non sbaglia per difetto"
        if f == 7:
            assert e < 2, f"{nome}: errore {e:.2f}% con la finestra di default, oltre il 2%"
    assert err[("Moderato", 7)] < 1, f"Moderato: errore {err[('Moderato', 7)]:.2f}%, documentato sotto l'1%"
    frasi = ("sotto l'1% con Moderato",
             f"{_pct_it(err[('Crescita', 7)])}% con Crescita",
             f"{_pct_it(err[('Azionario totale', 7)])}% con Azionario totale",
             f"{_pct_it(err[('Crescita massima', 10)])}% con Crescita massima",
             f"{_pct_it(err[('Azionario totale', 10)])}% con Azionario totale")
    motore_txt = open(asset("it-pianificazione-obiettivi", "motore_bootstrap.py"), encoding="utf-8").read()
    for rel, testo in (("references/glide-path.md", _testo_obiettivi("references/glide-path.md")),
                       ("references/matematica-accumulo.md", _testo_obiettivi("references/matematica-accumulo.md")),
                       ("motore_bootstrap.py", " ".join(motore_txt.split()))):
        testo = " ".join(testo.split())
        assert "meno dell'1,5%" not in testo, f"{rel}: «meno dell'1,5%» vale solo fino al profilo Crescita"
        for frase in frasi:
            assert frase in testo, f"{rel}: manca «{frase}» (errore misurato)"


@test("glidepath e Monte Carlo: «come la colonna I» solo coi parametri del foglio, avviso sui pianificatori vecchi")
def t_glide_come_foglio():
    # Revisione T8: su un pianificatore compilato prima della 0.10.0 (senza Parametri!B7, colonna I
    # a quota fissa) o con ATTERRAGGIO_AZ/CURVATURA/FINESTRA_DERISK cambiati, gli script davano un
    # versamento diverso dal foglio dichiarandolo «come la colonna I».
    from openpyxl import load_workbook
    gp = asset("it-pianificazione-obiettivi", "glidepath-obiettivi.py")
    mc = asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py")
    rc, out = run_script(gp)
    assert rc == 0 and "come la colonna I del foglio" in out, "default: i versamenti sono quelli della colonna I"
    assert "0.10.0" not in out, "default: avviso sul pianificatore vecchio fuori luogo"
    for args in (("--set", "CURVATURA=0.5"), ("--set", "ATTERRAGGIO_AZ=0.2"), ("--set", "FINESTRA_DERISK=5")):
        rc, out = run_script(gp, *args)
        assert rc == 0, out[-400:]
        assert "come la colonna I" not in out and "DIVERSI dalla colonna I" in out, f"{args}: versamenti spacciati per la colonna I"
        rc, out = run_script(mc, *args)
        assert rc == 0 and "non è quello della colonna I" in out, f"Monte Carlo {args}: manca l'avviso"
    xlsx = _excel_test([("Casa", 40000, 6, "Moderato", 5000, "importante", "terminante")])
    try:
        wb = load_workbook(xlsx)
        wb["Parametri"]["A7"] = None
        wb["Parametri"]["B7"] = None
        wb.save(xlsx)
        for script in (gp, mc, asset("it-pianificazione-obiettivi", "allocatore-capitale.py")):
            rc, out = run_script(script, xlsx)
            assert rc == 0, out[-400:]
            assert "prima della 0.10.0" in out, f"{os.path.basename(script)}: nessun avviso sul pianificatore vecchio"
            assert "come la colonna I" not in out, f"{os.path.basename(script)}: «come la colonna I» su un foglio a quota fissa"
    finally:
        shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    assert "prima della 0.10.0" in _testo_obiettivi("SKILL.md"), "SKILL.md: manca la nota sui pianificatori vecchi"
    assert "ATTERRAGGIO_AZ" in _testo_obiettivi("references/glide-path.md").split("## Il glide path nel versamento")[0] \
        and "colonna I" in _testo_obiettivi("references/glide-path.md").split("## Il glide path nel versamento")[0], \
        "glide-path.md: dire che atterraggio e curvatura diversi dal foglio fanno divergere dalla colonna I"


@test("montecarlo-obiettivi: il rendimento del profilo è la media aritmetica annua (volatilità → 0 ⇒ 50%)")
def t_convenzione_media():
    # F041: il modello normale centra la MEDIA del capitale sul versamento deterministico; con
    # volatilità quasi nulla la probabilità tende al 50%. Il mutante «media mensile = rend/12»
    # (media aritmetica annua > rend) la porta al 100%.
    import numpy as np
    mod = _modulo_obiettivi("montecarlo-obiettivi.py")
    mod.MODELLO = "normale"
    np.random.seed(3)
    target, rend, anni = 100_000.0, 0.05, 10
    pmt = mod.pmt_deterministico(target, 0.0, rend, anni)
    p, _ = mod.prob_successo(target, 0.0, pmt, rend, 0.001, anni)
    assert 0.45 <= p <= 0.55, f"volatilità 0,1%: probabilità {p:.1%}, attesa ~50%"
    # le etichette dicevano «geometrica»; serve la nota σ²/2 per tarare su rendimenti annualizzati
    assert "(geometrica)" not in open(asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"),
                                      encoding="utf-8").read(), "etichetta «(geometrica)» nello script"
    mc = _testo_obiettivi("references/monte-carlo.md")
    assert "composizione geometrica" not in mc, "monte-carlo.md: «composizione geometrica»"
    for rel in ("references/monte-carlo.md", "references/mercati-e-correlazione.md"):
        assert "σ²/2" in _testo_obiettivi(rel), f"{rel}: manca la conversione σ²/2 tra media aritmetica e annualizzata"


@test("obiettivi: la ricetta DMS → profilo è svolta sui Parametri, confronti tra grandezze omogenee")
def t_ricetta_dms():
    # Revisione T8: la ricetta aggiungeva σ²/2 al lordo e poi tassava, mentre il Monte Carlo applica la
    # volatilità del profilo al rendimento NETTO (e il suo esempio 4,65% ↔ 5,5% era già netto); mancava
    # un esempio svolto fino a un default del foglio, che sta sopra la ricetta mondiale; monte-carlo.md
    # chiamava «prudenti» i profili confrontando il loro netto con la media storica lorda.
    import numpy as np
    foglio, profili, _ = _obiettivi_foglio()
    infl = foglio["B6"].value
    al = _modulo_obiettivi("asset-location-optimizer.py")
    costi, tassa = al.TER_DOSSIER + al.BOLLO, al.TASSA_PLUS_DOSSIER
    rend, vol, _ = profili["Azionario totale"]
    dms = 0.052  # azionario mondiale reale annualizzato 1900-2024 (DMS), citato nella reference
    profilo_dms = ((1 + dms) * (1 + infl) - 1 - costi) * (1 - tassa) + vol ** 2 / 2  # nominale → netto → +σ²/2
    reale_foglio = (1 + (rend - vol ** 2 / 2) / (1 - tassa) + costi) / (1 + infl) - 1  # il default, a ritroso
    mer = " ".join(_testo_obiettivi("references/mercati-e-correlazione.md").split())
    ricetta = mer.split("**Come tradurre un dato DMS")[1].split("**(b)")[0]
    assert ricetta.index("(1 − 26%)") < ricetta.index("aggiungere σ²/2"), \
        "ricetta DMS: σ²/2 si aggiunge al netto, dopo costi e imposte (il Monte Carlo applica σ al netto)"
    for cifra in (f"{_pct_it(profilo_dms * 100)}%", f"{_pct_it(rend * 100)}%", f"{_pct_it(reale_foglio * 100)}% reale"):
        assert cifra in ricetta, f"mercati-e-correlazione.md: manca {cifra} nell'esempio svolto DMS → Azionario totale"
    crescita, vol_c, peso_c = profili["Crescita"]
    m = re.search(r"(\d+,\d+)% netto annualizzato ↔ " + re.escape(_pct_it(crescita * 100)) + "% del foglio", ricetta)
    assert m and abs(float(m.group(1).replace(",", ".")) - (crescita - vol_c ** 2 / 2) * 100) <= 0.006, \
        "ricetta DMS: l'esempio σ²/2 della Crescita va sul netto annualizzato (rendimento − σ²/2)"
    motore = _modulo_obiettivi("motore_bootstrap.py")
    az, ob = motore.carica_storico("rendimenti-storici.csv")
    storico = float(np.mean(peso_c * az + (1 - peso_c) * ob)) * 100
    mc = " ".join(_testo_obiettivi("references/monte-carlo.md").split())
    for cifra in (f"~{_pct_it(storico)}% nominale lordo", f"{_pct_it(al.lordo_da_netto(crescita) * 100)}% lordo"):
        assert cifra in mc, f"monte-carlo.md (ricentratura): manca {cifra}, il confronto va fatto lordo con lordo"


@test("montecarlo-obiettivi (bootstrap): la coda congiunta di due essenziali supera il prodotto delle marginali")
def t_rischio_sistemico():
    # F022: simula_congiunta non era mai esercitata dai test. Con un mercato condiviso la
    # probabilità che falliscano insieme è più alta che sotto indipendenza.
    xlsx = _excel_test([("Casa", 60000, 8, "Moderato", 0, "essenziale", "terminante"),
                        ("Università", 80000, 12, "Crescita", 0, "essenziale", "terminante")])
    rc, out = run_script(asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"), xlsx,
                         "--set", "MODELLO=bootstrap")
    shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)
    assert rc == 0, out[-400:]
    sez = out.split("RISCHIO SISTEMICO")[-1]
    assert "essenziali): Casa, Università" in sez, sez[:300]
    coda = re.search(r"falliscono.*?prodotto\):\s*([\d,]+)%.*?congiunta\):\s*([\d,]+)%", sez, re.S)
    insieme = re.search(r"riescono insieme.*?prodotto\):\s*([\d,]+)%.*?congiunta\):\s*([\d,]+)%", sez, re.S)
    ind, cong = _da_it(coda.group(1)), _da_it(coda.group(2))
    assert cong >= ind + 2, f"P(falliscono insieme): congiunta {cong}% vs prodotto {ind}%: shock non condivisi?"
    assert _da_it(insieme.group(2)) >= _da_it(insieme.group(1)), "P(riescono insieme) congiunta < prodotto"


@test("montecarlo-obiettivi: accanto all'OPZIONE la probabilità sul percorso glide, più bassa che a quota fissa")
def t_opzione_glide():
    # F017 (resto): la OPZIONE a quota fissa (58%) ignorava il de-risking che it-rebalance eseguirà
    # (sul glide path la stessa leva dà molto meno). Si mostrano entrambi i numeri, in entrambi i modelli.
    xlsx = _excel_test([("Barca", 400000, 20, "Moderato", 0, "desiderabile", "terminante")])
    try:
        for args in ((), ("--set", "GLIDE=False")):
            rc, out = run_script(asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"), xlsx, *args)
            assert rc == 0, out[-400:]
            testo = _sezioni_obiettivi(out)["Barca"][2]
            assert "OPZIONE:" in testo, f"{args}: OPZIONE attesa per un desiderabile a 20 anni"
            glide = re.search(r"percorso glide[^\n]*dal (\d+)% al (\d+)%", testo)
            fissa = re.search(r"quota fissa[^\n]*dal (\d+)% al (\d+)%", testo)
            assert glide and fissa, f"{args}: mancano le probabilità glide e quota fissa accanto all'OPZIONE"
            assert int(glide.group(2)) < int(fissa.group(2)), "col de-risking la leva non può rendere di più"
            assert int(glide.group(2)) > int(glide.group(1)), "OPZIONE senza guadagno sul percorso glide"
    finally:
        shutil.rmtree(os.path.dirname(xlsx), ignore_errors=True)


@test("asset-location: l'ottimizzatore riceve un rendimento lordo esplicito, «netto» definito ovunque")
def t_rendimento_lordo():
    # F018: i rendimenti dei profili sono netti di costi e imposte; l'ottimizzatore toglie TER,
    # bollo e imposte da sé, quindi deve partire dal LORDO (prima: 5,5% netto trattato come lordo,
    # verdetto PIR ribaltato). Conversione dichiarata: lordo = netto / (1 − 26%) + TER + bollo.
    mod = _modulo_obiettivi("asset-location-optimizer.py")
    lordo = 0.055 / (1 - 0.26) + 0.002 + 0.002
    def dossier(r, gia, anni):  # replica: ETF ad accumulazione, plus 26% alla fine
        m = gia * (1 + r - 0.002 - 0.002) ** anni
        return m - max(0.0, m - gia) * 0.26
    def pareggio(r):
        lo, hi = 0.0, 0.10
        for _ in range(60):  # TER del PIR a cui PIR (plus esente, bollo) = dossier
            mid = (lo + hi) / 2
            lo, hi = (mid, hi) if 10_000 * (1 + r - mid - 0.002) ** 15 > dossier(r, 10_000, 15) else (lo, mid)
        return lo
    lo = pareggio(lordo)
    rc, out = run_script(asset("it-pianificazione-obiettivi", "asset-location-optimizer.py"))
    assert rc == 0, out[-400:]
    m = re.search(r"TER di pareggio del PIR: ([\d,]+)%", out)
    assert m and abs(_da_it(m.group(1)) - lo * 100) <= 0.006, \
        f"TER di pareggio {m and m.group(1)}%, atteso {lo*100:.2f}% col lordo {lordo:.2%}"
    assert abs(mod.lordo_da_netto(0.055) - lordo) < 1e-12
    assert re.search(r"lordo[^\n]*7[.,]8%", out), "l'output deve dichiarare il rendimento lordo usato"
    # Revisione T8: la conversione tratta l'imposta come annua, quindi dà il lordo più alto compatibile
    # col netto e favorisce i contenitori esenti. Con l'imposta alla vendita (ETF ad accumulazione) il
    # lordo equivalente è più basso e il pareggio del PIR scende: verso e variante vanno detti.
    differito = ((1.055 ** 15 - 0.26) / 0.74) ** (1 / 15) - 1 + 0.004
    m = re.search(r"lordo equivalente a 15 anni è ([\d,]+)%.*?pareggio scende a ([\d,]+)%", out, re.S)
    assert m, "l'output deve dare la variante con l'imposta alla vendita (lordo equivalente e pareggio)"
    assert abs(_da_it(m.group(1)) - differito * 100) <= 0.006, f"lordo con imposta differita {m.group(1)}%, atteso {differito:.2%}"
    assert abs(_da_it(m.group(2)) - pareggio(differito) * 100) <= 0.006, \
        f"pareggio con imposta differita {m.group(2)}%, atteso {pareggio(differito):.2%}"
    assert "Prudente:" not in (mod.lordo_da_netto.__doc__ or ""), "lordo_da_netto: nell'ottimizzatore non è prudente"
    skill = _testo_obiettivi("SKILL.md")
    for cifra in (f"{_pct_it(pareggio(differito) * 100, 2)}%", "favorisce i contenitori esenti"):
        assert cifra in skill, f"SKILL.md (asset location): manca «{cifra}»"
    rc, out = run_script(asset("it-pianificazione-obiettivi", "asset-location-optimizer.py"),
                         "--set", "RENDIMENTO_LORDO=0.07")
    assert rc == 0 and "@ 7,0%" in out, "RENDIMENTO_LORDO esplicito non usato nello scenario pensione"
    par, _, _ = _obiettivi_foglio()
    note = " ".join(str(par[f"A{r}"].value or "") for r in range(1, 30))
    formula = "(lordo − costi) × (1 − 26%)"
    assert "netti di costi e imposte" in str(par["A8"].value), f"Parametri!A8: {par['A8'].value!r}"
    for dove, testo in (("Parametri", note), ("SKILL.md", _testo_obiettivi("SKILL.md")),
                        ("matematica-accumulo.md", _testo_obiettivi("references/matematica-accumulo.md"))):
        assert formula in testo, f"{dove}: manca la definizione di «netto» ({formula})"


@test("obiettivi: le cifre citate in SKILL.md e nelle reference sono l'output sui default")
def t_cifre_documentate():
    # F043: gli esempi usavano 3,5%/5% (€487/€366/€1.353) contro il 4,0%/5,5% dell'Excel, e le
    # cifre dell'asset location venivano da un «test sui dati reali» non riproducibile.
    from openpyxl import load_workbook
    skill = _testo_obiettivi("SKILL.md")
    mat = _testo_obiettivi("references/matematica-accumulo.md")
    mc = _testo_obiettivi("references/monte-carlo.md")
    dec = _testo_obiettivi("references/decumulo.md")
    mer = _testo_obiettivi("references/mercati-e-correlazione.md")
    # piano di accumulo: i versamenti del foglio d'esempio
    foglio, profili, obiettivi = _obiettivi_foglio()
    totale = sum(g["versamento"] for g in obiettivi)
    for v in [g["versamento"] for g in obiettivi] + [totale]:
        assert _euro_it(v) in mat, f"matematica-accumulo.md non cita {_euro_it(v)} del foglio d'esempio"
    for vecchio in ("€487", "€366", "€1.353"):
        assert vecchio not in mat, f"matematica-accumulo.md: {vecchio} non è l'output dei Parametri"
    glide = {g["nome"]: _euro_it(g["versamento"]) for g in obiettivi}
    for rel, testo, nomi in (("references/glide-path.md", _testo_obiettivi("references/glide-path.md"),
                              ("Anticipo casa", "Lungo termine / libertà")),
                             ("SKILL.md", skill, ("Lungo termine / libertà",))):
        for nome in nomi:
            assert glide[nome] in testo, f"{rel} non cita il versamento col glide path di {nome} ({glide[nome]})"
    tmp = tempfile.mkdtemp()
    try:
        rc, out = run_script(asset("it-pianificazione-obiettivi", "allocatore-capitale.py"),
                             "--json", os.path.join(tmp, "a.json"))
        assert rc == 0, out[-400:]
        righe = {r["nome"]: r for r in json.load(open(os.path.join(tmp, "a.json"), encoding="utf-8"))["obiettivi"]}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    richiesto = _euro_it(righe["Lungo termine / libertà"]["richiesto"])
    assert richiesto in _testo_obiettivi("references/allocazione-capitale.md"), \
        f"allocazione-capitale.md non cita il capitale richiesto col glide path ({richiesto})"
    # Revisione T8: anche le cifre a quota fissa e quelle che dipendono dai default o dalla serie
    # (aggiornata ogni gennaio) devono essere l'output, non numeri scritti a mano.
    infl = foglio["B6"].value
    per_nome = {g["nome"]: g for g in obiettivi}
    gp = _testo_obiettivi("references/glide-path.md")
    alloc = _testo_obiettivi("references/allocazione-capitale.md")
    fissi = {}
    for nome in ("Anticipo casa", "Lungo termine / libertà"):
        g = per_nome[nome]
        rend = profili[g["profilo"]][0]
        target_nom = g["target"] * (1 + infl) ** g["anni"]
        fissi[nome], _ = _pmt_anno_per_anno(target_nom, g["gia"] or 0, rend, rend, g["anni"], 7)
        for rel, testo in (("references/glide-path.md", gp), ("references/matematica-accumulo.md", mat)):
            assert _euro_it(fissi[nome]) in testo, \
                f"{rel}: manca il versamento a quota fissa di {nome} ({_euro_it(fissi[nome])})"
    g = per_nome["Lungo termine / libertà"]
    rend, fisso = profili[g["profilo"]][0], fissi["Lungo termine / libertà"]
    target_nom = g["target"] * (1 + infl) ** g["anni"]
    assert _euro_it(fisso) in skill, f"SKILL.md: manca il versamento a quota fissa a 20 anni ({_euro_it(fisso)})"
    aumento = f"(+{round((g['versamento'] / fisso - 1) * 100)}%)"
    assert aumento in gp, f"glide-path.md: il glide path alza il versamento a 20 anni di {aumento}"
    rend_l = min(profili.values(), key=lambda p: p[2])[0]
    cap = 0.0  # capitale finale col versamento a quota fissa, seguendo il glide path
    for r in rend_l + (rend - rend_l) * _modulo_obiettivi("motore_bootstrap.py").frazioni_glide(g["anni"], 7):
        cap = cap * (1 + r) + fisso * r / ((1 + r) ** (1 / 12) - 1)
    quota = f"circa il {round(cap / target_nom * 100)}% del target"
    assert quota in gp, f"glide-path.md: a quota fissa il capitale si ferma a «{quota}»"
    pv_fisso = g["target"] * ((1 + infl) / (1 + rend)) ** g["anni"]
    for rel, testo in (("references/allocazione-capitale.md", alloc), ("SKILL.md", skill)):
        for cifra in (_euro_it(pv_fisso), richiesto):
            assert cifra in testo, f"{rel}: manca {cifra} (present value a rendimento pieno e col glide path)"
    aumento = f"(+{round((righe[g['nome']]['richiesto'] / pv_fisso - 1) * 100)}%)"
    assert aumento in alloc, f"allocazione-capitale.md: il glide path alza il capitale richiesto di {aumento}"
    # tabella dei profili di monte-carlo.md = Parametri (sei profili)
    righe = re.findall(r"^\| ([A-Z][a-z ]+) \| ([\d,]+)% \| ([\d,]+)% \|", mc, re.M)
    doc = {n: (float(r.replace(",", ".")) / 100, float(v.replace(",", ".")) / 100) for n, r, v in righe}
    assert set(doc) == set(profili), f"monte-carlo.md elenca {sorted(doc)}, il foglio {sorted(profili)}"
    for n, (r, v, _) in profili.items():
        assert abs(doc[n][0] - r) < 1e-9 and abs(doc[n][1] - v) < 1e-9, f"monte-carlo.md, {n}: {doc[n]} ≠ {(r, v)}"
    # Monte Carlo sui default
    rc, out = run_script(asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"))
    assert rc == 0, out[-400:]
    sez = _sezioni_obiettivi(out)
    casa = sez["Anticipo casa"][2]
    pmt = re.search(r"Versamento deterministico:\s*(€[\d.]+)/mese", casa).group(1)
    p = re.search(r"questo versamento: (\d+)%", casa).group(1)
    pmt85 = re.search(r"confidenza:\s*(€[\d.]+)/mese", casa).group(1)
    frase = f"con {pmt}/mese hai il {p}% di arrivarci; per l'85% te ne servono {pmt85}"
    assert frase in mc, f"monte-carlo.md deve citare l'output sui default: «{frase}»"
    prob = _prob_obiettivi(out)
    assert f"solo il ~{prob['Lungo termine / libertà']}% di probabilità" in mc, "monte-carlo.md: probabilità a 20 anni"
    assert f"vicino al {prob['Fondo di emergenza']}%" in mc, "monte-carlo.md: probabilità del fondo di emergenza"
    # asset location sui default
    rc, out = run_script(asset("it-pianificazione-obiettivi", "asset-location-optimizer.py"))
    assert rc == 0, out[-400:]
    vant = re.search(r"VANTAGGIO FONDO PENSIONE: (€[\d.]+)\s+\(\+(\d+)% vs dossier\)", out)
    pir = re.search(r"(VANTAGGIO|SVANTAGGIO) PIR: (-?€[\d.]+)", out)
    par = re.search(r"TER di pareggio del PIR: ([\d,]+)%", out)
    for cifra in (vant.group(1), f"+{vant.group(2)}% vs dossier", pir.group(2).replace("-", "−"),
                  par.group(1).replace(".", ",") + "%"):
        assert cifra in skill, f"SKILL.md non cita {cifra} dell'ottimizzatore sui default"
    al = _modulo_obiettivi("asset-location-optimizer.py")
    deduzione = _euro_it(al.CONTRIBUTO_ANNUO_FP * al.ALIQUOTA_IRPEF_MARGINALE * al.ANNI_A_PENSIONE)
    reinvestita = re.search(r"Driver: la deduzione IRPEF \((€[\d.]+)\)", out).group(1)
    for cifra in (deduzione, reinvestita):
        assert cifra in skill, f"SKILL.md non cita {cifra} (risparmio IRPEF e suo reinvestimento)"
    # decumulo sui default
    rc, out = run_script(asset("it-pianificazione-obiettivi", "decumulo-obiettivi.py"))
    assert rc == 0, out[-400:]
    sens = dict(re.findall(r"^\s+€\s*[\d.]+\s+([\d,]+)%\s+([\d,]+)%$", out, re.M))
    for tasso in ("3,0", "4,0", "5,0", "6,0"):
        assert f"{sens[tasso]}% al {tasso[0]}%" in dec, \
            f"decumulo.md: manca «{sens[tasso]}% al {tasso[0]}%» (output sui default)"
    anno = re.search(r"peggiori rendimenti PRIMA:\s+ESAURITO all'anno (\d+)", out).group(1)
    assert f"all'anno {anno}" in dec, f"decumulo.md: la sequenza peggiore esaurisce all'anno {anno}"
    iniziale = _da_it(re.search(r"capitale €([\d.]+)", out).group(1))
    finale = _da_it(re.search(r"migliori rendimenti PRIMA:\s+capitale finale €([\d.]+)", out).group(1))
    assert 1.75 <= finale / iniziale <= 2.25 and "circa il doppio" in dec, \
        f"decumulo.md: «circa il doppio» con i migliori prima, l'output dà {finale / iniziale:.2f} volte"
    # rendimento reale storico dell'azionario USA (dalla serie, deflazionata con la CPI)
    motore = _modulo_obiettivi("motore_bootstrap.py")
    az, _ = motore.carica_storico("rendimenti-storici.csv")
    cpi = motore.carica_inflazione("rendimenti-storici.csv")
    import numpy as np
    reale = (np.prod((1 + az) / (1 + cpi)) ** (1 / len(az)) - 1) * 100
    for rel, testo in (("mercati-e-correlazione.md", mer), ("monte-carlo.md", mc)):
        assert f"~{_pct_it(reale)}%" in testo, f"{rel}: l'azionario USA ha reso ~{_pct_it(reale)}% reale annuo"
    costante = (np.prod((1 + az) / 1.02) ** (1 / len(az)) - 1) * 100
    frase = f"inflazione costante del 2% l'azionario USA renderebbe ~{_pct_it(costante)}% reale invece di ~{_pct_it(reale)}%"
    assert frase in mer, f"mercati-e-correlazione.md: «{frase}»"
    i = 1974 - 1928
    meno = lambda x: x.replace("-", "−")
    frase = (f"1974: {meno(_pct_it(az[i] * 100))}% nominale con inflazione al {_pct_it(cpi[i] * 100)}%, "
             f"{meno(str(round(((1 + az[i]) / (1 + cpi[i]) - 1) * 100)))}% reale")
    assert frase in mer, f"mercati-e-correlazione.md: «{frase}» (dalla serie)"
    # numero di formule del pianificatore
    wb = load_workbook(asset("it-pianificazione-obiettivi", "pianificatore-obiettivi.xlsx"))
    n = sum(1 for ws in wb.worksheets for riga in ws.iter_rows() for c in riga
            if isinstance(c.value, str) and c.value.startswith("="))
    citati = re.findall(r"(\d+) formule", skill)
    assert citati and all(int(c) == n for c in citati), f"SKILL.md cita {citati} formule, il foglio ne ha {n}"


@test("obiettivi: nessun rimando superato nelle reference (decumulo, correlazione, glide path, --json)")
def t_rimandi_superati():
    # F043: rimandi a «roadmap» per cose già fatte e limiti non più veri.
    # Revisione T8: l'allocatore non usa pmt_glide ma crescita_glide; il flag Tipo cambia anche la
    # vista per obiettivo; la finestra si legge da Parametri!B7; i profili non vanno confrontati
    # (netti) con la media storica (lorda) come «ipotesi prudenti»; SKILL.md descriveva l'allocatore
    # a rendimento pieno.
    vietati = {
        "references/monte-carlo.md": ("roadmap di `it-rebalance`", "roadmap di it-rebalance",
                                      "Nessuna correlazione tra obiettivi", "o glide path**",
                                      "7 anni come in `glidepath-obiettivi.py`", "ipotesi prudenti dei profili"),
        "references/glide-path.md": ("roadmap di it-rebalance", "sovrastima leggermente",
                                     "e l'allocatore (per il capitale richiesto oggi)",
                                     "La vista **per obiettivo** non dipende dal flag"),
        "references/matematica-accumulo.md": ("glide path e allocatore usano la stessa",),
        "SKILL.md": ("correlazione sistemica tra obiettivi nel Monte Carlo", "alimenta il tracker e il glide path",
                     "il capitale che al rendimento del profilo cresce fino al target",
                     "finestra di de-risking di 7 anni"),
        "assets/montecarlo-obiettivi.py": ("Parametri!A10:C14",),
    }
    for rel, frasi in vietati.items():
        testo = _testo_obiettivi(rel)
        for frase in frasi:
            assert frase not in testo, f"{rel}: «{frase}»"
    gp = _testo_obiettivi("references/glide-path.md")
    assert re.search(r"circa (il )?10%", gp), "glide-path.md: l'ordine di grandezza del de-risking ignorato (~10%)"
    # Step 4 di SKILL.md: lo pseudocodice era solo a quota fissa (chi lo segue in chat sottostima il
    # versamento dei terminanti); il report deve dichiarare le convenzioni.
    skill = _testo_obiettivi("SKILL.md")
    passo4 = skill.split("### Step 4")[1].split("### Step 5")[0]
    for formula in (r"m\s*=\s*min\(finestra, n\) − 1", r"r2\s*=\s*\(r \+ r_atterraggio\) / 2",
                    r"Quota fissa[^\n]*permanenti"):
        assert re.search(formula, passo4), f"SKILL.md, Step 4: manca «{formula}» (versamento in due fasi)"
    sezione3 = skill.split("### 3. Piano di accumulo")[1].split("### 4.")[0]
    assert "glide path" in sezione3 and "netto" in sezione3, "SKILL.md, report: dichiarare le convenzioni"


@test("serie storica: nominale in USD dichiarata, il decumulo deflaziona anno per anno con la CPI USA")
def t_serie_storica():
    # Critica dell'audit: la serie è nominale in dollari ma il motore la chiamava «reale», e il
    # decumulo la deflazionava con un 2% costante (cancellando la stagflazione degli anni '70).
    csv_path = asset("it-pianificazione-obiettivi", "rendimenti-storici.csv")
    righe = open(csv_path, encoding="utf-8").read().splitlines()
    commenti = " ".join(r for r in righe if r.startswith("#")).lower()
    assert "nominali" in commenti and "dollari" in commenti, "intestazione: dichiarare rendimenti nominali in dollari"
    assert "anno,azionario,obbligazionario,inflazione" in righe, "manca la colonna inflazione"
    dati = {int(r.split(",")[0]): [float(x) for x in r.split(",")[1:]] for r in righe
            if r and r[0].isdigit()}
    assert sorted(dati) == list(range(1928, 2026)) and all(len(v) == 3 for v in dati.values())
    # CPI-U dicembre su dicembre pubblicata dal BLS (tabella storica; release CPI dicembre 2025)
    for anno, bls in ((1946, 18.1), (1974, 12.3), (1979, 13.3), (1980, 12.5), (2008, 0.1), (2022, 6.5), (2025, 2.7)):
        assert abs(dati[anno][2] - bls) <= 0.05 + 1e-9, f"inflazione {anno}: {dati[anno][2]} vs BLS {bls}"
    motore = open(asset("it-pianificazione-obiettivi", "motore_bootstrap.py"), encoding="utf-8").read()
    assert "serie annuali reali" not in motore, "motore_bootstrap: la serie non è reale"
    mod = _modulo_obiettivi("decumulo-obiettivi.py")
    mod.PESO_AZIONARIO, mod.INFLAZIONE = 1.0, None
    reali = mod.carica_rendimenti_reali(0.0)
    atteso = (1 + dati[1974][0] / 100) / (1 + dati[1974][2] / 100) - 1
    assert abs(reali[1974 - 1928] - atteso) < 1e-12, f"1974 reale {reali[1974 - 1928]:.4f}, atteso {atteso:.4f}"
    mod.INFLAZIONE = 0.02  # override esplicito: inflazione costante come prima
    assert abs(mod.carica_rendimenti_reali(0.0)[1974 - 1928] - ((1 + dati[1974][0] / 100) / 1.02 - 1)) < 1e-12


# ══════════════════════ motore bootstrap condiviso ══════════════════════

@test("sincronia: motore_bootstrap.py e rendimenti-storici.csv identici tra le due skill")
def t_sincronia():
    import hashlib
    def h(p):
        return hashlib.md5(open(p, "rb").read()).hexdigest()
    for nome in ("motore_bootstrap.py", "rendimenti-storici.csv"):
        a = h(asset("it-pianificazione", nome))
        b = h(asset("it-pianificazione-obiettivi", nome))
        assert a == b, f"{nome} DIVERGE tra le skill: riallineare le copie"


@test("motore bootstrap: blocchi circolari, valori dalla serie, riproducibile col seed")
def t_motore():
    import numpy as np
    m = carica_modulo(asset("it-pianificazione", "motore_bootstrap.py"))
    port = np.array([0.10, -0.05, 0.03, 0.20, -0.15, 0.07, 0.01])
    np.random.seed(7)
    idx = m.indici_bootstrap(len(port), n_sim=50, anni=12, blocco=5)
    assert idx.shape == (50, 12) and idx.min() >= 0 and idx.max() < len(port)
    # dentro un blocco gli indici sono consecutivi (circolari): idx[:,k+1] == (idx[:,k]+1) % n
    for k in range(4):  # il primo blocco copre le colonne 0..4
        assert (idx[:, k + 1] == (idx[:, k] + 1) % len(port)).all(), "blocco non consecutivo"
    np.random.seed(7)
    idx2 = m.indici_bootstrap(len(port), 50, 12, 5)
    assert (idx == idx2).all(), "stesso seed → indici diversi (non riproducibile)"
    np.random.seed(7)
    rend = m.blocchi_bootstrap(port, 50, 12, 5)
    assert (rend == port[idx]).all(), "blocchi_bootstrap incoerente con indici_bootstrap"
    assert set(np.unique(rend)) <= set(port), "il bootstrap ha inventato valori"


# ══════════════════════ it-revisione (generatore IPS) ══════════════════════

@test("genera-ips: documento completo (sezioni, tabelle, nota di rischio calibrata)")
def t_genera_ips():
    try:
        from docx import Document
    except ImportError:
        raise AssertionError("python-docx non installato: pip install python-docx")
    tmp = tempfile.mkdtemp()
    out = os.path.join(tmp, "ips.docx")
    rc, log = run_script(asset("it-revisione", "genera-ips.py"), out)
    assert rc == 0, log[-400:]
    doc = Document(out)
    testo = "\n".join(p.text for p in doc.paragraphs)
    shutil.rmtree(tmp, ignore_errors=True)
    assert len(doc.tables) >= 3, f"solo {len(doc.tables)} tabelle"
    for chiave in ("INVESTMENT POLICY STATEMENT", "Obiettivi", "Glide path",
                   "ribilanciamento", "Appendice A", "Onestà sul rischio"):
        assert chiave in testo, f"sezione mancante: {chiave}"
    # la nota di rischio è CALCOLATA dalla quota azionaria (57% nell'esempio → ~28%)
    m = re.search(r"perdere circa il (\d+)%", testo)
    assert m and abs(int(m.group(1)) - 28) <= 1, "nota di rischio non calibrata sull'allocazione"
    # T14: titolo e lingua del documento (prima: titolo vuoto e controllo ortografico in inglese)
    from docx.oxml.ns import qn
    assert doc.core_properties.title.startswith("Investment Policy Statement — "), doc.core_properties.title
    assert doc.core_properties.language == "it-IT", doc.core_properties.language
    lang = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:lang"))
    assert lang is not None and lang.get(qn("w:val")) == "it-IT", "docDefaults senza w:lang it-IT"


@test("genera-ips: quota azionaria sommata su tutte le classi, mese in italiano in copertina")
def t_ips_quota():
    # REGRESSIONE F060: contava solo la prima classe che inizia per "azion" e scriveva il mese in inglese.
    from docx import Document
    mod = carica_modulo(asset("it-revisione", "genera-ips.py"))
    mesi = ("gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio",
            "agosto", "settembre", "ottobre", "novembre", "dicembre")
    tmp = tempfile.mkdtemp()
    try:
        for allocazione, attesa in (
                ([("Azionario sviluppati", 45_000, 45.0), ("Azionario emergenti", 11_700, 11.7),
                  ("Obbligazionario", 31_300, 31.3), ("Liquidità", 12_000, 12.0)], 57),
                ([("ETF azionario globale", 60_000, 60.0), ("Obbligazioni", 40_000, 40.0)], 60)):
            out = os.path.join(tmp, "ips.docx")
            # i dati dell'esempio con l'allocazione del caso (genera() con un dict usa solo quel dict)
            mod.genera(dict(mod.DATI, allocazione=allocazione, capitale=100_000,
                            obiettivi=[o[:4] + (100_000 * o[4] / mod.DATI["capitale"],) + o[5:]
                                       for o in mod.DATI["obiettivi"]]), out)
            testo = "\n".join(p.text for p in Document(out).paragraphs)
            m = re.search(r"allocazione al (\d+)% azionario", testo)
            assert m and int(m.group(1)) == attesa, f"quota azionaria {m.group(1) if m else '—'}%, attesa {attesa}%"
        assert any(f"{mese.capitalize()} " in testo for mese in mesi), "mese della copertina non in italiano"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ══════════════════════ it-tax-optimization (T1-T3) ══════════════════════

TAX = os.path.join(SKILLS, "it-tax-optimization")


def _testi_regole():
    """{percorso relativo: testo} di skill, comandi, README e descrizione del plugin.json,
    con il grassetto Markdown tolto. CHANGELOG.md racconta le versioni passate e resta
    fuori: gli errata li aggiunge chi rilascia, non questi controlli."""
    testi = {}
    for base in (SKILLS, os.path.join(ROOT, "commands")):
        for cartella, _, files in os.walk(base):
            for f in files:
                if f.endswith((".md", ".py")):
                    p = os.path.join(cartella, f)
                    testi[os.path.relpath(p, ROOT).replace(os.sep, "/")] = open(p, encoding="utf-8").read()
    testi["README.md"] = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    plugin = json.load(open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8"))
    testi[".claude-plugin/plugin.json"] = plugin["description"]
    return {k: v.replace("*", "") for k, v in testi.items()}


def _blocchi(testo):
    """Paragrafi, voci di elenco e righe di tabella."""
    return [b for b in re.split(r"\n\s*\n|\n(?=\s*[-*|] )|\n(?=\|)", testo) if b.strip()]


def _frasi(testo):
    """Spezza in frasi (punto/punto e virgola seguiti da maiuscola), righe e celle di tabella.
    Le abbreviazioni come «art. 67» o «c. 1-bis» restano unite perché segue una cifra."""
    return [f for f in re.split(r"(?<=[.!?;])\s+(?=[A-ZÀ-Ý«(`\[])|\n|\|", testo) if f.strip()]


_LIFO = re.compile(r"\bL\.?I\.?F\.?O\b\.?", re.I)
_REGIME = re.compile(r"amministrat|dichiarativ", re.I)
_CONFINE = re.compile(r"[,;:]|\bmentre\b|\binvece\b", re.I)
# LIFO legato subito al dichiarativo: «LIFO in dichiarativo», «LIFO (dichiarativo)», «LIFO solo in
# dichiarativo», «il LIFO vale solo in dichiarativo»; oppure «in dichiarativo è il LIFO».
_LIFO_DICH_DOPO = re.compile(r"\s*\)?\s*\(?\s*(?:vale\s+)?(?:solo\s+)?(?:in|nel|per\s+il)?\s*(?:regime\s+)?dichiarativ", re.I)
_LIFO_DICH_PRIMA = re.compile(r"dichiarativ\w*[^.;:,()]{0,25}$", re.I)
_ELENCO_USA = re.compile(r"HIFO\s*/\s*L\.?I\.?F\.?O|L\.?I\.?F\.?O\.?\s*/\s*Spec", re.I)


def _lifo_errato(frase):
    """True se la frase riferisce il LIFO all'amministrato, o lo nomina senza alcun regime fuori
    dall'elenco dei metodi USA (HIFO/LIFO/SpecID). Ogni occorrenza (anche «Lifo», «L.I.F.O.») si
    lega al regime della propria clausola (confini: , ; : «mentre» «invece»); se la clausola non
    ne nomina, al regime nominato più vicino prima di lei nella frase, o al primo dopo."""
    for m in _LIFO.finditer(frase):
        if _LIFO_DICH_DOPO.match(frase, m.end()) or _LIFO_DICH_PRIMA.search(frase, 0, m.start()):
            continue
        inizio = max([0] + [c.end() for c in _CONFINE.finditer(frase, 0, m.start())])
        dopo = _CONFINE.search(frase, m.end())
        clausola = frase[inizio:dopo.start() if dopo else len(frase)]
        if re.search(r"amministrat", clausola, re.I):
            return True
        if re.search(r"dichiarativ", clausola, re.I):
            continue
        prima = list(_REGIME.finditer(frase, 0, m.start()))
        regime = prima[-1] if prima else _REGIME.search(frase, m.end())
        if regime is None:
            if not _ELENCO_USA.search(frase[max(0, m.start() - 8):m.end() + 8]):
                return True
        elif regime.group(0).lower().startswith("amministrat"):
            return True
    return False


@test("T1: nessun LIFO riferito al regime amministrato (costo medio ponderato, art. 6 c. 4 D.Lgs. 461/97)")
def t_t1_lifo_amministrato():
    # REGRESSIONE F012: in amministrato la base è il costo medio ponderato; il LIFO
    # (art. 67 c. 1-bis TUIR) individua i titoli ceduti solo in dichiarativo. Ogni LIFO deve
    # quindi essere legato al dichiarativo (o all'elenco dei metodi USA), non all'amministrato.
    for frase in ("- ❌ Assumere HIFO o Specific ID per la scelta del lotto: in amministrato è sempre LIFO.",
                  "In regime amministrato la plus si calcola con il LIFO, mentre in dichiarativo si usa il costo medio ponderato.",
                  "- ❌ Ignorare il LIFO: la scelta del lotto NON è libera come negli USA, e influenza la plus realizzata.",
                  "LIFO obbligatorio in regime amministrato",
                  "In regime amministrato il metodo di valorizzazione è Lifo per default.",
                  "In amministrato i lotti si scaricano con il L.I.F.O."):
        assert _lifo_errato(frase), f"il controllo non riconosce: «{frase}»"
    for frase in ("Costo medio ponderato in amministrato, LIFO in dichiarativo: il lotto non si sceglie",
                  "costo medio ponderato (LIFO solo in dichiarativo)",
                  "la base è il costo medio ponderato (il LIFO vale solo in dichiarativo).",
                  "Cambiano il metodo del costo, che in dichiarativo è il LIFO e non il costo medio ponderato dell'amministrato",
                  "Scelta investitore (HIFO/LIFO/SpecID)"):
        assert not _lifo_errato(frase), f"falso positivo: «{frase}»"
    errori = []
    for rel, testo in _testi_regole().items():
        for frase in _frasi(testo):
            if _lifo_errato(frase):
                errori.append(f"{rel}: «{frase.strip()[:110]}»")
    assert not errori, "LIFO riferito all'amministrato o senza regime:\n  " + "\n  ".join(errori[:12])
    skill = open(os.path.join(TAX, "SKILL.md"), encoding="utf-8").read()
    assert re.search(r"costo medio ponderato[^\n]*art\. 6 c\. 4", skill), "SKILL.md: manca la regola del costo medio (art. 6 c. 4)"
    dich = open(os.path.join(TAX, "references", "regime-dichiarativo.md"), encoding="utf-8").read()
    assert re.search(r"LIFO[^\n]*art\. 67 c\. 1-bis", dich), "regime-dichiarativo.md: manca il LIFO con l'art. 67 c. 1-bis"
    assert "identiche all'amministrato" not in dich, "regime-dichiarativo.md: regole «identiche» all'amministrato, ma il metodo del costo cambia"
    # Cripto in dichiarativo: la legge non richiama la lettera c-sexies, ma la Circ. 30/E/2023
    # (par. 3.2.1) prescrive il LIFO: non è una prassi «per analogia» da concordare.
    metodo = open(os.path.join(TAX, "references", "criptovalute.md"), encoding="utf-8").read().split("Metodo di valorizzazione")[1].split("\n")[0]
    assert re.search(r"LIFO[^\n]*Circ\. 30/E/2023, par\. 3\.2\.1", metodo), "criptovalute.md: manca il LIFO della Circ. 30/E/2023 par. 3.2.1"
    assert "analogia" not in metodo and "da concordare" not in metodo, "criptovalute.md: il LIFO cripto presentato come scelta da concordare"


_NA = r"(?:non[ -]armonizzat|non[ -]UCITS|US-listed)"
_NA_COMPENSABILE = (
    re.compile(_NA + r"[^→\n|]*→\s*`?reddito[ _]divers", re.I),          # «ETF non armonizzato → `reddito_diverso`»
    re.compile(r"reddit[oi][ _]divers[oi]`?\s*\([^)\n]*" + _NA, re.I),  # «`reddito_diverso` (azioni, ETF non armonizzati)»
    re.compile(_NA + r"[^\n]{0,40}?(?<!non )(?<!mai )compensabil", re.I),  # «ETF non armonizzati (compensabili)»
    re.compile(_NA + r"[^\n→]*→\s*SÌ"),                                   # «plus su ... ETF non armonizzati → SÌ, consumato»
)


@test("T1: nessun ETF o fondo non armonizzato tra gli strumenti compensabili con lo zainetto")
def t_t1_non_armonizzati():
    # REGRESSIONE F003: la plus di un OICR non armonizzato è reddito di capitale (art. 44
    # c. 1 lett. g TUIR; art. 10-ter L. 77/1983), mai compensabile; solo la minus è reddito diverso.
    # Su tutti i testi, riga per riga (gli anti-pattern «❌» descrivono l'errore e si saltano).
    errori = []
    for rel, testo in _testi_regole().items():
        for riga in testo.splitlines():
            if "❌" not in riga and any(p.search(riga) for p in _NA_COMPENSABILE):
                errori.append(f"{rel}: «{riga.strip()[:110]}»")
    assert not errori, "non armonizzati tra i compensabili:\n  " + "\n  ".join(errori[:12])
    joint = open(os.path.join(SKILLS, "it-pianificazione", "references", "integrazione-fiscale-joint.md"),
                 encoding="utf-8").read()
    mappa = joint.split("```")[1].splitlines()
    col = next(r for r in mappa if "Imposta sostitutiva" in r).index("Imposta sostitutiva")
    for r in mappa:
        assert not re.search(_NA, r[col:], re.I), f"integrazione-fiscale-joint: non armonizzati nella colonna della sostitutiva: {r.strip()}"
    assert re.search(r"regime dichiarativo[^\n]*non armonizzat[^\n]*quadro RL|regime dichiarativo[^\n]*quadro RL[^\n]*non armonizzat", joint.replace("*", "")), \
        "integrazione-fiscale-joint: in dichiarativo manca l'eccezione IRPEF dei non armonizzati"
    etf = open(os.path.join(TAX, "references", "etf-trap-redditi-capitale.md"), encoding="utf-8").read()
    amici = etf.split('## Strumenti "amici"')[1].split('## Strumenti "nemici"')[0]
    assert "armonizzat" not in amici, "etf-trap: un OICR non armonizzato nella tabella «amici»"
    assert re.search(r'## Strumenti "nemici".*non armonizzat', etf, re.S), "etf-trap: non armonizzati assenti dai «nemici»"
    skill = open(os.path.join(TAX, "SKILL.md"), encoding="utf-8").read()
    riga = next(r for r in skill.splitlines() if r.startswith("| **Reddito diverso** (art. 67"))
    assert "armonizzat" not in riga, f"SKILL.md: riga reddito diverso: {riga[:120]}"
    assert "reddito_capitale_irpef" in skill, "SKILL.md: manca la categoria reddito_capitale_irpef per gli OICR extra UE/SEE"
    assert re.search(r"non armonizzat[^\n]*→\s*`reddito_capitale`[^\n]*`reddito_capitale_irpef`", skill), \
        "SKILL.md: la deduzione della categoria non manda i non armonizzati a reddito_capitale/_irpef"
    reb = open(os.path.join(SKILLS, "it-rebalance", "SKILL.md"), encoding="utf-8").read()
    assert re.search(r"`categoria_fiscale`[^\n]*`reddito_capitale`\s*\([^)\n]*non armonizzat", reb), \
        "it-rebalance: la riga categoria_fiscale non mette i non armonizzati sotto reddito_capitale"
    trovate = 0
    for r in reb.splitlines():
        if "reddito diverso" in r and "**può**" in r:
            trovate += 1
            assert "armonizzat" not in r, f"it-rebalance: {r[:120]}"
    assert trovate >= 1, "it-rebalance: manca la riga su cosa il reddito diverso può compensare"
    mb = open(os.path.join(SKILLS, "it-rebalance", "references", "ribilanciamento-multibroker.md"), encoding="utf-8").read()
    m = re.search(r"compensa solo i \*\*redditi diversi\*\* \(([^)]*)\)", mb)
    assert m and "armonizzat" not in m.group(1), "ribilanciamento-multibroker.md: elenco dei redditi diversi"


@test("T1: alla chiusura del rapporto le minus si certificano (art. 6 c. 5 D.Lgs. 461/97), non «si perdono»")
def t_t1_chiusura_rapporto():
    # REGRESSIONE F011: chiudendo il rapporto (o revocando l'opzione) l'intermediario certifica
    # le minus residue, utilizzabili su un altro rapporto con gli stessi intestatari o in
    # dichiarazione entro la scadenza originaria. Il grassetto è tolto prima del confronto.
    for rel, testo in _testi_regole().items():
        assert not re.search(r"si\s+perde\s+definitivamente", testo), f"{rel}: «si perde definitivamente»"
        assert "unico modo per consolidare" not in testo, f"{rel}: il dichiarativo presentato come unica via"
    for rel in ("references/multi-intermediario.md", "SKILL.md"):
        testo = open(os.path.join(TAX, rel), encoding="utf-8").read()
        assert "certificazione" in testo and "art. 6 c. 5" in testo, f"{rel}: manca la certificazione ex art. 6 c. 5"
    skill = open(os.path.join(TAX, "SKILL.md"), encoding="utf-8").read()
    assert "Pattern E" in skill, "SKILL.md: manca il Pattern E (consolidamento via certificazione)"
    multi = open(os.path.join(TAX, "references", "multi-intermediario.md"), encoding="utf-8").read()
    pratico = multi.split("## Caso pratico")[1].split("\n---\n")[0]   # senza il marcatore in fondo
    assert "Verificato" not in pratico, "caso pratico: il taglio include il marcatore"
    for lettera, chiave in (("a", "consumo in loco"), ("b", "certificazione"), ("c", "dichiarativo")):
        assert re.search(rf"^\*\*Scenario \({lettera}\)[^\n]*{chiave}", pratico, re.M), \
            f"caso pratico: manca l'intestazione dello scenario ({lettera}) «{chiave}»"
        assert re.search(rf"^\| \({lettera}\) ", pratico, re.M), f"caso pratico: manca la riga ({lettera}) della tabella riassuntiva"
    # I totali seguono le regole dei Pattern A, D ed E applicate alla tabella dei rapporti: il
    # dichiarativo vale dall'anno dopo e non salva le minus che scadono il 31/12/N.
    rapporti = re.findall(r"^\| (\w[\w ]*?) \| ([\d.]+) \| (31/12/N(?:\+\d)?|—) \| ([\d.]+) \|$", pratico, re.M)
    assert len(rapporti) == 3, f"caso pratico: tabella dei rapporti illeggibile ({rapporti})"
    num = lambda s: int(s.replace(".", ""))
    in_loco = sum(min(num(z), num(p)) for _, z, _, p in rapporti)
    residuo = {b: num(z) - min(num(z), num(p)) for b, z, _, p in rapporti}
    capienza = sum(num(p) - min(num(z), num(p)) for _, z, _, p in rapporti)
    scade_ora = [b for b, _, s, _ in rapporti if s == "31/12/N" and residuo[b] > 0]
    assert scade_ora, "caso pratico: nessun rapporto con minus in scadenza il 31/12/N"
    eur = lambda x: "€" + f"{round(x * 0.26):,}".replace(",", ".")
    certificate = min(sum(residuo[b] for b in scade_ora), capienza)
    tutte = min(sum(residuo.values()), capienza)
    dichiarativo = min(sum(r for b, r in residuo.items() if b not in scade_ora), capienza)
    for lettera, attesi in (("a", [in_loco]), ("b", [in_loco + certificate, in_loco + tutte]),
                            ("c", [in_loco + dichiarativo])):
        riga = re.search(rf"^\| \({lettera}\) [^\n]*", pratico, re.M).group(0)
        for x in attesi:
            assert eur(x) in riga, f"caso pratico, riga ({lettera}): atteso {eur(x)} — {riga}"
    # Pattern D senza le minus in scadenza nell'anno; Pattern E con la capienza della destinazione
    # divisa tra le origini (le stesse plus non si contano due volte).
    pattern_d = skill.split("**Pattern D")[1].split("**Pattern E")[0]
    assert re.search(r"anno_scadenza\s*>\s*anno_corrente", pattern_d), "Pattern D: lo zainetto consolidato conta anche le minus che scadono nell'anno"
    assert "Pattern E" in pattern_d.split("```")[1], "Pattern D: la proposta non passa dal confronto con il Pattern E"
    pattern_e = skill.split("**Pattern E")[1].split("### Step 4")[0]
    assert re.search(r"capienza\s*[−-]=\s*recuperabile", pattern_e), "Pattern E: la capienza della destinazione non si riduce tra un'origine e l'altra"


@test("T1: scarto di emissione e coefficiente 48,08% con fonte corretta, marcatori delle reference")
def t_t1_scarto_marcatori():
    # REGRESSIONE F052/F035/F042: BOT e zero coupon tenuti a scadenza non consumano zainetto;
    # il 48,08% viene dall'art. 3 c. 5 DL 66/2014, non dall'art. 68 c. 6 TUIR.
    aliq = open(os.path.join(TAX, "references", "aliquote-12-5-26.md"), encoding="utf-8").read()
    assert "art. 68 c. 6 TUIR, art. 7 DLgs 461/97" not in aliq, "aliquote: fonte del coefficiente superata"
    assert "art. 3 c. 5 DL 66/2014" in aliq, "aliquote: manca la fonte del 48,08%"
    assert re.search(r"### Esempio D[^\n]*(BOT|zero coupon)", aliq), "aliquote: manca l'esempio BOT/zero coupon"
    assert "prezzo di emissione" in aliq and "scarto di emissione" in aliq, "aliquote: esempio senza prezzo di emissione"
    for rel in ("it-tax-optimization/references/etf-trap-redditi-capitale.md",
                "it-tax-optimization/references/multi-intermediario.md",
                "it-tax-optimization/references/aliquote-12-5-26.md",
                "it-pianificazione-obiettivi/references/asset-location.md"):
        testo = open(os.path.join(SKILLS, rel), encoding="utf-8").read()
        assert "Verificato: 2026-10" in testo, f"{rel}: manca il marcatore della verifica di ottobre 2026"
    al = open(os.path.join(SKILLS, "it-pianificazione-obiettivi", "references", "asset-location.md"), encoding="utf-8").read()
    assert "rischio demografico" in al, "asset-location: polizze esenti IRPEF solo per il rischio demografico"


@test("T1: le citazioni del D.Lgs. 461/97 e dell'art. 10-ter L. 77/1983 portano il riferimento dal 2027")
def t_t1_doppie_citazioni():
    # Regola della roadmap «Citazioni dal 2027»: l'art. 6 D.Lgs. 461/97 confluisce nell'art. 306
    # del TU D.Lgs. 117/2026 (c. 2 → c. 3, c. 4 → c. 5, c. 5 → c. 6, c. 7 → c. 8), l'art. 5 nel 304,
    # l'art. 10-ter L. 77/1983 nell'art. 59 dell'All. 1 al D.Lgs. 33/2025. Nei file del tema T1
    # ogni riga che cita la norma vecchia nomina anche quella nuova.
    file_t1 = ["README.md",
               "skills/it-pianificazione/references/integrazione-fiscale-joint.md",
               "skills/it-rebalance/SKILL.md",
               "skills/it-rebalance/assets/ribilanciamento-multibroker.py",
               "skills/it-rebalance/references/ribilanciamento-multibroker.md",
               "skills/it-rebalance/references/sequenza-liquidazione.md",
               "skills/it-tax-optimization/SKILL.md"] + [
               f"skills/it-tax-optimization/references/{f}" for f in
               ("aliquote-12-5-26.md", "criptovalute.md", "etf-trap-redditi-capitale.md",
                "multi-intermediario.md", "regime-dichiarativo.md")]
    testi = _testi_regole()
    errori = []
    for rel in file_t1:
        for riga in testi[rel].splitlines():
            if re.search(r"D\.\s?Lgs\.\s?461|10-ter", riga) and not re.search(r"30[4-7]|TU 2027|117/2026|33/2025", riga):
                errori.append(f"{rel}: «{riga.strip()[:110]}»")
    assert not errori, "citazioni senza il riferimento dal 2027:\n  " + "\n  ".join(errori)


IMPORTATORE = os.path.join(TAX, "assets", "importa-estratto-broker.py")


def _importa(*args):
    """Lancia l'importatore con --json - e ritorna (rc, dati JSON o None, output grezzo)."""
    rc, out = run_script(IMPORTATORE, *args, "--json", "-")
    try:
        return rc, json.loads(out), out
    except ValueError:
        return rc, None, out


@test("T2 importatore: sugli esempi distribuiti nessuna plus «a imposta zero» (ETF = reddito di capitale)")
def t_t2_esempi():
    # REGRESSIONE F001: con il default «diverso» la plus di €1.650 su iShares Core MSCI World
    # risultava compensabile; la minus 2021 scaduta finiva nello zainetto.
    es = lambda f: os.path.join(TAX, "assets", f)
    rc, d, out = _importa(es("esempio-estratto.csv"), es("esempio-zainetto.csv"), "--anno", "2027")
    assert rc == 0 and d, out[-400:]
    a = d["analisi"]
    assert a["plus_latenti_compensabili"] == 0, a
    assert a["zainetto_totale"] == 800, a                 # solo la minus 2023, valida fino al 31/12/2027
    assert a["zainetto_scaduto"] == 1200, a               # la 2022 è scaduta il 31/12/2026
    assert a["plus_realizzabili_a_costo_zero"] == 0, a
    ishares = next(p for p in d["posizioni"] if p["isin"] == "IE00B4L5Y983")
    assert ishares["categoria_fiscale"] == "reddito_capitale" and ishares["plus_minus"] == 1650, ishares
    btp = next(p for p in d["posizioni"] if p["isin"] == "IT0005454241")
    assert btp["aliquota"] == 12.5, btp
    rc, d, out = _importa(es("esempio-estratto.csv"), es("esempio-zainetto.csv"), "--anno", "2026")
    assert rc == 0 and d["analisi"]["zainetto_totale"] == 2000 and d["analisi"]["zainetto_in_scadenza"] == 1200, out[-400:]
    assert all(z["stato"] != "scaduta" for z in d["zainetto"]), "l'esempio distribuito non deve contenere minus già scadute"


@test("T2 importatore: casi limite (costo zero, migliaia italiane, azione DE, ETF USA, minus scaduta, 12,5/26)")
def t_t2_casi_limite():
    # REGRESSIONE F034/F001/F003: costo zero mandava in crash, «1.000» diventava 1,0 e
    # «1.234.567» spariva, gli ETF USA risultavano compensabili, il 12,5% era sommato nominale.
    rc, d, out = _importa(os.path.join(FIXTURES, "estratto-casi-limite.csv"),
                          os.path.join(FIXTURES, "zainetto-casi-limite.csv"), "--anno", "2026")
    assert rc == 0 and d, out[-400:]
    pos = {p["isin"]: p for p in d["posizioni"]}
    sap = pos["DE0007164600"]
    assert sap["categoria_fiscale"] == "reddito_diverso", sap          # azione tedesca, non un ETF
    assert sap["costo"] == 120_000 and sap["valore"] == 150_000, sap   # «1.000» e «150.000,00»
    assert pos["IE00BK5BQT80"]["valore"] == 1_234_567, pos["IE00BK5BQT80"]
    assert pos["US9229083632"]["categoria_fiscale"] == "reddito_capitale_irpef", pos["US9229083632"]
    assert pos["JE00B1VS3770"]["categoria_fiscale"] == "reddito_diverso", "ETC: reddito diverso"
    assert pos["IT0004965148"]["categoria_fiscale"] == "da_classificare", "nessun indizio → da classificare"
    assert pos["IT0005611659"]["categoria_fiscale"] == "da_classificare", "BOT: la plus è scarto di emissione"
    assert pos["IT0005436693"]["aliquota"] == 12.5 and pos["IT0005436693"]["categoria_fiscale"] == "reddito_diverso"
    gratuite = pos["IT0005218380"]
    assert gratuite["costo"] == 0 and gratuite["pct"] is None and gratuite["plus_minus"] == 250, gratuite
    assert "XS0000000000" not in pos and any("Riga senza prezzi" in av for av in d["avvisi"]), d["avvisi"]
    a = d["analisi"]
    assert a["plus_latenti_compensabili"] == 30_650, a                 # SAP + BTP + gratuite + ETC
    assert a["plus_irpef"] == 1_000 and a["plus_da_classificare"] == 125, a
    assert a["zainetto_totale"] == 3_000 and a["zainetto_scaduto"] == 1_000, a
    k = 12.5 / 26
    zeq = 2_000 + 1_000 * k
    assert abs(a["zainetto_26_equivalente"] - zeq) < 0.01, a
    # minus 26% assorbe 2,08 € di plus al 12,5% per euro: si coprono prima i €100 di BTP
    assert abs(a["plus_realizzabili_a_costo_zero"] - (100 + zeq - 100 * k)) < 0.02, a
    assert abs(a["imposta_risparmiabile"] - 0.26 * zeq) < 0.02, a
    rc, testo = run_script(IMPORTATORE, os.path.join(FIXTURES, "estratto-casi-limite.csv"),
                           os.path.join(FIXTURES, "zainetto-casi-limite.csv"), "--anno", "2026")
    assert rc == 0 and "Traceback" not in testo, testo[-400:]
    for atteso in ("31/12/2026", "scaduta", "da classificare", "non costituisce consulenza"):
        assert atteso in testo, f"report senza «{atteso}»"


@test("T2 importatore: zainetto per broker, formato inglese, --json su file completo, errori leggibili")
def t_t2_broker_json_errori():
    # REGRESSIONE F034/F001: la compensazione non attraversa i broker; il JSON era troncato
    # a 1.200 caratteri; file mancanti davano un traceback.
    rc, d, out = _importa(os.path.join(FIXTURES, "estratto-multibroker.csv"),
                          os.path.join(FIXTURES, "zainetto-multibroker.csv"), "--anno", "2026")
    assert rc == 0 and d, out[-400:]
    a = d["analisi"]
    assert a["plus_latenti_compensabili"] == 1_700, a                 # «12.40» letto all'inglese
    assert a["plus_realizzabili_a_costo_zero"] == 800, a               # 300 su Directa + 500 su Fineco
    assert set(a["per_broker"]) == {"Directa", "Fineco"}, a
    tmp = tempfile.mkdtemp()
    try:
        dest = os.path.join(tmp, "estratto.json")
        rc, out = run_script(IMPORTATORE, os.path.join(TAX, "assets", "esempio-estratto.csv"),
                             "--broker", "Fineco", "--anno", "2026", "--json", dest)
        assert rc == 0, out[-400:]
        dati = json.load(open(dest, encoding="utf-8"))
        assert len(dati["posizioni"]) == 4 and all(p["broker"] == "Fineco" for p in dati["posizioni"]), dati
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    rc, out = run_script(IMPORTATORE, os.path.join(FIXTURES, "non-esiste.csv"))
    assert rc == 1 and "Traceback" not in out and "non trovato" in out, out[-300:]
    rc, out = run_script(IMPORTATORE, os.path.join(FIXTURES, "estratto-multibroker.csv"),
                         os.path.join(TAX, "assets", "esempio-zainetto.csv"))
    assert rc == 1 and "Traceback" not in out and "broker" in out, out[-300:]


@test("T3: quadro RW senza soglia per i titoli, stato dei broker esteri con data di verifica")
def t_t3_rw_broker():
    # REGRESSIONE F006/F053/F036: la soglia di €15.000 vale solo per conti e depositi; Trade
    # Republic e Scalable non sono più broker «senza sostituto d'imposta» per tutti.
    # Revisione T3: il controllo lavora sul paragrafo (o sulla riga di tabella), non sulla frase,
    # e riconosce anche la frase originale di F053 («non ha un sostituto d'imposta italiano»).
    senza_sostituto = re.compile(r"(senza|nessun|non ha(nno)?( un| alcun)?|privi? di)\s+(un |alcun )?sostitut", re.I)
    for rel, testo in _testi_regole().items():
        for blocco in _blocchi(testo):
            frasi = [blocco] if blocco.lstrip().startswith("|") else _frasi(blocco)
            if re.search(r"\bRW\b|monitoraggio", blocco):
                for frase in frasi:
                    if "15.000" in frase:
                        assert re.search(r"\bconti\b|conti? corrent|deposit", frase), f"{rel}: «{frase.strip()[:120]}»"
            for frase in frasi:
                if senza_sostituto.search(frase):
                    assert not re.search(r"Trade Republic|Scalable", frase), f"{rel}: «{frase.strip()[:120]}»"
                for clausola in re.split(r"[();]", frase):     # la description è una frase sola
                    if re.search(r"broker ester|dichiarativ", clausola) and re.search(r"Trade Republic|Scalable", clausola):
                        assert re.search(r"amministrat|succursale|verific", clausola), f"{rel}: «{clausola.strip()[:120]}»"
                if re.search(r"cripto", frase) and re.search(r"Trade Republic", frase):
                    assert re.search(r"verific", frase), f"{rel}: cripto su Trade Republic date per certe: «{frase.strip()[:120]}»"
                assert not re.search(r"broker ester\w*[^.\n]{0,40}\bobbligatori|regime di chi usa (un |i )?broker ester|"
                                     r"dichiarativ\w*\**\s*\(broker ester", frase), \
                    f"{rel}: broker estero = dichiarativo: «{frase.strip()[:120]}»"
    dich = open(os.path.join(TAX, "references", "regime-dichiarativo.md"), encoding="utf-8").read()
    rw = next(b for b in _blocchi(dich) if b.startswith("**Quadro RW"))
    assert re.search(r"[Pp]er titoli, ETF, fondi e cripto[^.]*qualunque sia l'importo", rw), "quadro RW: obbligo senza soglia per i titoli"
    stato = dich.split("## Quali broker")[1].split("\n## ")[0] if "## Quali broker" in dich else ""
    regimi = {}
    for riga in stato.splitlines():
        celle = [c.strip() for c in riga.strip().strip("|").split("|")]
        if len(celle) >= 3:
            regimi[celle[0]] = celle[1]
    for broker, regime in (("Interactive Brokers", "Dichiarativo"), ("Degiro", "Dichiarativo"),
                           ("Trade Republic", "Amministrato"), ("Scalable Capital", "Amministrato")):
        assert regimi.get(broker, "").startswith(regime), f"regime-dichiarativo.md: {broker} → «{regimi.get(broker)}»"
    assert "ottobre 2026" in stato, "tabella dei broker senza data di verifica"
    assert "non entrano nello zainetto del nuovo rapporto in amministrato" in dich, "minus RT e passaggio all'amministrato"
    assert not re.search(r"revoca[^.\n]{0,40}\(Pattern E\)", dich), "la revoca dell'amministrato è il Pattern D, la chiusura il Pattern E"
    # un conto Interactive Brokers o Degiro nell'importatore: avviso che i numeri per broker non valgono
    tmp = tempfile.mkdtemp()
    try:
        estratto = os.path.join(tmp, "ibkr.csv")
        with open(estratto, "w", encoding="utf-8") as f:
            f.write("Broker;ISIN;Descrizione;Quantità;Prezzo medio di carico;Controvalore\n"
                    "Interactive Brokers;US0378331005;Apple Inc;10;100,00;1.700,00\n")
        rc, d, out = _importa(estratto)
        assert rc == 0 and any("dichiarativo" in av and "Interactive Brokers" in av for av in d["avvisi"]), out[-300:]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    riga_tr = next((r for r in stato.splitlines() if r.startswith("| Trade Republic")), "")
    frasi_cripto = [f for f in re.split(r"(?<=\.)\s+", riga_tr) if "cripto" in f]
    assert frasi_cripto and all("verific" in f for f in frasi_cripto), \
        f"cripto su Trade Republic: va verificato se rientrano nell'amministrato: {frasi_cripto}"
    imp = open(os.path.join(TAX, "references", "import-dati-broker.md"), encoding="utf-8").read()
    assert "Scalable" in imp and "Interactive Brokers" in imp and "Verificato: 2026-10" in imp, \
        "import-dati-broker.md: stato dei broker esteri o marcatore mancanti"
    skill = open(os.path.join(TAX, "SKILL.md"), encoding="utf-8").read()
    assert "soglie RW" not in skill, "SKILL.md: «soglie RW» verificate (non esistono soglie per i titoli)"


@test("T3: dichiarativo con cambio del giorno, conti in valuta, minus in RT e plus di ETF nel quadro RM")
def t_t3_dichiarativo():
    # REGRESSIONE F053: plus in euro al cambio del giorno (art. 9 c. 2 TUIR), soglia di
    # €51.645,69 per 7 giorni lavorativi sui conti in valuta, minus da indicare in RT nell'anno.
    dich = open(os.path.join(TAX, "references", "regime-dichiarativo.md"), encoding="utf-8").read()
    for atteso in ("art. 9 c. 2", "51.645,69", "sette giorni lavorativi", "60/E", "quadro RM",
                   "anno di realizzo", "4043"):
        assert atteso in dich, f"regime-dichiarativo.md senza «{atteso}»"
    # l'esempio in dollari deve tornare con i numeri scritti nel testo (cambi, importi e segno)
    num = lambda x: float(x.replace(".", "").replace(",", "."))
    m = re.search(r"(\d+) azioni comprate a \$(\d+) con il cambio EUR/USD a (\d,\d+) \(costo €([\d.]+,\d\d)\) "
                  r"e vendute a \$(\d+) con il cambio a (\d,\d+) \(corrispettivo €([\d.]+,\d\d)\)"
                  r".*?in euro è una (minus|plus) di €([\d.]+,\d\d)", dich, re.S)
    assert m, "esempio USD non trovato"
    q, p0, c0, costo, p1, c1, ricavo, segno, diff = m.groups()
    q, costo_ok, ricavo_ok = int(q), int(q) * int(p0) / num(c0), int(q) * int(p1) / num(c1)
    assert abs(num(costo) - costo_ok) < 0.01 and abs(num(ricavo) - ricavo_ok) < 0.01, m.group(0)
    assert segno == ("minus" if ricavo_ok < costo_ok else "plus") and abs(num(diff) - abs(ricavo_ok - costo_ok)) < 0.01, m.group(0)
    # le minus di ETF e fondi sono redditi diversi: quadro RT, anche se le plus vanno nel quadro RM
    assert re.search(r"minus[^.\n]*(ETF|fondi)[^.\n]*quadro RT|quadro RT[^.\n]*minus[^.\n]*(ETF|fondi)", dich), "minus di ETF nel quadro RT"
    assert "Nel quadro RT vanno anche le minus da vendita di ETF e fondi, che sono redditi diversi" in dich.replace("**", "")
    assert not re.search(r"non vanno[^.\n]*minus|minus[^.\n]*(ETF|fondi)[^.\n]*redditi di capitale", dich), "minus di ETF negate nel quadro RT"
    for rel in ("skills/it-tax-optimization/references/regime-dichiarativo.md", "skills/it-tax-optimization/SKILL.md"):
        assert "sette giorni lavorativi continui" in open(os.path.join(ROOT, rel), encoding="utf-8").read(), rel
    checklist = dich.split("## Checklist")[1].split("\n## ")[0]
    for atteso in ("730", "30/9", "4047", "4048", "30/11"):
        assert atteso in checklist, f"checklist del dichiarativo senza «{atteso}»"
    assert re.search(r"(RT|quadro T)[^.\n]*(senza|non ha|nessun) acconti?", checklist), "l'imposta del quadro RT si versa a saldo, senza acconti"
    multi = open(os.path.join(TAX, "references", "multi-intermediario.md"), encoding="utf-8").read()
    assert "senza acconti" in multi and "e in acconto" not in multi and "(non 730)" not in multi and "730" in multi, \
        "multi-intermediario: 730 ammesso, imposta del quadro RT senza acconti"
    for rel, testo in _testi_regole().items():
        for m in re.finditer(r"quadro RT del Modello Redditi PF", testo):
            assert "730" in testo[m.end():m.end() + 60], f"{rel}: «{testo[m.start():m.end() + 60]}» senza il 730"
        for vecchio, nuovo in ((r"68 c\. 9-bis", "77 c. 11"), (r"68 c\. 5", "77 c. 6"),
                               (r"6 c\. 2 D\.Lgs\. 461/97", "306 c. 3"), (r"6 c\. 5 D\.Lgs\. 461/97", "306 c. 6")):
            for m in re.finditer(vecchio + r"[^\n]{0,40}?(?:→|dal 2027) art\. (\d+ c\. \d+)", testo):
                assert m.group(1) == nuovo, f"{rel}: «{m.group(0)}» (atteso art. {nuovo})"
    skill = open(os.path.join(TAX, "SKILL.md"), encoding="utf-8").read()
    pattern_d = skill.split("**Pattern D")[1].split("**Pattern E")[0]
    assert "cambio" in pattern_d and "entrambi i regimi" in pattern_d, "Pattern D: il cambio del giorno vale in entrambi i regimi"
    assert "continui" in pattern_d and "51.645,69" in pattern_d, "Pattern D: €51.645,69 per sette giorni lavorativi continui"
    calendario = skill.split("### Step 4")[1].split("## Formato output")[0]
    assert "dichiarativo" in calendario and "Redditi PF" in calendario, "calendario senza le scadenze del dichiarativo"
    assert "730" in calendario and "30/11" in calendario, "calendario: 730 e secondo acconto IVAFE"
    assert "quadro RT, senza acconti" in calendario, "calendario: l'imposta del quadro RT non ha acconti"
    # Revisione T3: il workflow non può trattare un conto in dichiarativo come un dossier in amministrato
    analisi = skill.split("### Step 2")[1].split("### Step 3")[0]
    ramo = analisi.split("**Broker in dichiarativo**")[1].split("\n\n")[0] if "**Broker in dichiarativo**" in analisi else ""
    for atteso in ("lotti", "cambio del giorno", "non calcolare", "Pattern A è cross-broker",
                   "Il Pattern D si valuta solo per i rapporti in amministrato",
                   "Pattern D (revoca dell'amministrato", "Pattern E (chiusura", "sezione 5"):
        assert atteso in ramo, f"Step 2, ramo dei broker in dichiarativo senza «{atteso}»"
    assert "indicative" not in ramo, "plus e minus dei conti in dichiarativo non sono «indicative» sul PMC in valuta"
    assert "### 5. (Solo con conti in dichiarativo) Adempimenti del dichiarativo" in skill
    assert re.search(r"### 4\. \(Solo se multi-broker[^)\n]*amministrato\)", skill), "sezione 4 solo con un rapporto in amministrato"
    ingest = skill.split("### Step 1")[1].split("### Step 2")[0]
    assert re.search(r"quadro RT[^.\n]*(per anno|non per broker)", ingest), "Step 1: minus riportabili dal quadro RT"
    disclaimer = skill.split("### Disclaimer")[1].split("\n## ")[0]
    assert "dichiarativo" in disclaimer, "disclaimer limitato all'amministrato"
    # F057: /minusvalenze è un alias della skill; i punti che chiedeva stanno nello Step 1
    assert re.search(r"(?i)regime[^.\n]*(ciascun|ogni) broker|(ciascun|ogni) broker[^.\n]*regime", ingest), \
        "Step 1 non chiede il regime di ogni broker"
    assert re.search(r"minus riportabili dal quadro RT", ingest), "Step 1: minus del quadro RT"
    assert re.search(r"[Nn]on (lanciare|usare) l'importatore[^.\n]*dichiarativo", ingest), "Step 1: importatore solo per l'amministrato"


@test("T3: cripto con leva ETN/ETC al 26% (con cautela) e reference marcate")
def t_t3_cripto():
    # REGRESSIONE F036: manca il confronto tra 33% della detenzione diretta e 26% degli
    # strumenti finanziari con sottostante cripto; pene e ravvedimento restano quelli verificati.
    cripto = open(os.path.join(TAX, "references", "criptovalute.md"), encoding="utf-8").read()
    leve = cripto.split("## Leve di pianificazione")[1].split("\n## ")[0]
    voce = next((r for r in leve.splitlines() if re.search(r"ETN|ETC|ETP", r)), "")
    assert re.search(r"al 26% invece del 33%", voce), f"leva ETN/ETC: {voce[:120]}"
    assert "zainetto ordinario" in voce and "non con le minus cripto" in voce, f"leva ETN/ETC: {voce[:120]}"
    assert "verificare" in voce.lower() and "rischio emittente" in voce, "leva ETN/ETC senza cautele"
    assert re.search(r"[Nn]é ETN[^.]*né ETP|ETC, ETN ed ETP[^.]*non sono citati", voce), \
        "la Circ. 30/E/2023 non cita ETN, ETC ed ETP"
    assert not re.search(r"circolare[^.]*(lo )?conferm", voce), "la circolare non conferma il trattamento di ETN ed ETC"
    assert re.search(r"\bW\b[^.\n]{0,30}dal 730/2024", cripto.replace("*", "")), "criptovalute: il quadro W c'è dal 730/2024"
    assert "l'AdE applica il LIFO" not in cripto
    avanzate = open(os.path.join(TAX, "references", "criptovalute-avanzate.md"), encoding="utf-8").read()
    assert "da 2 a 5 anni" in avanzate and "1-3 anni" not in avanzate, "pena dell'omessa dichiarazione"
    for f in ("criptovalute.md", "criptovalute-avanzate.md", "criptovalute-casi-speciali.md",
              "regime-dichiarativo.md", "import-dati-broker.md"):
        assert "Verificato: 2026-10" in open(os.path.join(TAX, "references", f), encoding="utf-8").read(), f


def _fx(nome):
    return os.path.join(FIXTURES, nome)


@test("T2 importatore: 12,5% a Stati white list, enti territoriali e sovranazionali; ETC e nomi di emittenti senza default rischiosi")
def t_t2_pubblici_etc():
    # REGRESSIONE revisione T2: Bund, Treasury e BEI uscivano al 26% senza avviso (art. 6 c. 1
    # D.Lgs. 461/97 → art. 306 c. 1 TU 2027); «WisdomTree Physical Gold» (JE) finiva a IRPEF
    # con quadro RL; «Vanguard … Inc» e «ISHARES III PLC» erano azioni con la plus compensabile.
    rc, d, out = _importa(_fx("estratto-pubblici-etc.csv"), _fx("zainetto-un-broker.csv"), "--anno", "2026")
    assert rc == 0 and d, out[-400:]
    pos = {p["descrizione"]: p for p in d["posizioni"]}
    for nome in ("Bund 0% 15/02/2032", "BEI 1,5% 15/11/2030", "US Treasury 4,5% 15/11/2033",
                 "Comune di Roma 5,345% 2048", "Land Berlin 1% 2030", "BTP 2,5% 01/12/2032"):
        assert pos[nome]["categoria_fiscale"] == "reddito_diverso" and pos[nome]["aliquota"] == 12.5, pos[nome]
    enel = pos["Enel Finance 3% 2030"]
    assert enel["categoria_fiscale"] == "reddito_diverso" and enel["aliquota"] == 26, enel
    avvisi_aliquota = [av for av in d["avvisi"] if "12,5%" in av]
    assert any("Enel Finance" in av for av in avvisi_aliquota), d["avvisi"]
    assert not any(re.search(r"Bund|BEI|Treasury|Comune|Land Berlin|BTP", av) for av in avvisi_aliquota), avvisi_aliquota
    for nome in ("WisdomTree Physical Gold", "iShares Physical Gold"):
        assert pos[nome]["categoria_fiscale"] == "reddito_diverso", pos[nome]
    assert pos["Vanguard FTSE All-World High Dividend Yield Inc"]["categoria_fiscale"] == "reddito_capitale"
    for nome in ("ISHARES III PLC", "WisdomTree Brent Crude Oil", "Fondo Esempio non armonizzato",
                 "Fondo Immobiliare Esempio", "HSBC MSCI World Inc", "SPDR S&P 500 Trust"):
        assert pos[nome]["categoria_fiscale"] == "da_classificare", pos[nome]
    assert not any(p["categoria_fiscale"] == "reddito_capitale_irpef" for p in d["posizioni"])
    a = d["analisi"]
    assert a["plus_latenti_compensabili"] == 1_050 and a["plus_non_compensabili"] == 100, a
    assert a["plus_da_classificare"] == 270 and a["plus_irpef"] == 0, a
    # 1.000 di minus al 26% coprono prima i 550 al 12,5% (264,42 in 26%-equivalente), poi i 500 al 26%
    assert a["plus_realizzabili_a_costo_zero"] == 1_050, a
    assert abs(a["imposta_risparmiabile"] - (550 * 0.125 + 500 * 0.26)) < 0.02, a
    assert set(a["per_broker"]) == {"Fineco"}, a
    # Revisione di T2: nomi che somigliano a emittenti pubblici, ETC o indici senza esserlo
    rc, d, out = _importa(_fx("estratto-falsi-positivi.csv"), "--anno", "2026")
    assert rc == 0 and d, out[-400:]
    pos = {p["descrizione"]: p for p in d["posizioni"]}
    for nome in ("UBS Bloomberg US Treasury Bond 1-10", "JPM BetaBuilders US Treasury Bond 0-1 yr",
                 "OATLY GROUP AB ADR", "ETFS Brent Crude", "Bundesrepublik Deutschland"):
        assert pos[nome]["categoria_fiscale"] == "da_classificare", pos[nome]
    for nome in ("BUNDESIMMOBILIENGESELLSCHAFT 3,375% 2030", "OESTERREICHISCHE BUNDESBAHNEN 3% 2033",
                 "OATLY GROUP 9,25% 2028"):
        assert pos[nome]["categoria_fiscale"] == "reddito_diverso" and pos[nome]["aliquota"] == 26, pos[nome]
    for nome in ("ETFS Physical Gold", "MSCI Inc", "S&P Global Inc", "World Kinect Corp"):
        assert pos[nome]["categoria_fiscale"] == "reddito_diverso", pos[nome]
    assert pos["Xtrackers FTSE Developed Europe Real Estate UCITS ETF"]["categoria_fiscale"] == "reddito_capitale"
    assert not any(p["aliquota"] == 12.5 or p["categoria_fiscale"] == "reddito_capitale_irpef" for p in d["posizioni"])


@test("T2 importatore: zainetto per anno di scadenza, residuo accanto al lordo, righe anomale")
def t_t2_zainetto_colonne():
    # REGRESSIONE revisione T2: «Anno di scadenza» era letto come anno di realizzo (+4 anni);
    # con «Minusvalenza» e «Residuo» vinceva il lordo; «-» nell'aliquota fermava tutto l'import.
    estratto = _fx("estratto-migliaia-it.csv")
    rc, d, out = _importa(estratto, _fx("zainetto-scadenza.csv"), "--anno", "2026")
    assert rc == 0 and d, out[-400:]
    z = sorted(d["zainetto"], key=lambda r: r["anno"])
    assert [r["anno"] for r in z] == [2022, 2023], z
    assert z[0]["scade_il"] == "31/12/2026" and z[0]["stato"] == "in_scadenza", z
    assert d["analisi"]["zainetto_in_scadenza"] == 1_000 and d["analisi"]["zainetto_totale"] == 1_800, d["analisi"]
    assert any("scadenza" in av for av in d["avvisi"]), d["avvisi"]
    rc, testo = run_script(IMPORTATORE, estratto, _fx("zainetto-scadenza.csv"), "--anno", "2026")
    assert rc == 0 and "Priorità" in testo, testo[-400:]
    # lordo 5.000, utilizzato 1.000 e residuo 500: vale il residuo, non lordo − utilizzato
    for nome in ("zainetto-lordo-residuo.csv", "zainetto-plurale.csv", "zainetto-utilizzato.csv",
                 "zainetto-compensabili.csv"):
        rc, d, out = _importa(estratto, _fx(nome), "--anno", "2026")
        assert rc == 0 and d, f"{nome}: {out[-300:]}"
        assert d["analisi"]["zainetto_totale"] == 500, f"{nome}: {d['analisi']}"
        if nome != "zainetto-plurale.csv":
            assert any("residu" in av.lower() for av in d["avvisi"]), f"{nome}: {d['avvisi']}"
    rc, d, out = _importa(estratto, _fx("zainetto-esercizio-scadenza.csv"), "--anno", "2026")
    assert rc == 0 and d["zainetto"][0]["anno"] == 2022 and d["zainetto"][0]["stato"] == "in_scadenza", out[-300:]
    # CSV con la virgola come delimitatore e importi interi all'italiana: «1.500» sono millecinquecento
    rc, d, out = _importa(estratto, _fx("zainetto-virgola-interi.csv"), "--anno", "2026")
    assert rc == 0 and d["analisi"]["zainetto_totale"] == 3_500, out[-300:]
    assert any(av.startswith("zainetto: formato numeri") for av in d["avvisi"]), d["avvisi"]
    rc, d, out = _importa(estratto, _fx("zainetto-anomalie.csv"), "--anno", "2026")
    assert rc == 0 and d, out[-400:]
    assert d["analisi"]["zainetto_totale"] == 1_000, d["analisi"]          # la riga senza anno resta fuori
    assert [r["stato"] for r in d["zainetto"]].count("anno_mancante") == 1, d["zainetto"]
    assert any("riga 3" in av and "aliquota" in av for av in d["avvisi"]), d["avvisi"]


@test("T2 importatore: numeri senza indizi di formato, colonna tipo fiscale, broker unico, errori di input")
def t_t2_numeri_colonne():
    # REGRESSIONE revisione T2: «1,000» in un CSV con la virgola come delimitatore diventava 1;
    # «Redditi diversi (capital gain)» diventava reddito di capitale; «Natura strumento» era
    # scambiata per il tipo fiscale; errori su aliquota e cartelle senza file né riga.
    for nome in ("estratto-migliaia-it.csv", "estratto-migliaia-en.csv"):
        rc, d, out = _importa(_fx(nome))
        assert rc == 0 and d, f"{nome}: {out[-300:]}"
        sap = d["posizioni"][0]
        assert sap["costo"] == 120_000 and sap["valore"] == 150_000, f"{nome}: {sap}"
        assert any("formato numeri" in av for av in d["avvisi"]), f"{nome}: {d['avvisi']}"
    rc, d, out = _importa(_fx("estratto-tipo-fiscale.csv"))
    assert rc == 0 and d, out[-400:]
    pos = {p["isin"]: p for p in d["posizioni"]}
    assert pos["IE00B4L5Y983"]["categoria_fiscale"] == "reddito_capitale_irpef", pos["IE00B4L5Y983"]
    assert pos["IE00B4L5Y983"]["costo"] == 10_000 and pos["IE00B4L5Y983"]["valore"] == 12_000, pos["IE00B4L5Y983"]
    assert pos["US0378331005"]["categoria_fiscale"] == "reddito_diverso", pos["US0378331005"]
    assert pos["IE00B5BMR087"]["categoria_fiscale"] == "reddito_capitale", pos["IE00B5BMR087"]
    assert any("boh" in av for av in d["avvisi"]), d["avvisi"]
    rc, d, out = _importa(_fx("estratto-migliaia-it.csv"), _fx("zainetto-un-broker.csv"), "--anno", "2026")
    assert rc == 0 and d["analisi"]["plus_realizzabili_a_costo_zero"] == 1_000, out[-400:]
    rc, out = run_script(IMPORTATORE, _fx("estratto-migliaia-it.csv"), _fx("zainetto-multibroker.csv"))
    assert rc == 1 and "Traceback" not in out and "broker" in out, out[-300:]
    tmp = tempfile.mkdtemp()
    try:
        natura = os.path.join(tmp, "natura.csv")
        with open(natura, "w", encoding="utf-8") as f:
            f.write("ISIN;Descrizione;Natura strumento;Quantità;Prezzo medio di carico;Controvalore\n"
                    "DE0007164600;SAP SE;Azione;10;100,00;1.500,00\n")
        rc, d, out = _importa(natura)
        assert rc == 0 and d["posizioni"][0]["categoria_fiscale"] == "reddito_diverso", out[-300:]
        assert d["posizioni"][0]["fonte_categoria"] == "descrizione", d["posizioni"][0]
        assert not any("tipo fiscale" in av for av in d["avvisi"]), d["avvisi"]
        ambiguo = os.path.join(tmp, "ambiguo.csv")
        with open(ambiguo, "w", encoding="utf-8") as f:
            f.write("Anno realizzo;Importo;Ammontare;Aliquota\n2024;5.000,00;500,00;26%\n")
        rc, out = run_script(IMPORTATORE, _fx("estratto-migliaia-it.csv"), ambiguo)
        assert rc == 1 and "Traceback" not in out and "ambigu" in out, out[-300:]
        vuota = os.path.join(tmp, "vuota.csv")
        with open(vuota, "w", encoding="utf-8") as f:
            f.write("Anno realizzo;Minusvalenza residua;Aliquota\n2024;1.000,00;\n")
        rc, d, out = _importa(_fx("estratto-migliaia-it.csv"), vuota)
        assert rc == 0 and any("riga 2" in av and "aliquota" in av for av in d["avvisi"]), out[-300:]
        zain = os.path.join(tmp, "zainetto.csv")
        with open(zain, "w", encoding="utf-8") as f:
            f.write("Anno realizzo;Minusvalenza residua;Aliquota\n2024;1.000,00;30%\n")
        rc, out = run_script(IMPORTATORE, _fx("estratto-migliaia-it.csv"), zain)
        assert rc == 1 and "Traceback" not in out and "zainetto.csv" in out and "riga 2" in out, out[-300:]
        rc, out = run_script(IMPORTATORE, tmp)
        assert rc == 1 and "Traceback" not in out, out[-300:]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ══════════════════════ T13: dividendi esteri, successione e donazione ══════════════════════

def _perc_t13(x):
    return float(x.replace(",", "."))


def _euro_t13(x):
    return float(x.replace(".", "").replace(",", "."))


@test("reference: ognuna è nell'indice della propria SKILL.md")
def t_reference_indicizzate():
    # REGRESSIONE revisione T13: due reference nuove non erano citate da nessun file, quindi
    # la skill non le caricava mai.
    mancanti = []
    for skill in sorted(os.listdir(SKILLS)):
        cartella = os.path.join(SKILLS, skill, "references")
        if not os.path.isdir(cartella):
            continue
        indice = open(os.path.join(SKILLS, skill, "SKILL.md"), encoding="utf-8").read().split("\n## Riferimenti")[1]
        indice = indice.split("\n## ")[0]
        mancanti += [f"{skill}/{f}" for f in sorted(os.listdir(cartella))
                     if f.endswith(".md") and f"references/{f}" not in indice]
    assert not mancanti, f"reference fuori dall'indice: {mancanti}"


@test("T13: dividendi esteri, prelievo r + 26% × (1 − r) in tabella ed esempi ricalcolati dal testo")
def t_t13_dividendi():
    # REGRESSIONE F021: i numeri della reference non erano vincolati da nessun test.
    div = open(os.path.join(TAX, "references", "dividendi-esteri.md"), encoding="utf-8").read()
    prelievo = lambda r: round((r + 0.26 * (1 - r)) * 100, 1)
    tabella = div.split("## Le ritenute nei principali Paesi")[1].split("\n\n(*)")[0]
    righe = [[c.strip() for c in r.strip().strip("|").split("|")] for r in tabella.splitlines() if r.startswith("| ")]
    paesi = {}
    for paese, interna, tetto, senza, con, _ in righe[1:]:
        pct = lambda cella: [_perc_t13(x) for x in re.findall(r"(\d+(?:,\d+)?)%", cella)]
        r = pct(interna)[0] / 100 if pct(interna) else 0.0
        paesi[paese] = r
        assert prelievo(r) in pct(senza), f"{paese}: senza rimborso {senza} ≠ {prelievo(r)}%"
        if pct(con):
            r_eff = 0.0 if "esent" in tetto else min(r, pct(tetto)[0] / 100)
            assert pct(con) == [prelievo(r_eff)], f"{paese}: con rimborso {con} ≠ {prelievo(r_eff)}%"
    assert {"Stati Uniti", "Francia", "Germania", "Svizzera", "Irlanda", "Paesi Bassi", "Regno Unito"} <= set(paesi), paesi
    assert paesi["Francia"] == 0.128, "Francia: ritenuta per le persone fisiche al 12,8%"
    nota = div.split("\n(*)")[1].split("\n\n")[0]
    assert "tutti i Paesi con tetto al 15%" not in nota and "Francia" in nota and "Irlanda" in nota, nota[:200]
    # esempio in amministrato: netto = lordo − r·lordo − 26% del netto frontiera
    rit = [_perc_t13(x) / 100 for x in re.findall(r"Ritenuta USA[^\n]*?\((\d+)%\)[^\n]*?\((\d+)%\)", div)[0]]
    netti = [_euro_t13(x) for x in re.search(r"Netto accreditato\s+([\d.]+,\d\d)\s+([\d.]+,\d\d)", div).groups()]
    assert netti == [round(1000 - 1000 * r - 0.26 * 1000 * (1 - r), 2) for r in rit], netti
    perdita = re.search(r"Senza W-8BEN si perdono €(\d+)", div).group(1)
    assert int(perdita) == round(netti[0] - netti[1]), perdita
    # esempio in dichiarativo: 26% sul lordo
    netto_dich = _euro_t13(re.search(r"Imposta sostitutiva 26% sul lordo[^\n]*\n\s*Netto\s+([\d.]+,\d\d)", div).group(1))
    assert netto_dich == 1000 - 150 - 260 and "41,0%" in div, netto_dich
    assert re.search(r"lordo[^.\n]*netto frontiera|netto frontiera[^.\n]*lordo", div.split("## Cosa resta da verificare")[1]), \
        "dichiarativo: la base del 26% (lordo o netto frontiera) tra le cose da verificare"
    francia = next(r for r in righe if r[0] == "Francia")
    assert "25%" in francia[5] and "rimborso" in francia[5], "Francia: il caso del 25% se il beneficiario non è documentato"


@test("T13: successione e donazione, esempi ricalcolati e costo fiscale degli eredi")
def t_t13_successione():
    # REGRESSIONE F054: esempi «vendo o lascio» non vincolati; zainetto e fondi/ETF.
    suc = open(os.path.join(TAX, "references", "successione-donazione.md"), encoding="utf-8").read()
    for pct, base, valore in re.findall(r"(\d+)% × ([\d.]+)\s+([\d.]+)", suc):
        assert int(pct) * _euro_t13(base) / 100 == _euro_t13(valore), f"{pct}% × {base} ≠ {valore}"
    for pct, a, b, valore in re.findall(r"(\d+)% × \(([\d.]+) − ([\d.]+)\)\s+([\d.]+)", suc):
        assert int(pct) * (_euro_t13(a) - _euro_t13(b)) / 100 == _euro_t13(valore), f"{pct}% × ({a} − {b}) ≠ {valore}"
    es2 = suc.split("**Esempio 2")[1].split("**I limiti")[0]
    netto_vendita = re.search(r"4% × \(([\d.]+) −", es2).group(1)
    assert _euro_t13(netto_vendita) == 3_000_000 - 260_000, netto_vendita
    totali = [_euro_t13(x) for x in re.findall(r"Totale imposte\s+([\d.]+)", es2)]
    differenza = _euro_t13(re.search(r"Differenza\s+([\d.]+)", es2).group(1))
    assert totali == [260_000 + 69_600, 80_000] and differenza == totali[0] - totali[1], (totali, differenza)
    es1 = suc.split("**Esempio 1")[1].split("**Esempio 2")[0]
    assert _euro_t13(re.search(r"Differenza\s+([\d.]+)", es1).group(1)) == 78_000
    assert "€1.000.000" in suc and "zainetto si estingue" in suc.replace("**", "")
    assert "da verificare sulla fonte ufficiale" not in suc and "va verificato sulla fonte ufficiale" not in suc, \
        "citazione non verificata dentro un file marcato «Verificato»"
    assert "`it-pianificazione-obiettivi/references/asset-location.md`" in suc and "(vedi `asset-location.md`)" not in suc
    skill = open(os.path.join(TAX, "SKILL.md"), encoding="utf-8").read()
    indice = skill.split("## Riferimenti")[1].split("\n## ")[0]
    for ref in ("dividendi-esteri.md", "successione-donazione.md"):
        assert f"- `references/{ref}` — " in indice, f"SKILL.md: {ref} non è nell'elenco dei Riferimenti"
    casi_limite = next(r for r in skill.splitlines() if "Casi limite" in r)
    assert "references/successione-donazione.md" in casi_limite, casi_limite
    costi = skill.split("### Costi fiscali non compensabili")[1].split("## Workflow")[0]
    assert "references/dividendi-esteri.md" in costi and "W-8BEN" in costi, "costi non compensabili senza le ritenute estere"
    # l'intervallo del prelievo nel bullet = minimo e massimo della tabella della reference
    div = open(os.path.join(TAX, "references", "dividendi-esteri.md"), encoding="utf-8").read()
    tabella = div.split("## Le ritenute nei principali Paesi")[1].split("\n\n(*)")[0]
    valori = [_perc_t13(x) for r in tabella.splitlines()[2:] for x in re.findall(r"(\d+,\d)%", "|".join(r.split("|")[4:6]))]
    bullet = next(r for r in costi.splitlines() if "Ritenute estere" in r)
    nel_bullet = [_perc_t13(x) for x in re.findall(r"(\d+,\d)%", bullet)]
    assert min(nel_bullet) == min(valori) and max(nel_bullet) == max(valori), (bullet[:200], min(valori), max(valori))
    # aliquote e franchigie: tabella della reference e copia cripto coincidono
    tab = suc.split("| Beneficiario |")[1].split("\n\n")[0]
    righe = {r.split("|")[1].strip(): [c.strip() for c in r.split("|")[2:4]] for r in tab.splitlines()[2:] if r.startswith("|")}
    assert righe["Coniuge e parenti in linea retta"] == ["€1.000.000", "4%"], righe
    assert righe["Fratelli e sorelle"] == ["€100.000", "6%"], righe
    assert righe["Altri soggetti"] == ["nessuna", "8%"], righe
    assert righe["Persona con disabilità grave (art. 3 c. 3 L. 104/1992)"][0] == "€1.500.000", righe
    assert any(k.startswith("Altri parenti fino al quarto grado, affini in linea retta, affini in linea collaterale fino al terzo grado")
               and v == ["nessuna", "6%"] for k, v in righe.items()), righe
    cripto_cs = open(os.path.join(TAX, "references", "criptovalute-casi-speciali.md"), encoding="utf-8").read()
    for atteso in ("4% per coniuge e parenti in linea retta oltre la franchigia di €1.000.000 ciascuno",
                   "6% per fratelli e sorelle oltre €100.000", "8% per gli altri soggetti",
                   "franchigia di €1.500.000 per i disabili gravi"):
        assert atteso in cripto_cs, f"criptovalute-casi-speciali: «{atteso}»"
    assert re.search(r"nella donazione non godono dell'esenzione: l'art\. 59", suc), "titoli di Stato in donazione"
    assert re.search(r"il testo vigente \*\*non esclude\*\* successioni e donazioni", suc), "OICR: l'eccezione per successione non c'è più"
    seq = open(os.path.join(SKILLS, "it-rebalance", "references", "sequenza-liquidazione.md"), encoding="utf-8").read()
    assert "netto frontiera" in seq and "it-tax-optimization/references/dividendi-esteri.md" in seq
    assert "Finalità successoria" in seq and "it-tax-optimization/references/successione-donazione.md" in seq
    obiettivi = os.path.join(SKILLS, "it-pianificazione-obiettivi")
    for rel in ("references/allocazione-capitale.md", "references/asset-location.md", "SKILL.md"):
        assert "it-tax-optimization/references/successione-donazione.md" in open(os.path.join(obiettivi, rel), encoding="utf-8").read(), rel
    alloc = open(os.path.join(obiettivi, "references", "allocazione-capitale.md"), encoding="utf-8").read()
    assert "€1.000.000" in alloc and re.search(r"[Ff]ondi ed ETF[^.]*non è garantit", alloc), "bucket Eredità senza la riserva sugli ETF"
    dich = open(os.path.join(TAX, "references", "regime-dichiarativo.md"), encoding="utf-8").read()
    assert re.search(r"dividendi[^.\n]*lordo[^\n]*dividendi-esteri\.md", dich), "regime-dichiarativo: dividendi esteri sul lordo nel quadro RM"
    cripto = open(os.path.join(TAX, "references", "criptovalute-casi-speciali.md"), encoding="utf-8").read()
    assert "successione-donazione.md" in cripto and not re.search(r"per gli altri parenti fino al quarto grado e per gli affini\b", cripto), \
        "criptovalute-casi-speciali: aliquote di successione diverse dalla reference"
    assert "per ciascun beneficiario" in alloc, "bucket Eredità: la franchigia è per ciascun beneficiario"
    assert re.search(r"lordo[^.\n]*dividendi esteri|dividendi esteri[^.\n]*lordo", dich.split("## In sintesi")[1]), \
        "regime-dichiarativo, sintesi: la base dei dividendi esteri cambia"


@test("T13: le regole su eredità di ETF, BTP in donazione e dividendi in dichiarativo valgono in ogni copia")
def t_t13_propagazione():
    # REGRESSIONE revisione T13: le regole erano vincolate solo nella reference; una copia
    # che le contraddice (ETF che «azzerano la plus» in eredità, BTP esenti anche in donazione,
    # dividendi in dichiarativo sul netto) passava.
    errori = []
    for rel, testo in _testi_regole().items():
        for blocco in _blocchi(testo):
            if re.search(r"\b(ETF|fondi)\b", blocco) and re.search(r"\bered|successione", blocco, re.I) \
                    and re.search(r"plusvalenz[ae][^.\n]*azzer|azzer[^.\n]*plusvalenz", blocco):
                if not re.search(r"\bcessione|non è garantit|da verificare|può (quindi )?non esistere|non vale con certezza", blocco):
                    errori.append(f"{rel}: ETF in eredità senza riserva: «{blocco.strip()[:120]}»")
            if re.search(r"\bBTP\b|titoli di Stato|PIR", blocco) and re.search(r"esent", blocco) and re.search(r"donazion", blocco):
                for frase in _frasi(blocco):
                    if re.search(r"esent[^.\n]*donazion|donazion[^.\n]*esent", frase) and not re.search(r"\bnon\b", frase):
                        errori.append(f"{rel}: esenzione in donazione: «{frase.strip()[:120]}»")
            if re.search(r"dividend", blocco) and re.search(r"lordo[^.\n]{0,40}come in amministrato|netto frontiera,? come in amministrato|in dichiarativo[^.\n]{0,40}sul netto frontiera", blocco):
                errori.append(f"{rel}: dividendi in dichiarativo: «{blocco.strip()[:120]}»")
    assert not errori, "\n  ".join(errori[:10])


# ══════════════════════ T13: riscatto della laurea (F055) ══════════════════════

@test("T13 riscatto della laurea: onere agevolato e ordinario, netto, rata, pensione in più e anticipo, valori chiusi")
def t_riscatto_laurea():
    # F055: la description prometteva il riscatto e la skill rimandava al patronato senza l'ordine di grandezza.
    # Agevolato: €18.808 × 33% = €6.206,64 l'anno (art. 2 c. 5-quater D.Lgs. 184/1997; Circ. INPS 14/2026)
    P = asset("it-pianificazione", "riscatto-laurea.py")
    tmp = tempfile.mkdtemp()
    try:
        j = os.path.join(tmp, "r.json")
        rc, out = run_script(P, "--json", j)
        assert rc == 0, out[-400:]
        r = json.load(open(j, encoding="utf-8"))
        assert r["onere_agevolato"] == 31033.2 and r["onere_ordinario"] == 59400.0, r      # 5 × 6.206,64; 5 × 36.000 × 33%
        assert r["rata_mensile"] == 258.61 and r["costo_netto"] == 20792.24, r              # 120 rate; deduzione al 33%
        assert r["delta_montante_al_pensionamento"] == 36403.28, r                          # × 1,005^32
        assert r["delta_pensione_lorda_annua"] == 1751.46 and r["anni_per_rientrare"] == 19.1, r   # × 4,8113%; netto/(× 0,62)
        assert r["coefficiente_trasformazione"] == 0.04811 and r["eta_pensionamento"] == 69.0, r
        assert r["eta_anticipata_senza_riscatto"] == 68.33 and r["eta_anticipata_con_riscatto"] == 63.33, r
        assert r["montante_attuale"] == 85000 + 31033.2 and r["modalita"] == "agevolato", r
        for atteso in ("€ 6.206,64 (€ 18.808 × 33%)", "€ 31.033", "€ 59.400", "€ 258,61", "€ 20.792", "€ 1.751",
                       "senza riscatto: a 68,3 anni", "con il riscatto: a 63,3 anni", "rendono allo stesso modo",
                       "Deduzione dell'anno con 120 rate: € 3.103", "requisito uomini dal 2028"):
            assert atteso in out, f"riscatto: manca «{atteso}»"
        # le due modalità rendono uguale per euro versato: stessi anni per rientrare
        assert re.search(r"rientrare\s+19,1\s+19,1", out), out
        # ordinario sotto il minimale: costa meno, e lo script lo dice
        rc, out = run_script(P, "--set", "RAL_ATTUALE=15_000")
        assert rc == 0 and "sotto il minimale" in out, out[-300:]
        # unica soluzione: rata = onere, deduzione dell'anno = onere
        rc, out = run_script(P, "--set", "RATE_MENSILI=1", "--json", j)
        r = json.load(open(j, encoding="utf-8"))
        assert rc == 0 and r["rata_mensile"] == r["onere_riscatto"] == 31033.2 and "unica soluzione" in out, out[-300:]
        # revisione T13: la deduzione di un anno esce dallo scaglione del 33% e lo script lo dice
        assert "Deduzione dell'anno con l'unica soluzione: € 31.033" in out and "il costo netto vero è più alto" in out, out[-600:]
        # le donne: requisito di 42 anni e 1 mese, l'anticipata un anno prima
        rc, out = run_script(P, "--set", "SESSO='F'", "--json", j)
        r = json.load(open(j, encoding="utf-8"))
        assert rc == 0 and r["eta_anticipata_senza_riscatto"] == 67.33 and r["eta_anticipata_con_riscatto"] == 62.33, r
        assert "requisito donne" in out, out[-900:]
        # anni di corso prima del 1996: l'onere conta solo gli anni dal 1996, l'anzianità tutti
        rc, out = run_script(P, "--set", "ANNI_PRIMA_DEL_1996=2", "--json", j)
        r = json.load(open(j, encoding="utf-8"))
        assert rc == 0 and r["onere_agevolato"] == 18619.92 and r["anzianita_con_riscatto"] == 17.0, r
        assert "cadono prima del 1996" in out and "riserva matematica" in out and "3 anni nel contributivo" in out, out[:900]
        # modalità ordinaria nella catena: il montante porta l'onere ordinario
        rc, out = run_script(P, "--set", "MODALITA='ordinario'", "--json", j)
        r = json.load(open(j, encoding="utf-8"))
        assert rc == 0 and r["onere_riscatto"] == 59400.0 and r["montante_attuale"] == 85000 + 59400.0, r
        # posizione mista: l'agevolato solo per gli anni dal 1996, con l'avviso sull'opzione
        rc, out = run_script(P, "--set", "CONTRIBUTI_SOLO_DAL_1996=True")
        assert rc == 0 and "posizione mista" in out and "art. 1 c. 23 L. 335/1995" in out, out[:700]
        # anche l'ordinario degli anni prima del 1996 non è stimabile qui (riserva matematica)
        assert "art. 2 c. 4 D.Lgs. 184/1997, art. 13 L. 1338/1962" in out, out[:700]
        # chi senza riscatto arriverebbe alla vecchiaia sotto i 20 anni di contributi
        rc, out = run_script(P, "--set", "ANZIANITA_ATTUALE=0", "--set", "ETA_ATTUALE=52", "--set", "ETA_PENSIONAMENTO=67")
        assert rc == 0 and "sotto i 20 della vecchiaia" in out, out[-500:]
        # errori nei dati: rc 1 con ERRORE su stderr
        for espr, atteso in (("ANNI_RISCATTO=0", "ANNI_RISCATTO"), ("RATE_MENSILI=121", "RATE_MENSILI"),
                             ("MODALITA='gratis'", "MODALITA"), ("ALIQUOTA_MARGINALE_IRPEF=33", "frazione"),
                             ("ETA_PENSIONAMENTO=30", "ETA_PENSIONAMENTO"), ("SESSO='X'", "SESSO"),
                             ("ANNI_PRIMA_DEL_1996=6", "ANNI_PRIMA_DEL_1996")):
            rc, so, se = run_script_split(P, "--set", espr)
            assert rc == 1 and se.lstrip().startswith("ERRORE:") and atteso in se and "Traceback" not in se, (espr, rc, se[-200:])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T13 riscatto della laurea: catena importa-previdenza → riscatto → proiezione-inps, stessi numeri")
def t_riscatto_catena():
    P = lambda n: asset("it-pianificazione", n)
    tmp = tempfile.mkdtemp()
    try:
        imp, inps, ris = (os.path.join(tmp, f) for f in ("imp.json", "inps.json", "ris.json"))
        rc, out = run_script(P("importa-previdenza.py"), os.path.join(FIXTURES, "estratto-inps-test.xml"), "--json", imp)
        assert rc == 0 and "ANZIANITA_ATTUALE = 2.0" in out, out[-400:]
        assert json.load(open(imp, encoding="utf-8"))["anzianita_contributiva"] == 2.0     # 104 settimane / 52
        # revisione T13: due rapporti nello stesso anno contano una volta per il diritto (al massimo 52 settimane)
        xml = open(os.path.join(FIXTURES, "estratto-inps-test.xml"), encoding="utf-8").read()
        riga = re.findall(r"      <RigaContributi>.*?</RigaContributi>\n", xml, re.S)[-1]
        doppio = os.path.join(tmp, "doppio.xml")
        open(doppio, "w", encoding="utf-8").write(xml.replace(riga, riga + riga.replace("AZIENDA TEST SRL", "SECONDO DATORE SRL")))
        rc, out = run_script(P("importa-previdenza.py"), doppio, "--json", os.path.join(tmp, "doppio.json"))
        assert rc == 0 and "settimane utili: 156" in out and "ANZIANITA_ATTUALE = 2.0" in out, out[-600:]
        rc, out = run_script(P("proiezione-inps.py"), "--from", imp, "--json", inps)
        assert rc == 0, out[-300:]
        base = json.load(open(inps, encoding="utf-8"))
        rc, out = run_script(P("riscatto-laurea.py"), "--from", imp, "--from", inps, "--json", ris)
        assert rc == 0 and "ANZIANITA_ATTUALE=2.0" in out and "COEFF_TRASFORMAZIONE=" in out, out[:600]
        r = json.load(open(ris, encoding="utf-8"))
        m0 = json.load(open(imp, encoding="utf-8"))["montante_attuale"]
        assert r["montante_attuale"] == round(m0 + r["onere_riscatto"], 2) and r["anzianita_con_riscatto"] == 7.0, r
        # l'anticipata arriva comunque dopo l'uscita ipotizzata: il riscatto non anticipa
        assert "il riscatto non anticipa" in out and "posizione mista" not in out, out[-900:]
        # revisione T13: la posizione mista attraversa la catena (importatore → riscatto → proiezione)
        mista = os.path.join(tmp, "mista.json")
        json.dump(dict(json.load(open(imp, encoding="utf-8")), contributi_solo_dal_1996=True), open(mista, "w", encoding="utf-8"))
        rm = os.path.join(tmp, "ris-mista.json")
        rc, out_m = run_script(P("riscatto-laurea.py"), "--from", mista, "--from", inps, "--json", rm)
        assert rc == 0 and out_m.index("posizione mista") < out_m.index("Onere lordo"), out_m[:700]
        assert json.load(open(rm, encoding="utf-8"))["contributi_solo_dal_1996"] is True
        rc, out_m = run_script(P("proiezione-inps.py"), "--from", rm)
        assert rc == 0 and out_m.index("posizione mista") < out_m.index("PENSIONE LORDA ANNUA"), out_m[:700]
        # un --set di età dopo il JSON della proiezione: anni e coefficiente non valgono più, errore con il rimedio
        for espr in ("ETA_PENSIONAMENTO=64", "ETA_ATTUALE=50"):
            rc, so, se = run_script_split(P("riscatto-laurea.py"), "--from", imp, "--from", inps, "--set", espr)
            assert rc == 1 and se.lstrip().startswith("ERRORE:") and "proiezione-inps.py" in se and "Traceback" not in se, (espr, se[-300:])
        rc, out = run_script(P("proiezione-inps.py"), "--from", ris, "--json", inps)
        assert rc == 0, out[-300:]
        con = json.load(open(inps, encoding="utf-8"))
        # la proiezione completa con il riscatto e la pensione in più dello script coincidono (arrotondamenti)
        delta = con["pensione_inps_lorda_annua"] - base["pensione_inps_lorda_annua"]
        assert abs(delta - r["delta_pensione_lorda_annua"]) < 2, (delta, r["delta_pensione_lorda_annua"])
        # senza il JSON della proiezione il coefficiente è quello di proiezione-inps.py alla stessa età di uscita,
        # e la catena con un'uscita diversa dai default resta coerente (revisione T13: era fisso a 69 anni)
        b64, r64, c64 = (os.path.join(tmp, f) for f in ("b64.json", "r64.json", "c64.json"))
        rc, out = run_script(P("proiezione-inps.py"), "--from", imp, "--set", "ETA_PENSIONAMENTO=64", "--json", b64)
        rc2, out2 = run_script(P("riscatto-laurea.py"), "--from", imp, "--set", "ETA_PENSIONAMENTO=64", "--json", r64)
        assert rc == rc2 == 0, (out[-300:], out2[-300:])
        r = json.load(open(r64, encoding="utf-8"))
        assert r["coefficiente_trasformazione"] == json.load(open(b64, encoding="utf-8"))["coefficiente_trasformazione"], r
        rc, out = run_script(P("proiezione-inps.py"), "--from", r64, "--json", c64)
        assert rc == 0 and "ETA_PENSIONAMENTO=64" in out, out[:400]      # l'età di uscita viaggia nel JSON
        delta = json.load(open(c64, encoding="utf-8"))["pensione_inps_lorda_annua"] - json.load(open(b64, encoding="utf-8"))["pensione_inps_lorda_annua"]
        assert abs(delta - r["delta_pensione_lorda_annua"]) < 2, (delta, r["delta_pensione_lorda_annua"])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T13 riscatto della laurea: reference e skill con lo stesso onere, fonti nel marcatore, description")
def t_riscatto_testi():
    ref = open(os.path.join(SKILLS, "it-pianificazione", "references", "previdenza-obbligatoria.md"), encoding="utf-8").read()
    sezione = ref.split("\n## Riscatto della laurea\n")[1].split("\n## ")[0]
    for atteso in ("€18.808 × 33% = **€6.206,64**", "120 rate mensili senza interessi", "art. 10 c. 1 lett. e TUIR",
                   "dal 2027 art. 10 c. 1 lett. f D.Lgs. 117/2026", "c. 5-quater", "art. 20 c. 6 DL 4/2019",
                   "riserva matematica", "art. 1 c. 23 L. 335/1995", "€31.033", "€20.792", "€258,61", "€1.751",
                   "circa 19 anni", "del 38%", "Circ. INPS 14/2026", "Verificato: 2026-10", "(c. 5-ter) e, poiché",
                   "dal 2027 art. 10 c. 1 lett. u e art. 14 c. 1 lett. i D.Lgs. 117/2026", "per le donne 42 anni e 1 mese",
                   "ripetuti"):
        assert atteso in sezione, f"previdenza-obbligatoria.md, Riscatto della laurea: manca «{atteso}»"
    # lo stesso minimale nella reference, nello script e nella Gestione Separata degli autonomi
    sorgente = open(asset("it-pianificazione", "riscatto-laurea.py"), encoding="utf-8").read()
    assert re.search(r"^MINIMALE_ARTIGIANI_COMMERCIANTI = 18_808\.0\b", sorgente, re.M)
    gs = open(os.path.join(SKILLS, "it-pianificazione-autonomi", "references", "previdenza-gestione-separata.md"), encoding="utf-8").read()
    assert "€18.808" in gs, "minimale 2026 diverso tra le reference"
    skill = open(os.path.join(SKILLS, "it-pianificazione", "SKILL.md"), encoding="utf-8").read()
    assert "Pattern P7 — Riscattare la laurea" in skill and "assets/riscatto-laurea.py" in skill
    assert "Non gestisce ricongiunzione, totalizzazione, riscatto laurea" not in skill
    p7 = skill.split("**Pattern P7")[1].split("\n")[0]
    for atteso in ("ANNI_RISCATTO", "ANNI_PRIMA_DEL_1996", "SESSO='F'", "RATE_MENSILI", "ANZIANITA_ATTUALE"):
        assert atteso in p7, f"Pattern P7: manca «{atteso}» tra i dati da chiedere"
    assert 'riscatto-laurea.py"     --from dati.json --from inps.json' in skill, "Step 1: il riscatto nella catena"
    mappa = open(os.path.join(SKILLS, "it-tax-optimization", "references", "mappa-nuovo-tuir.md"), encoding="utf-8").read()
    assert "art. 2 c. 5-bis D.Lgs. 184/1997, ultimo periodo" in mappa and "art. 10 c. 1 lett. u e art. 14 c. 1 lett. i" in mappa


def _it(x, dec=2):
    """Importo in formato italiano, come nelle reference: 17760.92 → «17.760,92»."""
    return f"{x:,.{dec}f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".")


@test("T13 freelance: rivalsa del 4%, saldo e acconti, il secondo anno ricalcolato, eccedenza LM49 → RP21")
def t_freelance_flussi():
    # F063: rivalsa, shock del secondo anno ed eccedenza dei contributi non comparivano nella skill.
    # L'esempio si ricalcola da capo: compensi €40.000, coefficiente 78%, Gestione Separata 26,07%,
    # acconti dei contributi 40% + 40% (art. 1 c. 212 L. 662/1996), dell'imposta 50% + 50% (art. 58 DL 124/2019)
    A = os.path.join(SKILLS, "it-pianificazione-autonomi")
    gs = open(os.path.join(A, "references", "previdenza-gestione-separata.md"), encoding="utf-8").read()
    sezione = gs.split("\n## Flussi di cassa: rivalsa del 4%, saldo e acconti\n")[1].split("\n## ")[0]
    lordo = 40_000 * 0.78
    contributi = round(lordo * 0.2607, 2)
    rata_c = round(contributi * 0.40, 2)
    versati_c = contributi + 2 * rata_c
    imponibile2 = lordo - versati_c
    attesi = {f"€{_it(contributi)}", f"€{_it(versati_c)}", f"2 × €{_it(rata_c)}", f"€{_it(imponibile2)}"}
    totali = {}
    for aliquota in (0.05, 0.15):
        imposta1 = round(lordo * aliquota, 2)
        giugno = contributi + rata_c + imposta1 * 1.5
        novembre = rata_c + imposta1 * 0.5
        totali[aliquota] = round(giugno + novembre, 2)
        regime = contributi + round((lordo - contributi) * aliquota, 2)
        if aliquota == 0.05:
            imposta2 = round(imponibile2 * aliquota, 2)
            attesi |= {f"€{_it(giugno)}", f"€{_it(novembre)}", f"€{_it(imposta2)}", f"€{_it(imposta1 - imposta2)}",
                       f"€{_it(round((lordo - contributi) * aliquota, 2))}", f"**€{_it(totali[aliquota])}**",
                       f"circa {_it(totali[aliquota] / regime, 1)} volte", f"il {round(totali[aliquota] / 40_000 * 100)}% dei compensi"}
        else:
            attesi |= {f"€{_it(totali[aliquota])}", f"il {round(totali[aliquota] / 40_000 * 100)}% dei compensi"}
    for atteso in sorted(attesi):
        assert atteso in sezione, f"previdenza-gestione-separata.md, Flussi di cassa: manca «{atteso}»"
    for atteso in ("art. 1 c. 212 L. 662/1996", "due rate del 40%", "due rate del 50%", "Ris. AdE 93/E/2019",
                   "dal 2027 art. 235 c. 1 D.Lgs. 117/2026", "dal 2027 art. 72 c. 4 D.Lgs. 33/2025", "100%",
                   "Circ. INPS 8/2026", "abrogato dalla L. 113/2026", "parte integrante dei compensi",
                   "integrativo delle casse", "metodo previsionale", "Verificato: 2026-10"):
        assert atteso in sezione, f"previdenza-gestione-separata.md, Flussi di cassa: manca «{atteso}»"
    assert "prorogate al 20 luglio con D.L. 89/2026" not in gs, "il rinvio del 2026 presentato come regola"
    # nel testo, non solo nel marcatore delle fonti: il rinvio del 2026 è datato e il decreto abrogato
    corpo = sezione.split("\n*Rivalsa del 4%")[0]
    assert "abrogato dalla L. 113/2026 che ne ha fatto salvi gli effetti" in corpo and "non la regola" in corpo, corpo[-600:]
    ff = open(os.path.join(A, "references", "fiscalita-forfettario.md"), encoding="utf-8").read()
    for atteso in ("rivalsa INPS del 4%", "è un compenso, non un contributo", "integrativo", "**versati**",
                   "rigo LM49", "rigo RP21", "art. 1 c. 64 L. 190/2014, dal 2027 art. 235 c. 1 D.Lgs. 117/2026",
                   "per il forfettario puro quell'eccedenza va persa"):
        assert atteso in ff, f"fiscalita-forfettario.md: manca «{atteso}»"
    # la skill riporta gli stessi numeri della reference
    skill = open(os.path.join(A, "SKILL.md"), encoding="utf-8").read()
    a6 = skill.split("**Pattern A6")[1].split("\n")[0]
    for atteso in (f"€{_it(totali[0.05])}", f"€{_it(totali[0.15])}", "1,9 volte", "44%", "rivalsa del 4%"):
        assert atteso in a6, f"Pattern A6: manca «{atteso}»"
    assert "pattern A1-A6" in skill and "rigo LM49" in skill and "Trattare la rivalsa INPS del 4% come un contributo" in skill
    # revisione T13: i contributi si versano nella sezione INPS dell'F24 con le causali PXX o P10, non con il 1299
    for testo, nome in ((skill, "SKILL.md"), (gs, "previdenza-gestione-separata.md")):
        assert "PXX" in testo and "P10" in testo and "1299" not in testo, f"{nome}: causali F24 della Gestione Separata"

# ══════════════════════ T11: testi unici dal 2027 ══════════════════════

# Norme abrogate o rinumerate dal 1/1/2027 (mappa in it-tax-optimization/references/mappa-nuovo-tuir.md).
# Restano in vigore, e non servono qui: art. 24 c. 11 DL 201/2011, gli altri articoli del D.Lgs.
# 252/2005, art. 34 DPR 601/1973, art. 3 DL 66/2014.
_ART = r"(?i:\bartt?)\."
_NORME_2027 = re.compile(
    _ART + r"\s?\d+[\w-]*(?:[^;\n()]{0,30}?)(?<!nuovo)\s(?:TUIR|DPR 917/1986)\b"
    r"|" + _ART + r"\s?6[78]\b[^;\n()]{0,15}?\bc-(?:bis|ter|quater|quinquies|sexies)\b|\((?i:art)\.\s?44\)"
    r"|D\.\s?Lgs\.\s?461/(?:19)?97|L\.\s?77/1983|26-quinquies|DL 167/1990|D\.\s?Lgs\.\s?472/(?:19)?97"
    r"|D\.\s?Lgs\.\s?74/2000|L\.\s?228/2012|DPR 642/(?:19)?72|" + _ART + r"\s?19\b[^;\n]{0,20}?DL 201/2011|D\.\s?Lgs\.\s?346/(?:19)?90"
    r"|" + _ART + r"\s?(?:17|11 c\.\s?(?:6|4-ter|4-quater|4-quinquies))\b[^;\n()]{0,40}?D\.\s?Lgs\.\s?252/2005"
    r"|" + _ART + r"\s?27\b[^;\n()]{0,30}?DPR 600/(?:19)?73"
    r"|L\.\s?232/2016,?\s?c\.\s?(?:153 e )?158\b|\bc\.\s?158\b[^;\n]{0,10}L\.\s?232/2016"
    r"|L\.\s?190/2014|L\.\s?383/2001|L\.\s?207/2024[^.;\n]{0,10}c\.\s?24\b|1 c\.\s?(?:24|6) L\.\s?207/2024"
    r"|" + _ART + r"\s?11 c\.\s?3 D\.\s?Lgs\.\s?47/2000|D\.\s?Lgs\.\s?47/2000,? art\.\s?11 c\.\s?3\b"
    r"|lett(?:era|\.)\s?c-(?:bis|ter|quater|quinquies|sexies)"
    r"|c\.\s?5-bis,?(?: ultimo periodo,?)? D\.\s?Lgs\.\s?184/1997")
# il riferimento dal 2027: «dal 2027», una freccia o il numero di uno dei testi unici (non l'anno nudo,
# che è anche una data)
_SUCCESSORE_2027 = re.compile(r"(?i:dal (?:1/1/|1° gennaio )?2027)|→|TU 2027|117/2026|141/2026|173/2024|174/2024"
                              r"|33/2025|123/2025")


_ABBREVIAZIONI = {"art", "artt", "co", "lett", "conv", "circ", "risp", "par", "cfr", "es", "all", "mod", "fasc",
                   "ris", "lgs", "dpr", "tab", "sez", "cass", "nt", "pag", "ss"}


def _fine_frase(testo, da, a):
    """Posizione della fine della prima frase in testo[da:a] (un punto che non chiude un'abbreviazione)."""
    for f in re.finditer(r"\.\s+(?=[A-Z«])", testo[da:a]):
        parola = re.search(r"(\w+)[)»\]]*$", testo[da:da + f.start()])
        if parola and len(parola.group(1)) > 1 and parola.group(1).lower() not in _ABBREVIAZIONI:
            return da + f.start()
    return a


def _citazioni_scoperte(testo, fino=None):
    """Citazioni di norme che cambiano dal 2027 senza il successore, citazione per citazione.
    Citazioni contigue («art. 67 TUIR, lettera c-sexies») formano un gruppo. Il successore va
    cercato nella parentesi che contiene il gruppo (elenchi delle fonti: «(A, B; dal 2027 a, b)»)
    o, fuori parentesi, tra il gruppo e la citazione successiva, entro la frase. Con `fino`, solo i
    gruppi che iniziano prima di quella posizione (testo di una riga unito alla successiva)."""
    chiusa, pila = {}, []          # apertura → chiusura delle parentesi della riga
    for k, ch in enumerate(testo):
        if ch == "(":
            pila.append(k)
        elif ch == ")" and pila:
            chiusa[pila.pop()] = k
    for k in pila:
        chiusa[k] = len(testo)

    def parentesi(pos):            # la parentesi più interna che contiene pos
        dentro = [(a, c) for a, c in chiusa.items() if a < pos <= c]
        return max(dentro) if dentro else None

    m = list(_NORME_2027.finditer(testo))
    scoperte, i = [], 0
    while i < len(m):
        j = i
        while j + 1 < len(m):
            tra = testo[m[j].end():m[j + 1].start()]
            if len(tra) > 12 or "(" in tra or ")" in tra or _SUCCESSORE_2027.search(tra):
                break
            j += 1
        par = parentesi(m[i].start())
        if par:
            fine = par[1]
        else:
            fine = _fine_frase(testo, m[j].end(), m[j + 1].start() if j + 1 < len(m) else len(testo))
        if (fino is None or m[i].start() < fino) and not _SUCCESSORE_2027.search(testo, m[i].start(), fine):
            scoperte.append(testo[m[i].start():m[j].end()])
        i = j + 1
    return scoperte


def _righe_scoperte(testo, unisci=False):
    """[(numero di riga, citazione)] delle citazioni senza successore. Con `unisci` (commenti e
    docstring degli script, che vanno a capo) ogni riga si legge unita alla successiva: la
    citazione spezzata su due righe si riconosce, e il successore può stare a capo."""
    righe = testo.splitlines()
    errori = []
    for n, riga in enumerate(righe):
        if unisci and n + 1 < len(righe):
            unite = riga.rstrip() + " " + re.sub(r"^\s*#?\s*", "", righe[n + 1])
            errori += [(n + 1, c) for c in _citazioni_scoperte(unite, fino=len(riga.rstrip()))]
        else:
            errori += [(n + 1, c) for c in _citazioni_scoperte(riga)]
    return errori


_RINVIATI_T11 = ()


def _celle_scoperte(path):
    """Le celle di testo di un Excel con citazioni senza il riferimento dal 2027."""
    import openpyxl
    errori = []
    for ws in openpyxl.load_workbook(path).worksheets:
        for row in ws.iter_rows():
            for cella in row:
                if isinstance(cella.value, str) and not cella.value.startswith("="):
                    errori += [f"{os.path.basename(path)}!{ws.title}!{cella.coordinate}: «{c}»"
                               for _, c in _righe_scoperte(cella.value)]
    return errori


@test("T11: ogni norma che cambia dal 2027 è citata anche col testo unico che la sostituisce")
def t_t11_doppie_citazioni():
    # REGRESSIONE F035: dal 1/1/2027 i testi unici abrogano e rinumerano le norme citate; la
    # roadmap chiede la doppia citazione. Riga per riga, su skill, comandi e README.
    # REVISIONE T11: il controllo era per riga (un successore copriva tutte le citazioni della riga)
    # e accettava l'anno nudo; ora è per citazione, anche spezzata su due righe e negli Excel.
    for caso, atteso in (("art. 68 c. 5 TUIR, comunicati entro il 30/6/2027", 1),
                         ("art. 44 TUIR per i proventi, e l'art. 67 TUIR (dal 2027 art. 76)", 1),
                         ("art. 44 e art. 67 TUIR (dal 2027 artt. 46 e 76)", 0),
                         ("Art. 13 co. 3 TUIR, struttura vigente", 1), ("artt. 44-45 e 68 c. 7 TUIR", 1),
                         ("(art. 67 c-sexies)", 1), ("art. 27 c. 4 lett. b DPR 600/73", 1), ("L. 232/2016, c. 158", 1)):
        assert len(_citazioni_scoperte(caso)) == atteso, (caso, _citazioni_scoperte(caso))
    assert _righe_scoperte("coefficiente 48,08%, art. 6 c. 1 D.Lgs.\n    461/97): una minus", unisci=True), "citazione spezzata"
    assert not _righe_scoperte("# fondo pensione (art. 17 D.Lgs. 252/2005; L. 190/2014\n# c. 622, dal 2027 art. 257 c. 10)", unisci=True)
    assert len(_citazioni_scoperte("l'art. 6 c. 1 D.Lgs. 461/97 (conv. L. 89/2014). Dal 2027 art. 306 c. 1")) == 1, "la frase finisce"
    errori = []
    for rel, testo in sorted(_testi_regole().items()):
        if rel.startswith(_RINVIATI_T11) or rel.endswith("mappa-nuovo-tuir.md"):
            continue
        errori += [f"{rel}:{n}: «{c}»" for n, c in _righe_scoperte(testo, unisci=rel.endswith(".py"))]
    import openpyxl
    tmp = tempfile.mkdtemp()
    try:                           # il controllo delle celle trova davvero una cella scoperta
        wb = openpyxl.Workbook()
        wb.active["B3"] = "Detrazione pensionati (art. 13 c. 3 TUIR)"
        wb.save(os.path.join(tmp, "prova.xlsx"))
        assert _celle_scoperte(os.path.join(tmp, "prova.xlsx")) == ["prova.xlsx!Sheet!B3: «art. 13 c. 3 TUIR»"]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    for cartella, _, files in os.walk(SKILLS):
        for f in files:
            if f.endswith(".xlsx"):
                errori += _celle_scoperte(os.path.join(cartella, f))
    assert not errori, f"{len(errori)} citazioni senza il riferimento dal 2027:\n  " + "\n  ".join(errori[:25])
    # il +€50 dei pensionati: dal 2027 è il c. 4 dell'art. 13, non il 3-bis
    pian = os.path.join(SKILLS, "it-pianificazione")
    for rel in ("SKILL.md", "references/previdenza-obbligatoria.md"):
        for riga in open(os.path.join(pian, rel), encoding="utf-8").read().splitlines():
            if "+€50" in riga and "3-bis" in riga:
                assert "c. 4" in riga, f"it-pianificazione/{rel}: «{riga.strip()[:90]}»"
    tax = os.path.join(SKILLS, "it-tax-optimization")
    mappa = open(os.path.join(tax, "references", "mappa-nuovo-tuir.md"), encoding="utf-8").read()
    assert "## Da verificare" in mappa and mappa.rstrip().endswith("Verificato: 2026-10.*"), "mappa: sezione «Da verificare» e marcatore"
    assert "GU n. 152 del 3/7/2026" in mappa, "mappa: riferimento alla Gazzetta Ufficiale"
    # la regex conosce tutte le norme della colonna «fino al 31/12/2026» della mappa
    for riga in mappa.splitlines():
        celle = [c.strip() for c in riga.strip().strip("|").split("|")]
        if riga.startswith("| ") and len(celle) == 2 and not celle[0].startswith(("Fino al", "---")) \
                and not celle[1].startswith("invariat"):
            assert _NORME_2027.search(celle[0].replace("*", "")), f"mappa: _NORME_2027 non riconosce «{celle[0]}»"
    for riga in ("| art. 68 c. 9-bis TUIR (cripto: costo zero, proventi da detenzione) | art. 77 c. 11 |",
                 "| art. 6 D.Lgs. 461/1997 (amministrato): c. 1 (48,08%) / 1-bis / 2 / 4 / 5 / 6 / 7 | art. 306: c. 1 / 2 / 3 / 5 / 6 / 7 / 8 |"):
        assert riga in mappa, f"mappa: «{riga}»"
    skill = open(os.path.join(tax, "SKILL.md"), encoding="utf-8").read()
    assert "- `references/mappa-nuovo-tuir.md` — " in skill.split("## Riferimenti")[1], "SKILL.md: mappa fuori dall'elenco dei Riferimenti"
    regole = skill.split("## Regole fiscali fondamentali")[1].split("\n### ")[0]
    assert "references/mappa-nuovo-tuir.md" in regole, "SKILL.md: le regole fondamentali non rimandano alla mappa del 2027"


# ══════════════════════ fiscalità e tracker ══════════════════════

@test("netto-pensione: IRPEF 2026, detrazioni e aliquota rendita FP (valori chiusi)")
def t_netto():
    d = os.path.join(SKILLS, "it-pianificazione", "assets")
    sys.path.insert(0, d)
    try:
        m = carica_modulo(os.path.join(d, "netto-pensione.py"))
    finally:
        sys.path.remove(d)
    assert abs(m.irpef_lorda(30_000) - 7_100) < 0.01      # 6.440 + 660 (scaglione 33%, L.199/2025)
    assert abs(m.irpef_lorda(50_000) - 13_700) < 0.01
    assert abs(m.detrazione_pensione(8_500) - 1_955) < 0.01
    assert m.detrazione_pensione(60_000) == 0.0
    assert m.aliquota_rendita_fp(15) == 0.15
    assert abs(m.aliquota_rendita_fp(20) - 0.135) < 1e-9
    assert m.aliquota_rendita_fp(37) == 0.09              # floor
    # valore chiuso: pensione 27.138 → IRPEF 6.241,74 − detr 755,48; addizionali 2,23%
    netto, irpef, addiz = m.netto_pensionato(27_138, 0.0173, 0.005)
    att_irpef = 27_138 * 0.23 - (700 + 1_255 * (28_000 - 27_138) / 19_500)
    assert abs(irpef - att_irpef) < 0.01 and abs(addiz - 27_138 * 0.0223) < 0.01
    assert abs(netto - (27_138 - att_irpef - 27_138 * 0.0223)) < 0.01


@test("netto-pensione: ritenuta del fondo sulla sola quota imponibile della rendita (TFR e dedotti)")
def t_netto_quota_imponibile():
    # Revisione T4: lo script tassava l'intera rendita mentre le reference dicevano «come in netto-pensione.py».
    m = _modulo_previdenza("netto-pensione.py")
    assert m.QUOTA_IMPONIBILE_FP == 1.0, "default prudente: tutta la rendita"
    assert abs(m.rendita_netta_fp(10_000, 37, 0.6) - 10_000 * (1 - 0.09 * 0.6)) < 1e-9
    assert abs(m.rendita_netta_fp(10_000, 20, 1.0) - 10_000 * (1 - 0.135)) < 1e-9
    tmp = tempfile.mkdtemp()
    try:
        out_json = os.path.join(tmp, "n.json")
        rc, out = run_script(asset("it-pianificazione", "netto-pensione.py"), "--set", "QUOTA_IMPONIBILE_FP=0.6",
                             "--set", "ANNI_ISCRIZIONE_FP=37", "--json", out_json)
        assert rc == 0, out[-300:]
        d = json.load(open(out_json, encoding="utf-8"))
        assert d["quota_imponibile_fp"] == 0.6
        for sc in d["scenari"]:
            assert abs(sc["rendita_netta"] - round(sc["rendita_lorda"] * (1 - 0.09 * 0.6), 2)) < 0.011, sc
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    tv = " ".join(open(os.path.join(SKILLS, "it-pianificazione", "references", "proiezione-tenore-vita.md"),
                       encoding="utf-8").read().split())
    assert "come in netto-pensione.py" not in tv and "QUOTA_IMPONIBILE_FP" in tv, \
        "proiezione-tenore-vita.md: il rimando allo script deve nominare il parametro della quota imponibile"


@test("proiezione-pensione.xlsx: detrazione pensionati = netto-pensione.py (art. 13 c. 3 TUIR)")
def t_excel_detrazione():
    # REGRESSIONE F008: Pubblica!B64 usava una formula inventata (1.880 € fissi oltre 28.000).
    from openpyxl import load_workbook
    orig = asset("it-pianificazione", "proiezione-pensione.xlsx")
    ws = load_workbook(xlsx_ricalcolati()[orig], data_only=True)["Pubblica"]
    cartella = os.path.dirname(orig)
    sys.path.insert(0, cartella)  # netto-pensione importa motore_bootstrap
    try:
        netto = carica_modulo(asset("it-pianificazione", "netto-pensione.py"))
    finally:
        sys.path.remove(cartella)
    lorda, detrazione = ws["B61"].value, ws["B64"].value
    attesa = netto.detrazione_pensione(lorda)
    assert abs(detrazione - attesa) < 0.01, f"B64 = {detrazione:.2f} su lorda {lorda:,.0f}, attesa {attesa:.2f}"


@test("parametri previdenziali: nessun valore superato né dato personale dell'autore")
def t_parametri_superati():
    # REGRESSIONE F002/F009/F051: massimale 2025 spacciato per 2026, plafond "era €5.300",
    # "Quota 104" (mai esistita), riferimenti all'autore nelle reference.
    vietati = ("~€120.607", "era €5.300", "Quota 104", "Eugenio", "Calabria", "no, attenzione")
    for base in (SKILLS, os.path.join(ROOT, "commands")):
        for cartella, _, files in os.walk(base):
            for f in files:
                if f.endswith((".md", ".py")):
                    testo = open(os.path.join(cartella, f), encoding="utf-8").read()
                    for v in vietati:
                        assert v not in testo, f"{os.path.relpath(os.path.join(cartella, f), ROOT).replace(os.sep, '/')}: «{v}»"


@test("tracker-secchielli: fogli, formule vive, semaforo demo funzionante")
def t_tracker():
    from openpyxl import load_workbook
    p = asset("it-revisione", "tracker-secchielli.xlsx")
    wb = load_workbook(p)
    for foglio in ("Leggimi", "Impostazioni", "Glide", "Posizioni", "Controllo", "Versamenti"):
        assert foglio in wb.sheetnames, f"foglio mancante: {foglio}"
    assert str(wb["Controllo"]["B5"].value).startswith("=IF") and "SUMIFS" in str(wb["Controllo"]["B5"].value)
    assert len(wb["Posizioni"].data_validations.dataValidation) >= 2, "tendine mancanti"
    wv = load_workbook(p, data_only=True)  # il file è recalcolato: i valori demo devono esserci
    c = wv["Controllo"]
    assert c["G7"].value == "RIBILANCIA", "il semaforo demo (Lungo termine oltre banda) non scatta"
    assert c["G6"].value == "OK" and c["B18"].value and c["B18"].value > 100_000


@test("ponte-pensione: proprietà limite (capitale enorme, spesa zero, RITA copre tutto)")
def t_ponte():
    sc = asset("it-pianificazione", "ponte-pensione.py")
    def prob(*sets):
        rc, out = run_script(sc, *sets)
        assert rc == 0, out[-300:]
        m = re.search(r"PROBABILITÀ di coprire il ponte: ([\d,]+)%", out)
        return _da_it(m.group(1))
    assert prob("--set", "CAPITALE_PONTE=10_000_000.0") == 100.0
    assert prob("--set", "SPESA_ANNUA_PONTE=0.0", "--set", "CAPITALE_PONTE=0.0") == 100.0
    # RITA dall'uscita con rata netta > spesa: il capitale non serve
    assert prob("--set", "CAPITALE_PONTE=0.0", "--set", "ANNI_MAX_RITA=10",
                "--set", "SPESA_ANNUA_PONTE=20000.0") == 100.0
    # e il capitale necessario per confidenza è monotono nel senso giusto: con RITA
    # ordinaria (5 anni) e capitale 0, la probabilità NON può essere 100
    assert prob("--set", "CAPITALE_PONTE=0.0", "--set", "ANNI_MAX_RITA=5") < 100.0
    # REGRESSIONE F056: rendimento reale del fondo a zero → ZeroDivisionError. A r = 0 i
    # versamenti mancati valgono versamento × anni (annuità lineare).
    rc, out = run_script(sc, "--set", "RENDIMENTO_FP_REALE=0.0", "--set", "VERSAMENTO_FP_ANNUO=1000.0",
                         "--set", "COEFF_TRASFORMAZIONE=0.05")
    assert rc == 0, out[-300:]
    anni = int(re.search(r"\((\d+) anni di contributi in meno", out).group(1))
    assert euro(out, "Rendita fondo:") == 50 * anni, "rendita persa a r = 0 ≠ versamento × anni × coeff."


@test("obiettivi: il fondo pensione al plafond di €5.300 presuppone nessun contributo del datore (plafond condiviso)")
def t_obiettivi_plafond_condiviso():
    # Nota del § 4 dell'handoff (dopo T4 e T8): «€5.300 al 33%» vale solo senza contributo del datore,
    # che conta nel plafond insieme al dovuto del lavoratore (art. 8 c. 4 D.Lgs. 252/2005).
    cartella = os.path.join(SKILLS, "it-pianificazione-obiettivi")
    for rel in ("SKILL.md", "references/asset-location.md", "assets/asset-location-optimizer.py"):
        testo = open(os.path.join(cartella, rel), encoding="utf-8").read()
        for par in testo.split("\n\n"):
            if re.search(r"5\.300|CONTRIBUTO_ANNUO_FP = 5300", par):
                assert "datore" in par, f"{rel}: €5.300 senza dire che il plafond è condiviso col datore: {par.strip()[:90]}"
    rc, out = run_script(asset("it-pianificazione-obiettivi", "asset-location-optimizer.py"))
    assert rc == 0 and re.search(r"plafond[^\n]*datore", out), "lo stdout deve ricordare il plafond condiviso"
    # oltre il plafond niente deduzione: con 8.000 €/anno il risparmio IRPEF è quello di 5.300 (revisione)
    al = _modulo_obiettivi("asset-location-optimizer.py")
    def driver(*sets):
        rc, out = run_script(asset("it-pianificazione-obiettivi", "asset-location-optimizer.py"),
                             "--set", "REINVESTI_RISPARMIO_IRPEF=False", *sets)
        assert rc == 0, out[-300:]
        return out, float(re.search(r"Driver: la deduzione IRPEF \(€([\d.]+)\)", out).group(1).replace(".", ""))
    out8, d8 = driver("--set", "CONTRIBUTO_ANNUO_FP=8000")
    _, d53 = driver("--set", "CONTRIBUTO_ANNUO_FP=5300")
    atteso = 5_300 * al.ALIQUOTA_IRPEF_MARGINALE * al.ANNI_A_PENSIONE
    assert abs(d8 - atteso) <= 1 and abs(d53 - atteso) <= 1, f"deduzione oltre il plafond: {d8} e {d53}, attesi {atteso:,.0f}"
    assert "dedotti solo €5.300/anno" in out8


@test("ponte-pensione: serie storica deflazionata anno per anno con la CPI USA, come il decumulo")
def t_ponte_reale():
    # Nota del § 4 dell'handoff (dopo T8): il ponte confrontava rendimenti NOMINALI in dollari con una
    # spesa in euro di oggi. Sui default la probabilità scende da circa 44% a circa 21%.
    righe = open(asset("it-pianificazione", "rendimenti-storici.csv"), encoding="utf-8").read().splitlines()
    dati = {int(r.split(",")[0]): [float(x) / 100 for x in r.split(",")[1:]] for r in righe if r and r[0].isdigit()}
    mod = _modulo_previdenza("ponte-pensione.py")
    assert mod.INFLAZIONE is None and mod.INFLAZIONE_RIPIEGO == 0.02
    peso = mod.PESO_AZIONARIO_PONTE
    port = mod.rendimenti_reali_ponte()
    for anno in (1974, 1980, 2008, 2022):
        az, ob, inf = dati[anno]
        atteso = peso * ((1 + az) / (1 + inf) - 1) + (1 - peso) * ((1 + ob) / (1 + inf) - 1)
        assert abs(port[anno - 1928] - atteso) < 1e-12, f"{anno}: reale {port[anno - 1928]:.4f}, atteso {atteso:.4f}"
    mod.INFLAZIONE = 0.02  # override esplicito: inflazione costante
    az, ob, _ = dati[1974]
    assert abs(mod.rendimenti_reali_ponte()[1974 - 1928] - (peso * ((1 + az) / 1.02 - 1) + (1 - peso) * ((1 + ob) / 1.02 - 1))) < 1e-12
    # valore chiuso sui default (seed 42): la serie reale, non quella nominale
    rc, out = run_script(asset("it-pianificazione", "ponte-pensione.py"))
    assert rc == 0, out[-300:]
    prob = _da_it(re.search(r"PROBABILITÀ di coprire il ponte: ([\d,]+)%", out).group(1))
    assert prob == PROB_PONTE_DEFAULT, f"probabilità sui default {prob}% ≠ {PROB_PONTE_DEFAULT}%"
    assert "CPI" in out, "lo stdout deve dire con quale inflazione è deflazionata la serie"


PROB_PONTE_DEFAULT = 21.1


@test("pensione: la serie storica è dichiarata nominale in dollari, non «reale», in script, SKILL e reference")
def t_serie_pensione_nominale():
    # Nota del § 4 dell'handoff (critica dell'audit a T8): montecarlo-pensione e it-pianificazione
    # chiamavano «reale» una serie nominale in USD.
    for rel in ("assets/montecarlo-pensione.py", "SKILL.md", "references/montecarlo-pensione.md"):
        testo = " ".join(open(os.path.join(SKILLS, "it-pianificazione", rel), encoding="utf-8").read().split())
        assert "serie storica reale" not in testo and "FORMA reale" not in testo, f"{rel}: serie storica detta reale"
        assert re.search(r"nominale in dollari", testo, re.I) and "ricentrat" in testo, \
            f"{rel}: dire che la serie è nominale in dollari e che se ne usa solo la forma (media ricentrata)"
    # senza ricentratura la serie si deflaziona: mai la media nominale in dollari contro euro di oggi
    import numpy as np
    mc = _modulo_previdenza("montecarlo-pensione.py")
    righe = open(asset("it-pianificazione", "rendimenti-storici.csv"), encoding="utf-8").read().splitlines()
    dati = np.array([[float(x) / 100 for x in r.split(",")[1:]] for r in righe if r and r[0].isdigit()])
    az, ob, inf = dati[:, 0], dati[:, 1], dati[:, 2]
    w = mc.FP_PESO_AZIONARIO
    assert abs(mc.serie_portafoglio(az, ob).mean() - mc.FP_RENDIMENTO_ATTESO_REALE) < 1e-12, "media ricentrata"
    mc.RICENTRA_MEDIA = False
    reale = w * ((1 + az) / (1 + inf) - 1) + (1 - w) * ((1 + ob) / (1 + inf) - 1)
    assert np.allclose(mc.serie_portafoglio(az, ob), reale, atol=1e-12), "RICENTRA_MEDIA=False: serie non deflazionata"


@test("cruscotto TWR: ultima riga dinamica (5 o 14 rilevazioni), righe vuote senza −100%")
def t_twr_dinamico():
    # REGRESSIONE F049: il foglio funzionava solo con 13 righe esatte (meno → −100%,
    # di più → TWR, XIRR e flusso terminale fermi alla riga 18).
    import datetime
    from openpyxl import load_workbook
    mod = ricalcolatore()
    tmp = tempfile.mkdtemp()
    try:
        orig = asset("it-revisione", "cruscotto-patrimoniale.xlsx")
        corta, lunga = os.path.join(tmp, "corta.xlsx"), os.path.join(tmp, "lunga.xlsx")
        wb = load_workbook(orig)
        for r in range(11, 19):  # compilati solo i primi 5 mesi
            for col in "ABC":
                wb["Rendimento TWR"][f"{col}{r}"] = None
        wb.save(corta)
        wb = load_workbook(orig)
        ws = wb["Rendimento TWR"]
        ws["A19"], ws["B19"], ws["C19"] = datetime.datetime(2026, 2, 1), 90500, 750
        wb.save(lunga)
        out_corta, out_lunga = mod.ricalcola([corta, lunga], os.path.join(tmp, "out"))
        c = load_workbook(out_corta, data_only=True)["Rendimento TWR"]
        assert c["I5"].value == 5, f"rilevazioni contate: {c['I5'].value!r}"
        assert abs(c["I6"].value - (c["E10"].value - 1)) < 1e-9, "TWR non legge l'ultima riga compilata"
        assert isinstance(c["I9"].value, float), f"XIRR su 5 righe: {c['I9'].value!r}"
        for r in range(11, 19):
            for col in "DEF":
                assert c[f"{col}{r}"].value in (None, ""), f"{col}{r} = {c[f'{col}{r}'].value!r} su riga vuota"
        lw = load_workbook(out_lunga, data_only=True)["Rendimento TWR"]
        assert lw["I5"].value == 14 and abs(lw["I6"].value - (lw["E19"].value - 1)) < 1e-9, "la 14ª riga non entra"
        assert lw["F18"].value == -750 and lw["F19"].value == -750 + 90500, "flusso terminale non spostato"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("cruscotto: Leggimi, Sintesi con soglia per priorità, bande 5/25, costi e riga ISC")
def t_cruscotto_revisione():
    # REGRESSIONE F064 (soglie fisse 70/50), F030 (banda ±5 uguale per tutte le classi),
    # F050 (nessun controllo dei costi), F049 (limiti non dichiarati).
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_ricalcolati()[asset("it-revisione", "cruscotto-patrimoniale.xlsx")], data_only=True)
    assert wb.sheetnames[0] == "Leggimi", f"fogli: {wb.sheetnames}"
    assert any("120" in str(c.value) for r in wb["Leggimi"].iter_rows() for c in r), "Leggimi senza i limiti"
    s = wb["Sintesi"]
    assert (s["G5"].value, s["H5"].value) == ("Priorità", "Soglia"), (s["G5"].value, s["H5"].value)
    assert s["F6"].value == "On track" and s["F7"].value == "A rischio", (s["F6"].value, s["F7"].value)
    a = wb["Allocazione"]
    assert a["F5"].value == 0.05 and abs(a["F7"].value - 0.0375) < 1e-12, (a["F5"].value, a["F7"].value)
    assert str(a["E9"].value).startswith("Sottopeso"), f"alternativi 0% vs target 5%: {a['E9'].value!r}"
    an = wb["Andamento"]
    etichette = {str(an[f"A{r}"].value): r for r in range(1, 40) if an[f"A{r}"].value}
    assert any(e.startswith("COSTI DELL'ANNO") for e in etichette), "blocco costi assente"
    bollo = next(r for e, r in etichette.items() if e.startswith("Bollo"))
    assert abs(an[f"B{bollo}"].value - 88000 * 0.002) < 1e-9, f"bollo: {an[f'B{bollo}'].value!r}"
    assert any("ISC" in str(wb["Azioni"][f"A{r}"].value) for r in range(5, 12)), "manca la riga demo sull'ISC"


@test("bande di ribilanciamento: regola 5/25 in cruscotto, tracker e IPS")
def t_bande_5_25():
    # REGRESSIONE F030: stesso portafoglio «in banda» nel cruscotto e fuori banda nello script.
    from openpyxl import load_workbook
    def formula(nome, foglio, cella):
        return str(load_workbook(asset("it-revisione", nome))[foglio][cella].value)
    assert "MIN($B$12" in formula("cruscotto-patrimoniale.xlsx", "Allocazione", "F5")
    assert "MIN(Impostazioni!$B$14" in formula("tracker-secchielli.xlsx", "Controllo", "F15")
    assert "MIN(0.05" in formula("investment-policy-statement.xlsx", "IPS", "C21")


@test("revisione: costi dell'anno nel metodo e nel template, nessun percorso di sandbox")
def t_revisione_testi():
    cartella = os.path.join(SKILLS, "it-revisione")
    for rel in ("references/metodo-revisione.md", "assets/template-revisione.md"):
        assert "Costi dell'anno" in open(os.path.join(cartella, rel), encoding="utf-8").read(), rel
    for nome in sorted(os.listdir(SKILLS)):
        f = os.path.join(SKILLS, nome, "SKILL.md")
        if os.path.exists(f):
            assert "/mnt/skills" not in open(f, encoding="utf-8").read(), f"{nome}/SKILL.md cita recalc.py del sandbox claude.ai"


@test("revisione: spazio deducibile residuo del fondo (plafond − datore − dovuto) e perdita sui non dedotti non comunicati")
def t_revisione_spazio_deducibile():
    # Nota del § 4 dell'handoff (T10 dopo T4): la revisione di fine anno deve dire quanto resta
    # deducibile e quanto costa non comunicare i non dedotti.
    for rel in ("skills/it-revisione/SKILL.md", "skills/it-revisione/references/metodo-revisione.md",
                "skills/it-revisione/assets/template-revisione.md"):
        testo = " ".join(open(os.path.join(ROOT, rel), encoding="utf-8").read().split())
        assert "spazio deducibile residuo" in testo, f"{rel}: manca lo spazio deducibile residuo"
        assert "max(0, €5.300 − datore − dovuto − volontari già versati nell'anno)" in testo, \
            f"{rel}: lo spazio residuo è max(0, plafond − datore − dovuto − volontari già versati)"
        assert "forfettario" in testo and "altri redditi IRPEF" in testo, f"{rel}: manca il caso del forfettario"
        assert re.search(r"non dedott[^.]*9-15%[^.]*23%", testo), \
            f"{rel}: perdita sui non dedotti non comunicati: 9-15%, 23% con anticipazioni e riscatti per altre cause"


@test("scadenze di fine anno: ovunque c'è il fondo pensione c'è anche la comunicazione dei non dedotti")
def t_non_dedotti():
    # REGRESSIONE F020: senza comunicazione al fondo entro il 31/12 dell'anno successivo,
    # i contributi non dedotti vengono tassati di nuovo in uscita.
    from openpyxl import load_workbook
    for rel in ("skills/it-revisione/SKILL.md", "skills/it-revisione/references/metodo-revisione.md",
                "skills/it-revisione/assets/template-revisione.md",
                "skills/it-revisione/references/investment-policy-statement.md",
                "skills/it-revisione/assets/genera-ips.py",
                "skills/it-pianificazione-autonomi/references/fiscalita-forfettario.md"):
        assert "non dedott" in open(os.path.join(ROOT, rel), encoding="utf-8").read(), f"{rel}: manca la comunicazione"
    for nome, foglio, cella in (("cruscotto-patrimoniale.xlsx", "Azioni", "A2"),
                                ("investment-policy-statement.xlsx", "IPS", "A47")):
        testo = load_workbook(asset("it-revisione", nome))[foglio][cella].value
        assert "non dedott" in testo, f"{nome} {foglio}!{cella}: manca la comunicazione"


@test("autonomi: artigiani e commercianti fuori perimetro, casse CIPAG ed ENPAM aggiornate")
def t_autonomi_perimetro():
    # REGRESSIONE F016/F038/F009: artigiani-commercianti accettati e trattati come Gestione
    # Separata; CIPAG integrativo al 4%; ENPAM senza aliquota di Quota B; soglie superate.
    cartella = os.path.join(SKILLS, "it-pianificazione-autonomi")
    skill = open(os.path.join(cartella, "SKILL.md"), encoding="utf-8").read()
    assert "/ artigiani-commercianti" not in skill, "artigiani-commercianti ancora tra i tipi di cassa in input"
    assert re.search(r"description:[^\n]*artigian", skill), "description senza l'esclusione di artigiani e commercianti"
    comando = open(os.path.join(ROOT, "commands", "pensione-autonomi.md"), encoding="utf-8").read()
    assert "artigian" in comando, "il comando non esclude artigiani e commercianti"
    casse = open(os.path.join(cartella, "references", "casse-professionali.md"), encoding="utf-8").read()
    cipag = next(r for r in casse.splitlines() if r.startswith("| **CIPAG**"))
    enpam = next(r for r in casse.splitlines() if r.startswith("| **ENPAM**"))
    assert "5%" in cipag and "Verificato: 2026-10" in cipag, f"riga CIPAG: {cipag[:120]}"
    assert "19,5" in enpam and "Verificato: 2026-10" in enpam, f"riga ENPAM: {enpam[:120]}"
    ff = open(os.path.join(cartella, "references", "fiscalita-forfettario.md"), encoding="utf-8").read()
    for testo, rel in ((skill, "SKILL.md"), (ff, "fiscalita-forfettario.md")):
        avvisi = [r for p in testo.split("\n\n") for r in p.splitlines() if "artigian" in r and re.search(r"coefficient\w* (?:di redditività )?del \d+%", r)]
        assert avvisi, f"{rel}: manca l'avviso su artigiani e commercianti accanto ai coefficienti"
        for avviso in avvisi:
            for coeff in ("40%", "54%", "62%", "86%"):
                assert coeff in avviso, f"{rel}: l'avviso «{avviso.strip()[:60]}…» non cita il coefficiente {coeff}"
    gs = open(os.path.join(cartella, "references", "previdenza-gestione-separata.md"), encoding="utf-8").read()
    assert "Σ (reddito_anno × 0,25)" in gs and "0,2607) rivalutato" not in gs, "Gestione Separata: a montante va il 25% IVS"
    assert "× 0,25 fino al pensionamento" in skill, "SKILL.md: montante al 26,07% invece del 25% di IVS"
    for rel, testo in _file_testo(cartella, os.path.join(SKILLS, "it-pianificazione"), os.path.join(ROOT, "commands")):
        for riga in testo.splitlines():
            assert not re.search(r"aliquota 26,07% vs 33%|separata ~0,2607", riga), f"{rel}: 26,07% come aliquota del montante: {riga.strip()[:100]}"
    for superato in ("1,5× assegno sociale", "considerabile anche la rendita della complementare"):
        assert superato not in gs, f"previdenza-gestione-separata.md: «{superato}»"


# ═══════════════════════════ T15: capitale umano ═══════════════════════════

CU = ("it-pianificazione-obiettivi", "capitale-umano.py")


@test("T15 capitale umano: valori chiusi della valutazione")
def t_cu_valutazione():
    m = carica_modulo(asset(*CU))
    # rendita: g = 0, frazione 1, ν = 0 → formula chiusa
    assert abs(m.valore_capitale_umano(32000, 0, 32, 0.01, 0, 1) - 32000 * (1 - 1.01 ** -32) / 0.01) < 0.01
    assert m.valore_capitale_umano(32000, 0.01, 0, 0.01, 0.03, 1) == 0
    assert abs(m.frazione_lavorata(0.08, 0.5, 0.6) - 0.984) < 1e-12
    assert abs(m.valore_pensione(18000, 32, 55, 0.01, 0.0075) - 194248.56) < 0.01
    assert m.quota_azionaria_effettiva(0.3, 1) == 1 and abs(m.quota_azionaria_effettiva(0.3, 0.5) - 0.65) < 1e-12
    # invarianti: il capitale umano decresce con ν e con p_u, cresce con g
    base = m.valore_capitale_umano(35000, 0.01, 32, 0.01, 0.03, 0.984)
    assert abs(base - 705036.01) < 0.01, base
    assert m.valore_capitale_umano(35000, 0.01, 32, 0.01, 0.04, 0.984) < base
    assert m.valore_capitale_umano(35000, 0.01, 32, 0.01, 0.03, m.frazione_lavorata(0.16, 0.5, 0.6)) < base
    assert m.valore_capitale_umano(35000, 0.02, 32, 0.01, 0.03, 0.984) > base


@test("T15 capitale umano: profili con fonte, natura e valori plausibili")
def t_cu_profili():
    m = carica_modulo(asset(*CU))
    p = m.carica_profili()
    assert p == m.carica_profili(asset("it-pianificazione-obiettivi", "profili-capitale-umano.csv"))
    assert set(p) == {"pubblico", "privato_stabile", "privato_ciclico", "finanza", "autonomo", "imprenditore"}, set(p)
    assert p["privato_ciclico"]["quota_azionaria_hc"] == 0.30 and p["privato_ciclico"]["nu"] == 0.03
    assert p["finanza"]["settore_default"] == "Financials"
    for nome, r in p.items():
        assert r["fonte"].strip() and r["natura"].strip(), nome
        assert 0 <= r["quota_azionaria_hc"] <= 1 and 0 <= r["nu"] <= 0.10 and -1 <= r["rho"] <= 1, nome
        assert 0 <= r["sigma_h"] <= 0.5 and 0 <= r["p_u"] <= 0.5 and 0 <= r["sostituzione_naspi"] <= 1, nome
        assert r["settore_default"] in (None, *m.SETTORI_GICS), nome
    e = m.profilo_effettivo({"profilo": "privato_ciclico", "override": {"p_u": 0.08}}, p)
    assert e["p_u"] == 0.08 and p["privato_ciclico"]["p_u"] == 0.05   # l'override non tocca la tabella


def _cu(*args, cwd=None):
    return run_script_split(asset(*CU), *args, cwd=cwd)


def _cu_esempio():
    return json.load(open(asset("it-pianificazione-obiettivi", "esempio-capitale-umano.json"), encoding="utf-8"))


def _scrivi_json(path, dati):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(dati, f)
    return path


@test("T15 capitale umano: bilancio, esposizione ed equivalenza sull'esempio")
def t_cu_bilancio():
    tmp = tempfile.mkdtemp()
    try:
        out_json = os.path.join(tmp, "cu.json")
        rc, so, se = _cu("--json", out_json, cwd=tmp)
        assert rc == 0 and "dati d'esempio" in so, se[-300:]
        j = json.load(open(out_json, encoding="utf-8"))
        b = j["bilancio"]
        assert abs(b["HC"] - 705036.01) < 0.01 and abs(b["PP"] - 194248.56) < 0.01 and b["F"] == 113000, b
        assert abs(j["esposizione_finanziaria"] - 0.610619) < 1e-5 and abs(j["esposizione_totale"] - 0.277107) < 1e-5
        assert j["equivalenza"] is None
        m = carica_modulo(asset(*CU))
        assert m.equivalenza(0.5, b)["limite"] == "max" and m.equivalenza(0.5, b)["x"] == 1
        assert m.equivalenza(0.1, b)["limite"] == "min" and m.equivalenza(0.1, b)["x"] == 0
        assert abs(m.equivalenza(0.25, b)["x"] - 0.361781) < 1e-5
        assert "calcolo, non target" in m.equivalenza(0.25, b)["messaggio"]
        # s = 0 con un portafoglio azionario: esposizione totale < finanziaria
        dati = _cu_esempio()
        dati["persona"]["override"]["quota_azionaria_hc"] = 0
        rc, so, se = _cu("--dati", _scrivi_json(os.path.join(tmp, "s0.json"), dati), "--json", out_json, cwd=tmp)
        j = json.load(open(out_json, encoding="utf-8"))
        assert rc == 0 and j["esposizione_totale"] < j["esposizione_finanziaria"], se[-300:]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T15 capitale umano: dati incoerenti respinti con rc 1, casi limite gestiti")
def t_cu_errori():
    tmp = tempfile.mkdtemp()
    def con(modifica):
        d = _cu_esempio()
        modifica(d)
        return _scrivi_json(os.path.join(tmp, "d.json"), d)
    try:
        for modifica, atteso in [
            (lambda d: d["persona"].update(profilo="astronauta"), "profilo"),
            (lambda d: d["persona"].update(eta=95), "eta_fine"),
            (lambda d: d["persona"].update(eta_pensione=30), "eta_pensione"),
            (lambda d: d["fondo_pensione"].update(quota_azionaria=60), "quota"),
            (lambda d: d.update(liquidita=-1), "negativ"),
            (lambda d: d.update(persone=[d.pop("persona")]), "persona"),
            (lambda d: d["persona"].update(tfr_nel_fondo="sì"), "tfr_nel_fondo"),
        ]:
            rc, so, se = _cu("--dati", con(modifica), cwd=tmp)
            assert rc == 1 and se.lstrip().startswith("ERRORE:") and atteso in se, (atteso, rc, se[-200:])
        # portafoglio vuoto con l'equivalenza richiesta: niente divisione per zero
        def vuoto(d):
            d.update(posizioni=[], liquidita=0, esposizione_totale_desiderata=0.3)
        rc, so, se = _cu("--dati", con(vuoto), "--json", os.path.join(tmp, "o.json"), cwd=tmp)
        assert rc == 0, se[-300:]
        j = json.load(open(os.path.join(tmp, "o.json"), encoding="utf-8"))
        assert j["equivalenza"]["x"] is None and "nessun portafoglio" in j["equivalenza"]["messaggio"]
        # asset_class sconosciuta: non azionaria, con un avviso che la nomina
        rc, so, se = _cu("--dati", con(lambda d: d["posizioni"].append({"asset_class": "oro", "valore": 5000})), cwd=tmp)
        assert rc == 0 and "oro" in so and "non azionari" in so, so[-400:]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T15 capitale umano: proiezione, fascia Monte Carlo e caso di Ibbotson")
def t_cu_proiezione():
    m = carica_modulo(asset(*CU))
    dati = _cu_esempio()
    prof = m.profilo_effettivo(dati["persona"], m.carica_profili())
    c = {"R_F": 0.01, "NU_PENSIONE": 0.0075, "ETA_FINE": 90, "ETA_PENSIONE": 67, "PENSIONE_NETTA": 18000}
    p = m.proiezione(dati, prof, c, None, 500, 7)
    righe = {r["eta"]: r for r in p["anni"]}
    assert min(righe) == 35 and max(righe) == 72, sorted(righe)[:3]
    hc = [righe[e]["HC"] for e in range(35, 68)]
    assert all(a > b for a, b in zip(hc, hc[1:])) and righe[67]["HC"] == 0 and righe[70]["HC"] == 0
    assert righe[66]["PP"] > righe[35]["PP"] and righe[72]["PP"] < righe[67]["PP"]
    assert all(r["p10"] <= r["p50"] <= r["p90"] for r in p["anni"])
    assert p == m.proiezione(dati, prof, c, None, 500, 7)                      # riproducibile col seme
    # ρ = σ_h = 0 e σ_S = 0: nessuna dispersione, fascia sul deterministico
    m.SIGMA_AZ = 0.0
    piatto = dict(prof, rho=0.0, sigma_h=0.0)
    for r in m.proiezione(dati, piatto, c, None, 200, 7)["anni"]:
        assert abs(r["p10"] - r["p90"]) < 1e-9 and abs(r["p50"] - r["esposizione"]) < 1e-9, r
    m.SIGMA_AZ = 0.18
    # glide path più corto della proiezione: si tiene l'ultima quota, dicendolo
    g = m.proiezione(dati, prof, c, [0.6, 0.5, 0.4], 200, 7)
    assert any("più corto" in a for a in g["avvisi"])
    # già in pensione
    pens = _cu_esempio()
    pens["persona"].update(eta=68, netto_annuo=0, tfr_annuo=0)
    r = m.proiezione(pens, prof, dict(c), None, 200, 7)
    assert r["eta_sorpasso"] is None and all(x["HC"] == 0 for x in r["anni"])
    # Ibbotson et al. 2007, fig. 2.3: (HC + PP)/W ≈ 94% a 25 anni (tolleranza larga: letto da un grafico)
    ib = {"persona": {"profilo": "pubblico", "eta": 25, "eta_pensione": 65, "netto_annuo": 50000, "crescita_reale": 0,
                      "tfr_annuo": 0, "contributo_datore_fp": 0, "quota_reddito_azioni_datore": 0,
                      "override": {"p_u": 0, "nu": 0.04}},
          "pensione": {"netta_annua": 10000, "eta_fine": 85}, "posizioni": [], "liquidita": 50000,
          "fondo_pensione": {"valore": 0, "quota_azionaria": 0}, "tfr_maturato_azienda": 0, "risparmio_annuo": 5000}
    pi = m.profilo_effettivo(ib["persona"], m.carica_profili())
    b = m.bilancio(ib, pi, {"R_F": 0.02, "NU_PENSIONE": 0.04, "ETA_FINE": 85, "ETA_PENSIONE": 65, "PENSIONE_NETTA": 10000})
    assert 0.92 <= (b["HC"] + b["PP"]) / b["W"] <= 0.96, (b["HC"], b["PP"], b["W"])


@test("T15 capitale umano: sovrapposizione settoriale e CSV da Excel italiano")
def t_cu_settore():
    m = carica_modulo(asset(*CU))
    assert set(m.PESI_MONDO) == set(m.SETTORI_GICS) and 0.99 <= sum(m.PESI_MONDO.values()) <= 1.01
    tmp = tempfile.mkdtemp()
    try:
        p = os.path.join(tmp, "comp.csv")
        with open(p, "w", encoding="utf-8-sig") as f:              # BOM, «;» e percento con la virgola
            f.write("strumento;settore;peso\nA;Information Technology;50,0\nA;Financials;50,0\n"
                    "B;Health Care;1,0\n")
        comp = m.leggi_composizione(p)
        assert comp == {"A": {"Information Technology": 0.5, "Financials": 0.5}, "B": {"Health Care": 1.0}}, comp
        pos = [{"asset_class": "azionario", "valore": 30000, "strumento": "A"},
               {"asset_class": "azionario", "valore": 10000, "strumento": "B"},
               {"asset_class": "azionario", "valore": 10000, "strumento": "C"}]
        s = m.sovrapposizione(pos, comp, "Information Technology", 40000, 60000,
                              {**m.PESI_MONDO, "Information Technology": 0.25})
        assert abs(s["quota_settore_fin"] - 15000 / 40000) < 1e-12 and s["mancanti"] == ["C"], s
        assert abs(s["esposizione_combinata"] - (40000 * 0.375 + 60000) / 100000) < 1e-12
        assert abs(s["rapporto"] - 1.5) < 1e-12 and s["sovrappeso"] is False
        # l'esempio del plugin passa da riga di comando
        rc, so, se = run_script_split(asset(*CU), "--composizione",
                                      asset("it-pianificazione-obiettivi", "esempio-composizione.csv"), cwd=tmp)
        assert rc == 0 and "4. SOVRAPPOSIZIONE" in so, se[-300:]
        with open(p, "w", encoding="utf-8") as f:
            f.write("strumento;settore;peso\nA;Tecnologia;100\n")
        rc, so, se = run_script_split(asset(*CU), "--composizione", p, cwd=tmp)
        assert rc == 1 and "Information Technology" in se, se[-300:]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T15 capitale umano: sensibilità ordinata e catena --from dagli script previdenziali")
def t_cu_sensibilita_catena():
    m = carica_modulo(asset(*CU))
    dati = _cu_esempio()
    prof = m.profilo_effettivo(dati["persona"], m.carica_profili())
    c = {"R_F": 0.01, "NU_PENSIONE": 0.0075, "ETA_FINE": 90, "ETA_PENSIONE": 67, "PENSIONE_NETTA": 18000}
    righe = m.sensibilita(dati, prof, c)
    assert {r["parametro"] for r in righe} == {"s", "nu", "g", "p_u", "r_f"} and len(righe) == 9, righe
    ampiezze = []
    for r in righe:
        if not ampiezze or ampiezze[-1][0] != r["parametro"]:
            ampiezze.append([r["parametro"], 0])
        ampiezze[-1][1] = max(ampiezze[-1][1], abs(r["delta"]))
    assert [a[1] for a in ampiezze] == sorted((a[1] for a in ampiezze), reverse=True), ampiezze
    assert len(ampiezze) == 5, "le varianti dello stesso parametro devono essere adiacenti"
    assert next(r for r in righe if r["parametro"] == "nu" and r["variante"] == "+1,5 punti")["HC"] < m.bilancio(dati, prof, c)["HC"]
    tmp = tempfile.mkdtemp()
    try:
        A = "it-pianificazione"
        assert run_script(asset(A, "proiezione-inps.py"), "--json", os.path.join(tmp, "inps.json"), cwd=tmp)[0] == 0
        assert run_script(asset(A, "netto-pensione.py"), "--json", os.path.join(tmp, "netto.json"), cwd=tmp)[0] == 0
        senza = _cu_esempio()
        senza["persona"].pop("eta_pensione")
        senza.pop("pensione")
        p = _scrivi_json(os.path.join(tmp, "d.json"), senza)
        rc, so, se = run_script_split(asset(*CU), "--dati", p, "--from", os.path.join(tmp, "inps.json"),
                                      "--from", os.path.join(tmp, "netto.json"), "--json", os.path.join(tmp, "cu.json"), cwd=tmp)
        assert rc == 0, se[-300:]
        j = json.load(open(os.path.join(tmp, "cu.json"), encoding="utf-8"))
        inps = json.load(open(os.path.join(tmp, "inps.json"), encoding="utf-8"))
        netto = json.load(open(os.path.join(tmp, "netto.json"), encoding="utf-8"))
        assert j["profilo_usato"]["eta_pensione"] == int(inps["eta_pensionamento"]), j["profilo_usato"]
        assert j["profilo_usato"]["pensione_netta"] == netto["netto_pensione_inps"]
        assert j["sensibilita"] and len(j["sensibilita"]) == 9
        # --set prevale su --dati
        rc, so, se = run_script_split(asset(*CU), "--set", "ETA_PENSIONE=70", "--json", os.path.join(tmp, "cu.json"), cwd=tmp)
        assert rc == 0 and json.load(open(os.path.join(tmp, "cu.json"), encoding="utf-8"))["profilo_usato"]["eta_pensione"] == 70
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T15 capitale umano: input realistici dalla revisione finale (età frazionarie, CSV da Excel, --from scartato)")
def t_cu_revisione_finale():
    m = carica_modulo(asset(*CU))
    tmp = tempfile.mkdtemp()
    def con(modifica, nome="d.json"):
        d = _cu_esempio()
        modifica(d)
        return _scrivi_json(os.path.join(tmp, nome), d)
    try:
        # età frazionarie (proiezione-inps.py le accetta): arrotondate per difetto, senza traceback
        rc, so, se = _cu("--dati", con(lambda d: d["persona"].update(eta=35.5, eta_pensione=67.5)), cwd=tmp)
        assert rc == 0 and "Traceback" not in se and "per difetto" in so, (rc, se[-300:])
        # la frase della diagnosi 2 nomina il blocco più grande: per un pensionato è la pensione
        def pensionato(d):
            d["persona"].update(eta=68, netto_annuo=0, tfr_annuo=0)
        rc, so, se = _cu("--dati", con(pensionato), cwd=tmp)
        assert rc == 0 and "gli somiglia" not in so and "capitale pensionistico" in so.split("2. ESPOSIZIONE")[1].split("3. ")[0], so[-900:]
        # CSV da Excel italiano: righe «;;», «;» finale, cp1252 con accenti, celle in percento
        comp = os.path.join(tmp, "comp.csv")
        with open(comp, "w", encoding="cp1252") as f:
            f.write("strumento;settore;peso;\nETF Società;Information Technology;60,0%;\n"
                    "ETF Società;Financials;40,0%;\n;;;\n")
        assert m.leggi_composizione(comp) == {"ETF Società": {"Information Technology": 0.6, "Financials": 0.4}}
        # --dati che scavalca un valore letto con --from: lo dice
        A = "it-pianificazione"
        assert run_script(asset(A, "netto-pensione.py"), "--json", os.path.join(tmp, "netto.json"), cwd=tmp)[0] == 0
        rc, so, se = _cu("--from", os.path.join(tmp, "netto.json"), cwd=tmp)
        assert rc == 0 and "togli la chiave" in so, so[:900]
        # nessuna composizione che combaci: niente esposizione combinata fittizia; nomi confrontati senza spazi
        pos = [{"asset_class": "azionario", "valore": 1000, "strumento": " A "}]
        z = m.sovrapposizione(pos, {"B": {"Financials": 1.0}}, "Financials", 0, 500, m.PESI_MONDO)
        assert z["quota_settore_fin"] is None and z["esposizione_combinata"] is None, z
        z = m.sovrapposizione(pos, {"A": {"Financials": 1.0}}, "Financials", 1000, 500, m.PESI_MONDO)
        assert z["quota_settore_fin"] == 1.0, z
        # override fuori dominio: ERRORE pulito, e nessun pezzo di report prima dell'errore
        rc, so, se = _cu("--dati", con(lambda d: d["persona"]["override"].update(rho=1.5)), cwd=tmp)
        assert rc == 1 and se.lstrip().startswith("ERRORE:") and "rho" in se and "1. BILANCIO" not in so, (rc, so[-200:], se[-200:])
        rc, so, se = _cu("--dati", con(lambda d: d["persona"].update(crescita_reale="1%")), cwd=tmp)
        assert rc == 1 and "crescita_reale" in se and "Traceback" not in se, se[-300:]
        with open(comp, "w", encoding="utf-8") as f:
            f.write("strumento;settore;peso\nA;Tecnologia;100\n")
        rc, so, se = _cu("--composizione", comp, cwd=tmp)
        assert rc == 1 and "1. BILANCIO" not in so, so[-300:]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T15 capitale umano: proiezione a tre contenitori, TFR e contributi trasferiti, orizzonte dei pensionati")
def t_cu_contenitori():
    m = carica_modulo(asset(*CU))
    c = {"R_F": 0.01, "NU_PENSIONE": 0.0075, "ETA_FINE": 90, "ETA_PENSIONE": 67, "PENSIONE_NETTA": 18000}
    dati = _cu_esempio()
    prof = m.profilo_effettivo(dati["persona"], m.carica_profili())
    b = m.bilancio(dati, prof, c)
    # l'anno 0 è la fotografia di oggi, anche con un glide path: coincide con la diagnosi 2
    for glide in (None, [0.9, 0.9, 0.5]):
        r0 = m.proiezione(dati, prof, c, glide, 50, 1)["anni"][0]
        assert abs(r0["esposizione"] - b["esposizione_totale"]) < 1e-12, (glide, r0["esposizione"], b["esposizione_totale"])
    # il glide path agisce solo sul portafoglio, dall'anno 1: fondo pensione e TFR tengono la loro quota
    g1 = m.proiezione(dati, prof, c, [0.0, 0.0], 50, 1)["anni"][1]
    assert g1["az_fondo"] > 0 and g1["az_portafoglio"] == 0, g1
    # con rendimenti, crescita e risparmio a zero il TFR fa crescere il patrimonio di tfr_annuo l'anno
    m.REND_AZ_REALE, m.REND_OBB_REALE, m.SIGMA_AZ = 0.0, 0.0, 0.0
    try:
        d = _cu_esempio()
        d["persona"].update(crescita_reale=0, tfr_annuo=3000, contributo_datore_fp=1000)
        d["risparmio_annuo"] = 0
        righe = m.proiezione(d, prof, c, None, 20, 1)["anni"]
        assert abs((righe[1]["F"] - righe[0]["F"]) - 4000) < 1e-6, righe[1]["F"] - righe[0]["F"]
        assert abs((righe[1]["tfr"] - righe[0]["tfr"]) - 3000) < 1e-6 and abs((righe[1]["fondo"] - righe[0]["fondo"]) - 1000) < 1e-6
        d["persona"]["tfr_nel_fondo"] = True
        righe = m.proiezione(d, prof, c, None, 20, 1)["anni"]
        assert righe[1]["tfr"] == righe[0]["tfr"] and abs((righe[1]["fondo"] - righe[0]["fondo"]) - 4000) < 1e-6
        # dopo la pensione non entra più niente
        ultimi = [r for r in righe if r["eta"] >= 67]
        assert all(abs(x["F"] - ultimi[0]["F"]) < 1e-6 for x in ultimi)
    finally:
        m.REND_AZ_REALE, m.REND_OBB_REALE, m.SIGMA_AZ = 0.03, 0.005, 0.18
    # un pensionato di 80 anni: cinque anni di proiezione, senza superare l'età finale
    pens = _cu_esempio()
    pens["persona"].update(eta=80, netto_annuo=0, tfr_annuo=0)
    eta = [r["eta"] for r in m.proiezione(pens, prof, c, None, 20, 1)["anni"]]
    assert eta == list(range(80, 86)), eta
    pens["persona"]["eta"] = 87
    eta = [r["eta"] for r in m.proiezione(pens, prof, c, None, 20, 1)["anni"]]
    assert eta[0] == 87 and eta[-1] <= 89 and len(eta) > 1, eta


@test("T15 capitale umano: report leggibile (sensibilità in punti, P50, stato del sorpasso, chiavi sconosciute)")
def t_cu_report():
    m = carica_modulo(asset(*CU))
    c = {"R_F": 0.01, "NU_PENSIONE": 0.0075, "ETA_FINE": 90, "ETA_PENSIONE": 67, "PENSIONE_NETTA": 18000}
    dati = _cu_esempio()
    prof = m.profilo_effettivo(dati["persona"], m.carica_profili())
    varianti = {(r["parametro"], r["variante"]) for r in m.sensibilita(dati, prof, c)}
    assert varianti == {("s", "+10 punti"), ("s", "-10 punti"), ("nu", "+1,5 punti"), ("nu", "-1,5 punti"),
                        ("g", "0%"), ("g", "2%"), ("p_u", "× 2"), ("r_f", "+1 punto"), ("r_f", "-1 punto")}, varianti
    assert m.punti(-0.0002) == "0,0 punti" and m.punti(0.0123) == "+1,2 punti" and m.punti(-0.07) == "-7,0 punti"
    tmp = tempfile.mkdtemp()
    try:
        out = os.path.join(tmp, "cu.json")
        rc, so, se = _cu("--json", out, cwd=tmp)
        assert rc == 0 and "P10 · P50 · P90" in so and "-0,0 punti" not in so, so[-1500:]
        j = json.load(open(out, encoding="utf-8"))
        assert j["stato_sorpasso"] == "entro_la_proiezione" and j["eta_sorpasso"] is not None
        pens = _cu_esempio()
        pens["persona"].update(eta=68, netto_annuo=0, tfr_annuo=0)
        rc, so, se = _cu("--dati", _scrivi_json(os.path.join(tmp, "p.json"), pens), "--json", out, cwd=tmp)
        assert rc == 0 and json.load(open(out, encoding="utf-8"))["stato_sorpasso"] == "gia_superato"
        povero = _cu_esempio()
        povero.update(posizioni=[], liquidita=0, risparmio_annuo=0, tfr_maturato_azienda=0)
        povero["fondo_pensione"]["valore"] = 0
        povero["persona"].update(tfr_annuo=0, contributo_datore_fp=0)
        rc, so, se = _cu("--dati", _scrivi_json(os.path.join(tmp, "n.json"), povero), "--json", out, cwd=tmp)
        j = json.load(open(out, encoding="utf-8"))
        assert rc == 0 and j["stato_sorpasso"] == "non_prima_della_fine" and j["eta_sorpasso"] is None, j["stato_sorpasso"]
        # chiavi sconosciute: avviso che le nomina; «_nota» e i campi extra delle posizioni sono ammessi
        refuso = _cu_esempio()
        refuso["liquidta"] = 5000
        refuso["persona"]["netto_anuo"] = 1
        refuso["posizioni"][0]["broker"] = "Fineco"
        rc, so, se = _cu("--dati", _scrivi_json(os.path.join(tmp, "r.json"), refuso), cwd=tmp)
        assert rc == 0 and "liquidta" in so and "persona.netto_anuo" in so, so[:1200]
        assert "_nota" not in so and "broker" not in so
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ═══════════════════════════ T16: composizione degli ETF ═══════════════════════════

CE = ("it-pianificazione-obiettivi", "importa-composizione-etf.py")


@test("T16 composizione ETF: lettura della tabella (intestazione, preambolo, fine, numeri, formati di file)")
def t_ce_lettura():
    m = carica_modulo(asset(*CE))
    for nome, sep in (("composizione-ishares-en.csv", ","), ("composizione-ishares-it.csv", ";")):
        righe, foglio = m.leggi_righe(os.path.join(FIXTURES, nome))
        i, col, testi = m.trova_intestazione(righe)
        assert i == 5 and foglio is None, (nome, i)
        assert set(col) == {"peso", "settore", "paese", "asset_class", "nome"}, col
        dati = m.righe_dati(righe, i, col["peso"])
        assert len(dati) == 8, (nome, len(dati))                 # il disclaimer in fondo non entra
        v = m.virgola_decimale([r[col["peso"]] for r in dati])
        assert v is (sep == ";")
        assert [m.peso_numerico(r[col["peso"]], v) for r in dati] == [30, 20, 15, 15, 10, 9, 0.8, 0.2]
    assert testi["peso"] == "Ponderazione (%)" and testi["paese"] == "Area geografica"
    assert m.peso_numerico("1.234,5", True) == 1234.5 and m.peso_numerico("1,234.5", False) == 1234.5
    assert m.peso_numerico("4,46%", True) == 4.46 and m.peso_numerico("N/A", False) is None
    assert m.peso_numerico(0.0565, False) == 0.0565
    # precedenza tra due candidate: «Sector Classification» batte «Industry Classification»
    _, col, testi = m.trova_intestazione([["ISIN", "Percent of Fund", "Trade Country Name",
                                           "Industry Classification", "Sector Classification"]])
    assert testi["settore"] == "Sector Classification" and col["settore"] == 4
    tmp = tempfile.mkdtemp()
    try:
        # CSV risalvato da Excel italiano: cp1252, «;», righe di soli «;»
        p = os.path.join(tmp, "excel.csv")
        with open(p, "w", encoding="cp1252") as f:
            f.write("Fondo;;;\nPosizioni al;30/09/2026;;\n;;;\nNome;Settore;Area geografica;Ponderazione (%)\n"
                    "ALPHA;Informatica;Stati Uniti;60,0\nBETA;Sanità;Giappone;40,0\n;;;\n")
        righe, _ = m.leggi_righe(p)
        i, col, _ = m.trova_intestazione(righe)
        assert i == 3 and len(m.righe_dati(righe, i, col["peso"])) == 2
        assert righe[5][1] == "Sanità"
        # XLSX con l'intestazione alla riga 4 e il titolo del foglio
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2026-10-25"
        ws.append([None])
        ws.append(["Although information in this document…"])
        ws.append([None])
        ws.append([None, "Name", "ISIN", "Country", "Type of Security", "Industry Classification", "Weighting"])
        ws.append([1, "ALPHA", "US0000000001", "United States", "Equities", "Information Technology", 0.6])
        px = os.path.join(tmp, "x.xlsx")
        wb.save(px)
        righe, foglio = m.leggi_righe(px)
        i, col, _ = m.trova_intestazione(righe)
        assert foglio == "2026-10-25" and i == 3 and col["peso"] == 6
        # errori: pagina HTML, file mancante, scheda prodotto, estensione
        ph = os.path.join(tmp, "pagina.csv")
        with open(ph, "w", encoding="utf-8") as f:
            f.write("  <!DOCTYPE html><html><head><title>iShares</title></head></html>")
        for percorso, atteso in ((ph, "pagina HTML"), ("/non/esiste.csv", "non trovato")):
            rc, so, se = run_script_split(asset(*CE), "--file", percorso, "--strumento", "X", cwd=tmp)
            assert rc == 1 and se.startswith("ERRORE:") and atteso in se and "Traceback" not in se, se[-300:]
        ps = os.path.join(tmp, "scheda.csv")
        with open(ps, "w", encoding="utf-8") as f:
            f.write("Caratteristiche principali\nTER,0.20%\nValuta,USD\n")
        rc, so, se = run_script_split(asset(*CE), "--file", ps, "--strumento", "X", cwd=tmp)
        assert rc == 1 and "file delle posizioni" in se and "TER" in se, se[-300:]
        pp = os.path.join(tmp, "f.pdf")
        with open(pp, "w", encoding="ascii") as f:    # non «wb»: t_cli_utf8 vuole encoding= o «rb» in ogni open()
            f.write("%PDF-1.4")
        rc, so, se = run_script_split(asset(*CE), "--file", pp, "--strumento", "X", cwd=tmp)
        assert rc == 1 and "XLSX o CSV" in se, se[-300:]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _xlsx_emittente(tipo, cartella):
    """File delle posizioni sintetico con il layout reale dell'emittente (preambolo, riga d'intestazione,
    disclaimer, pesi in frazione o percento, righe non classificate): i file veri non si ridistribuiscono.
    Il nome del file non contiene l'emittente, così il riconoscimento passa dal preambolo."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    if tipo == "xtrackers":
        ws.title = "2026-10-25"
        ws.append([None])
        ws.append(["Although information in this document has been obtained from sources believed to be reliable, "
                   "DWS does not guarantee its accuracy, completeness or fairness."])
        ws.append([None])
        ws.append([None, "Name", "ISIN", "Country", "Currency", "Type of Security", "Industry Classification", "Weighting"])
        for n, (nome, settore, paese, tipo_titolo, peso) in enumerate((
                ("A", "Information Technology", "United States", "Equities", 0.40),
                ("B", "Financials", "Japan", "Equities", 0.25),
                ("C", "Health Care", "Germany", "Equities", 0.20),
                ("D", "Materials", "Switzerland", "Equities", 0.12),
                ("E", "unknown", "Germany", "Equities", 0.02),
                ("F", "Health Care", "United Arab Emirates", "Equities", "N/A"),
                ("G", "unknown", "Germany", "Future", 0.01)), 1):
            ws.append([n, nome, f"XX000000000{n}", paese, "USD", tipo_titolo, settore, peso])
    elif tipo == "spdr":
        ws.title = "holdings"
        for riga in (("Fund Name:", "State Street SPDR MSCI ACWI IMI UCITS ETF"), ("ISIN:", "IE00B3YLTY66"),
                     ("Ticker Symbol:", "SPYI GY"), ("Holdings As Of:", "22-Oct-2026"), (None,)):
            ws.append(list(riga))
        ws.append(["ISIN", "Security Name", "Percent of Fund", "Trade Country Name", "Sector Classification",
                   "Industry Classification"])
        for riga in (("US0000000001", "A", 50.0, "United States", "Information Technology", "Semiconductors"),
                     ("TW0000000002", "B", 30.0, "Taiwan", "Consumer Discretionary", "Broadline Retail"),
                     ("HK0000000003", "C", 15.0, "Hong Kong", "Communication Services", "Interactive Media"),
                     ("Unassigned", "D", 5.04, "Unassigned", "Unassigned", "Unassigned"),
                     ("Unassigned", "CNY:HKD 20250930", -0.04, "Unassigned", "Unassigned", "Unassigned")):
            ws.append(list(riga))
        ws.append([None])
        ws.append(["Marketing communication. The performance data quoted represents past performance."])
    elif tipo == "vanguard":
        ws.append(["Vanguard FTSE All-World UCITS ETF"])
        ws.append(["As at 30 Sep 2026"])
        ws.append([None])
        ws.append(["Holding name", "Sector", "Country", "% of net assets"])
        for riga in (("A", "Technology", "United States", 40), ("B", "Financials", "Japan", 30),
                     ("C", "Telecommunications", "France", 10), ("D", "Basic Materials", "Australia", 20)):
            ws.append(list(riga))
    elif tipo == "amundi":
        ws.append(["Amundi MSCI Europe UCITS ETF"])
        ws.append(["Dati al", "30/09/2026"])
        ws.append([None])
        ws.append(["Nome", "Settore", "Paese", "Peso (%)"])
        for riga in (("A", "Beni di prima necessità", "Francia", 50), ("B", "Servizi di pubblica utilità", "Italia", 30),
                     ("C", "Immobiliare", "Paesi Bassi", 20)):
            ws.append(list(riga))
    else:
        raise ValueError(tipo)
    path = os.path.join(cartella, f"posizioni-{('xtrackers', 'spdr', 'vanguard', 'amundi').index(tipo)}.xlsx")
    wb.save(path)
    return path


@test("T16 composizione ETF: valori chiusi per emittente, riconoscimento ed esclusioni")
def t_ce_composizione():
    m = carica_modulo(asset(*CE))
    r2 = lambda d: {k: round(v * 100, 2) for k, v in d.items()}  # noqa: E731
    en = m.componi(os.path.join(FIXTURES, "composizione-ishares-en.csv"), "ETF MSCI World")
    it = m.componi(os.path.join(FIXTURES, "composizione-ishares-it.csv"), "ETF MSCI World")
    for c in (en, it):
        assert c["emittente"] == "ishares" and c["riconosciuto_da"] == "preambolo" and c["classificazione"] == "GICS", c
        assert c["data_posizioni"] == "2026-09-30" and c["titoli"] == 6, c
        assert r2(c["settori"]) == {"Information Technology": 40.40, "Financials": 20.20, "Health Care": 15.15,
                                    "Industrials": 15.15, "Energy": 9.09}, c["settori"]
        assert r2(c["paesi"]) == {"US": 50.51, "JP": 15.15, "GB": 15.15, "KR": 10.10, "CA": 9.09}, c["paesi"]
        assert abs(c["esclusi"]["liquidita_derivati"] - 0.01) < 1e-9 and c["avvisi"] == [], c
    assert list(en["settori"]) == list(it["settori"]) and en["settori"] == it["settori"]   # stesso ordine e valori
    assert list(en["settori"])[0] == "Information Technology"                             # peso decrescente
    tmp = tempfile.mkdtemp()
    try:
        x = m.componi(_xlsx_emittente("xtrackers", tmp), "X")
        assert x["emittente"] == "xtrackers" and x["data_posizioni"] == "2026-10-25" and x["titoli"] == 5, x
        assert r2(x["settori"]) == {"Information Technology": 41.24, "Financials": 25.77, "Health Care": 20.62,
                                    "Materials": 12.37}, x["settori"]
        assert r2(x["paesi"]) == {"US": 40.40, "JP": 25.25, "DE": 22.22, "CH": 12.12}, x["paesi"]
        assert abs(x["esclusi"]["liquidita_derivati"] - 0.01) < 1e-9
        assert abs(x["esclusi"]["settore_non_classificato"] - 0.02) < 1e-9 and x["esclusi"]["paese_non_classificato"] == 0
        s = m.componi(_xlsx_emittente("spdr", tmp), "S")
        assert s["emittente"] == "spdr" and s["data_posizioni"] == "2026-10-22", s
        assert s["colonne"]["settore"] == "Sector Classification" and s["colonne"]["paese"] == "Trade Country Name"
        assert r2(s["settori"]) == {"Information Technology": 52.63, "Consumer Discretionary": 31.58,
                                    "Communication Services": 15.79}, s["settori"]
        assert r2(s["paesi"]) == {"US": 52.63, "TW": 31.58, "HK": 15.79}, s["paesi"]
        assert abs(s["esclusi"]["liquidita_derivati"] - 0.0004) < 1e-9, s["esclusi"]
        assert any("non classificat" in a and "5%" in a for a in s["avvisi"]), s["avvisi"]   # 5,04% > 5%
        v = m.componi(_xlsx_emittente("vanguard", tmp), "V")
        assert v["emittente"] == "vanguard" and v["classificazione"] == "ICB" and v["settori"] == {}, v
        assert r2(v["paesi"]) == {"US": 40.0, "JP": 30.0, "AU": 20.0, "FR": 10.0} and v["data_posizioni"] == "2026-09-30", v
        assert any("ICB" in a for a in v["avvisi"]), v["avvisi"]
        a = m.componi(_xlsx_emittente("amundi", tmp), "A")
        assert a["emittente"] == "amundi" and a["data_posizioni"] == "2026-09-30", a
        assert r2(a["settori"]) == {"Consumer Staples": 50.0, "Utilities": 30.0, "Real Estate": 20.0}, a["settori"]
        assert r2(a["paesi"]) == {"FR": 50.0, "IT": 30.0, "NL": 20.0}, a["paesi"]
        # riconoscimento: --emittente vince sul preambolo; dal nome del file; nessuno con colonne note → avviso
        f = m.componi(os.path.join(FIXTURES, "composizione-ishares-en.csv"), "Z", "spdr")
        assert f["riconosciuto_da"] == "--emittente" and f["emittente"] == "spdr"
        g = os.path.join(tmp, "posizioni.csv")
        with open(g, "w", encoding="utf-8") as fh:
            fh.write("Name,Sector,Country,Weight\nA,Energy,Italy,100\n")
        c = m.componi(g, "G")
        assert c["emittente"] is None and c["riconosciuto_da"] == "nessuno", c
        assert any("regole generiche" in w for w in c["avvisi"]), c["avvisi"]
        os.rename(g, os.path.join(tmp, "SWDA_holdings_ishares.csv"))
        c = m.componi(os.path.join(tmp, "SWDA_holdings_ishares.csv"), "G")
        assert c["riconosciuto_da"] == "nome del file" and c["emittente"] == "ishares", c
        # ETF obbligazionario, settore vuoto ovunque → solo paesi, con l'avviso
        with open(g, "w", encoding="utf-8") as fh:
            fh.write("Name,Sector,Asset Class,Location,Weight (%)\nBTP 2030,-,Fixed Income,Italy,60\n"
                     "OAT 2031,-,Fixed Income,France,40\n")
        c = m.componi(g, "B")
        assert c["settori"] == {} and r2(c["paesi"]) == {"IT": 60.0, "FR": 40.0}, c
        assert any("settore" in w for w in c["avvisi"]), c["avvisi"]
        # liquidità sopra il 5% → avviso; Telecommunication Services → Communication Services con avviso
        with open(g, "w", encoding="utf-8") as fh:
            fh.write("Name,Sector,Asset Class,Location,Weight (%)\nA,Telecommunication Services,Equity,Italy,94\n"
                     "CASH,Cash and/or Derivatives,Cash,Italy,6\n")
        c = m.componi(g, "C")
        assert c["settori"] == {"Communication Services": 1.0}, c["settori"]
        assert any("liquidit" in w for w in c["avvisi"]) and any("Telecommunication" in w for w in c["avvisi"]), c["avvisi"]
        # dai file reali: paesi minori (Zambia in un MSCI World), «European Union» e «NA» come non classificati,
        # quote escluse minime scritte «< 0,01%» e non «0,00%»
        with open(g, "w", encoding="utf-8") as fh:
            fh.write("Name,Sector,Location,Weight (%)\nA,Materials,Zambia,0.02\nB,Financials,European Union,0.003\n"
                     "C,Energy,NA,0.001\nD,Energy,Italy,99.976\n")
        c = m.componi(g, "R")
        assert set(c["paesi"]) == {"ZM", "IT"}, c["paesi"]
        assert abs(c["esclusi"]["paese_non_classificato"] - 0.00004) < 1e-12, c["esclusi"]
        assert all(len(m.PAESI[k]) == 2 for k in m.PAESI) and len(m.PAESI) == 249     # ISO 3166-1 completo
        assert "paese non classificato < 0,01%" in m.righe_strumento(c, "25/10/2026")[0], m.righe_strumento(c, "x")[0]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    assert m.SETTORI_GICS == carica_modulo(asset(*CU)).SETTORI_GICS


@test("T16 composizione ETF: nomi sconosciuti, somme incoerenti e classificazione mista respinti")
def t_ce_errori():
    tmp = tempfile.mkdtemp()
    try:
        out = os.path.join(tmp, "composizione.csv")
        with open(out, "w", encoding="utf-8") as f:
            f.write("strumento;settore;peso\nA;Energy;100\n")
        prima = (open(out, encoding="utf-8").read(), os.path.getmtime(out))
        casi = (
            ("Name,Sector,Location,Weight (%)\nA,Spazio,Atlantide,30\nB,Energy,Lemuria,20\nC,Energy,Italy,50\n",
             ("Atlantide", "Lemuria", "Spazio", "segnala")),
            ("Name,Sector,Location,Weight (%)\nA,Energy,Italy,40\nB,Energy,France,20\n", ("sommano a 60",)),
            ("Name,Sector,Location,Weight (%)\nA,Energy,Italy,abc\nB,Energy,France,100\n", ("abc", "riga")),
            ("Name,Sector,Location,Weight (%)\nA,Technology,Italy,50\nB,Information Technology,France,50\n", ("mista",)),
            ("Name,Sector,Asset Class,Location,Weight (%)\nCASH,Cash and/or Derivatives,Cash,Italy,100\n", ("nessuna riga",)),
        )
        for n, (testo, attesi) in enumerate(casi):
            p = os.path.join(tmp, f"caso{n}.csv")
            with open(p, "w", encoding="utf-8") as f:
                f.write(testo)
            rc, so, se = run_script_split(asset(*CE), "--file", p, "--strumento", "X", "--out", out, cwd=tmp)
            assert rc == 1 and se.lstrip().startswith("ERRORE:") and "Traceback" not in se, (n, rc, se[-400:])
            assert "ERRORE:" not in so, (n, so[-200:])
            for parola in attesi:
                assert parola in se, (n, parola, se[-400:])
            assert (open(out, encoding="utf-8").read(), os.path.getmtime(out)) == prima, n
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T16 composizione ETF: CSV fuso con l'esistente, formato, --json e report")
def t_ce_uscita():
    tmp = tempfile.mkdtemp()
    try:
        out = os.path.join(tmp, "composizione.csv")
        with open(out, "w", encoding="utf-8") as f:     # vecchio formato, scritto a mano, con un commento libero
            f.write("# compilato a mano\nstrumento;settore;peso\nETF Europa;Financials;0,6\nETF Europa;Industrials;0,4\n"
                    "ETF MSCI World;Energy;100\n")
        en = os.path.join(FIXTURES, "composizione-ishares-en.csv")
        uscite = []
        for _ in range(2):                             # reimportare due volte non duplica
            rc, so, se = run_script_split(asset(*CE), "--file", en, "--strumento", "ETF MSCI World",
                                          "--out", out, "--json", os.path.join(tmp, "c.json"), cwd=tmp)
            assert rc == 0, se[-400:]
            uscite.append(so)
        assert "convertito" in uscite[0] and "convertito" not in uscite[1], uscite   # solo il primo trova il vecchio formato
        testo = open(out, encoding="utf-8-sig").read()
        with open(out, "rb") as f:
            assert f.read().startswith(b"\xef\xbb\xbf")
        righe = testo.splitlines()
        assert righe[0] == "# compilato a mano" and righe[1] == "strumento;dimensione;voce;peso", righe[:3]
        assert righe[2:4] == ["ETF Europa;settore;Financials;60,00", "ETF Europa;settore;Industrials;40,00"], righe[2:4]
        assert sum(1 for r in righe if r.startswith("# ETF MSCI World:")) == 1, righe
        commento = next(r for r in righe if r.startswith("# ETF MSCI World:"))
        assert re.fullmatch(r"# ETF MSCI World: iShares, composizione-ishares-en\.csv, posizioni al 30/09/2026, "
                            r"importato il \d\d/\d\d/\d{4}; esclusi liquidità e derivati 1,00%", commento), commento
        mw = [r for r in righe if r.startswith("ETF MSCI World;")]
        assert mw[0] == "ETF MSCI World;settore;Information Technology;40,40" and "ETF MSCI World;paese;KR;10,10" in mw, mw
        assert len(mw) == 10 and not any(";Energy;100" in r for r in righe), mw
        for dim in ("settore", "paese"):
            tot = sum(float(r.split(";")[3].replace(",", ".")) for r in mw if r.split(";")[1] == dim)
            assert abs(tot - 100) <= 0.05, (dim, tot)
        so = uscite[1]
        assert "ETF Europa" in so, so                                        # il report dice cosa ha conservato
        assert "Tecnologia dell'informazione" in so and "Corea del Sud" in so and "Weight (%)" in so, so
        d = json.load(open(os.path.join(tmp, "c.json"), encoding="utf-8"))
        assert d["out"] == out and [s["strumento"] for s in d["strumenti"]] == ["ETF MSCI World"], d
        s = d["strumenti"][0]
        assert {"strumento", "emittente", "riconosciuto_da", "file", "data_posizioni", "classificazione", "colonne",
                "titoli", "esclusi", "settori", "paesi", "avvisi"} == set(s), set(s)
        assert abs(sum(s["settori"].values()) - 1) < 1e-9 and abs(sum(s["paesi"].values()) - 1) < 1e-9
        # --json - : solo JSON su stdout, --out scritto lo stesso
        rc, so, se = run_script_split(asset(*CE), "--file", en, "--strumento", "Y", "--out", out, "--json", "-", cwd=tmp)
        assert rc == 0 and json.loads(so)["strumenti"][0]["strumento"] == "Y", (so[:300], se[-300:])
        assert "Y;paese;US;50,51" in open(out, encoding="utf-8-sig").read()
        # due file in una chiamata; riga di comando spaiata o ripetuta → rc 2
        rc, so, se = run_script_split(asset(*CE), "--file", en, "--strumento", "A", "--file",
                                      os.path.join(FIXTURES, "composizione-ishares-it.csv"), "--strumento", "B",
                                      "--out", out, cwd=tmp)
        assert rc == 0 and "A;settore;Energy;9,09" in open(out, encoding="utf-8-sig").read(), se[-300:]
        assert "B;paese;KR;10,10" in open(out, encoding="utf-8-sig").read()
        for args in (("--file", en, "--file", en, "--strumento", "A"),
                     ("--file", en, "--strumento", "A", "--file", en, "--strumento", "A")):
            rc, so, se = run_script_split(asset(*CE), *args, "--out", out, cwd=tmp)
            assert rc == 2 and "usage:" in se, (args, rc, se[-300:])
        # un --out con colonne sconosciute non si sovrascrive
        strano = os.path.join(tmp, "strano.csv")
        with open(strano, "w", encoding="utf-8") as f:
            f.write("isin;quota\nIE00;1\n")
        rc, so, se = run_script_split(asset(*CE), "--file", en, "--strumento", "A", "--out", strano, cwd=tmp)
        assert rc == 1 and open(strano, encoding="utf-8").read() == "isin;quota\nIE00;1\n", se[-300:]
        # il CSV scritto si rilegge con csv senza perdite: quattro campi per riga dati
        dati = [r for r in csv.reader(io.StringIO(open(out, encoding="utf-8-sig").read()), delimiter=";")
                if r and not r[0].startswith("#")]
        assert all(len(r) == 4 for r in dati), dati
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ═══════════════════════════ T17: cruscotto grafico ═══════════════════════════

def _righe_csv_ordini(cartella):
    """(broker, operazione, classe, importo) di ogni riga dei CSV degli ordini, senza la riga NOTA finale."""
    import csv
    righe = []
    for nome in sorted(os.listdir(cartella)):
        with open(os.path.join(cartella, nome), encoding="utf-8-sig", newline="") as f:
            r = list(csv.reader(f, delimiter=";"))
        for riga in r[1:]:
            if riga and riga[1] != "NOTA":
                righe.append((riga[0], riga[1], riga[2], float(riga[5].replace(".", "").replace(",", "."))))
    return righe


@test("T17 ribilanciamento --json: chiavi del piano, ordini uguali ai CSV, stdout invariato")
def t_t17_ribilancia_json():
    script = asset("it-rebalance", "ribilanciamento-multibroker.py")
    a, b = tempfile.mkdtemp(), tempfile.mkdtemp()
    try:
        rc, out = run_script(script, "--out", "ordini", cwd=a)
        rc_json, out_json = run_script(script, "--out", "ordini", "--json", "piano.json", cwd=b)
        assert rc == 0 and rc_json == 0, (rc, rc_json, out_json[-400:])
        riga_json = "JSON del piano: piano.json\n"
        assert out_json.count(riga_json) == 1 and out == out_json.replace(riga_json, ""), out_json[-600:]
        d = json.load(open(os.path.join(b, "piano.json"), encoding="utf-8"))
        chiavi = {"totale", "totale_post", "imposta", "prelievo", "entrate", "liquidato", "drift", "risultante",
                  "zainetto", "ordini"}
        assert chiavi <= set(d), set(d)
        assert d["totale"] == 100000 and d["imposta"] == 0, (d["totale"], d["imposta"])
        az = d["drift"]["azionario"]
        assert az["fuori_banda"] is True and abs(az["peso"] - 0.68) < 1e-9, az
        assert {"peso", "target", "drift", "tight", "wide", "fuori_banda", "esito", "motivo"} <= set(az), set(az)
        assert abs(d["risultante"]["azionario"] - 0.60) < 1e-6, d["risultante"]
        assert round(d["zainetto"]["Fineco"]) == 1800, d["zainetto"]
        dal_json = sorted((o["broker"], o["operazione"], o["classe"], round(o["importo_eur"], 2)) for o in d["ordini"])
        assert dal_json and dal_json == sorted(_righe_csv_ordini(os.path.join(b, "ordini"))), dal_json
        # file di output non scrivibile → ERRORE e rc 1
        rc, so, se = run_script_split(script, "--out", "ordini", "--json", "/non/esiste/x.json", cwd=b)
        assert rc == 1 and se.lstrip().startswith("ERRORE:") and "impossibile scrivere" in se, (rc, se[-300:])
        assert "ERRORE:" not in so and "Traceback" not in se, so[-300:]
    finally:
        shutil.rmtree(a, ignore_errors=True)
        shutil.rmtree(b, ignore_errors=True)
    # Review Focus 4: zainetto Fineco a zero → nessun ordine, il drift resta fuori banda, JSON serializzabile
    m, p = _ribilancia(ZAINETTO={"Fineco": 0, "Directa": 0})
    d = m.piano_json(p, m.genera_ordini(p))
    assert d["ordini"] == [] and d["drift"]["azionario"]["fuori_banda"] is True, d["ordini"]
    json.dumps(d)


_JSON_ESEMPIO = {}
NOME_PERICOLOSO = "</script><img src=x onerror=alert(1)> & l'«anticipo»"


def _json_esempio():
    """I JSON degli script del piano sugli esempi del plugin, generati una volta per esecuzione della suite:
    {nome del file: percorso}. Sono gli input reali del cruscotto."""
    if not _JSON_ESEMPIO:
        tmp = tempfile.mkdtemp(prefix="cruscotto-")
        pian = asset("it-pianificazione-obiettivi", "pianificatore-obiettivi.xlsx")
        OB, PE = (lambda n: asset("it-pianificazione-obiettivi", n)), (lambda n: asset("it-pianificazione", n))
        lanci = [
            ("alloc.json", OB("allocatore-capitale.py"), [pian]),
            ("glide.json", OB("glidepath-obiettivi.py"), [pian]),
            ("mco.json", OB("montecarlo-obiettivi.py"), [pian, "--set", "N_SIM=500"]),
            ("cu.json", OB("capitale-umano.py"), ["--composizione", OB("esempio-composizione.csv")]),
            ("dati.json", PE("importa-previdenza.py"), [os.path.join(FIXTURES, "estratto-inps-test.xml"),
                                                       os.path.join(FIXTURES, "storico-fondo-test.csv"), "--anonimo"]),
            ("inps.json", PE("proiezione-inps.py"), ["--from", "dati.json"]),
            ("mc.json", PE("montecarlo-pensione.py"), ["--from", "dati.json", "--from", "inps.json", "--set", "N_SIM=500"]),
            ("netto.json", PE("netto-pensione.py"), ["--from", "mc.json"]),
            ("broker.json", asset("it-tax-optimization", "importa-estratto-broker.py"),
             [os.path.join(FIXTURES, "estratto-multibroker.csv"), os.path.join(FIXTURES, "zainetto-multibroker.csv"), "--anno", "2026"]),
            ("piano.json", asset("it-rebalance", "ribilanciamento-multibroker.py"), ["--out", "ordini"]),
        ]
        for nome, script, args in lanci:
            rc, out = run_script(script, *args, "--json", nome, cwd=tmp)
            assert rc == 0, f"{os.path.basename(script)}: rc {rc}, {out[-300:]}"
            _JSON_ESEMPIO[nome] = os.path.join(tmp, nome)
    return _JSON_ESEMPIO


def _cruscotto(*sorgenti, cartella=None, extra=()):
    """Lancia genera-cruscotto.py sui file dati; ritorna (rc, stdout, stderr, bundle o None, html o None)."""
    cartella = cartella or tempfile.mkdtemp(prefix="cruscotto-out-")
    args = [a for f in sorgenti for a in ("--from", f)]
    html, bundle = os.path.join(cartella, "c.html"), os.path.join(cartella, "b.json")
    rc, so, se = run_script_split(asset("it-revisione", "genera-cruscotto.py"), *args, "-o", html, "--json", bundle,
                                  "--data", "2026-10-25", *extra, cwd=cartella)
    b = json.load(open(bundle, encoding="utf-8")) if rc == 0 else None
    h = open(html, encoding="utf-8").read() if rc == 0 else None
    return rc, so, se, b, h


def _json_variante(nome, modifica):
    """Copia di un JSON d'esempio con `modifica(dati)` applicata, nella stessa cartella con un prefisso."""
    d = json.load(open(_json_esempio()[nome], encoding="utf-8"))
    modifica(d)
    path = os.path.join(os.path.dirname(_json_esempio()[nome]), "variante-" + nome)
    json.dump(d, open(path, "w", encoding="utf-8"), ensure_ascii=False)
    return path


@test("T17 cruscotto: ogni output reale riconosciuto con il suo tipo, nessuno con due firme")
def t_t17_riconosci():
    mod = carica_modulo(asset("it-revisione", "genera-cruscotto.py"))
    js = _json_esempio()
    attesi = {"alloc.json": "allocatore", "glide.json": "glide_path", "mco.json": "mc_obiettivi", "inps.json": "inps",
              "mc.json": "mc_pensione", "netto.json": "netto_pensione", "cu.json": "capitale_umano",
              "broker.json": "broker", "piano.json": "ribilanciamento", "dati.json": "ignorato:importa-previdenza"}
    for nome, tipo in attesi.items():
        d = json.load(open(js[nome], encoding="utf-8"))
        assert mod.riconosci(d) == tipo, (nome, mod.riconosci(d))
        assert sum(1 for _, firma in mod.FIRME if firma(d)) <= 1, nome
    assert {t for t, _ in mod.FIRME} == {t for t in attesi.values() if not t.startswith("ignorato")}
    assert mod.riconosci({"a": 1}) is None
    # ponte e riscatto si riconoscono e si ignorano: la catena previdenziale si passa intera
    tmp = os.path.dirname(js["inps.json"])
    PE = lambda n: asset("it-pianificazione", n)
    rc, out = run_script(PE("riscatto-laurea.py"), "--from", "dati.json", "--from", "inps.json", "--json", "ris.json", cwd=tmp)
    assert rc == 0, out[-300:]
    rc, out = run_script(PE("montecarlo-pensione.py"), "--from", "dati.json", "--from", "inps.json", "--set", "ANNI=25",
                         "--set", "N_SIM=500", "--json", "mc-uscita.json", cwd=tmp)
    assert rc == 0, out[-300:]
    rc, out = run_script(PE("ponte-pensione.py"), "--from", "inps.json", "--from", "mc-uscita.json", "--set", "ETA_USCITA=62",
                         "--json", "ponte.json", cwd=tmp)
    assert rc == 0, out[-300:]
    assert mod.riconosci(json.load(open(os.path.join(tmp, "ris.json"), encoding="utf-8"))) == "ignorato:riscatto-laurea"
    assert mod.riconosci(json.load(open(os.path.join(tmp, "ponte.json"), encoding="utf-8"))) == "ignorato:ponte-pensione"


@test("T17 cruscotto: il bundle porta i valori dei JSON di origine e nessuna posizione dell'estratto")
def t_t17_bundle():
    js = _json_esempio()
    L = lambda n: json.load(open(js[n], encoding="utf-8"))
    tutti = [js[n] for n in ("alloc.json", "glide.json", "mco.json", "inps.json", "mc.json", "netto.json", "cu.json",
                             "broker.json", "piano.json")]
    rc, so, se, b, html = _cruscotto(*tutti)
    assert rc == 0, (so[-300:], se[-300:])
    mco, mc, netto, cu, broker, piano = (L(n) for n in ("mco.json", "mc.json", "netto.json", "cu.json", "broker.json", "piano.json"))
    assert b["generato_il"] == "2026-10-25" and b["versione"] == 1
    assert b["plugin"] == json.load(open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8"))["version"]
    assert [f["tipo"] for f in b["fonti"]] == ["allocatore", "glide_path", "mc_obiettivi", "inps", "mc_pensione",
                                              "netto_pensione", "capitale_umano", "broker", "ribilanciamento"], b["fonti"]
    assert b["obiettivi"]["obiettivi"][0]["probabilita_successo"] == mco["obiettivi"][0]["probabilita_successo"]
    assert b["obiettivi"]["obiettivi"][1]["capitale_finale"] == mco["obiettivi"][1]["capitale_finale"]
    assert b["pensione"]["tasso_lordo"]["mediana"] == mc["tasso_sostituzione_lordo"]["mediana"]
    assert b["pensione"]["target"] == mc["target"]
    assert b["glide_path"]["portafoglio"] == L("glide.json")["portafoglio"]
    assert b["allocazione"]["target_aggregato"] == L("alloc.json")["aggregato"]
    ind = {i["id"]: i for i in b["indicatori"]}
    assert list(ind) == ["capitale_allocato", "valore_portafoglio", "probabilita_essenziali", "sostituzione_netta",
                         "zainetto_in_scadenza"], list(ind)
    assert ind["capitale_allocato"]["valore"] == L("alloc.json")["capitale"]
    assert ind["sostituzione_netta"]["valore"] == next(s for s in netto["scenari"] if s["nome"] == "mediana")["tasso_sostituzione_netto"]
    assert ind["probabilita_essenziali"]["nota"] == "Fondo di emergenza"
    assert ind["probabilita_essenziali"]["valore"] == mco["obiettivi"][0]["probabilita_successo"]
    assert b["capitale_umano"]["eta_sorpasso"] == cu["eta_sorpasso"] and b["capitale_umano"]["stato_sorpasso"] == cu["stato_sorpasso"]
    assert ind["zainetto_in_scadenza"]["valore"] == broker["analisi"]["zainetto_in_scadenza"]
    assert ind["valore_portafoglio"]["valore"] == broker["analisi"]["totale_valore"]
    assert b["ribilanciamento"]["drift"]["azionario"]["drift"] == piano["drift"]["azionario"]["drift"]
    assert len(b["ribilanciamento"]["ordini"]) == len(piano["ordini"]) > 0
    for p in broker["posizioni"]:
        assert p["isin"] not in html and p["descrizione"] not in html and p["isin"] not in json.dumps(b), p["isin"]
    # stdout: percorso, tipi riconosciuti, sezioni e indicatori in formato italiano
    assert "Cruscotto:" in so and "mc.json → mc_pensione" in so and "Sezioni:" in so and "€130.000" in so, so[-800:]


@test("T17 cruscotto: sezioni assenti, indicatori senza fonte, errori con rc 1 e rc 2")
def t_t17_errori():
    js = _json_esempio()
    rc, so, se, b, _ = _cruscotto(js["alloc.json"])
    assert rc == 0, se[-300:]
    assert b["allocazione"] and all(b[k] is None for k in ("obiettivi", "glide_path", "pensione", "capitale_umano",
                                                            "fiscalita", "ribilanciamento")), b
    assert [i["id"] for i in b["indicatori"]] == ["capitale_allocato"], b["indicatori"]
    # Review Focus 3: nessun obiettivo essenziale → niente indicatore della probabilità minima
    def senza_essenziali(d):
        for o in d["obiettivi"]:
            o["priorita"] = "importante"
    rc, so, se, b, _ = _cruscotto(_json_variante("mco.json", senza_essenziali))
    assert rc == 0 and [i["id"] for i in b["indicatori"]] == [] and b["obiettivi"], (se[-300:], b and b["indicatori"])
    # dati.json di importa-previdenza: riconosciuto e ignorato, con un avviso
    rc, so, se, b, _ = _cruscotto(js["alloc.json"], js["dati.json"])
    assert rc == 0 and "ℹ" in so and "ignorato" in so and "dati.json" in so, so[-400:]

    def errore(*args, atteso, extra=()):
        rc, so, se, _, _ = _cruscotto(*args, extra=extra)
        assert rc == 1, (args, rc, so[-200:], se[-300:])
        assert se.lstrip().startswith("ERRORE:") and "ERRORE:" not in so and "Traceback" not in se, (so[-200:], se[-300:])
        for a in (atteso if isinstance(atteso, tuple) else (atteso,)):
            assert a in se, (a, se[-300:])
    tmp = tempfile.mkdtemp()
    try:
        non_json = os.path.join(tmp, "rotto.json")
        open(non_json, "w", encoding="utf-8").write("{non è json")
        lista = os.path.join(tmp, "lista.json")
        open(lista, "w", encoding="utf-8").write("[1, 2]")
        ignoto = os.path.join(tmp, "ignoto.json")
        open(ignoto, "w", encoding="utf-8").write('{"a": 1}')
        copia = os.path.join(tmp, "mc-uscita.json")
        shutil.copy(js["mc.json"], copia)
        errore("/non/esiste.json", atteso="non trovato")
        errore(non_json, atteso=("rotto.json", "JSON"))
        errore(lista, atteso="lista.json")
        errore(ignoto, atteso=("ignoto.json", "non è un output riconosciuto", "mc_pensione"))
        errore(js["mc.json"], copia, atteso=("mc.json", "mc-uscita.json", "stesso tipo"))
        # Review Focus 1: un JSON di una versione precedente senza un campo usato
        errore(_json_variante("mc.json", lambda d: d["tasso_sostituzione_lordo"].pop("p90")),
               atteso=("variante-mc.json", "tasso_sostituzione_lordo.p90"))
        errore(js["alloc.json"], extra=("--data", "2026-13-01"), atteso="data")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    rc, so, se = run_script_split(asset("it-revisione", "genera-cruscotto.py"), "--from", js["alloc.json"],
                                  "-o", "/non/esiste/c.html")
    assert rc == 1 and se.lstrip().startswith("ERRORE:") and "impossibile scrivere" in se, (rc, se[-300:])
    rc, out = run_script(asset("it-revisione", "genera-cruscotto.py"))
    assert rc == 2 and "usage:" in out, (rc, out[-300:])
    rc, out = run_script(asset("it-revisione", "genera-cruscotto.py"), "--from")
    assert rc == 2 and "usage:" in out, (rc, out[-300:])


@test("T17 cruscotto: HTML autosufficiente, segnaposto unico, nomi pericolosi innocui")
def t_t17_html():
    mod = carica_modulo(asset("it-revisione", "genera-cruscotto.py"))
    js = _json_esempio()
    tpl = open(asset("it-revisione", "cruscotto.html"), encoding="utf-8").read()
    assert tpl.count("__DATI_CRUSCOTTO__") == 1
    def pericoloso(d):
        d["obiettivi"][0]["nome"] = NOME_PERICOLOSO
    rc, so, se, b, html = _cruscotto(js["alloc.json"], _json_variante("mco.json", pericoloso))
    assert rc == 0, se[-300:]
    for testo in (tpl, html):
        assert not re.search(r"https?://|(src|href)\s*=\s*[\"']//|@import|fetch\(", testo)
    assert "__DATI_CRUSCOTTO__" not in html
    # Review Focus 5: il nome non chiude il tag e il JSON incorporato si rilegge identico
    assert "</script><img" not in html and html.count("</script>") == tpl.count("</script>")
    dati = json.loads(re.search(r'<script id="dati" type="application/json">(.*?)</script>', html, re.S).group(1))
    assert dati["obiettivi"]["obiettivi"][0]["nome"] == NOME_PERICOLOSO
    assert dati == b
    for tpl_rotto in ("<p>senza</p>", "__DATI_CRUSCOTTO__ __DATI_CRUSCOTTO__"):
        try:
            mod.incorpora(tpl_rotto, {})
            raise AssertionError("incorpora deve respingere il template")
        except mod.ErroreDati:
            pass


# Il rendering si verifica in Chromium con Playwright per Python. Senza Playwright i test sono saltati,
# salvo RICHIEDI_BROWSER=1 (la CI). PLAYWRIGHT_CHROMIUM_EXECUTABLE indica un Chromium già installato
# fuori dalla cache di Playwright.

def browser():
    """Il modulo sync di Playwright e una funzione che lancia Chromium; Salta se Playwright manca."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        if os.environ.get("RICHIEDI_BROWSER"):
            raise AssertionError("RICHIEDI_BROWSER è impostato ma Playwright non è installato")
        raise Salta("Playwright non installato")
    def lancia(p):
        eseguibile = os.environ.get("PLAYWRIGHT_CHROMIUM_EXECUTABLE")
        return p.chromium.launch(executable_path=eseguibile) if eseguibile else p.chromium.launch()
    return sync_playwright, lancia


def formato_it_percento(x):
    """Una frazione come la mostra il cruscotto: un decimale, virgola, «%» attaccato (0,4901 → «49,0%»)."""
    return f"{x * 100:.1f}".replace(".", ",") + "%"


def _apri_cruscotto(pagina, html, schema="light", larghezza=1200):
    """Scrive l'HTML in un file e lo apre in una pagina nuova; ritorna (page, errori, eccezioni, richieste)."""
    cartella = tempfile.mkdtemp(prefix="cruscotto-browser-")
    path = os.path.join(cartella, "c.html")
    open(path, "w", encoding="utf-8").write(html)
    contesto = pagina.new_context(color_scheme=schema, viewport={"width": larghezza, "height": 900})
    page = contesto.new_page()
    page.set_default_timeout(5000)
    errori, eccezioni, richieste = [], [], []
    page.on("console", lambda m: errori.append(m.text) if m.type == "error" else None)
    page.on("pageerror", lambda e: eccezioni.append(str(e)))
    page.on("request", lambda r: richieste.append(r.url))
    page.goto("file://" + path)
    page.wait_for_load_state("load")
    return page, errori, eccezioni, richieste


SEZIONI_CRUSCOTTO = ("allocazione", "obiettivi", "glide_path", "pensione", "capitale_umano", "fiscalita", "ribilanciamento")


def _tutti_i_json():
    js = _json_esempio()
    return [js[n] for n in ("alloc.json", "glide.json", "mco.json", "inps.json", "mc.json", "netto.json", "cu.json",
                            "broker.json", "piano.json")]


@test("T17 cruscotto nel browser: si apre senza errori, ogni sezione col suo SVG, nessuna richiesta di rete")
def t_t17_browser_rendering():
    sync_playwright, lancia = browser()
    rc, so, se, b, html = _cruscotto(*_tutti_i_json())
    assert rc == 0, se[-300:]
    rc, _, se, _, solo_alloc = _cruscotto(_json_esempio()["alloc.json"])
    assert rc == 0, se[-300:]
    with sync_playwright() as p:
        br = lancia(p)
        try:
            page, errori, eccezioni, richieste = _apri_cruscotto(br, html)
            assert errori == [] and eccezioni == [], (errori, eccezioni)
            assert all(r.startswith(("file://", "data:")) for r in richieste), richieste
            for k in SEZIONI_CRUSCOTTO:
                assert page.locator(f"#sez-{k} svg").count() >= 1, k
            assert page.locator("#sez-indicatori").count() == 1
            page, errori, eccezioni, _ = _apri_cruscotto(br, solo_alloc)
            assert errori == [] and eccezioni == [], (errori, eccezioni)
            assert page.locator("section[id^=sez-]").count() == 2, page.locator("section[id^=sez-]").count()
        finally:
            br.close()


@test("T17 cruscotto nel browser: numeri in italiano, nomi intatti, null senza NaN, tema scuro, 375px")
def t_t17_browser_contenuti():
    sync_playwright, lancia = browser()
    def pericoloso(d):
        d["obiettivi"][0]["nome"] = NOME_PERICOLOSO
    def senza_sorpasso(d):
        d["eta_sorpasso"], d["stato_sorpasso"] = None, "non_prima_della_fine"
    def senza_ordini(d):
        d["ordini"] = []
    sorgenti = [f for f in _tutti_i_json() if not f.endswith(("mco.json", "cu.json", "piano.json"))]
    sorgenti += [_json_variante("mco.json", pericoloso), _json_variante("cu.json", senza_sorpasso),
                 _json_variante("piano.json", senza_ordini)]
    rc, so, se, b, html = _cruscotto(*sorgenti)
    assert rc == 0, se[-300:]
    rc, _, se, b_tutti, html_tutti = _cruscotto(*_tutti_i_json())
    assert rc == 0, se[-300:]
    ind = {i["id"]: i for i in b_tutti["indicatori"]}
    with sync_playwright() as p:
        br = lancia(p)
        try:
            page, errori, eccezioni, _ = _apri_cruscotto(br, html_tutti)
            assert page.inner_text('[data-valore="prob-0"]') == formato_it_percento(b_tutti["obiettivi"]["obiettivi"][0]["probabilita_successo"])
            assert page.inner_text('[data-valore="sostituzione_netta"]') == formato_it_percento(ind["sostituzione_netta"]["valore"])
            assert page.inner_text('[data-valore="capitale_allocato"]') == "€130.000", page.inner_text('[data-valore="capitale_allocato"]')
            assert page.inner_text('[data-valore="valore_portafoglio"]') == "€8.900"
            page, errori, eccezioni, _ = _apri_cruscotto(br, html)
            assert errori == [] and eccezioni == [], (errori, eccezioni)
            corpo = page.inner_text("body")
            for brutto in ("NaN", "undefined", "null", "Infinity"):
                assert brutto not in corpo, brutto
            # Review Focus 5: il nome pericoloso si legge identico, nessun markup interpretato
            assert NOME_PERICOLOSO in corpo and page.locator("img").count() == 0
            # Review Focus 2 e 4: nessun sorpasso, nessun ordine, detti a parole
            assert "nessun sorpasso" in page.inner_text("#sez-capitale_umano").lower()
            assert "nessun ordine" in page.inner_text("#sez-ribilanciamento").lower()
            sfondo = lambda pg: pg.evaluate("getComputedStyle(document.body).backgroundColor")
            chiaro = sfondo(page)
            scura, errori, eccezioni, _ = _apri_cruscotto(br, html, schema="dark")
            assert sfondo(scura) != chiaro and errori == [] and eccezioni == [], (chiaro, sfondo(scura))
            stretta, errori, eccezioni, _ = _apri_cruscotto(br, html_tutti, larghezza=375)
            larghezza = stretta.evaluate("document.documentElement.scrollWidth")
            assert larghezza <= 375 and errori == [] and eccezioni == [], larghezza
        finally:
            br.close()
    src = open(asset("it-revisione", "cruscotto.html"), encoding="utf-8").read()
    for vietato in ("innerHTML", "outerHTML", "insertAdjacentHTML", "document.write", "eval(", "new Function"):
        assert vietato not in src, vietato


@test("T17 cruscotto: la skill chiede conferma prima di pubblicare, comando e README allineati")
def t_t17_skill():
    skill_md = lambda s: os.path.join(SKILLS, s, "SKILL.md")
    skill = open(skill_md("it-revisione"), encoding="utf-8").read()
    assert "### 6. Cruscotto grafico" in skill, "it-revisione: manca il passo 6 «Cruscotto grafico»"
    passo = skill.split("### 6. Cruscotto grafico")[1].split("\n## ")[0]
    assert "${CLAUDE_SKILL_DIR}/assets/genera-cruscotto.py" in passo, passo[:300]
    assert "chiedi" in passo.lower() and "sì esplicito" in passo and "claude.ai" in passo and "Artifact" in passo, passo[:600]
    anti = skill.split("## Anti-pattern")[1].split("\n## ")[0]
    assert "senza conferma" in anti and "a mano" in anti, anti[-600:]
    assert os.path.isfile(os.path.join(ROOT, "commands", "cruscotto.md"))
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    assert "/cruscotto" in readme and "genera-cruscotto.py" in readme
    dati = readme.split("### Dove vanno i tuoi dati")[1].split("\n### ")[0]
    assert "Artifact" in dati and "conferma" in dati, dati[-600:]
    for s in ("it-pianificazione-obiettivi", "it-pianificazione", "it-rebalance"):
        assert "/cruscotto" in open(skill_md(s), encoding="utf-8").read(), s
    # revisione finale T17: i file di input degli script stanno nella stessa cartella e non si passano;
    # due file dello stesso tipo sono una scelta dell'utente, non un file da rigenerare
    for nome in ("piano.json", "capitale-umano.json", "ips.json"):
        assert nome in passo, f"il passo 6 deve nominare l'input {nome} da non passare"
    assert "stesso tipo" in passo and "chiedi all'utente" in passo.lower(), passo[:1500]
    comando = open(os.path.join(ROOT, "commands", "cruscotto.md"), encoding="utf-8").read()
    assert "--json" in comando and "input" in comando, comando


@test("T17 cruscotto: gli avvisi dell'importatore che nominano le posizioni non entrano nella pagina")
def t_t17_avvisi_posizioni():
    # revisione finale T17: gli avvisi dell'importatore nominano la posizione (descrizione o ISIN), e il
    # testo del consenso dice che le posizioni non ci sono. Restano solo gli avvisi di zainetto e broker.
    tmp = tempfile.mkdtemp()
    try:
        rc, out = run_script(asset("it-tax-optimization", "importa-estratto-broker.py"),
                             os.path.join(FIXTURES, "estratto-pubblici-etc.csv"), "--broker", "Fineco",
                             "--anno", "2026", "--json", "b.json", cwd=tmp)
        assert rc == 0, out[-300:]
        broker = json.load(open(os.path.join(tmp, "b.json"), encoding="utf-8"))
        nomi = [p[k] for p in broker["posizioni"] for k in ("isin", "descrizione") if p.get(k)]
        assert any(n in a for a in broker["avvisi"] for n in nomi), "la fixture deve produrre avvisi sulle posizioni"
        rc, so, se, b, html = _cruscotto(os.path.join(tmp, "b.json"))
        assert rc == 0, se[-300:]
        for n in nomi:
            assert n not in html and n not in json.dumps(b, ensure_ascii=False), n
        assert b["fiscalita"]["avvisi_omessi"] >= 1, b["fiscalita"]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T17 cruscotto: il ribilanciamento sui dati d'esempio lo dice nel JSON e nella pagina")
def t_t17_dati_esempio():
    js = _json_esempio()
    assert json.load(open(js["piano.json"], encoding="utf-8"))["dati_esempio"] is True
    rc, so, se, b, html = _cruscotto(js["piano.json"])
    assert rc == 0 and b["ribilanciamento"]["dati_esempio"] is True, se[-300:]
    # un piano da --dati non è d'esempio
    m, piano = _ribilancia(DATI_ESEMPIO=False)
    assert m.piano_json(piano, m.genera_ordini(piano))["dati_esempio"] is False
    sync_playwright, lancia = browser()
    with sync_playwright() as p:
        br = lancia(p)
        try:
            page, errori, eccezioni, _ = _apri_cruscotto(br, html)
            assert errori == [] and eccezioni == [], (errori, eccezioni)
            for sez in ("#sez-ribilanciamento", "#sez-allocazione"):
                assert "dati di esempio" in page.inner_text(sez).lower(), sez
        finally:
            br.close()


@test("T17 cruscotto nel browser: il template aperto senza dati spiega come generarlo, senza errori")
def t_t17_template_vuoto():
    # Segnalato dal proprietario: aperto così com'è, il template andava in SyntaxError sul segnaposto e la
    # pagina restava vuota. Ora dice cosa fare; lo stesso vale per un blocco dati non leggibile.
    sync_playwright, lancia = browser()
    tpl = open(asset("it-revisione", "cruscotto.html"), encoding="utf-8").read()
    rotto = tpl.replace("__DATI_CRUSCOTTO__", '{"versione": 1, "probabilita": NaN}')
    with sync_playwright() as p:
        br = lancia(p)
        try:
            for html in (tpl, rotto):
                page, errori, eccezioni, _ = _apri_cruscotto(br, html)
                assert errori == [] and eccezioni == [], (errori, eccezioni)
                testo = page.inner_text("#cruscotto")
                assert "genera-cruscotto.py" in testo and "/cruscotto" in testo, testo
                assert page.locator("section[id^=sez-]").count() == 0
        finally:
            br.close()


# ══════════════════════ freshness · json · guardia ══════════════════════

# File con dati normativi che DEVONO portare almeno un marcatore «Verificato: AAAA-MM»:
# senza, i loro numeri resterebbero fuori dal controllo mensile di freshness.yml (F024).
OBBLIGATI_FRESHNESS = (
    "it-tax-optimization/references/aliquote-12-5-26.md",
    "it-tax-optimization/references/etf-trap-redditi-capitale.md",
    "it-tax-optimization/references/regime-dichiarativo.md",
    "it-tax-optimization/references/criptovalute.md",
    "it-tax-optimization/references/criptovalute-avanzate.md",
    "it-tax-optimization/references/criptovalute-casi-speciali.md",
    "it-tax-optimization/references/dividendi-esteri.md",
    "it-tax-optimization/references/successione-donazione.md",
    "it-tax-optimization/references/multi-intermediario.md",
    "it-tax-optimization/references/import-dati-broker.md",
    "it-tax-optimization/references/mappa-nuovo-tuir.md",
    "it-pianificazione/SKILL.md",
    "it-pianificazione/references/previdenza-obbligatoria.md",
    "it-pianificazione/references/previdenza-complementare.md",
    "it-pianificazione/assets/netto-pensione.py",
    "it-pianificazione/assets/proiezione-inps.py",
    "it-pianificazione/assets/ponte-pensione.py",
    "it-pianificazione/assets/importa-previdenza.py",
    "it-pianificazione/assets/riscatto-laurea.py",
    "it-pianificazione-autonomi/references/casse-professionali.md",
    "it-pianificazione-autonomi/references/previdenza-gestione-separata.md",
    "it-pianificazione-autonomi/references/fiscalita-forfettario.md",
    "it-pianificazione-obiettivi/references/asset-location.md",
    "it-pianificazione-obiettivi/references/capitale-umano.md",
    "it-rebalance/SKILL.md",
    "it-revisione/assets/genera-ips.py",
)


def marcatori_freshness():
    """[(file relativo a skills/, anno, mese)] per ogni «Verificato: AAAA-MM» nei .py e .md."""
    trovati = []
    for base, _, files in os.walk(SKILLS):
        for f in sorted(files):
            if f.endswith((".py", ".md")):
                rel = os.path.relpath(os.path.join(base, f), SKILLS).replace(os.sep, "/")
                testo = open(os.path.join(base, f), encoding="utf-8", errors="replace").read()
                for m in re.finditer(r"Verificato: (\d{4})-(\d{2})", testo):
                    trovati.append((rel, int(m.group(1)), int(m.group(2))))
    return trovati


def tabella_freshness():
    """Tabella Markdown per il README: per ogni file la verifica più vecchia e fin quando vale."""
    piu_vecchia = {}
    for rel, a, m in marcatori_freshness():
        piu_vecchia[rel] = min(piu_vecchia.get(rel, (a, m)), (a, m))
    righe = ["| File | Verificato | Valido fino a |", "|---|---|---|"]
    for rel, (a, m) in sorted(piu_vecchia.items(), key=lambda kv: (kv[1], kv[0])):
        righe.append(f"| `{rel}` | {a}-{m:02d} | {a + 1}-{m:02d} |")
    return "\n".join(righe)


@test("freshness fiscale: i file normativi portano 'Verificato: AAAA-MM' e nessuno supera i 12 mesi")
def t_freshness():
    import datetime
    trovati = marcatori_freshness()
    marcati = {rel for rel, _, _ in trovati}
    mancanti = [rel for rel in OBBLIGATI_FRESHNESS if rel not in marcati]
    assert not mancanti, ("file normativi senza marcatore (verificarli alle fonti primarie e "
                          f"aggiungere «Verificato: AAAA-MM»): {mancanti}")
    oggi = datetime.date.today()
    for f, a, m in trovati:
        eta = (oggi.year - a) * 12 + (oggi.month - m)
        assert 0 <= eta <= 12, (f"{f}: dato normativo verificato {a}-{m:02d}, {eta} mesi fa. "
                                "È il design: i dati fiscali scadono — riverificare alle fonti "
                                "e aggiornare il marcatore.")


@test("pipeline --json: proiezione-inps scrive i risultati chiave su file")
def t_json():
    import json as _json
    tmp = tempfile.mkdtemp()
    out = os.path.join(tmp, "p.json")
    rc, log = run_script(asset("it-pianificazione", "proiezione-inps.py"),
                         "--set", "MONTANTE_ATTUALE=100_000.0", "--set", "RAL_ATTUALE=30_000.0",
                         "--set", "ETA_ATTUALE=37.0", "--set", "ETA_PENSIONAMENTO=67.0",
                         "--set", "CRESCITA_REALE_RAL=0.0", "--set", "PIL_REALE=0.02",
                         "--set", "COEFF_ESPLICITO=0.05", "--json", out)
    assert rc == 0, log[-300:]
    d = _json.load(open(out, encoding="utf-8"))
    shutil.rmtree(tmp, ignore_errors=True)
    atteso = (100_000 * 1.02 ** 30 + 30_000 * 0.33 * ((1.02 ** 30 - 1) / 0.02)) * 0.05
    assert abs(d["pensione_inps_lorda_annua"] - atteso) <= 2
    for k in ("montante_finale", "ral_finale", "tasso_sostituzione_lordo_vs_ral_finale"):
        assert k in d, f"chiave JSON mancante: {k}"


@test("T9 catena --json → --from: importa-previdenza → proiezione INPS → Monte Carlo → netto e ponte, stessi numeri")
def t_catena_previdenziale():
    # F027: la catena era promessa in SKILL.md e nel changelog ma nessuno script leggeva un JSON.
    import json as _json
    P = lambda n: asset("it-pianificazione", n)
    tmp = tempfile.mkdtemp()
    j = lambda n: os.path.join(tmp, n)
    leggi = lambda n: _json.load(open(j(n), encoding="utf-8"))
    try:
        rc, out = run_script(P("importa-previdenza.py"), os.path.join(FIXTURES, "estratto-inps-test.xml"),
                             os.path.join(FIXTURES, "storico-fondo-test.csv"), "--anonimo", "--json", j("dati.json"))
        assert rc == 0, out[-300:]
        dati = leggi("dati.json")
        assert set(dati) >= {"montante_attuale", "ral_attuale", "eta_attuale", "fp_montante_attuale", "fp_versamento_annuo"}, dati
        assert "ROSSI" not in open(j("dati.json"), encoding="utf-8").read() and dati["eta_attuale"] == int(dati["eta_attuale"]), dati
        assert dati["ral_attuale"] == 31000 and dati["fp_montante_attuale"] == 2000, dati
        rc, out = run_script(P("proiezione-inps.py"), "--from", j("dati.json"), "--json", j("inps.json"))
        assert rc == 0 and "[from] dati.json: MONTANTE_ATTUALE=" in out and "RAL_ATTUALE=31000" in out, out[-400:]
        inps = leggi("inps.json")
        assert abs(euro(out, "PENSIONE LORDA ANNUA") - inps["pensione_inps_lorda_annua"]) <= 1
        assert inps["eta_pensionamento"] == 69 and "RAL: € 31.000 oggi" in out, out[:400]
        rc, out = run_script(P("montecarlo-pensione.py"), "--from", j("dati.json"), "--from", j("inps.json"),
                             "--set", "N_SIM=2000", "--json", j("mc.json"))
        assert rc == 0 and "[from] inps.json: PENSIONE_INPS_ANNUA=" in out and "[from] dati.json: FP_POT_ATTUALE=2000" in out, out[-500:]
        mc = leggi("mc.json")
        assert mc["pensione_inps_lorda"] == inps["pensione_inps_lorda_annua"] and mc["ral_finale"] == inps["ral_finale"], (mc, inps)
        assert mc["anni"] == inps["anni_al_pensionamento"]
        rc, out = run_script(P("netto-pensione.py"), "--from", j("mc.json"), "--json", j("netto.json"))
        assert rc == 0 and "[from] mc.json: PENSIONE_INPS_LORDA=" in out and "RENDITE_FP=" in out, out[-400:]
        netto = leggi("netto.json")
        mod = _modulo_previdenza("netto-pensione.py")
        atteso, _, _ = mod.netto_pensionato(inps["pensione_inps_lorda_annua"], mod.ADDIZIONALE_REGIONALE, mod.ADDIZIONALE_COMUNALE)
        assert abs(netto["netto_pensione_inps"] - atteso) < 0.01, (netto["netto_pensione_inps"], atteso)
        assert [s["rendita_lorda"] for s in netto["scenari"]] == [mc["rendita_fp"]["p10"], mc["rendita_fp"]["mediana"], mc["rendita_fp"]["p90"]]
        # revisione T9: mc.json è il montante alla vecchiaia (69), non all'uscita (62). Il ponte lo
        # rifiuta e dice come rigenerarlo; con il Monte Carlo fino all'uscita la catena regge.
        rc, out = run_script(P("ponte-pensione.py"), "--from", j("inps.json"), "--from", j("mc.json"),
                             "--set", "ETA_USCITA=62", "--set", "N_SIM=2000")
        eta_oggi = inps["eta_pensionamento"] - inps["anni_al_pensionamento"]
        assert rc == 1 and "ERRORE:" in out and f"--set ANNI={int(62 - eta_oggi)}" in out and "Traceback" not in out, (rc, out[-500:])
        rc, out = run_script(P("montecarlo-pensione.py"), "--from", j("dati.json"), "--from", j("inps.json"),
                             "--set", f"ANNI={int(62 - eta_oggi)}", "--set", "N_SIM=2000", "--json", j("mc-uscita.json"))
        assert rc == 0, out[-300:]
        mc_uscita = leggi("mc-uscita.json")
        rc, out = run_script(P("ponte-pensione.py"), "--from", j("inps.json"), "--from", j("mc-uscita.json"),
                             "--set", "ETA_USCITA=62", "--set", "N_SIM=2000", "--json", j("ponte.json"))
        assert rc == 0 and "[from] mc-uscita.json: MONTANTE_FP_ALL_USCITA=" in out, out[-400:]
        ponte = leggi("ponte.json")
        assert ponte["montante_fp_all_uscita"] == mc_uscita["montante_fp"]["mediana"] < mc["montante_fp"]["mediana"], ponte
        assert ponte["eta_pensione_vecchiaia"] == 69, ponte
        assert abs(ponte["rita"]["rata_lorda"] - mc_uscita["montante_fp"]["mediana"] / 5) < 0.01, ponte["rita"]
        # il montante forzato a mano con --set non passa dal controllo
        rc, out = run_script(P("ponte-pensione.py"), "--from", j("inps.json"), "--from", j("mc.json"),
                             "--set", "ETA_USCITA=62", "--set", "MONTANTE_FP_ALL_USCITA=100000", "--set", "N_SIM=500")
        assert rc == 0, out[-300:]
        # precedenza: --set prevale su --from; un JSON di un altro script → rc 1; file mancante → rc 1
        rc, out = run_script(P("montecarlo-pensione.py"), "--from", j("inps.json"), "--set", "PENSIONE_INPS_ANNUA=1.0",
                             "--set", "N_SIM=500", "--json", j("mc2.json"))
        assert rc == 0 and leggi("mc2.json")["pensione_inps_lorda"] == 1.0, out[-300:]
        rc, out = run_script(P("proiezione-inps.py"), "--from", j("mc.json"))
        assert rc == 1 and "nessuna delle chiavi attese" in out and "Traceback" not in out, (rc, out[-300:])
        # contestato in revisione T9: la posizione mista (retribuzioni prima del 1996) si perdeva nella
        # catena; ora la proiezione la legge dal JSON dell'importatore e lo dice prima del numero
        _json.dump(dict(dati, contributi_solo_dal_1996=True), open(j("mista.json"), "w", encoding="utf-8"))
        rc, out = run_script(P("proiezione-inps.py"), "--from", j("mista.json"))
        assert rc == 0 and "posizione mista" in out and "La mia pensione futura" in out, out[-400:]
        assert out.index("posizione mista") < out.index("PENSIONE LORDA ANNUA"), "avviso dopo il numero"
        rc, out = run_script(P("proiezione-inps.py"), "--from", j("dati.json"))
        assert "posizione mista" not in out, out[-300:]
        rc, out = run_script(P("proiezione-inps.py"), "--from", j("manca.json"))
        assert rc == 1 and "non trovato" in out, (rc, out[-300:])
        rc, out = run_script(P("netto-pensione.py"), "--help")
        assert "chiavi lette da --from" in out and "rendita_fp → RENDITE_FP" in out, out[-500:]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T9 contratto di input della revisione: ogni chiave della tabella esiste nel JSON del suo script")
def t_contratto_revisione():
    # F031 e revisione T9: la tabella dice quale campo va in quale cella; una chiave rinominata in uno
    # script la romperebbe in silenzio. B12 e B13 della pensione stanno sulla stessa base (lorda).
    import json as _json
    skill = open(os.path.join(SKILLS, "it-revisione", "SKILL.md"), encoding="utf-8").read()
    tabella = skill.split("### Contratto di input")[1].split("\n## ")[0]
    tmp = tempfile.mkdtemp()
    pian = asset("it-pianificazione-obiettivi", "pianificatore-obiettivi.xlsx")
    lanci = {
        "mc.json": (asset("it-pianificazione", "montecarlo-pensione.py"), ["--set", "N_SIM=500"]),
        "netto.json": (asset("it-pianificazione", "netto-pensione.py"), []),
        "mco.json": (asset("it-pianificazione-obiettivi", "montecarlo-obiettivi.py"), [pian, "--set", "N_SIM=500"]),
        "alloc.json": (asset("it-pianificazione-obiettivi", "allocatore-capitale.py"), [pian]),
        "glide.json": (asset("it-pianificazione-obiettivi", "glidepath-obiettivi.py"), [pian]),
        "cu.json": (asset("it-pianificazione-obiettivi", "capitale-umano.py"),
                    ["--composizione", asset("it-pianificazione-obiettivi", "esempio-composizione.csv")]),
        "imp.json": (asset("it-tax-optimization", "importa-estratto-broker.py"),
                     [os.path.join(FIXTURES, "estratto-multibroker.csv"), os.path.join(FIXTURES, "zainetto-multibroker.csv"), "--anno", "2026"]),
    }
    try:
        js = {}
        for nome, (script, argomenti) in lanci.items():
            rc, out = run_script(script, *argomenti, "--json", os.path.join(tmp, nome), cwd=tmp)
            assert rc == 0, f"{os.path.basename(script)}: {out[-300:]}"
            js[nome] = _json.load(open(os.path.join(tmp, nome), encoding="utf-8"))
        def valori(dato, percorso):
            """I valori al percorso «a.b», «a[].b» o «a[nome=x].b» (tutti gli elementi della lista o quello col nome)."""
            correnti = [dato]
            for pezzo in percorso.split("."):
                m = re.fullmatch(r"(\w+)(?:\[(?:(\w+)=\"?([^\]\"]*)\"?)?\])?", pezzo)
                chiave, filtro_k, filtro_v = m.groups()
                prossimi = []
                for c in correnti:
                    assert isinstance(c, dict) and chiave in c, f"manca «{chiave}» in {percorso}"
                    v = c[chiave]
                    if "[" in pezzo:
                        elementi = list(v.values()) if isinstance(v, dict) else v
                        if filtro_k:
                            elementi = [e for e in elementi if str(e.get(filtro_k)) == filtro_v]
                        assert elementi, f"{percorso}: nessun elemento"
                        prossimi.extend(elementi)
                    else:
                        prossimi.append(v)
                correnti = prossimi
            return correnti
        attese = (("mc.json", "tasso_sostituzione_lordo.mediana"), ("mc.json", "target"),
                  ("netto.json", 'scenari[nome="mediana"].tasso_sostituzione_netto'),
                  ("mco.json", "obiettivi[].probabilita_successo"), ("mco.json", "obiettivi[].gia"),
                  ("mco.json", "obiettivi[].target_oggi"), ("mco.json", "obiettivi[].priorita"),
                  ("alloc.json", "aggregato.azioni"), ("alloc.json", "aggregato.totale"), ("glide.json", "portafoglio"),
                  ("imp.json", "posizioni[].valore"), ("imp.json", "analisi.per_broker[].zainetto_in_scadenza"),
                  ("imp.json", "analisi.per_broker[].plus_realizzabili_a_costo_zero"),
                  ("cu.json", "esposizione_finanziaria"), ("cu.json", "esposizione_totale"),
                  ("cu.json", "settore.quota_settore_fin"), ("cu.json", "settore.sovrappeso"), ("cu.json", "eta_sorpasso"),
                  ("cu.json", "stato_sorpasso"))
        for nome, percorso in attese:
            assert percorso.replace("[]", "[") .split("[")[0].split(".")[-1] in tabella or percorso in tabella, \
                f"la tabella del contratto non nomina «{percorso}»"
            assert valori(js[nome], percorso), percorso
        riga = next(r for r in tabella.splitlines() if r.startswith("| Pensione, dipendenti"))
        assert "`tasso_sostituzione_lordo.mediana`" in riga and "`target`" in riga and "mai in B12" in riga, riga
        assert js["mc.json"]["target"] == 0.70 and 0 < js["mc.json"]["tasso_sostituzione_lordo"]["mediana"] < 2
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T9 --json di glide path e Monte Carlo degli obiettivi: gli stessi numeri dello stdout, chiavi del contratto")
def t_json_obiettivi():
    import json as _json
    OB = "it-pianificazione-obiettivi"
    tmp = tempfile.mkdtemp()
    try:
        rc, out = run_script(asset(OB, "glidepath-obiettivi.py"), "--json", os.path.join(tmp, "g.json"))
        assert rc == 0, out[-300:]
        g = _json.load(open(os.path.join(tmp, "g.json"), encoding="utf-8"))
        nomi = {o["nome"]: o for o in g["obiettivi"]}
        assert set(nomi) == {"Fondo di emergenza", "Anticipo casa", "Lungo termine / libertà"} and g["finestra"] == 7
        lungo = nomi["Lungo termine / libertà"]
        assert lungo["quota_azionaria"][:14] == [0.65] * 14 and lungo["quota_azionaria"][20] == 0.0 and len(lungo["quota_azionaria"]) == 21
        assert nomi["Anticipo casa"]["quota_azionaria"][7] is None, "terminante concluso → None"
        assert nomi["Fondo di emergenza"]["permanente"] and nomi["Fondo di emergenza"]["quota_azionaria"][20] == 0.0
        assert f"Lungo termine / libertà: {_euro_it(lungo['versamento_mensile'])}/mese" in out, lungo["versamento_mensile"]
        assert len(g["portafoglio"]) == 21 and all(0 <= x <= 1 for x in g["portafoglio"]) and g["portafoglio"][20] == 0.0
        assert _pct_it(g["portafoglio"][3] * 100, 0) + "%" in out.split("PORTAFOGLIO")[1], g["portafoglio"][3]
        rc, out = run_script(asset(OB, "montecarlo-obiettivi.py"), "--set", "N_SIM=3000", "--set", "MODELLO='bootstrap'",
                             "--json", os.path.join(tmp, "m.json"))
        assert rc == 0, out[-300:]
        m = _json.load(open(os.path.join(tmp, "m.json"), encoding="utf-8"))
        prob = _prob_obiettivi(out)
        for o in m["obiettivi"]:
            assert round(o["probabilita_successo"] * 100) == prob[o["nome"]], (o["nome"], o["probabilita_successo"], prob)
            assert o["capitale_finale"]["p10"] <= o["capitale_finale"]["mediana"] <= o["capitale_finale"]["p90"]
            assert o["versamento_confidenza"] >= o["versamento_deterministico"] >= 0 and o["confidenza"] == 0.85
        assert abs(sum(o["versamento_deterministico"] for o in m["obiettivi"]) - m["totale_versamento_deterministico"]) < 0.05
        assert m["modello"] == "bootstrap" and m["n_sim"] == 3000
        rs = m["rischio_sistemico"]
        assert rs and rs["tutti_falliscono"]["congiunta"] >= rs["tutti_falliscono"]["indipendenti"], rs
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("guardia intestazioni: colonna spostata → errore esplicito, non dati sbagliati")
def t_guardia():
    from openpyxl import load_workbook
    tmp = tempfile.mkdtemp()
    rotto = os.path.join(tmp, "rotto.xlsx")
    shutil.copy(asset("it-pianificazione-obiettivi", "pianificatore-obiettivi.xlsx"), rotto)
    wb = load_workbook(rotto)
    wb["Obiettivi"]["B4"] = "Colonna a caso"
    wb.save(rotto)
    rc, out = run_script(asset("it-pianificazione-obiettivi", "glidepath-obiettivi.py"), rotto)
    shutil.rmtree(tmp, ignore_errors=True)
    assert rc != 0, "lo script deve FALLIRE su layout non riconosciuto"
    assert "NON riconosciuta" in out, out[-300:]


# ═══════════════════════════ EXCEL RICALCOLATI ═══════════════════════════
# openpyxl legge solo le cache salvate nei file: questi test ricalcolano gli
# xlsx con LibreOffice in locale italiano (scripts/ricalcola-excel.py). Senza
# LibreOffice i test sono saltati, salvo RICHIEDI_LIBREOFFICE=1 (la CI).

_RICALCOLATI = {}
VOLATILI = ("TODAY(", "NOW(", "RAND(", "RANDBETWEEN(")


def ricalcolatore():
    mod = carica_modulo(os.path.join(ROOT, "scripts", "ricalcola-excel.py"))
    if not mod.trova_soffice():
        if os.environ.get("RICHIEDI_LIBREOFFICE"):
            raise AssertionError("RICHIEDI_LIBREOFFICE è impostato ma LibreOffice non è installato")
        raise Salta("LibreOffice non installato")
    return mod


def xlsx_ricalcolati():
    """{xlsx del plugin: copia ricalcolata}, un solo ricalcolo per esecuzione della suite."""
    if not _RICALCOLATI:
        mod = ricalcolatore()
        originali = mod.tutti_gli_xlsx()
        _RICALCOLATI.update(zip(originali, mod.ricalcola(originali, tempfile.mkdtemp(prefix="ricalcolo-"))))
    return _RICALCOLATI


@test("ricalcolo Excel: LibreOffice ricalcola al caricamento, in locale italiano")
def t_ricalcolo():
    import zipfile
    from openpyxl import Workbook, load_workbook
    mod = ricalcolatore()
    tmp = tempfile.mkdtemp()
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = 5, "=A1*3", "=FIXED(1.5,1)"
    nuovo = os.path.join(tmp, "nuovo.xlsx")
    wb.save(nuovo)
    # Cache vecchia: A2 dichiara 6 (come se A1 valesse ancora 2), A3 "1.5".
    vecchio = os.path.join(tmp, "vecchio.xlsx")
    with zipfile.ZipFile(nuovo) as zi, zipfile.ZipFile(vecchio, "w") as zo:
        for voce in zi.infolist():
            dati = zi.read(voce.filename)
            if voce.filename == "xl/worksheets/sheet1.xml":
                dati = re.sub(rb"(<c r=\"A2\"[^>]*>\s*<f>A1\*3</f>)\s*(<v\s*/>|<v>\s*</v>)?", rb"\1<v>6</v>", dati)
                dati = re.sub(rb"(<c r=\"A3\")([^>]*>\s*<f>[^<]*</f>)\s*(<v\s*/>|<v>\s*</v>)?",
                              rb'\1 t="str"\2<v>1.5</v>', dati)
            zo.writestr(voce, dati)
    assert load_workbook(vecchio, data_only=True).active["A2"].value == 6, "fixture: cache vecchia non scritta"
    (out,) = mod.ricalcola([vecchio], os.path.join(tmp, "out"))
    ws = load_workbook(out, data_only=True).active
    assert ws["A2"].value == 15, f"cache non ricalcolata: A2 = {ws['A2'].value}"
    assert ws["A3"].value == "1,5", f"locale non italiano: FIXED(1.5;1) = {ws['A3'].value!r}"


@test("Excel: ricalcolati in locale italiano, nessuna cella in errore")
def t_excel_errori():
    mod = ricalcolatore()
    for orig, ricalcolato in xlsx_ricalcolati().items():
        trovati = mod.errori(ricalcolato)
        assert not trovati, f"{os.path.basename(orig)}: {len(trovati)} celle in errore, es. {trovati[:3]}"


@test("Excel: le cache consegnate coincidono col ricalcolo (file aggiornati prima del commit)")
def t_excel_cache():
    import math
    from openpyxl import load_workbook
    problemi = []
    for orig, ricalcolato in xlsx_ricalcolati().items():
        formule = load_workbook(orig)
        if any(isinstance(c.value, str) and any(v in c.value.upper() for v in VOLATILI)
               for ws in formule.worksheets for riga in ws.iter_rows() for c in riga):
            continue  # con OGGI() la cache dipende dal giorno: per questi file contano solo gli errori
        cache, nuovo = load_workbook(orig, data_only=True), load_workbook(ricalcolato, data_only=True)
        diverse = []
        for ws in formule.worksheets:
            for riga in ws.iter_rows():
                for c in riga:
                    if not (isinstance(c.value, str) and c.value.startswith("=")):
                        continue
                    a, b = cache[ws.title][c.coordinate].value, nuovo[ws.title][c.coordinate].value
                    numeri = all(isinstance(x, (int, float)) for x in (a, b))
                    if not (a == b or (numeri and math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-9))):
                        diverse.append(f"{ws.title}!{c.coordinate}: {a!r} → {b!r}")
        if diverse:
            problemi.append(f"{os.path.basename(orig)}: {len(diverse)} cache diverse, es. {diverse[:2]}")
    assert not problemi, ("; ".join(problemi) + ". Se le formule sono giuste: "
                          "python scripts/ricalcola-excel.py FILE --in-place")


# ═══════════════════════ T12: test che vincolano i numeri ═══════════════════════

# Struttura attesa di ogni xlsx del plugin: fogli nell'ordine, celle di ancoraggio e soglia minima di
# formule (il conteggio di oggi, arrotondato per difetto: un foglio svuotato per sbaglio emerge qui).
STRUTTURA_XLSX = {
    "it-pianificazione/assets/proiezione-pensione.xlsx": {
        "fogli": ["Input", "Coefficienti", "Pubblica", "Complementare", "Scenari"], "formule": 1090,
        "ancore": {("Input", "A5"): "Età corrente", ("Pubblica", "A57"): "RISULTATI FINALI A PENSIONAMENTO",
                   ("Complementare", "A57"): "RISULTATI FINALI A PENSIONAMENTO", ("Scenari", "A4"): "Voce"}},
    "it-pianificazione-autonomi/assets/proiezione-pensione-autonomi.xlsx": {
        "fogli": ["Input", "Coefficienti", "Previdenza", "Fiscalità", "Complementare", "Riepilogo"], "formule": 690,
        "ancore": {("Input", "A5"): "Età attuale", ("Previdenza", "A57"): "RISULTATI A PENSIONAMENTO",
                   ("Riepilogo", "A13"): "Tasso di sostituzione — totale"}},
    "it-pianificazione-obiettivi/assets/pianificatore-obiettivi.xlsx": {
        "fogli": ["Parametri", "Obiettivi", "Riepilogo"], "formule": 110,
        "ancore": {("Parametri", "A9"): "Profilo", ("Obiettivi", "A4"): "Obiettivo",
                   ("Obiettivi", "A16"): "TOTALE versamenti mensili", ("Riepilogo", "A4"): "RICONCILIAZIONE"}},
    "it-revisione/assets/cruscotto-patrimoniale.xlsx": {
        "fogli": ["Leggimi", "Patrimonio", "Allocazione", "Andamento", "Sintesi", "Azioni", "Rendimento TWR"], "formule": 415,
        "ancore": {("Patrimonio", "A21"): "PATRIMONIO NETTO", ("Allocazione", "A4"): "Classe",
                   ("Sintesi", "A12"): "Tasso di sostituzione proiettato", ("Rendimento TWR", "A5"): "Data"}},
    "it-revisione/assets/investment-policy-statement.xlsx": {
        "fogli": ["IPS"], "formule": 14,
        "ancore": {("IPS", "A1"): "Investment Policy Statement (IPS)", ("IPS", "A25"): "Totale"}},
    "it-revisione/assets/tracker-secchielli.xlsx": {
        "fogli": ["Leggimi", "Impostazioni", "Glide", "Posizioni", "Controllo", "Versamenti"], "formule": 145,
        "ancore": {("Impostazioni", "A2"): "Anno del piano (0 = partenza):", ("Posizioni", "A3"): "Strumento",
                   ("Controllo", "A18"): "TOTALE", ("Versamenti", "A3"): "Data"}},
}


@test("T12 Excel: ogni xlsx ha i fogli attesi, le celle di ancoraggio e almeno le formule di oggi")
def t_excel_struttura():
    from openpyxl import load_workbook
    for rel, attesa in STRUTTURA_XLSX.items():
        wb = load_workbook(os.path.join(SKILLS, rel))
        nome = os.path.basename(rel)
        assert wb.sheetnames == attesa["fogli"], f"{nome}: fogli {wb.sheetnames} ≠ {attesa['fogli']}"
        formule = sum(1 for ws in wb.worksheets for riga in ws.iter_rows() for c in riga
                      if isinstance(c.value, str) and c.value.startswith("="))
        assert formule >= attesa["formule"], f"{nome}: {formule} formule, attese almeno {attesa['formule']}"
        for (foglio, cella), testo in attesa["ancore"].items():
            assert str(wb[foglio][cella].value).strip() == testo, f"{nome}: {foglio}!{cella} = {wb[foglio][cella].value!r}, atteso {testo!r}"
        for ws in wb.worksheets:
            for riga in ws.iter_rows():
                for c in riga:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        assert "#REF!" not in c.value and "#NAME?" not in c.value, f"{nome}: {ws.title}!{c.coordinate} = {c.value}"
    assert len(STRUTTURA_XLSX) == 6 == sum(1 for _ in __import__("glob").glob(os.path.join(SKILLS, "*", "assets", "*.xlsx"))), \
        "un xlsx nuovo o rimosso: aggiornare STRUTTURA_XLSX"


@test("T12 Excel: Riepilogo del pianificatore e Patrimonio del cruscotto letti dal ricalcolo, tracker con cache uguali al ricalcolo")
def t_excel_valori_sentinella():
    from openpyxl import load_workbook
    ricalcolati = xlsx_ricalcolati()
    # pianificatore: la riconciliazione del Riepilogo ricostruita dai fogli a monte
    orig = asset("it-pianificazione-obiettivi", "pianificatore-obiettivi.xlsx")
    wb = load_workbook(ricalcolati[orig], data_only=True)
    r, o, par = wb["Riepilogo"], wb["Obiettivi"], wb["Parametri"]
    versamenti = sum(o[f"I{k}"].value or 0 for k in range(5, 15))
    assert abs(r["B5"].value - versamenti) < 0.01 and abs(r["B5"].value - o["I16"].value) < 0.01, (r["B5"].value, versamenti)
    assert r["B6"].value == par["B5"].value == 1200 and abs(r["B7"].value - (1200 - versamenti)) < 0.01
    assert r["B8"].value == ("Sostenibile" if r["B7"].value >= 0 else "NON sostenibile")
    per_priorita = {p: sum((o[f"I{k}"].value or 0) for k in range(5, 15) if o[f"G{k}"].value == p)
                    for p in ("essenziale", "importante", "desiderabile")}
    assert abs(r["B12"].value - per_priorita["essenziale"]) < 0.01 and abs(r["B14"].value - per_priorita["desiderabile"]) < 0.01
    # revisione T12: anche il totale «importante», e le tre priorità che ricompongono i versamenti
    assert abs(r["B13"].value - per_priorita["importante"]) < 0.01, (r["B13"].value, per_priorita)
    assert abs(r["B12"].value + r["B13"].value + r["B14"].value - r["B5"].value) < 0.01
    assert abs(r["B18"].value - sum(o[f"F{k}"].value or 0 for k in range(5, 15))) < 0.01 and r["B18"].value == 5000
    # cruscotto, modello consegnato: senza debiti il netto coincide con le attività e le quote sommano a 1
    orig = asset("it-revisione", "cruscotto-patrimoniale.xlsx")
    pat = load_workbook(ricalcolati[orig], data_only=True)["Patrimonio"]
    attivita = sum(pat[f"B{k}"].value or 0 for k in range(6, 14))
    passivita = sum(pat[f"B{k}"].value or 0 for k in range(17, 19))
    assert pat["B14"].value == attivita == 64000 and pat["B19"].value == passivita and pat["B21"].value == attivita - passivita
    assert abs(sum(pat[f"C{k}"].value for k in range(6, 14)) - 1) < 1e-9, "le quote sul netto devono sommare a 1"
    assert abs(pat["C6"].value - 40000 / 64000) < 1e-9
    # revisione T12: con passività a zero il modello non distingue netto e attività. La variante con
    # un mutuo (B17), un altro debito (B18) e una voce in B12 vincola netto = attività − passività,
    # quote sul netto (sommano a più di 1) e il totale delle attività esteso a B6:B13.
    tmp = tempfile.mkdtemp()
    try:
        wb = load_workbook(orig)
        wb["Patrimonio"]["B17"], wb["Patrimonio"]["B18"], wb["Patrimonio"]["B12"] = 20000, 5000, 1000
        variante = os.path.join(tmp, "debiti.xlsx")
        wb.save(variante)
        (ricalcolata,) = ricalcolatore().ricalcola([variante], os.path.join(tmp, "out"))
        v = load_workbook(ricalcolata, data_only=True)["Patrimonio"]
        assert (v["B14"].value, v["B19"].value, v["B21"].value) == (65000, 25000, 40000), \
            (v["B14"].value, v["B19"].value, v["B21"].value)
        assert abs(v["C6"].value - 1) < 1e-9 and abs(v["C12"].value - 1000 / 40000) < 1e-9, (v["C6"].value, v["C12"].value)
        assert abs(sum(v[f"C{k}"].value for k in range(6, 14)) - 65000 / 40000) < 1e-9, "quote non calcolate sul netto"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    # tracker: le cache consegnate coincidono col ricalcolo, salvo le celle volatili (TODAY)
    import math
    orig = asset("it-revisione", "tracker-secchielli.xlsx")
    formule = load_workbook(orig)
    cache, nuovo = load_workbook(orig, data_only=True), load_workbook(ricalcolati[orig], data_only=True)
    confrontate, diverse = 0, []
    for ws in formule.worksheets:
        for riga in ws.iter_rows():
            for c in riga:
                if not (isinstance(c.value, str) and c.value.startswith("=")) or any(v in c.value.upper() for v in VOLATILI):
                    continue
                a, b = cache[ws.title][c.coordinate].value, nuovo[ws.title][c.coordinate].value
                confrontate += 1
                numeri = all(isinstance(x, (int, float)) for x in (a, b))
                if not (a == b or (numeri and math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-9))):
                    diverse.append(f"{ws.title}!{c.coordinate}: {a!r} → {b!r}")
    assert confrontate >= 139 and not diverse, (confrontate, diverse[:3])


# Valori chiusi sui default (seed 42). Cambiano se cambia la meccanica, ma anche se cambiano i parametri
# normativi di netto-pensione.py (scaglioni, detrazioni, addizionali, contributi) o la serie
# rendimenti-storici.csv con la CPI (montecarlo-pensione e ponte-pensione: vale anche per
# PROB_PONTE_DEFAULT e il capitale necessario). Si aggiornano insieme al dato, seguendo il «Calendario»
# del README, e la modifica va detta nel changelog.
MC_PENSIONE_MEDIANA_DEFAULT = 169_205
NETTO_TASSO_NETTO_MEDIANA_DEFAULT = 0.9693


@test("T12 --json previdenziali: chiavi e valori chiusi di montecarlo-pensione, netto-pensione e ponte-pensione")
def t_json_previdenziali():
    P = lambda n: asset("it-pianificazione", n)
    tmp = tempfile.mkdtemp()
    try:
        j = os.path.join(tmp, "x.json")
        rc, out = run_script(P("montecarlo-pensione.py"), "--json", j)
        assert rc == 0, out[-300:]
        mc = json.load(open(j, encoding="utf-8"))
        assert set(mc) >= {"montante_fp", "rendita_fp", "pensione_inps_lorda", "ral_finale", "anni", "tasso_sostituzione_lordo", "prob_target", "target"}, sorted(mc)
        assert mc["montante_fp"]["mediana"] == MC_PENSIONE_MEDIANA_DEFAULT, mc["montante_fp"]
        assert mc["montante_fp"]["p10"] < mc["montante_fp"]["mediana"] < mc["montante_fp"]["p90"]
        assert mc["rendita_fp"]["mediana"] == round(mc["montante_fp"]["mediana"] * 0.05), mc["rendita_fp"]
        assert abs(mc["tasso_sostituzione_lordo"]["mediana"] - (21000 + mc["rendita_fp"]["mediana"]) / 45000) < 1e-3
        assert euro(out, "mediana simulata") == MC_PENSIONE_MEDIANA_DEFAULT, "lo stdout e il JSON devono dire la stessa mediana"
        rc, out = run_script(P("netto-pensione.py"), "--json", j)
        assert rc == 0, out[-300:]
        ne = json.load(open(j, encoding="utf-8"))
        mediana = next(s for s in ne["scenari"] if s["nome"] == "mediana")
        assert mediana["tasso_sostituzione_netto"] == NETTO_TASSO_NETTO_MEDIANA_DEFAULT, mediana
        assert abs(mediana["pensione_netta_totale"] - (ne["netto_pensione_inps"] + mediana["rendita_netta"])) < 0.01
        assert mediana["tasso_sostituzione_netto"] > mediana["tasso_sostituzione_lordo"], "il netto batte il lordo"
        rc, out = run_script(P("ponte-pensione.py"), "--json", j)
        assert rc == 0, out[-300:]
        po = json.load(open(j, encoding="utf-8"))
        assert po["probabilita_successo"] == 0.2107 and round(po["probabilita_successo"] * 100, 1) == PROB_PONTE_DEFAULT, po
        assert po["capitale_necessario"]["85%"] < po["capitale_necessario"]["95%"], po["capitale_necessario"]
        assert po["capitale_necessario"] == {"85%": 53167.0, "95%": 58619.0}, po["capitale_necessario"]
        assert po["rita"]["rata_lorda"] == 50000 and po["rita"]["aliquota"] == 0.105 and po["anni_ponte"] == 7
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test("T12 rami avanzati: curvatura del glide path, buffer del decumulo, coefficiente a età frazionaria, sensibilità al PIL")
def t_rami_avanzati():
    motore = _modulo_obiettivi("motore_bootstrap.py")
    # CURVATURA 0,5: a metà finestra la frazione trattenuta supera quella lineare (sqrt(0,5) > 0,5)
    lin, curva = motore.frazioni_glide(10, 10, 1.0), motore.frazioni_glide(10, 10, 0.5)
    assert abs(lin[5] - 0.5) < 1e-9 and abs(curva[5] - 0.5 ** 0.5) < 1e-9 and curva[5] > lin[5], (lin[5], curva[5])
    assert all(c >= l for c, l in zip(curva, lin)) and curva[0] == lin[0] == 1.0
    gp = asset("it-pianificazione-obiettivi", "glidepath-obiettivi.py")
    tmp = tempfile.mkdtemp()
    try:
        g_lin = json.load(open(_json_di(gp, "--json", os.path.join(tmp, "l.json")), encoding="utf-8"))
        g_cur = json.load(open(_json_di(gp, "--set", "CURVATURA=0.5", "--json", os.path.join(tmp, "c.json")), encoding="utf-8"))
        lungo_l = next(o for o in g_lin["obiettivi"] if o["nome"].startswith("Lungo"))["quota_azionaria"]
        lungo_c = next(o for o in g_cur["obiettivi"] if o["nome"].startswith("Lungo"))["quota_azionaria"]
        # a 3 anni dalla scadenza (indice 17 su orizzonte 20, finestra di 7): lineare 0,65 × 3/7 ≈ 0,279, curva 0,65 × sqrt(3/7) ≈ 0,426
        assert abs(lungo_l[17] - 0.65 * 3 / 7) < 1e-3 and lungo_c[17] > lungo_l[17] and abs(lungo_c[17] - 0.65 * (3 / 7) ** 0.5) < 1e-3
        assert g_cur["curvatura"] == 0.5 and any("curvatura" in sc for sc in g_cur["scarti_dal_foglio"])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    # decumulo: il buffer cambia davvero i percorsi (0 contro 2 anni), probabilità in [0, 1]
    dec = _modulo_obiettivi("decumulo-obiettivi.py")
    p0, f0, _ = dec.monte_carlo(dec.PRELIEVO_ANNUO_REALE, 0)
    p2, f2, _ = dec.monte_carlo(dec.PRELIEVO_ANNUO_REALE, 2)
    assert 0 < p0 <= 1 and 0 < p2 <= 1 and (p0, f0) != (p2, f2), (p0, p2)
    # proiezione-inps: coefficiente interpolato sui dodicesimi e tabella di sensibilità monotona nel PIL
    pi = _modulo_previdenza("proiezione-inps.py")
    assert abs(pi.coeff_vigente(67.5) - (pi.COEFF_2025_2026[67] + pi.COEFF_2025_2026[68]) / 2) < 1e-12
    assert pi.coeff_vigente(67) == pi.COEFF_2025_2026[67] and pi.coeff_vigente(80) == pi.COEFF_2025_2026[71] and pi.coeff_vigente(50) == pi.COEFF_2025_2026[57]
    rc, out = run_script(asset("it-pianificazione", "proiezione-inps.py"))
    assert rc == 0 and "COEFF_ESPLICITO impostato" not in out, out[-300:]
    righe = [r for r in out.split("SENSIBILITÀ")[1].splitlines() if re.match(r"\s+\d+,\d%\s", r)]
    tabella = [[_da_it(x) for x in re.findall(r"(\d+,\d)%", r)] for r in righe]
    assert len(tabella) == 3 and all(len(r) == 4 for r in tabella), tabella
    for col in range(1, 4):
        assert tabella[0][col] < tabella[1][col] < tabella[2][col], f"colonna {col}: il tasso deve crescere col PIL {tabella}"
    for r in tabella:
        assert r[1] > r[2] > r[3], f"a parità di PIL, più riduzione dei coefficienti → tasso più basso {r}"


def _json_di(script, *args):
    """Lancia lo script e ritorna il percorso del JSON scritto (l'argomento dopo --json)."""
    rc, out = run_script(script, *args)
    assert rc == 0, out[-300:]
    return args[list(args).index("--json") + 1]


@test("T12 fixture realistiche: BOM, unità «(EUR)» e segno «-» nello storico del fondo, estratto INPS con xmlns")
def t_fixture_realistiche():
    m = carica_modulo(asset("it-pianificazione", "importa-previdenza.py"))
    quote, nav, vers = m.leggi_storico_fondo(os.path.join(FIXTURES, "storico-fondo-bom-unita-segno.csv"))
    assert quote == {"FONDO TEST COMPARTO AZIONARIO": 125.0, "FONDO TEST COMPARTO OBBLIGAZIONARIO": 150.0}, quote
    assert nav["FONDO TEST COMPARTO AZIONARIO"][1] == 10.4 and nav["FONDO TEST COMPARTO OBBLIGAZIONARIO"][1] == 5.2, nav
    assert vers[2025] == {"iscritto": 1000.0, "azienda": 0.0, "TFR": 1000.0}, dict(vers[2025])
    anag, righe = m.leggi_estratto_inps(os.path.join(FIXTURES, "estratto-inps-xmlns.xml"))
    assert anag["Cognome"] == "ROSSI" and len(righe) == 2 and sum(r["settimane"] for r in righe) == 104, (anag, righe)
    rc, out = run_script(asset("it-pianificazione", "importa-previdenza.py"), os.path.join(FIXTURES, "estratto-inps-xmlns.xml"),
                         os.path.join(FIXTURES, "storico-fondo-bom-unita-segno.csv"), "--anonimo")
    assert rc == 0 and "MONTANTE FONDO: ≈ € 2.080" in out and "\ufeff" not in out, out[-500:]
    # il BOM davanti alla prima intestazione non deve far perdere la colonna della data
    assert "colonne non riconosciute" not in out


def tabella_test():
    """Tabella Markdown per il README: per ogni strumento del plugin (script, Excel, CSV) i test che lo
    esercitano, cioè quelli che nominano il file, direttamente o tramite un helper o una costante."""
    import ast
    from collections import Counter
    sorgente = open(os.path.abspath(__file__), encoding="utf-8").read()
    albero = ast.parse(sorgente)
    funzioni, costanti = {}, {}   # nome → (nome del test o None, sorgente); costante → sorgente
    for nodo in albero.body:
        if isinstance(nodo, ast.FunctionDef):
            nome = next((d.args[0].value for d in nodo.decorator_list
                         if isinstance(d, ast.Call) and getattr(d.func, "id", "") == "test" and d.args), None)
            funzioni[nodo.name] = (nome, ast.get_source_segment(sorgente, nodo) or "")
        elif isinstance(nodo, ast.Assign):
            for t in nodo.targets:
                if isinstance(t, ast.Name):
                    costanti[t.id] = ast.get_source_segment(sorgente, nodo) or ""
    strumenti = []
    for skill in sorted(os.listdir(SKILLS)):
        cartella = os.path.join(SKILLS, skill, "assets")
        if os.path.isdir(cartella):
            for f in sorted(os.listdir(cartella)):
                if f.endswith((".py", ".xlsx", ".csv")) and (f != "motore_bootstrap.py" or skill == "it-pianificazione"):
                    strumenti.append((f"{skill}/assets/{f}", f))

    def nomina(src, simboli):
        return any(re.search(rf"\b{re.escape(x)}\b", src) for x in simboli)

    righe = ["| Strumento | Test che lo esercitano |", "|---|---|"]
    for rel, base in strumenti:
        simboli = {c for c, src in costanti.items() if base in src}
        toccano = {fn for fn, (nome, src) in funzioni.items() if base in src or nomina(src, simboli)}
        for _ in range(3):   # chiusura: chi chiama un helper che tocca il file tocca il file
            helper = {fn for fn in toccano if funzioni[fn][0] is None}
            toccano |= {fn for fn, (nome, src) in funzioni.items() if nomina(src, {h + "(" for h in helper}) or
                        any(re.search(rf"\b{h}\(", src) for h in helper)}
        nomi = Counter(funzioni[fn][0].split(":")[0] for fn in toccano if funzioni[fn][0])
        totale = sum(nomi.values())
        voci = [f"{n} ×{k}" if k > 1 else n for n, k in sorted(nomi.items())]
        righe.append(f"| `{rel}` | {totale}: " + "; ".join(voci[:6]) + (" …" if len(voci) > 6 else "") + " |")
    return "\n".join(righe)


@test("T12 README: la tabella strumenti ↔ test è quella generata dal runner (--tabella-test)")
def t_tabella_test():
    r = subprocess.run([sys.executable, os.path.abspath(__file__), "--tabella-test"],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
    assert r.returncode == 0, r.stderr[-300:]
    generata = r.stdout.strip()
    assert generata.startswith("| Strumento | Test che lo esercitano |"), generata[:200]
    righe = generata.splitlines()[2:]
    assert len(righe) >= 20, f"solo {len(righe)} strumenti"
    senza = [r for r in righe if "| 0: " in r]
    assert not senza, f"strumenti senza alcun test: {senza}"
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    marcatore = "<!-- generata con: python tests/run_tests.py --tabella-test -->"
    assert marcatore in readme, "README senza la tabella strumenti ↔ test"
    nel_readme = readme.split(marcatore)[1].strip().split("\n\n")[0].strip()
    assert nel_readme == generata, ("la tabella del README è diversa da quella generata: "
                                    "rigenerarla con python tests/run_tests.py --tabella-test")


@test("runner: --only/--skip filtrano i test per sottostringa del nome")
def t_filtro():
    me = os.path.abspath(__file__)
    r = subprocess.run([sys.executable, me, "--only", "freshness"],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
    assert "1/1 test passati" in r.stdout, r.stdout[-400:]
    assert "freshness fiscale" in r.stdout
    r = subprocess.run([sys.executable, me, "--skip", "freshness", "--only", "igiene"],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
    attesi = sum(1 for v in globals().values() if getattr(v, "__test__", False)
                 and "igiene" in v.__nome__.lower() and "freshness" not in v.__nome__.lower())
    assert f"{attesi}/{attesi} test passati" in r.stdout, r.stdout[-400:]
    assert "freshness" not in r.stdout
    r = subprocess.run([sys.executable, me, "--only", "nome-inesistente"],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
    assert r.returncode != 0, "nessun test selezionato deve essere un errore, non un verde vuoto"


@test("runner: la tabella dello stato dei dati normativi elenca i file obbligati con la scadenza")
def t_tabella_dati_normativi():
    r = subprocess.run([sys.executable, os.path.abspath(__file__), "--tabella-freshness"],
                       capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
    assert r.returncode == 0, r.stderr[-300:]
    assert r.stdout.startswith("| File | Verificato | Valido fino a |"), r.stdout[:200]
    for rel in OBBLIGATI_FRESHNESS:
        assert f"| `{rel}` |" in r.stdout, f"{rel} assente dalla tabella"


# ═══════════════════════ T14: authoring, documentazione e packaging ═══════════════════════

def _frontmatter(path):
    """(frontmatter, corpo) di un file Markdown: il frontmatter è tra i primi due «---»."""
    testo = open(path, encoding="utf-8").read()
    assert testo.startswith("---\n"), f"{path}: manca il frontmatter"
    _, fm, corpo = testo.split("---\n", 2)
    return fm, corpo


def _yaml_stretto(fm, nome):
    """Il sottoinsieme di YAML che usiamo, controllato senza dipendenze: chiave: valore, con i
    valori che contengono «: » o « #» tra virgolette doppie (un parser stretto come PyYAML
    altrimenti li rifiuta), e mappe indentate di due spazi. Restituisce {chiave: valore}."""
    out, padre = {}, None
    for riga in fm.splitlines():
        if not riga.strip():
            continue
        m = re.fullmatch(r"( {2})?([a-z][\w-]*):(?: (.*))?", riga)
        assert m, f"{nome}: riga non valida nel frontmatter «{riga[:60]}»"
        rientro, chiave, valore = m.groups()
        if valore is None:
            assert not rientro, f"{nome}: mappa annidata oltre il primo livello ({chiave})"
            padre = chiave
            out[chiave] = {}
            continue
        if valore.startswith('"'):
            assert re.fullmatch(r'"[^"\\]*"', valore), f"{nome}: {chiave} tra virgolette malformato"
            valore = valore[1:-1]
        else:
            assert ": " not in valore and " #" not in valore and not re.match(r"[\[\]{}>|*&!%@`'\"]", valore), \
                f"{nome}: il valore di {chiave} va tra virgolette doppie (YAML stretto)"
            # revisione T14: senza PyYAML in CI, «version: 0.20» passava come stringa; YAML la legge 0.2
            assert valore.lower() not in ("null", "~", "yes", "no", "on", "off"), f"{nome}: {chiave}: {valore!r} è ambiguo in YAML"
            if re.fullmatch(r"[-+]?\d+(\.\d*)?([eE][-+]?\d+)?", valore):
                valore = float(valore) if any(c in valore for c in ".eE") else int(valore)
            valore = {"true": True, "false": False}.get(valore, valore)
        if rientro:
            assert padre, f"{nome}: {chiave} indentato senza una mappa"
            out[padre][chiave] = valore
        else:
            padre = None
            out[chiave] = valore
    try:  # se PyYAML c'è, conferma con il parser vero
        import yaml
        assert yaml.safe_load(fm) == out, f"{nome}: PyYAML legge il frontmatter in modo diverso"
    except ImportError:
        pass
    return out


@test("T14 skill: description entro 1.024 caratteri, YAML stretto, metadata.version uguale a README e CHANGELOG")
def t_frontmatter_skill():
    # F023: sei description oltre 1.024 caratteri e «: » non quotati, che un parser YAML stretto
    # (quello dell'upload di skill fuori da Claude Code) rifiuta; nessuna versione per skill;
    # 202 righe di «Note per l'iterazione» caricate nel contesto a ogni uso della skill.
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    tabella = dict(re.findall(r"\| `(it-[\w-]+)` \| ✅ v([\d.]+)", readme))
    changelog = open(os.path.join(ROOT, "CHANGELOG.md"), encoding="utf-8").read()
    for skill in sorted(os.listdir(SKILLS)):
        fm, corpo = _frontmatter(os.path.join(SKILLS, skill, "SKILL.md"))
        d = _yaml_stretto(fm, skill)
        assert set(d) <= {"name", "description", "metadata", "when_to_use", "license"}, f"{skill}: chiavi {sorted(d)}"
        assert d["name"] == skill, f"{skill}: name {d['name']!r}"
        assert 200 <= len(d["description"]) <= 1024, f"{skill}: description di {len(d['description'])} caratteri"
        assert re.search(r"\bNON\b", d["description"]), f"{skill}: la description dichiara cosa non copre"
        ver = d["metadata"].get("version", "")
        assert re.fullmatch(r"\d+\.\d+", ver), f"{skill}: metadata.version {ver!r} (stringa «0.x»)"
        assert tabella.get(skill) == ver, f"{skill}: README dice v{tabella.get(skill)}, il frontmatter {ver}"
        sezione = changelog.split(f"\n### {skill}\n")[1]
        assert re.match(r"\s*#### v" + re.escape(ver) + r"\b", sezione), f"{skill}: CHANGELOG senza la voce v{ver} in testa"
        assert not re.search(r"^#{2,4} Note per|In quest'iterazione|Questa è la versione|^- \*?\*?v0\.\d+", corpo, re.M), \
            f"{skill}: note d'iterazione nel SKILL.md"
        # revisione T14: ogni asset della skill è spiegato nei Riferimenti (le note spostate in CHANGELOG
        # erano l'unico posto che descriveva ponte-pensione, netto-pensione, tracker e generatore dell'IPS)
        riferimenti = corpo.split("\n## Riferimenti")[1] if "\n## Riferimenti" in corpo else ""
        cartella_assets = os.path.join(SKILLS, skill, "assets")
        if os.path.isdir(cartella_assets):
            for f in sorted(os.listdir(cartella_assets)):
                if f.endswith((".py", ".xlsx")) and f != "motore_bootstrap.py":
                    assert f"assets/{f}" in riferimenti, f"{skill}: assets/{f} non è nei Riferimenti"
    # ricongiunzione e totalizzazione non sono gestite: fuori dai trigger, dichiarate escluse; il riscatto della
    # laurea dalla T13 sì (riscatto-laurea.py), quindi è un trigger e non più un'esclusione
    d = _yaml_stretto(_frontmatter(os.path.join(SKILLS, "it-pianificazione", "SKILL.md"))[0], "it-pianificazione")
    trigger, esclusioni = d["description"].split("NON copre")
    assert not re.search(r"ricongiunzione|totalizzazione", trigger), "it-pianificazione: trigger non gestiti"
    assert all(x in esclusioni for x in ("ricongiunzione", "totalizzazione")), esclusioni
    assert "riscatto della laurea" in trigger and "riscatto" not in esclusioni, (trigger[-300:], esclusioni)


@test("T14 comandi: alias sottili delle skill, uno per skill e come nel README")
def t_comandi_alias():
    # F057: i 6 comandi duplicavano a mano i workflow delle skill (già divergenti) ed entravano
    # tra gli strumenti che Claude può invocare da solo.
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    uso = dict(re.findall(r"^/([\w-]+)\s+→ (it-[\w-]+)", readme, re.M))
    cartella = os.path.join(ROOT, "commands")
    visti = {}
    for nome in sorted(os.listdir(cartella)):
        fm, corpo = _frontmatter(os.path.join(cartella, nome))
        d = _yaml_stretto(fm, f"commands/{nome}")
        comando = nome[:-3]
        assert d.get("disable-model-invocation") is True, \
            f"/{comando}: disable-model-invocation: true (Claude carica la skill, non il comando)"
        assert d.get("argument-hint", "").startswith("["), f"/{comando}: argument-hint"
        assert len(d.get("description", "")) <= 250, f"/{comando}: description lunga"
        skill = re.findall(r"`(it-[\w-]+)`", corpo)
        assert skill and len(set(skill)) == 1, f"/{comando}: deve rimandare a una sola skill, trovate {set(skill)}"
        assert os.path.isfile(os.path.join(SKILLS, skill[0], "SKILL.md")), f"/{comando}: skill {skill[0]} inesistente"
        assert "$ARGUMENTS" in corpo, f"/{comando}: senza $ARGUMENTS"
        assert len([r for r in corpo.splitlines() if r.strip()]) <= 4, f"/{comando}: non è un alias sottile"
        assert uso.get(comando) == skill[0], f"README: /{comando} → {uso.get(comando)}, il comando → {skill[0]}"
        visti[skill[0]] = comando
    assert set(visti) == set(os.listdir(SKILLS)), f"skill senza comando: {set(os.listdir(SKILLS)) - set(visti)}"


@test("T14 manifest: nessun campo sconosciuto in plugin.json, upstream in metadata, validate --strict in CI")
def t_manifest_strict():
    # F061: «upstream» al top level è un campo sconosciuto (warning, e rc 1 con --strict).
    # Campi della documentazione ufficiale, plugins-reference, tabella «Fields» (letta il 24/10/2026).
    noti = {"$schema", "name", "displayName", "version", "description", "author", "homepage", "repository",
            "license", "keywords", "metadata", "icon", "documentationUrl", "supportUrl", "privacyPolicyUrl",
            "termsOfServiceUrl", "defaultEnabled", "dependencies", "settings", "userConfig", "types", "channels",
            "skills", "commands", "agents", "hooks", "mcpServers", "lspServers", "outputStyles", "workflows",
            "experimental"}
    pl = json.load(open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8"))
    assert set(pl) <= noti, f"campi sconosciuti in plugin.json: {set(pl) - noti}"
    assert pl["metadata"]["upstream"]["repo"] == "anthropics/financial-services", pl.get("metadata")
    ci = open(os.path.join(ROOT, ".github", "workflows", "tests.yml"), encoding="utf-8").read()
    assert re.search(r"run: claude plugin validate \. --strict\b", ci), "job validate senza --strict"
    assert "CHANGELOG.md" in open(os.path.join(ROOT, "scripts", "pacchetto.py"), encoding="utf-8").read()


@test("T14 README e CHANGELOG: aggiornamenti spiegati, voce del changelog per la versione del manifest")
def t_changelog_aggiornamenti():
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    for atteso in ("claude plugin update wealth-management-it@wm-it", "--scope project", "Enable auto-update",
                   "/reload-plugins", "[CHANGELOG.md](CHANGELOG.md)"):
        assert atteso in readme, f"README: manca «{atteso}»"
    assert "aggiornamenti con un comando" not in readme, "README: l'auto-update è spento di default per questo marketplace"
    assert not re.search(r"^\*\*v0\.\d", readme, re.M), "README: lo storico delle versioni sta in CHANGELOG.md"
    ver = json.load(open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8"))["version"]
    changelog = open(os.path.join(ROOT, "CHANGELOG.md"), encoding="utf-8").read()
    versioni = re.findall(r"^## (v[\d.]+|Non rilasciato)", changelog, re.M)
    assert f"v{ver}" in versioni, f"CHANGELOG.md senza la voce v{ver} del manifest"
    assert versioni.index(f"v{ver}") <= 1 and versioni[0] in ("Non rilasciato", f"v{ver}"), \
        f"CHANGELOG.md: la voce più recente deve essere «Non rilasciato» o v{ver}, trovate {versioni[:3]}"
    assert "## Storia delle skill" in changelog


@test("T14 reference: indice delle sezioni nelle reference di almeno 100 righe")
def t_indici_reference():
    # F066: 21 reference da 100 righe in su senza indice; Claude ne legge l'inizio per decidere dove andare.
    for f in sorted(glob.glob(os.path.join(SKILLS, "*", "references", "*.md"))):
        testo = open(f, encoding="utf-8").read()
        if testo.count("\n") < 100:
            continue
        rel = os.path.relpath(f, SKILLS)
        assert "\n## Indice\n" in testo, f"{rel}: manca «## Indice»"
        h2, in_codice = [], False
        for r in testo.split("\n"):
            if r.startswith("```"):
                in_codice = not in_codice
            elif not in_codice and r.startswith("## ") and r != "## Indice":
                h2.append(r[3:].strip())
        voci = re.findall(r"^- (.+)$", testo.split("\n## Indice\n")[1].split("\n## ")[0], re.M)
        assert voci == h2, f"{rel}: l'indice non corrisponde alle sezioni ({len(voci)} voci, {len(h2)} sezioni)"
        assert testo.index("\n## Indice\n") < testo.index("\n## " + h2[0] + "\n"), f"{rel}: l'indice va prima delle sezioni"


@test("T14 postura e fonti: leve come scenari in skill e reference, certificati e compensazione come nella norma")
def t_postura_fonti_t14():
    # Revisione T14: la rilettura del § 2.6 aveva corretto le skill ma non le reference condivise, e
    # nessun test ne vincolava l'esito; idem la verifica del § 2.7 sui certificati.
    pian = open(os.path.join(SKILLS, "it-pianificazione", "SKILL.md"), encoding="utf-8").read()
    auto = open(os.path.join(SKILLS, "it-pianificazione-autonomi", "SKILL.md"), encoding="utf-8").read()
    for nome, testo in (("it-pianificazione", pian), ("it-pianificazione-autonomi", auto)):
        passo5 = testo.split("### Step 5")[1].split("\n## ")[0]
        assert "scenari" in passo5 and "Raccomandazioni" not in passo5.split("\n")[0], f"{nome}: Step 5 come scenari"
        assert "raccomandazioni d'azione" not in testo.lower() and "Raccomandazioni d'azione" not in testo, nome
    p1 = next(r for r in pian.splitlines() if "Pattern P1" in r)
    assert "non un importo consigliato" in p1, p1
    vietate = ("Massimizzare contributo deducibile", "fino a saturarlo", "Praticamente impossibile fare di meglio",
               "DEVE versare", "A: fondo, sempre", "quasi sempre la scelta", "opportunità prioritaria",
               "prima di raccomandare", "raddoppia tipicamente")
    for cartella, _, files in os.walk(SKILLS):
        for f in files:
            if f.endswith(".md"):
                testo = open(os.path.join(cartella, f), encoding="utf-8").read()
                for v in vietate:
                    assert v not in testo, f"{os.path.relpath(os.path.join(cartella, f), SKILLS)}: «{v}»"
                assert not re.search(r"skill v0\.\d|\bv0\.\d+ non ", testo), f"{f}: versione di skill superata nel testo"
    imp = open(asset("it-tax-optimization", "importa-estratto-broker.py"), encoding="utf-8").read()
    assert "vendere per caricare" not in imp and "venderle caricherebbe lo zainetto" in imp
    assert "ponte-pensione.py" in pian.split("### Step 5")[1] and "4-bis" in pian and "Uscita anticipata e RITA" in pian
    # § 2.7 e revisione T14: le cinque lettere dell'art. 68 c. 5 e l'eccezione dei certificati protetti
    tax = open(os.path.join(SKILLS, "it-tax-optimization", "SKILL.md"), encoding="utf-8").read()
    riga = next(r for r in tax.splitlines() if r.startswith("| **Reddito diverso**"))
    for atteso in ("lett. c, c-bis, c-ter, c-quater, c-quinquies", "art. 68 c. 4 e 5", "capitale protetto", "art. 44 c. 2 lett. c n. 2",
                   "dal 2027 art. 46 c. 2 lett. b n. 2"):
        assert atteso in riga, f"riga dei redditi diversi senza «{atteso}»"
    ips = open(asset("it-revisione", "genera-ips.py"), encoding="utf-8").read()
    marcatore = ips.split("# Verificato: 2026-10 — plus ETF")[1].split("POLICY_FISCALITA = [")[0]
    verificato, fuori = marcatore.split("# Fuori dal marcatore")
    assert "certificat" not in verificato and "Non verificata" not in ips, "certificati dentro il marcatore"
    assert "capitale protetto" in fuori and "c, c-bis, c-ter, c-quater e" in verificato, marcatore


@test("pacchetto: lo zip per Cowork contiene solo il plugin, nome = versione del manifest")
def t_pacchetto():
    import zipfile
    script = os.path.join(ROOT, "scripts", "pacchetto.py")
    ver = json.load(open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8"))["version"]
    tmp = tempfile.mkdtemp()
    try:
        rc, out = run_script(script, "--out", tmp, "--tag", f"v{ver}", cwd=ROOT)
        assert rc == 0, out[-400:]
        zpath = os.path.join(tmp, f"wealth-management-it-v{ver}.zip")
        assert os.path.exists(zpath), f"zip atteso: {zpath}; presenti: {os.listdir(tmp)}"
        nomi = zipfile.ZipFile(zpath).namelist()
        radice = "wealth-management-it/"
        assert all(n.startswith(radice) for n in nomi), "tutto deve stare sotto wealth-management-it/"
        for atteso in (radice + ".claude-plugin/plugin.json", radice + "README.md", radice + "CHANGELOG.md",
                       radice + "LICENSE", radice + "NOTICE",  # Apache-2.0 §4: licenza e attribuzione
                       radice + "commands/pensione.md",
                       radice + "skills/it-tax-optimization/SKILL.md",
                       radice + "skills/it-pianificazione/assets/proiezione-inps.py"):
            assert atteso in nomi, f"manca {atteso}"
        for vietato in ("/tests/", "/.github/", "/.remember/", "__pycache__", ".DS_Store",
                        "/.git/", "/scripts/", "/dist/", "/ordini/"):
            assert not any(vietato in n for n in nomi), f"nello zip non deve esserci {vietato}"
        # il tag deve coincidere con la versione del manifest
        rc, out = run_script(script, "--out", tmp, "--tag", "v9.9.9", cwd=ROOT)
        assert rc != 0 and "9.9.9" in out, "tag ≠ manifest deve fallire citando le versioni"
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ═══════════════════════════════ RUNNER ═══════════════════════════════

def seleziona(tests, only, skip):
    """Filtra i test per sottostringa (case-insensitive) del nome dichiarato."""
    def nome(fn):
        return fn.__nome__.lower()
    if only:
        tests = [fn for fn in tests if any(o.lower() in nome(fn) for o in only)]
    if skip:
        tests = [fn for fn in tests if not any(k.lower() in nome(fn) for k in skip)]
    return tests


def main():
    import argparse
    for flusso in (sys.stdout, sys.stderr):   # ✅ e ❌ anche su una console Windows in cp1252
        try:
            flusso.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description="Test suite del plugin wealth-management-it.")
    ap.add_argument("--only", action="append", metavar="NOME",
                    help="esegui solo i test il cui nome contiene NOME (ripetibile)")
    ap.add_argument("--skip", action="append", metavar="NOME",
                    help="salta i test il cui nome contiene NOME (ripetibile)")
    ap.add_argument("--tabella-freshness", action="store_true",
                    help="stampa la tabella Markdown dello stato dei dati normativi (per il README) ed esce")
    ap.add_argument("--tabella-test", action="store_true",
                    help="stampa la tabella Markdown strumenti ↔ test (per il README) ed esce")
    args = ap.parse_args()
    if args.tabella_freshness:
        print(tabella_freshness())
        sys.exit(0)
    if args.tabella_test:
        print(tabella_test())
        sys.exit(0)

    print("=" * 74)
    print("TEST SUITE — wealth-management-it")
    print("=" * 74)
    tests = [v for v in globals().values() if getattr(v, "__test__", False)]
    tests = seleziona(tests, args.only, args.skip)
    if not tests:
        print("Nessun test selezionato: controlla --only/--skip.")
        sys.exit(2)
    for fn in tests:
        fn()
    ok = sum(1 for _, esito, _ in RISULTATI if esito)
    saltati = sum(1 for _, esito, _ in RISULTATI if esito is None)
    eseguiti = len(RISULTATI) - saltati
    for nome, esito, dett in RISULTATI:
        if esito is None:
            print(f"  ⏭️  {nome} — saltato: {dett}")
            continue
        print(("  ✅ " if esito else "  ❌ ") + nome)
        if not esito:
            print("       → " + dett)
    print("-" * 74)
    print(f"{ok}/{eseguiti} test passati" + (f", {saltati} saltati" if saltati else ""))
    sys.exit(0 if ok == eseguiti else 1)


if __name__ == "__main__":
    main()
